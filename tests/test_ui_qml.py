"""The QML window, loaded offscreen with a stubbed bridge: it must load without
warnings, and the welcome screen, progress bar and validators must behave."""
import os
import re
from datetime import date, timedelta

import pytest

pytest.importorskip("PySide6")
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QObject, QUrl  # noqa: E402
from PySide6.QtQml import QQmlApplicationEngine, QQmlExpression, qmlContext  # noqa: E402
from PySide6.QtQuick import QQuickItem  # noqa: E402
from PySide6.QtTest import QTest  # noqa: E402

from jobhelper import config, notify, qt_bridge, schedule, store, ui_api  # noqa: E402
from tests.conftest import ROOT  # noqa: E402

READY = {"claude": True, "plan": "pro", "claude_email": "me@example.com", "paid_api": False,
         "email": True, "email_addr": "jobs@example.com", "cv": "cv.docx",
         "vault": "ok", "changed_files": []}


def qml_source() -> str:
    """The text of every QML file of the window (it is split into one file per window)."""
    return "\n".join(p.read_text(encoding="utf-8") for p in sorted(list((ROOT / "ui_qt").glob("*.qml")) + list((ROOT / "ui_qt").glob("*.js"))))


def _wait(ms):
    """Let the window run for `ms`. In short steps: under test the window's loop keeps Python to
    itself for the whole of one wait (the running app does not), and a background thread, such
    as the connection check, gets its turn only between two waits."""
    import time
    end = time.monotonic() + ms / 1000
    while True:
        left = end - time.monotonic()
        if left <= 0:
            break
        QTest.qWait(max(1, min(8, round(left * 1000))))
        time.sleep(0.002)                  # a real pause: this is when background threads run


class Ui:
    """A loaded window plus helpers to poke it."""

    def __init__(self, engine, bridge, conn, warnings):
        self.engine, self.bridge, self.conn, self.warnings = engine, bridge, conn, warnings
        self.win = engine.rootObjects()[0]

    def find(self, name):
        obj = self.win.findChild(QObject, name)
        assert obj is not None, f"no QML object named {name!r}"
        return obj

    def seen(self, root="statusPanel"):
        """Every item drawn under a part of the window, rows made by a Repeater included
        (those are children on screen but not children of the object tree, so find() misses them)."""
        todo, out = [self.win.findChild(QQuickItem, root)], []
        while todo:
            item = todo.pop()
            out.append(item)
            todo.extend(item.childItems())
        return out

    def drawn(self, name, root="statusPanel"):
        for item in self.seen(root):
            if item.objectName() == name:
                return item
        raise AssertionError(f"nothing named {name!r} is drawn in {root}")

    # each window lives in its own file; names inside a window are only known in that window's context
    WINDOWS = ["tableDialog", "detailDialog", "compareDialog", "formDialog", "assistantDialog", "prepEditDialog", "helpDialog", "prepDialog", "letterDialog",
               "responseDialog", "limitDialog", "welcome", "optDialog", "schedDialog",
               "navPanel", "stagePanel", "statusPanel", "pipelineBar"]      # the parts of the main screen

    def _inside(self, obj):
        """The context of the file an object was written in. For a window made in Main.qml
        that is not the object's own context (Main's) but the one its children were created in."""
        own = qmlContext(obj)
        if obj is self.win:
            return own
        for child in obj.findChildren(QObject):
            ctx = qmlContext(child)
            if ctx is not None and ctx != own and ctx.parentContext() == own:
                return ctx
        return own

    def js(self, code, scope=None):
        """Evaluate in the main window's context; if a name is only known inside one of
        the windows (its text fields, lists …), evaluate there instead."""
        targets = [scope] if scope else [self.win] + [self.win.findChild(QObject, name) for name in self.WINDOWS]
        error = ""
        for target in targets:
            if target is None:
                continue
            expr = QQmlExpression(self._inside(target), target, code)   # that file's context, so its ids resolve
            value = expr.evaluate()
            if not expr.hasError():
                return value[0] if isinstance(value, tuple) else value
            error = error or expr.error().toString()
            if "ReferenceError" not in expr.error().toString():
                break
        raise AssertionError(error)

    def wait(self, ms=60):
        _wait(ms)


@pytest.fixture(scope="module")
def ui(qapp, tmp_path_factory):
    data = tmp_path_factory.mktemp("ui_data")
    mp = pytest.MonkeyPatch()
    for name, value in {
        "DATA_DIR": data, "CV_DIR": data / "cv", "LETTERS_DIR": data / "cover_letters",
        "TAILORED_CV_DIR": data / "tailored_cvs", "JOBS_JSON": data / "jobs.json",
        "MY_LETTERS_DIR": data / "my_letters", "COVER_LETTERS_DIR": data / "my_letters",
        "EXCEL_PATH": data / "applications.xlsx", "NEW_JOBS_EXCEL": data / "new_jobs.xlsx",
        "PROFILE_JSON": data / "profile.json", "USAGE_JSON": data / "usage.json",
        "RESPONSES_DB": data / "responses.db", "RESPONSES_REPORT": data / "responses_report.xlsx",
        "ENV_PATH": data / ".env", "MODEL": config.MODEL, "MIN_LEAD_SCORE": config.MIN_LEAD_SCORE,
        "IMAP_USER": "", "IMAP_PASSWORD": "",
    }.items():
        mp.setattr(config, name, value)
    conn = dict(READY)
    mp.setattr(ui_api.Api, "connections", lambda self: dict(conn))
    mp.setattr(qt_bridge.Bridge, "_ping_usage", lambda self: self.usageReady.emit())   # no Claude call
    sent = []
    mp.setattr(notify, "send", lambda message, subtitle="": sent.append((message, subtitle)) or True)   # no real notification
    mp.setattr(schedule, "read_schedule",
               lambda: {"enabled": False, "weekday": 1, "hour": 9, "minute": 0, "label": "Off"})
    config.ensure_dirs()
    store.save_jobs([
        store.new_job(id="lead", title="Junior Python Developer", company="Acme", match_score=80,
                      match_tier="strong", status="letter_ready"),
        store.new_job(id="app", title="QA Engineer", company="Beta", status="waiting",
                      date_applied=(date.today() - timedelta(days=20)).isoformat(),
                      description="We use Python, Docker and Kubernetes.", match_score=64, match_tier="strong",
                      match_reason="Solid Python background."),
    ])

    warnings = []
    engine = QQmlApplicationEngine()
    engine.warnings.connect(lambda ws: warnings.extend(w.toString() for w in ws))
    bridge = qt_bridge.Bridge()
    engine.rootContext().setContextProperty("bridge", bridge)
    engine.load(QUrl.fromLocalFile(str(ROOT / "ui_qt" / "Main.qml")))
    assert engine.rootObjects(), "Main.qml failed to load: %s" % warnings
    handle = Ui(engine, bridge, conn, warnings)
    handle.sent = sent
    handle.win.show()
    _wait(2300)                            # splash → welcome card
    for _ in range(200):                   # the connection check is a thread; wait for it, not for a guessed time
        if handle.win.property("connChecked") and handle.find("welcome").property("phase") == "card":
            break
        _wait(30)
    _wait(700)                             # the card's own entrance
    yield handle
    # Closing as the app does: the window goes before the bridge, so no binding is left calling a bridge that is gone
    # (the terminal used to show "TypeError: Cannot call method … of null" on every exit).
    import app_qt
    from PySide6.QtCore import qInstallMessageHandler
    said = []
    earlier = qInstallMessageHandler(lambda kind, context, message: said.append(message))
    try:
        handle.win.close()
        app_qt.shut_down(engine)
        del engine, bridge
        QTest.qWait(50)
    finally:
        qInstallMessageHandler(earlier)
        mp.undo()
    assert not [m for m in said if "of null" in m or "TypeError" in m], said


def test_window_loads_without_qml_warnings(ui):
    assert ui.warnings == []
    assert ui.win.property("connChecked") is True


def test_counts_come_from_the_data(ui):
    assert ui.js("countsData.leads") == 1
    assert ui.js("countsData.applications") == 1


# ── welcome screen ──────────────────────────────────────────────────────────
def test_welcome_card_is_shown_ready_at_launch(ui):
    welcome = ui.find("welcome")
    assert welcome.property("phase") == "card"
    assert welcome.property("shown") is True and welcome.property("first") is True
    assert welcome.property("ready") is True
    assert ui.find("welcomeStart").property("enabled") is True


def test_start_enters_the_app_and_the_chip_reopens_the_card(ui):
    welcome = ui.find("welcome")
    ui.js("welcome.finish()")
    assert welcome.property("shown") is False
    ui.wait(700)
    assert welcome.property("visible") is False        # fully animated away
    ui.js("welcome.openCard()")
    assert welcome.property("shown") is True and welcome.property("first") is False
    assert ui.find("welcomeStart").property("text") == "Done"
    ui.js("welcome.finish()")
    ui.wait(700)


def test_welcome_is_not_ready_without_claude_or_cv(ui):
    welcome = ui.find("welcome")
    for missing in ({"claude": False}, {"cv": ""}, {"paid_api": True},
                    {"vault": "locked"}, {"changed_files": ["CV rules"]}):
        ui.win.setProperty("connData", {**READY, **missing})
        assert welcome.property("ready") is False, missing
    ui.win.setProperty("connData", dict(READY))
    assert welcome.property("ready") is True


def test_changed_files_block_start_until_approved(ui, monkeypatch):
    approved = []
    monkeypatch.setattr(qt_bridge.integrity, "approve", lambda names=None: approved.append(names) or ["cv_rules"])
    welcome = ui.find("welcome")
    ui.win.setProperty("connData", {**READY, "changed_files": ["CV rules", "Master CV"]})
    assert welcome.property("protectedOk") is False
    assert ui.find("welcomeStart").property("enabled") is (not welcome.property("first"))
    ui.js("bridge.approveChanges()")
    assert approved == [None]
    ui.win.setProperty("connData", dict(READY))
    assert welcome.property("protectedOk") is True and ui.warnings == []


def test_plain_storage_without_keychain_does_not_block(ui):
    welcome = ui.find("welcome")
    ui.win.setProperty("connData", {**READY, "vault": "off"})
    assert welcome.property("ready") is True
    ui.win.setProperty("connData", dict(READY))


def _email_answer(ui):
    """Saving a mailbox tries it in the background; wait for the card to get its answer."""
    welcome = ui.find("welcome")
    for _ in range(200):
        if not welcome.property("emailBusy"):
            return
        ui.wait(20)
    raise AssertionError("the mailbox check never answered")


def test_the_recovery_key_is_shown_and_restored_in_the_account_card(ui, monkeypatch):
    """Both needed a command in the terminal before; someone who only uses the window could not do either."""
    from jobhelper import vault
    monkeypatch.setattr(vault, "state", lambda: "ok")
    monkeypatch.setattr(vault, "recovery_key", lambda: "k" * 43 + "=")
    copied = []
    monkeypatch.setattr(qt_bridge.Bridge, "copyText", qt_bridge.Slot(str, result=bool)(lambda self, text: copied.append(text) or True))
    welcome = ui.find("welcome")
    ui.js("welcome.openCard()")
    ui.win.setProperty("connData", {**READY, "vault": "ok", "changed_files": []})
    ui.wait(80)
    assert welcome.property("keyText") == "" and ui.find("keyShow").property("text") == "Recovery key"   # never shown unasked
    ui.js("welcome.toggleKey()")
    assert welcome.property("keyText") == "k" * 43 + "=" and ui.find("keyShow").property("text") == "Hide key"
    assert ui.find("keyField").property("readOnly") is True and ui.find("keyField").property("text") == "k" * 43 + "="
    ui.js("welcome.keyCopied = bridge.copyText(welcome.keyText)")
    assert copied == ["k" * 43 + "="] and ui.find("keyCopy").property("text") == "Copied"
    ui.js("welcome.toggleKey()")                                   # Hide
    assert welcome.property("keyText") == ""
    ui.js("welcome.toggleKey(); welcome.finish()")                 # closing the card puts it away too
    ui.wait(700)
    assert welcome.property("keyText") == ""
    assert "interval: 60000; running: welcome.keyText" in qml_source(), "and it is put away by itself after a minute"
    # no key to show
    monkeypatch.setattr(vault, "state", lambda: "off")
    ui.js("welcome.openCard(); welcome.toggleKey()")
    assert welcome.property("keyText") == "" and "no key to show" in welcome.property("keyErr")

    # a new Mac: the saved key opens the data again
    used = []
    def restore(text):
        if text != "good-key":
            raise vault.VaultError("That key does not belong to this data." if text == "other-key" else "That is not a valid JobHelper key.")
        used.append(text)
    monkeypatch.setattr(vault, "restore_key", restore)
    ui.win.setProperty("connData", {**READY, "vault": "locked", "changed_files": []})
    field = ui.find("restoreField")
    for typed, part in (("", "Paste your recovery key first"), ("two words", "not a recovery key"), ("ключ", "not a recovery key"),
                        ("x" * 201, "not a recovery key"), ("bad-key", "not a recovery key for this data"), ("other-key", "does not belong")):
        field.setProperty("text", typed)
        ui.js("welcome.restoreKey()")
        assert part in welcome.property("restoreErr"), typed
        assert field.property("text") == "", "what was pasted is never left on screen"
        assert typed == "" or typed not in welcome.property("restoreErr")
    field.setProperty("text", "good-key")
    ui.js("welcome.restoreKey()")
    assert used == ["good-key"] and welcome.property("restoreErr") == ""
    ui.win.setProperty("connData", dict(READY))
    ui.js('welcome.keyErr = ""; welcome.finish()')
    ui.wait(700)
    assert "python status.py" not in (ROOT / "ui_qt" / "WelcomeOverlay.qml").read_text(encoding="utf-8")
    assert ui.warnings == []


def test_a_refused_app_password_is_not_saved(ui, monkeypatch):
    """Before, only the form of the password was checked, and a wrong one showed up at the first scan."""
    from jobhelper import email_client
    saved, tried = [], []
    monkeypatch.setattr(config, "set_imap", lambda user, pwd, lookback=None, host=None: saved.append((user, pwd)))
    monkeypatch.setattr(email_client, "try_login", lambda user, password, timeout=8.0, host="": tried.append((user, password)) or
                        {"ok": False, "msg": "The mail server refused this address and App Password. Check both."})
    welcome = ui.find("welcome")
    ui.js("welcome.openCard(); welcome.editBox = 1")
    welcome.setProperty("editEmail", True)
    welcome.setProperty("emailDraft", "jobs@example.com")
    welcome.setProperty("hostDraft", "imap.example.com")
    welcome.setProperty("pwDraft", "abcdabcdabcdabcd")
    ui.js("welcome.saveEmail()")
    ui.js("welcome.saveEmail()")                                   # a second press while it is checking does nothing
    _email_answer(ui)
    assert tried == [("jobs@example.com", "abcdabcdabcdabcd")] and saved == []
    assert "refused" in welcome.property("emailErr") and welcome.property("editEmail") is True
    assert welcome.property("pwDraft") == "" and ui.find("emailSave").property("text") == "Save"

    # no connection: saved, and the card says the password was not tried
    monkeypatch.setattr(email_client, "try_login", lambda user, password, timeout=8.0, host="":
                        {"ok": None, "msg": "The mail server could not be reached, so the password was not tried."})
    welcome.setProperty("pwDraft", "abcdabcdabcdabcd")
    ui.js("welcome.saveEmail()")
    _email_answer(ui)
    assert saved == [("jobs@example.com", "abcdabcdabcdabcd")] and welcome.property("editEmail") is False
    note = ui.find("emailNote")
    assert note.property("visible") is True and "not tried" in note.property("text")
    ui.js('welcome.emailNote = ""; welcome.finish()')
    ui.wait(700)
    assert ui.warnings == []


def test_welcome_rejects_a_bad_email_and_stays_open(ui, monkeypatch):
    saved = []
    monkeypatch.setattr(config, "set_imap", lambda user, pwd, lookback=None, host=None: saved.append((user, pwd)))
    welcome = ui.find("welcome")
    ui.js("welcome.openCard()")
    welcome.setProperty("editEmail", True)
    welcome.setProperty("hostDraft", "")
    for email, password in (("not-an-email", ""), ("jobs@gmail.com", "short"), ("jobs@gmail.com", "пароль"), ("jobs@gmail.com", "abcdabcdabcd1234"),
                            ("jobs@example.com", "secret123"), ("jobs@outlook.com", "secret123"), ("jobs@seznam.cz", "пароль")):
        welcome.setProperty("emailDraft", email)
        welcome.setProperty("pwDraft", password)
        ui.js("welcome.saveEmail()")
        assert welcome.property("emailErr") != "", (email, password)
        assert welcome.property("editEmail") is True
    assert saved == []
    welcome.setProperty("emailDraft", "jobs@gmail.com")
    assert welcome.property("needsServer") is False and "Google App Password" in ui.find("mailHint").property("text")
    welcome.setProperty("emailDraft", "jobs@example.com")          # a provider the app does not know: its server is asked for
    assert welcome.property("needsServer") is True and "imap.example.com" in ui.find("mailHint").property("text")
    assert "bridge.saveEmailChecked(editBox, emailDraft, pwDraft, needsServer ? hostDraft" in qml_source()
    welcome.setProperty("hostDraft", "imap.example.com")
    welcome.setProperty("pwDraft", "abcd abcd abcd abcd")          # spaces as Google shows it
    ui.js("welcome.saveEmail()")
    assert welcome.property("emailBusy") is True and ui.find("emailSave").property("text") == "Checking…"
    _email_answer(ui)
    assert saved == [("jobs@example.com", "abcd abcd abcd abcd")]
    assert welcome.property("emailErr") == "" and welcome.property("editEmail") is False
    ui.js("welcome.finish()")
    ui.wait(700)


# ── run progress ────────────────────────────────────────────────────────────
def _begin_search(ui):
    ui.js('resultFade.stop(); runMode = "search"; runResult = ""; runProgress = 0; stageKey = ""; resetRobots()')
    ui.bridge._set_running(True)


def test_progress_follows_stages_and_job_counts(ui):
    b = ui.bridge
    _begin_search(ui)
    b.agentChanged.emit("a1", "running", "")
    assert ui.win.property("stageKey") == "a1" and ui.find("a1").property("active") is True
    b.agentChanged.emit("a1", "done", "junior")
    b.agentChanged.emit("a2", "running", "")
    b.progressChanged.emit("a2", 50, 100)
    assert ui.win.property("runProgress") == pytest.approx(0.04 + 0.24 * 0.5)
    b.agentChanged.emit("a2", "done", "+14 new")
    b.agentChanged.emit("a3", "running", "")
    b.progressChanged.emit("a3", 7, 14)
    assert ui.win.property("runProgress") == pytest.approx(0.28 + 0.68 * 0.5)
    assert ui.find("a2").property("done") is True and ui.find("a3").property("active") is True
    b.progressChanged.emit("a2", 0, 100)                 # a late signal from a finished stage
    assert ui.win.property("runProgress") == pytest.approx(0.28 + 0.68 * 0.5)
    b.runFinished.emit("search", "Done · +14 new jobs, 6 letters + CVs ready.")
    b._set_running(False)
    assert ui.win.property("runProgress") == pytest.approx(1.0)
    assert ui.win.property("runResult") == "ok" and ui.win.property("stageKey") == ""
    assert ui.warnings == []


def test_failed_run_is_marked_as_error_not_complete(ui):
    b = ui.bridge
    _begin_search(ui)
    b.agentChanged.emit("a2", "running", "")
    b.runFinished.emit("search", "Error: No usable job source")
    b._set_running(False)
    assert ui.win.property("runResult") == "error"
    assert ui.win.property("runProgress") < 1
    ui.js('resultFade.stop(); runResult = ""; runMode = ""; runProgress = 0')


def test_elapsed_time_format(ui):
    assert ui.js("fmtElapsed(0)") == "00:00"
    assert ui.js("fmtElapsed(75)") == "01:15"
    assert ui.js("fmtElapsed(3605)") == "60:05"


# ── dialogs and views ───────────────────────────────────────────────────────
@pytest.mark.parametrize("name", ["optDialog", "schedDialog"])
def test_dialogs_open_and_close(ui, name):
    dialog = ui.find(name)
    ui.js("open()", dialog)
    ui.wait(500)
    assert dialog.property("opened") is True
    ui.js("close()", dialog)
    ui.wait(400)
    assert dialog.property("visible") is False
    assert ui.warnings == []


@pytest.mark.parametrize("view, rows", [("new", 1), ("app", 1), ("resp", 0), ("ana", 21)])       # 2 headings + 5 lines for the last 7 days + 14 all-time rows
def test_views_list_the_data(ui, view, rows):
    ui.js(f'openView("{view}")')
    ui.wait(450)
    assert ui.find("tableDialog").property("opened") is True
    assert ui.find("listView").property("count") == rows
    ui.js("close()", ui.find("tableDialog"))
    ui.wait(350)
    assert ui.warnings == []


def test_both_themes_render_without_warnings(ui):
    for dark in (False, True):
        ui.win.setProperty("dark", dark)
        ui.wait(120)
    assert ui.warnings == []


# ── input validation helpers ────────────────────────────────────────────────
@pytest.mark.parametrize("code, expected", [
    ('isEmail("you@gmail.com")', True),
    ('isEmail("you@gmail")', False),
    ('isEmail("you gmail.com")', False),
    ('isAppPassword("abcd efgh ijkl mnop")', True),
    ('isAppPassword("abcdefgh1jklmnop")', False),
    ('isAppPassword("short")', False),
    ('hasNonLatin("привет")', True),
    ('hasNonLatin("エンジニア")', True),
    ('hasNonLatin("Plzeň developer")', False),
    ('looksLikeRole("junior python developer")', True),
    ('looksLikeRole("asdfgh qwerty")', False),
    ('isSlug("proton-ag")', True),
    ('isSlug("no spaces")', False),
    ('cityEN("Praha")', "Prague"),
    ('cityEN("")', ""),
    ('cityEN("Berlin")', "Berlin"),
])
def test_validators(ui, code, expected):
    assert ui.js(code) == expected


# ── job detail, follow-up, stop ─────────────────────────────────────────────
def test_follow_up_count_is_shown_in_the_sidebar(ui):
    assert ui.js("countsData.followups") == 1


def test_detail_view_shows_score_keywords_and_follow_up(ui, monkeypatch):
    monkeypatch.setattr(ui_api.Api, "_cv_text", lambda self: "Python and Docker")
    dialog = ui.find("detailDialog")
    ui.js('detailDialog.show("app")')
    ui.wait(500)
    assert dialog.property("opened") is True
    assert ui.find("detailTitle").property("text") == "QA Engineer"
    assert ui.js("detailDialog.job.reason") == "Solid Python background."
    assert ui.js("detailDialog.job.matched.join(',')") == "Python,Docker"
    assert ui.js("detailDialog.job.missing.join(',')") == "Kubernetes"
    assert ui.js("detailDialog.fu.due") is True
    ui.js("detailDialog.draft = bridge.followupDraft(detailDialog.job.id)")
    ui.wait(100)
    assert "QA Engineer" in ui.find("draftBody").property("text")
    assert ui.js("bridge.markFollowedUp(detailDialog.job.id)") is True
    ui.js("detailDialog.refresh(); countsData = bridge.counts()")
    assert ui.js("detailDialog.fu.due") is False and ui.js("countsData.followups") == 0
    ui.js("detailDialog.close()")
    ui.wait(400)
    assert ui.warnings == []


def test_detail_view_ignores_an_unknown_job(ui):
    ui.js('detailDialog.show("does-not-exist")')
    ui.wait(150)
    assert ui.find("detailDialog").property("visible") is False


def test_run_button_becomes_stop_while_running(ui, monkeypatch):
    button = ui.find("runButton")
    assert button.property("text") == "▶ Run search"
    _begin_search(ui)
    assert button.property("text") == "■ Stop" and button.property("enabled") is True
    ui.js("stopRun()")
    assert ui.win.property("stopping") is True and ui.bridge._stop.is_set()
    assert button.property("text") == "Stopping…" and button.property("enabled") is False
    ui.bridge.agentChanged.emit("a3", "running", "")
    ui.bridge.runFinished.emit("search", "Stopped. Everything finished before that is saved.")
    ui.bridge._set_running(False)
    assert ui.win.property("stopping") is False and ui.win.property("runResult") == "error"
    assert ui.find("a3").property("active") is False          # no robot keeps working after a stop
    assert button.property("text") == "▶ Run search"
    ui.bridge._stop.clear()
    ui.js('resultFade.stop(); runResult = ""; runMode = ""; runProgress = 0')


def test_email_scan_shows_real_progress(ui):
    ui.js('resultFade.stop(); runMode = "scan"; runResult = ""; runProgress = 0; stageKey = ""')
    ui.bridge._set_running(True)
    ui.bridge.agentChanged.emit("a5", "running", "")
    ui.bridge.progressChanged.emit("a5", 3, 12)
    assert ui.win.property("runProgress") == pytest.approx(0.25)
    ui.bridge.runFinished.emit("scan", "Scanned 12 emails · 1 new (0 rejections).")
    ui.bridge._set_running(False)
    assert ui.win.property("runProgress") == pytest.approx(1.0)
    ui.js('resultFade.stop(); runResult = ""; runMode = ""; runProgress = 0')


def test_follow_up_days_option_is_loaded_and_saved(ui):
    dialog = ui.find("optDialog")
    ui.js("open()", dialog)
    ui.wait(450)
    spin = ui.find("fuSpin")
    assert spin.property("value") == 10
    ui.js("close()", dialog)
    ui.wait(350)
    ui.bridge.saveOptions({**ui.bridge.options(), "followup_days": 21})
    ui.js("optsData = bridge.options()")              # what Save does in the window
    ui.js("open()", dialog)
    ui.wait(450)
    assert spin.property("value") == 21
    ui.js("close()", dialog)
    ui.wait(350)


def test_scan_period_option_is_loaded_and_saved(ui):
    dialog = ui.find("optDialog")
    ui.bridge.saveOptions({**ui.bridge.options(), "scan_days": 45})
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    assert ui.find("scanSpin").property("value") == 45
    ui.js("close()", dialog)
    ui.wait(350)
    assert ui.warnings == []


# ── paused on Claude's limit ────────────────────────────────────────────────
def test_pause_shows_resume_warning_and_sound_then_resume_turns_blue(ui, monkeypatch):
    played = []
    monkeypatch.setattr(qt_bridge.subprocess, "Popen", lambda cmd, **k: played.append(cmd[0]))
    run, resume, dialog = ui.find("runButton"), ui.find("resumeButton"), ui.find("limitDialog")
    _begin_search(ui)
    ui.bridge.agentChanged.emit("a3", "running", "")
    info = {"paused": True, "can_resume": False, "window": "5-hour", "used": 99, "resets_at": "13:20",
            "resets_day": "07 Oct", "message": "limit"}
    ui.bridge.limitChanged.emit(info)
    ui.wait(450)
    assert ui.win.property("paused") is True and ui.win.property("working") is False
    assert run.property("visible") is False and resume.property("visible") is True
    assert ui.find("stopWhilePaused").property("visible") is True
    assert dialog.property("opened") is True and played == ["afplay"]
    assert "5-hour limit is used up (99%)" in ui.find("limitText").property("text")
    assert "13:20" in ui.find("limitText").property("text")
    assert ui.find("a3").property("active") is False                      # the robot rests while paused
    why = ui.win.property("limitWhy")
    assert "Can't resume yet" in why and "13:20" in why and "turns blue" in why
    grey = resume.property("color").name()
    ui.js("close()", dialog)
    ui.wait(350)

    ui.bridge.limitChanged.emit({**info, "can_resume": True})
    ui.wait(350)
    assert ui.win.property("canResume") is True and "limit is back" in ui.win.property("limitWhy")
    assert resume.property("color").name() != grey                       # grey -> blue
    assert dialog.property("visible") is False and played == ["afplay"]  # no second warning for the same pause

    ui.bridge.limitChanged.emit({**info, "paused": False, "can_resume": False})
    assert ui.win.property("paused") is False and run.property("visible") is True
    assert ui.find("a3").property("active") is True                       # the same agent continues
    ui.bridge.runFinished.emit("search", "Done · +0 new jobs, 0 letters + CVs ready.")
    ui.bridge._set_running(False)
    ui.js('resultFade.stop(); runResult = ""; runMode = ""; runProgress = 0')
    assert ui.warnings == []


# ── number fields and the Last run card ─────────────────────────────────────
def test_typed_number_is_saved_without_pressing_enter(ui):
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest as _QTest
    dialog, spin = ui.find("optDialog"), ui.find("scanSpin")
    ui.bridge.saveOptions({**ui.bridge.options(), "scan_days": 120})
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    assert spin.property("value") == 120
    assert spin.property("activeFocus") is False                 # nothing is pre-selected on open
    ui.js("contentItem.forceActiveFocus(); contentItem.selectAll()", spin)
    ui.win.requestActivate()
    for key in (Qt.Key_1, Qt.Key_0):                             # type over it, no Enter, no click elsewhere
        _QTest.keyClick(ui.win, key)
    ui.wait(80)
    assert spin.property("value") == 10
    assert spin.property("wheelEnabled") is False
    ui.js("close()", dialog)
    ui.wait(350)


def test_last_run_card_survives_a_restart(ui):
    ui.bridge._api.note_run(when="07 Oct 14:20", found="+12 new", letters=20, cvs=19)
    ui.bridge._api.note_run(scan="07 Oct · 2 new")
    assert ui.bridge.lastRun() == {"when": "07 Oct 14:20", "found": "+12 new", "letters": 20, "cvs": 19,
                                   "scan": "07 Oct · 2 new"}
    ui.js("reloadAll()")                                         # what the window does when it starts
    assert ui.js("lastRunData.letters") == 20 and ui.js("lastRunData.found") == "+12 new"
    # a long value must not break the layout (it is shortened inside the card)
    ui.js('lastRunData = ({ when: "07 Oct 14:20", found: "a very long value that would never fit in this narrow card at all" })')
    ui.wait(80)
    ui.js("lastRunData = bridge.lastRun()")
    assert ui.warnings == []


# ── status, apply date, merge and the second mailbox in the window ──────────
def test_status_buttons_and_apply_date_in_the_job_view(ui):
    ui.js('detailDialog.show("app")')
    ui.wait(450)
    assert ui.js("detailDialog.job.is_application") is True
    ui.js('detailDialog.setStatus("rejected")')
    assert ui.js("detailDialog.job.status") == "rejected"
    ui.js('detailDialog.setStatus("waiting")')
    field = ui.find("dateField")
    field.setProperty("text", "07.10.2026")
    ui.js("detailDialog.saveDate()")
    assert "YYYY-MM-DD" in ui.js("detailDialog.dateErr")                    # red message, window stays open
    assert ui.find("detailDialog").property("opened") is True
    field.setProperty("text", "2026-09-30")
    ui.js("detailDialog.saveDate()")
    assert ui.js("detailDialog.dateErr") == "" and ui.js("detailDialog.job.applied") == "2026-09-30"
    assert ui.js("detailDialog.job.applied_from_email") is False
    ui.js("detailDialog.close()")
    ui.wait(350)
    assert ui.warnings == []


def test_second_mailbox_row_follows_the_connection(ui):
    welcome = ui.find("welcome")
    ui.js("welcome.openCard()")
    ui.wait(300)
    ui.win.setProperty("connData", {**READY, "email2": False, "email2_addr": ""})
    assert ui.find("secondMailAdd").property("text") == "Add"
    ui.js("welcome.editBox = 2; welcome.editEmail = true")
    welcome.setProperty("emailDraft", READY["email_addr"])
    welcome.setProperty("pwDraft", "abcdabcdabcdabcd")
    ui.js("welcome.saveEmail()")
    _email_answer(ui)
    assert "already your main job email" in welcome.property("emailErr")    # refused, card stays open
    ui.js('welcome.editEmail = false; welcome.emailErr = ""; welcome.editBox = 1')
    ui.win.setProperty("connData", {**READY, "email2": True, "email2_addr": "second@example.com"})
    assert ui.find("secondMailAdd").property("text") == "Change"
    ui.win.setProperty("connData", dict(READY))
    ui.js("welcome.finish()")
    ui.wait(600)
    assert ui.warnings == []


# ── skip words, readable replies, analytics that add up ─────────────────────
def test_skip_words_and_exceptions_are_edited_loaded_and_saved(ui):
    dialog, skip, keep = ui.find("optDialog"), ui.find("skipTags"), ui.find("keepTags")
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    assert ui.js('all().join(",")', skip) == "intern,internship"          # the default
    assert ui.js('add("Student")', skip) is True and ui.js('add("student")', skip) is False
    assert "already" in skip.property("err")
    assert ui.js('add("стажёр")', skip) is False and "Latin" in skip.property("err")
    assert ui.js('add("7")', skip) is False
    assert ui.js('add("graduate program")', keep) is True
    assert ui.js('all().join(",")', skip) == "intern,internship,student"
    ui.js("close()", dialog)
    ui.wait(350)
    ui.bridge.saveOptions({**ui.bridge.options(), "skip_words": ["sales"], "keep_words": ["sales engineer"]})
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    assert ui.js('all().join(",")', skip) == "sales" and ui.js('all().join(",")', keep) == "sales engineer"
    ui.js("close()", dialog)
    ui.wait(350)
    ui.bridge.saveOptions({**ui.bridge.options(), "skip_words": ["intern", "internship"], "keep_words": []})
    assert ui.warnings == []


@pytest.mark.parametrize("roles, skips, keeps, clash", [
    (["software engineering intern"], ["intern"], [], True),
    (["Machine Learning Interns"], ["intern"], [], True),
    (["internal tools engineer"], ["intern"], [], False),
    (["junior python developer"], ["intern", "internship"], [], False),
    (["software engineering intern"], ["intern"], ["software engineering"], False),
    (["c++ developer"], ["c++"], [], True),
])
def test_a_search_role_may_not_contain_a_skipped_word(ui, roles, skips, keeps, clash):
    import json
    message = ui.js(f"roleClash({json.dumps(roles)}, {json.dumps(skips)}, {json.dumps(keeps)})", ui.find("optDialog"))
    assert (message != "") is clash
    if clash:
        assert roles[0] in message and "skip" in message


def test_click_a_reply_to_read_it_and_open_its_application(ui):
    from jobhelper import db
    db.upsert_response({"msg_id": "<r1@x>", "job_id": "app", "company": "Beta", "role": "QA Engineer", "stage": "rejection",
                        "received_date": "Mon, 5 Oct 2026 10:15:00 +0200", "subject": "Your application", "from_addr": "hr@beta.com",
                        "summary": "Beta chose another candidate.", "rejection_category": "went_with_other_candidate"})
    ui.js('openView("resp")')
    ui.wait(450)
    assert ui.js("listView.model[0].rid") == "<r1@x>"
    ui.js('responseDialog.show("<r1@x>")')
    ui.wait(450)
    assert ui.find("responseDialog").property("opened") is True
    assert ui.find("responseTitle").property("text") == "Beta"
    assert ui.find("responseSummary").property("text") == "Beta chose another candidate."
    assert ui.js("responseDialog.r.category") == "went with other candidate" and ui.js("responseDialog.r.date") == "2026-10-05 10:15"
    assert ui.js("responseDialog.r.job_title") == "QA Engineer"
    ui.js("responseDialog.close(); tableDialog.close()")
    ui.wait(400)
    ui.js('responseDialog.show("<missing>")')
    assert ui.find("responseDialog").property("visible") is False
    assert ui.warnings == []


def test_analytics_rows_add_up(ui):
    ui.js('openView("ana")')
    ui.wait(450)
    rows = {r["primary"]: int(r["value"]) for r in [
        {"primary": ui.js(f"listView.model[{i}].primary"), "value": ui.js(f"listView.model[{i}].value")}
        for i in range(ui.find("listView").property("count"))] if str(r["value"]).strip() != ""}
    assert rows["Applications"] == rows["↳ waiting for an answer"] + rows["↳ in interviews now"] + rows["↳ offers"] + rows["↳ rejected"]
    assert rows["Reply emails"] == (rows["↳ confirmations"] + rows["↳ interview emails"] + rows["↳ tests and tasks"]
                                    + rows["↳ rejections"] + rows["↳ other"])
    assert rows["Why the rejection emails said no"] == rows["↳ rejections"]
    ui.js("tableDialog.close()")
    ui.wait(350)


# ── search / filter / sort, not interested, letter editor ───────────────────
def _titles(ui):
    return [ui.js(f"listView.model[{i}].primary") for i in range(ui.find("listView").property("count"))]


def test_tables_can_be_searched_filtered_and_sorted(ui):
    from jobhelper import store as _store
    jobs = _store.load_jobs()
    extra = [_store.new_job(id="t1", title="Data Engineer", company="Zeta Corp", match_score=40, match_tier="stretch",
                            status="letter_ready", date_posted="2026-10-06", location="Brno"),
             _store.new_job(id="t2", title="Python Developer", company="Alpha", match_score=20, date_posted="2026-10-01"),
             _store.new_job(id="t3", title="QA Tester", company="Midway", date_posted="2026-10-07")]
    _store.save_jobs(jobs + extra)
    try:
        ui.js('openView("new")')
        ui.wait(450)
        dialog, search = ui.find("tableDialog"), ui.find("searchField")
        assert len(_titles(ui)) == 4 and ui.js("titleTxt.text", dialog).endswith("— 4")
        assert _titles(ui)[0].startswith("Junior Python Developer")              # best score first
        search.setProperty("text", "python")
        assert len(_titles(ui)) == 2 and ui.js("titleTxt.text", dialog).endswith("2 of 4")
        search.setProperty("text", "brno zeta")                                  # all words must match, place included
        assert _titles(ui) == ["Data Engineer  ·  Zeta Corp"]
        search.setProperty("text", "no such thing")
        assert _titles(ui) == []
        search.setProperty("text", "")
        ui.js('groupFilter = "stretch"; apply()', dialog)
        assert _titles(ui) == ["Data Engineer  ·  Zeta Corp"]
        ui.js('groupFilter = "not scored"; apply()', dialog)
        assert _titles(ui) == ["QA Tester  ·  Midway"]
        ui.js('groupFilter = ""; sortKey = "company"; apply()', dialog)
        assert [t.split("  ·  ")[1] for t in _titles(ui)] == ["Acme", "Alpha", "Midway", "Zeta Corp"]
        ui.js('sortKey = "date"; apply()', dialog)
        assert _titles(ui)[0].startswith("QA Tester") and _titles(ui)[1].startswith("Data Engineer")
        groups = ui.js("groups.map(function (g) { return g.key + ':' + g.n; }).join(',')", dialog)
        assert groups == "strong:1,stretch:1,no documents:1,not scored:1"
    finally:
        ui.js("tableDialog.close()")
        ui.wait(350)
        _store.save_jobs(jobs)
    assert ui.warnings == []


def test_not_interested_hides_the_lead_and_the_filter_restores_it(ui):
    from jobhelper import store as _store
    jobs = _store.load_jobs()
    _store.save_jobs(jobs + [_store.new_job(id="bye", title="Boring Role", company="Nope Ltd", match_score=10)])
    try:
        ui.js('openView("new")')
        ui.wait(450)
        assert any(t.startswith("Boring Role") for t in _titles(ui))
        ui.js('detailDialog.show("bye")')
        ui.wait(400)
        assert ui.find("notInterested").property("visible") is True
        assert ui.js('bridge.dismissJob("bye").ok') is True
        ui.js("detailDialog.close(); refreshOpenView()")
        ui.wait(400)
        assert not any(t.startswith("Boring Role") for t in _titles(ui)) and ui.js("countsData.leads") is not None
        ui.js('groupFilter = "not interested"; apply()', ui.find("tableDialog"))
        assert _titles(ui) == ["Boring Role  ·  Nope Ltd"] and ui.js("listView.model[0].actions[0].kind") == "restore"
        assert ui.js('bridge.restoreJob("bye").ok') is True
        ui.js("refreshOpenView()")
        ui.wait(300)
        assert ui.js("groupFilter", ui.find("tableDialog")) == ""                # that group is empty again
        assert any(t.startswith("Boring Role") for t in _titles(ui))
    finally:
        ui.js("tableDialog.close()")
        ui.wait(350)
        _store.save_jobs(jobs)
    assert ui.warnings == []


def test_letter_editor_edits_rewrites_and_saves(ui, monkeypatch):
    import shutil
    from jobhelper import cover_letter, llm, store as _store
    template = ROOT / "data" / "cover_letter_template.docx"
    if not template.exists():
        pytest.skip("letter template not present")
    monkeypatch.setattr(config, "LETTER_TEMPLATE", config.DATA_DIR / "cover_letter_template.docx")
    shutil.copy2(template, config.LETTER_TEMPLATE)
    monkeypatch.setattr(ui_api.Api, "_cv_text", lambda self: "Alex Novak\nPython, Docker.")
    paras = [("I am applying because the work matches what I build every day. " * 5).strip() for _ in range(5)]
    cover_letter.build_docx(config.LETTER_TEMPLATE, config.LETTERS_DIR / "L.docx", "QA Engineer", paras, date_text="7 October 2026")
    jobs = _store.load_jobs()
    lead = _store.find(jobs, "lead"); lead["cover_letter_file"] = "L.docx"
    _store.save_jobs(jobs)
    dialog = ui.find("letterDialog")
    ui.js('letterDialog.show("lead")')
    ui.wait(450)
    assert dialog.property("opened") is True and ui.js("letterDialog.texts.length") == 5
    assert ui.js("letterDialog.dirty") is False and ui.find("letterSave").property("enabled") is False

    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"paragraph": "I would like to bring the same drive to this team."})
    ui.js("letterDialog.rewrite(3)")
    assert ui.js("letterDialog.busyIndex") == 3
    for _ in range(100):
        ui.wait(30)
        if ui.js("letterDialog.busyIndex") == -1:
            break
    assert ui.js("letterDialog.err") == ""
    assert ui.js("letterDialog.texts[3]") == "I would like to bring the same drive to this team."
    assert ui.js("letterDialog.dirty") is True and "rewritten" in ui.js("letterDialog.info")
    assert cover_letter.read_letter(config.LETTERS_DIR / "L.docx")["paragraphs"][3] == paras[3]     # nothing saved yet

    ui.js("letterDialog.save()")
    assert ui.js("letterDialog.err") == "" and ui.js("letterDialog.dirty") is False
    assert cover_letter.read_letter(config.LETTERS_DIR / "L.docx")["paragraphs"][3] == "I would like to bring the same drive to this team."
    ui.js('letterDialog.setText(0, ""); letterDialog.save()')
    assert "5 paragraphs" in ui.js("letterDialog.err") and dialog.property("opened") is True        # red message, stays open
    ui.js("letterDialog.close()")
    ui.wait(350)
    lead["cover_letter_file"] = ""
    _store.save_jobs(jobs)
    assert ui.warnings == []


# ── insights and interview preparation ──────────────────────────────────────
def test_insights_view_lists_roles_keywords_and_skills(ui, monkeypatch):
    monkeypatch.setattr(ui_api.Api, "_cv_text", lambda self: "Python and Docker")
    ui.js('openView("ins")')
    ui.wait(450)
    rows = [(ui.js(f"listView.model[{i}].primary"), ui.js(f"listView.model[{i}].heading")) for i in range(ui.find("listView").property("count"))]
    heads = [text for text, heading in rows if heading is True]
    assert heads == ["What is working: kinds of role", "What is working: where the job came from", "Week by week",
                     "How long companies take to decide", "Search keywords", "Skills to learn", "Why jobs score low",
                     "What the postings pay"]
    assert any(str(text).startswith("Week of ") and str(text).endswith("(this week)") for text, _ in rows)
    assert any(text == "Kubernetes" for text, _ in rows)                           # asked for in the seeded posting, not in the CV
    assert ui.js("sorts.length", ui.find("tableDialog")) == 0                      # a plain list, no toolbar
    ui.js("tableDialog.close()")
    ui.wait(350)
    assert ui.warnings == []


def test_interview_preparation_from_the_job_window(ui, monkeypatch):
    from jobhelper import llm
    monkeypatch.setattr(ui_api.Api, "_cv_text", lambda self: "Alex Novak\nPython, Docker.")
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {
        "summary": "Mostly Python.", "ask_them": ["What is the team like?"], "revise": ["Docker"],
        "questions": [{"kind": "about you", "question": "Tell me about yourself.", "why": "Opening.", "answer": "I am a junior developer."},
                      {"kind": "technical", "question": "Kubernetes?", "why": "", "answer": "I have years of production experience with Kubernetes."}]})
    ui.js('detailDialog.show("app")')
    ui.wait(450)
    assert ui.js("detailDialog.job.has_prep") is False and ui.find("prepMake").property("text") == "Prepare"
    ui.js('detailDialog.prepErr = ""; detailDialog.prepBusy = true; detailDialog.prepJob = "app"; bridge.prepareInterview("app")')
    assert ui.find("prepMake").property("enabled") is False
    import time
    for _ in range(200):
        time.sleep(0.02)                                  # let the background worker run
        ui.wait(10)
        if ui.js("detailDialog.prepBusy") is False:
            break
    prep = ui.find("prepDialog")
    for _ in range(60):                                   # the window opens with an animation
        ui.wait(30)
        if prep.property("opened") is True:
            break
    assert ui.js("detailDialog.prepErr") == "", ui.js("detailDialog.prepErr")
    assert prep.property("opened") is True and ui.find("prepFor").property("text") == "Beta  ·  QA Engineer"
    assert ui.js("prepDialog.prep.questions.length") == 2
    assert ui.js("prepDialog.prep.questions[1].check") == "Kubernetes"             # an unbacked claim is flagged for you
    assert ui.js("detailDialog.job.has_prep") is True
    ui.js('prepDialog.note = "x"; prepDialog.close(); detailDialog.close()')
    ui.wait(400)
    ui.js('prepDialog.show("lead")')
    assert prep.property("visible") is False                                       # nothing saved for that job
    assert ui.warnings == []


# ── help, and the cost wording ──────────────────────────────────────────────
def test_help_opens_from_the_button_and_covers_the_app(ui):
    dialog = ui.find("helpDialog")
    assert ui.find("helpButton").property("text") == "Help"
    ui.js("open()", dialog)
    ui.wait(500)
    assert dialog.property("opened") is True
    titles = ui.js("topics.map(function (t) { return t.title; }).join('|')", dialog)
    for expected in ("What JobHelper does", "A normal week", "Scores and labels", "Claude and the limit", "Privacy and safety"):
        assert expected in titles
    text = ui.js("topics.map(function (t) { return t.lines.join(' '); }).join(' ')", dialog)
    assert "never sends" in text and "never invents experience" in text and "recovery key" in text
    assert "—" not in text                                             # en dash rule applies to the help too
    ui.js("close()", dialog)
    ui.wait(350)
    assert ui.warnings == []


def test_last_run_card_says_what_you_paid_and_the_api_price(ui):
    source = qml_source()
    assert '"You paid"' in source and '"Same work via API"' in source and '"subscription only"' in source
    assert "$0" not in source, "nothing was paid per run, so no price is shown next to 'You paid'"
    assert '"Total cost"' not in source and '"Cost (API-equiv.)"' not in source
    assert "usageData.cost" not in source, "the card shows the last run's cost, not the total of everything"


def _card_values(ui):
    return {t.property("text") for t in ui.seen() if isinstance(t.property("text"), str)}


def test_last_run_card_shows_the_cost_of_the_last_run_only(ui):
    ui.js('lastRunData = ({ when: "07 Oct 14:20", found: "+3 new" })')
    ui.wait(60)
    assert {"Same work via API", "after next run"} <= _card_values(ui)     # a summary saved before costs were noted
    ui.js('lastRunData = ({ when: "07 Oct 14:20", found: "+3 new", cost: 0.4, cost_of: "search" })')
    ui.wait(60)
    values = _card_values(ui)
    assert {"Same search via API", "≈ $0.40", "subscription only"} <= values and "Same work via API" not in values
    ui.js('lastRunData = ({ scan: "07 Oct · 0 new", cost: 0, cost_of: "email scan" })')
    ui.wait(60)
    assert {"Same scan via API", "≈ $0.00"} <= _card_values(ui)
    ui.js("lastRunData = bridge.lastRun()")
    assert ui.warnings == []


def test_last_run_card_says_how_much_mail_the_scan_went_through(ui):
    ui.js('lastRunData = ({ scan: "09 Oct · 0 new" })')
    ui.wait(60)
    assert "Emails read" not in _card_values(ui)                # a summary saved before this was noted: no empty row
    ui.js('lastRunData = ({ scan: "09 Oct · 0 new", scan_detail: "45 · 5 asked", cost: 0.08, cost_of: "email scan" })')
    ui.wait(60)
    assert {"Emails read", "45 · 5 asked", "09 Oct · 0 new"} <= _card_values(ui)
    ui.js("lastRunData = bridge.lastRun()")
    assert ui.warnings == []


def test_options_take_your_own_salary_figure_and_currency(ui):
    dialog = ui.find("optDialog")
    start = ui.bridge.options()
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    spin, combo, note = ui.find("salarySpin"), ui.find("currencyCombo"), ui.find("salaryNote")
    assert spin.property("value") == 0 and combo.property("currentText") == "CZK" and combo.property("count") >= 10
    ui.js('cityCombo.editText = "Prague"; countryField.text = ""', dialog)
    assert "software roles in Czechia only" in note.property("text")
    ui.js('cityCombo.editText = "Berlin"', dialog)
    assert "no estimate for this country" in note.property("text")
    spin.setProperty("value", 3200)
    ui.js('currencyCombo.currentIndex = currencyCombo.find("EUR")', dialog)
    assert "Form answers give this" in note.property("text")
    import json
    got = json.loads(ui.js("JSON.stringify(optDialog.collect().options)", dialog))
    assert (got["salary_target"], got["salary_currency"], got["where"]) == (3200, "EUR", "Berlin")
    assert ui.bridge.saveOptions(got)["ok"] is True
    ui.js("confirmClose = true; close()", dialog)
    ui.wait(350)

    # the job window and the list say it in that currency
    ui.bridge._api._change("lead", lambda jobs, j: j.update(salary="70 000 - 80 000 CZK", desired_monthly_czk=58000), render=False)
    detail = ui.bridge.jobDetail("lead")
    assert detail["desired_text"] == "3 200 EUR/mo target" and detail["pay_text"] == "posting says 2 800–3 200 EUR/mo"
    assert [r["desired"] for r in ui.bridge.leads() if r["id"] == "lead"] == ["3 200 EUR/mo"]
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    assert ui.find("salarySpin").property("value") == 3200 and ui.find("currencyCombo").property("currentText") == "EUR"   # loaded again
    ui.js("confirmClose = true; close()", dialog)
    ui.wait(350)
    # no figure of your own, in Berlin: nothing is shown, rather than a Prague figure
    assert ui.bridge.saveOptions({**ui.bridge.options(), "salary_target": 0})["ok"] is True
    assert ui.bridge.jobDetail("lead")["desired_text"] == ""
    assert ui.bridge.saveOptions({**ui.bridge.options(), "where": "Prague", "salary_currency": "CZK"})["ok"] is True
    detail = ui.bridge.jobDetail("lead")
    assert detail["desired_text"] == "58 000 CZK/mo target" and detail["pay_text"] == "posting says 70–80k CZK/mo"
    assert ui.bridge.saveOptions({**ui.bridge.options(), "where": start.get("where", ""), "country": start.get("country", "")})["ok"] is True
    ui.bridge._api._change("lead", lambda jobs, j: j.update(salary="", desired_monthly_czk=""), render=False)
    ui.js("optsData = bridge.options()")
    assert "CZK/mo" not in (ROOT / "ui_qt" / "DetailDialog.qml").read_text(encoding="utf-8") and "CZK/mo" not in (ROOT / "ui_qt" / "views.js").read_text(encoding="utf-8")
    assert ui.warnings == []


def test_options_has_a_country_for_any_city(ui):
    dialog = ui.find("optDialog")
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    field = ui.find("countryField")
    ui.js('cityCombo.editText = "Prague"', dialog)
    assert field.property("text") == "" and field.property("placeholderText") == "Czechia (from the city)"
    ui.js('cityCombo.editText = "Smallville"', dialog)
    assert field.property("placeholderText") == "not known for this city: type it"
    for wrong, part in (("Германия", "in English"), ("Germany 2", "no digits"), ("DE;", "only the country's name")):
        assert part in ui.js('optDialog.countryProblem("%s")' % wrong, dialog), wrong
    for fine in ("", "Germany", "United Kingdom", "Côte d'Ivoire", "St. Lucia", "Bosnia-Herzegovina"):
        assert ui.js('optDialog.countryProblem("%s")' % fine, dialog) == "", fine
    ui.js('cityCombo.editText = "Tbilisi"; countryField.text = "Georgia 1"', dialog)
    before = ui.bridge.options().get("country", "")
    assert "no digits" in ui.js("optDialog.countryProblem(countryField.text)", dialog)
    assert ui.js("optDialog.collect().options.country", dialog) == "Georgia 1"
    assert ui.bridge.saveOptions(ui.js("optDialog.collect().options", dialog) and {**ui.bridge.options(), "country": "Georgia 1"})["ok"] is False
    assert ui.bridge.options().get("country", "") == before             # the backend refuses the same thing
    ui.js('countryField.text = " Georgia "', dialog)
    assert ui.js("optDialog.collect().options.country", dialog) == "Georgia"
    ui.js("confirmClose = true; close()", dialog)
    ui.wait(350)
    ui.js("optsData = bridge.options()")
    assert ui.warnings == []


def test_options_can_start_one_longer_email_scan(ui, monkeypatch):
    started = []
    monkeypatch.setattr(qt_bridge.Bridge, "_scan", lambda self, days=None: started.append(days))
    monkeypatch.setattr(config, "IMAP_USER", "jobs@example.com")
    monkeypatch.setattr(config, "IMAP_PASSWORD", "abcdabcdabcdabcd")
    dialog = ui.find("optDialog")
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    saved_days = ui.find("scanSpin").property("value")
    spin, note = ui.find("backSpin"), ui.find("backNote")
    assert spin.property("value") == 90 and note.property("visible") is False
    spin.setProperty("value", 60)
    ui.js("optDialog.scanBack()", dialog)
    for _ in range(50):
        ui.wait(20)
        if started:
            break
    assert started == [60] and "last 60 days" in note.property("text") and dialog.property("backBad") is False
    assert ui.find("scanSpin").property("value") == saved_days and ui.bridge.options()["scan_days"] == saved_days   # the setting stays
    ui.js("optDialog.scanBack()", dialog)                          # the first one is still "running"
    assert dialog.property("backBad") is True and "running" in note.property("text") and started == [60]
    ui.bridge._set_running(False)
    ui.js('win.runMode = ""; win.resetRobots()')
    monkeypatch.setattr(config, "IMAP_PASSWORD", "")
    ui.js("optDialog.scanBack()", dialog)                          # no mailbox: said here, in red
    assert dialog.property("backBad") is True and "job email" in note.property("text") and started == [60]
    ui.js("confirmClose = false; close()", dialog)
    ui.wait(350)
    assert ui.warnings == []


def test_closing_buttons_ask_once_before_they_discard_unsaved_text(ui, monkeypatch):
    """Mark applied and Not interested close the job window; with text not saved they used to discard it silently."""
    done = []
    monkeypatch.setattr(qt_bridge.Bridge, "applyJob", qt_bridge.Slot(str, result=bool)(lambda self, job_id: done.append(("apply", job_id)) or True))
    monkeypatch.setattr(ui_api.Api, "dismiss_job", lambda self, job_id: done.append(("dismiss", job_id)) or {"ok": True})
    ui.js('detailDialog.show("lead")')
    ui.wait(450)
    note = ui.find("closeNote")
    ui.js('cnNotes.text = "They pay monthly."')
    ui.js("detailDialog.markApplied()")
    assert done == [] and ui.find("detailDialog").property("opened") is True
    assert note.property("text") == "Not saved yet: the notes about the company. Save it, or press Mark applied once more to discard it."
    ui.js("detailDialog.notInterested()")                       # another button: asked about again, by its own name
    assert done == [] and "press Not interested once more" in note.property("text")
    ui.js("detailDialog.notInterested()")                       # the second press of the same button goes ahead
    ui.wait(350)
    assert done == [("dismiss", "lead")] and ui.find("detailDialog").property("opened") is False
    ui.js('detailDialog.show("lead")')
    ui.wait(450)
    assert ui.js("detailDialog.confirmAction") == "" and note.property("visible") is False
    ui.js("detailDialog.markApplied()")                         # nothing typed: no question
    ui.wait(350)
    assert done[-1] == ("apply", "lead") and ui.find("detailDialog").property("opened") is False
    ui.js("tableDialog.close()")
    ui.wait(300)
    source = (ROOT / "ui_qt" / "DetailDialog.qml").read_text(encoding="utf-8")
    assert source.count("bridge.applyJob(") == 1 and source.count("bridge.dismissJob(") == 1, "both buttons go through the asking functions"
    assert ui.warnings == []


def test_usage_card_corrects_itself_when_a_reset_time_passes(ui):
    timer = ui.find("usageReset")
    keep = ui.js("JSON.stringify(usageData)")
    ui.js('usageData = ({ model: "sonnet", five_hour_used: 12, next_reset_in: 0 })')
    assert timer.property("running") is False
    ui.js('usageData = ({ model: "sonnet", five_hour_used: 100, five_hour_resets: "14:00", next_reset_in: 600 })')
    assert timer.property("running") is True and timer.property("interval") == 601500
    ui.js('usageData = ({ model: "sonnet", seven_day_used: 100, next_reset_in: 590000 })')      # days away: still a valid timer
    assert 0 < timer.property("interval") <= 2147483647
    ui.js('usageData = ({ model: "sonnet", five_hour_used: 100, five_hour_resets: "14:00", next_reset_in: 1 })')
    ui.wait(2900)                                               # the time passes: the card asks the app again
    assert ui.js("usageData.five_hour_used") != 100
    ui.js("usageData = JSON.parse('%s')" % keep.replace("\\", "\\\\").replace("'", "\\'"))
    assert ui.warnings == []


def test_a_long_choice_in_a_combo_box_wraps_instead_of_being_cut(ui):
    dialog = ui.find("optDialog")
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    ui.js("srcCombo.currentIndex = 1", dialog)                  # "LinkedIn only": short
    ui.wait(80)
    short = ui.js("srcCombo.height", dialog)
    assert short == 32 and ui.js("srcCombo.contentItem.lineCount", dialog) == 1
    ui.js("srcCombo.implicitWidth = 150", dialog)                # a box that cannot hold the longest choice on one line
    ui.js("srcCombo.currentIndex = 0", dialog)                  # "Czech boards, LinkedIn & careers"
    ui.wait(80)
    assert ui.js("srcCombo.contentItem.lineCount", dialog) >= 2 and ui.js("srcCombo.height", dialog) > short
    assert ui.js("srcCombo.contentItem.truncated", dialog) is False
    assert ui.js("srcCombo.contentItem.text", dialog) == "Czech boards, LinkedIn & careers"
    assert ui.js("srcCombo.contentItem.x + srcCombo.contentItem.width", dialog) <= ui.js("srcCombo.width", dialog)   # clear of the arrow's side
    # and it is really drawn: the box holds pixels in the text's colour (a blank box once passed every check above)
    line = ui.js("srcCombo.contentItem", dialog)
    assert line.property("visible") is True and line.property("opacity") == 1
    from PySide6.QtQuick import QQuickWindow
    corner = line.mapToScene(line.boundingRect().topLeft())
    picture = QQuickWindow.grabWindow(ui.win).copy(int(corner.x()), int(corner.y()), int(line.width()), int(line.height()))
    ink = ui.js("theme.ink")
    close = lambda c: abs(c.red() - ink.red()) + abs(c.green() - ink.green()) + abs(c.blue() - ink.blue()) < 90      # noqa: E731
    assert sum(close(picture.pixelColor(x, y)) for x in range(0, picture.width(), 2) for y in range(0, picture.height(), 2)) > 20
    assert ui.js("cityCombo.height", dialog) == 32               # an editable box keeps its one-line field
    ui.js("srcCombo.implicitWidth = 210; confirmClose = false; close()", dialog)
    ui.wait(350)
    assert ui.warnings == []


def test_usage_card_shows_nothing_left_and_when_it_resets(ui):
    keep = ui.js("JSON.stringify(usageData)")
    width = ui.find("statusPanel").property("width")
    ui.js('usageData = ({ model: "sonnet", five_hour_used: 100, five_hour_resets: "12 Oct 14:00", seven_day_used: 34, seven_day_resets: "12 Oct 09:00" })')
    ui.wait(60)
    five, seven = ui.drawn("usageLeft_five_hour"), ui.drawn("usageLeft_seven_day")
    assert five.property("text") == "0% left" and five.property("font").bold()
    resets = ui.drawn("usageResets_five_hour")
    assert resets.property("visible") is True and resets.property("text") == "resets 12 Oct 14:00"
    assert seven.property("text") == "66% left"
    assert ui.drawn("usageResets_seven_day").property("visible") is False      # the reset time matters only when little is left
    assert ui.find("statusPanel").property("width") == width, "a longer text must not widen the column past the window"
    assert resets.mapToItem(None, 0, 0).x() + resets.property("width") <= ui.win.width()
    ui.js('usageData = ({ model: "sonnet", five_hour_used: 85, five_hour_resets: "14:00" })')
    ui.wait(60)
    five, seven = ui.drawn("usageLeft_five_hour"), ui.drawn("usageLeft_seven_day")      # the rows are made again
    assert five.property("text") == "15% left" and not five.property("font").bold()
    assert ui.drawn("usageResets_five_hour").property("text") == "resets 14:00"
    assert seven.property("text") == "—"
    ui.js("usageData = JSON.parse('%s')" % keep.replace("\\", "\\\\").replace("'", "\\'"))
    assert ui.warnings == []


def test_the_right_column_fits_inside_the_window(ui):
    panel = ui.find("statusPanel")
    assert panel.mapToItem(None, 0, 0).x() + panel.property("width") <= ui.win.width()
    for item in ui.seen():
        if item.property("visible") and isinstance(item.property("text"), str) and item.property("text"):
            assert item.mapToItem(None, 0, 0).x() + item.property("width") <= ui.win.width() + 0.5, item.property("text")


# ── about you, notifications, backups, readable help ────────────────────────
def test_new_options_are_loaded_and_saved(ui):
    dialog = ui.find("optDialog")
    ui.bridge.saveOptions({**ui.bridge.options(), "about_right_to_work": "EU citizen", "about_start": "one month",
                           "notifications": False, "keep_backups": 25})
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    assert ui.find("rightField").property("text") == "EU citizen" and ui.find("startField").property("text") == "one month"
    assert ui.find("notifySw").property("checked") is False and ui.find("keepSpin").property("value") == 25
    ui.js("close()", dialog)
    ui.wait(350)
    ui.bridge.saveOptions({**ui.bridge.options(), "about_right_to_work": "", "about_start": "", "notifications": True, "keep_backups": 10})
    assert ui.warnings == []


def test_help_marks_key_words(ui):
    text = ui.js("topics.map(function (t) { return t.lines.join(' '); }).join(' ')", ui.find("helpDialog"))
    for word in ("Run search", "Scan email", "Mark applied", "strong", "stretch", "Not interested", "recovery key"):
        assert f"<b>{word}</b>" in text, word
    assert text.count("<b>") == text.count("</b>") and "<b><b>" not in text
    source = qml_source()
    assert "textFormat: Text.StyledText" in source


# ── dates, company notes, pasted replies, CV comparison, pay ────────────────
def _open_job(ui, job_id="app"):
    ui.js(f'detailDialog.show("{job_id}")')
    ui.wait(450)
    assert ui.find("detailDialog").property("opened") is True


def test_a_date_on_a_job_shows_on_the_main_screen_and_is_announced(ui):
    card = ui.find("upcomingCard")
    assert card.property("visible") is False                                       # nothing planned: no empty card
    _open_job(ui)
    ui.js('evDate.text = "next week"; detailDialog.saveEvent()')
    assert "YYYY-MM-DD" in ui.find("eventErr").property("text") and ui.find("eventErr").property("visible") is True
    assert ui.find("detailDialog").property("opened") is True and ui.js("upcomingData.length") == 0
    tomorrow = (date.today() + timedelta(days=1)).isoformat()
    ui.js(f'evDate.text = "{tomorrow}"; evTime.text = "14:30"; evNote.text = "Teams call"; detailDialog.eventKind = "call"; detailDialog.saveEvent()')
    assert ui.js("detailDialog.eventErr") == "" and ui.js("detailDialog.job.event.kind") == "call"
    assert ui.js("upcomingData.length") == 1 and ui.js("upcomingData[0].label") == "tomorrow"
    assert card.property("visible") is True
    ui.sent.clear()
    assert ui.js("bridge.remind()") == 1 and ui.js("bridge.remind()") == 0         # once a day
    assert ui.sent == [("Tomorrow at 14:30: call · Beta", "Reminder")]
    ui.js('detailDialog.show("app")')                                              # reopening shows what was saved
    ui.wait(200)
    assert ui.find("evDate").property("text") == tomorrow and ui.find("evNote").property("text") == "Teams call"
    ui.js("bridge.clearEvent(detailDialog.job.id); detailDialog.refresh(); upcomingData = bridge.upcoming()")
    assert card.property("visible") is False and ui.find("evDate").property("text") == ""
    ui.js("detailDialog.close()")
    ui.wait(350)
    assert ui.warnings == []


def test_company_notes_are_kept_for_the_company(ui):
    from jobhelper import planner
    _open_job(ui)
    ui.js('cnContact.text = "Jana, HR"; cnSalary.text = "55k"; cnAsked.text = "SQL joins"; cnNotes.text = "Nice team"; detailDialog.saveNotes()')
    assert ui.js("detailDialog.notesInfo") == "Saved for Beta."
    assert planner.company_notes("beta")["asked"] == "SQL joins"
    ui.js("detailDialog.close()")
    ui.wait(350)
    _open_job(ui)
    assert ui.find("cnContact").property("text") == "Jana, HR" and ui.find("cnNotes").property("text") == "Nice team"
    ui.js('cnContact.text = ""; cnSalary.text = ""; cnAsked.text = ""; cnNotes.text = ""; detailDialog.saveNotes()')
    assert planner.company_notes("Beta")["updated"] == ""
    ui.js("detailDialog.close()")
    ui.wait(350)
    assert ui.warnings == []


def test_a_pasted_reply_is_filed_from_the_job_window(ui, monkeypatch):
    import time

    from jobhelper import db, llm
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"stage": "acknowledged", "summary": "They got it."})
    _open_job(ui)
    assert ui.js("detailDialog.replyOpen") is False
    ui.js('detailDialog.replyOpen = true; replyText.text = "ok"; detailDialog.fileReply()')
    assert "too short" in ui.js("detailDialog.replyErr") and ui.js("detailDialog.replyBusy") is False
    ui.js('replyText.text = "Hello, thank you for applying. We received your application and will be in touch."; detailDialog.fileReply()')
    assert ui.js("detailDialog.replyBusy") is True
    for _ in range(200):
        time.sleep(0.02)                                  # let the background worker run
        ui.wait(10)
        if ui.js("detailDialog.replyBusy") is False:
            break
    assert ui.js("detailDialog.replyErr") == "", ui.js("detailDialog.replyErr")
    assert ui.js("detailDialog.replyInfo") == "Filed as: acknowledged." and ui.js("detailDialog.replyOpen") is False
    assert ui.find("replyText").property("text") == ""
    rows = [r for r in db.all_responses() if r["msg_id"].startswith("pasted-")]
    assert len(rows) == 1 and rows[0]["job_id"] == "app"
    ui.js("detailDialog.close()")
    ui.wait(350)
    assert ui.warnings == []


def test_cv_comparison_window(ui, monkeypatch):
    import docx

    from jobhelper import agent1_cv

    def make(path, lines):
        d = docx.Document()
        for line in lines:
            d.add_paragraph(line)
        d.save(str(path))
        return path

    master = make(config.DATA_DIR / "master.docx", ["Alex Novak", "PROFILE", "I know Python."])
    make(config.TAILORED_CV_DIR / "cv_lead.docx", ["Alex Novak", "PROFILE", "I know Python and SQL."])
    monkeypatch.setattr(agent1_cv, "find_cv", lambda: master)
    _open_job(ui, "lead")
    assert ui.find("compareCv").property("visible") is False                       # no adapted CV yet
    ui.js("compareDialog.show(detailDialog.job.id)")
    ui.wait(150)
    assert ui.find("compareDialog").property("visible") is False                   # and nothing opens
    ui.bridge._api._change("lead", lambda jobs, j: j.update(tailored_cv_file="cv_lead.docx", salary="50 000 - 60 000 CZK"), render=False)
    ui.js("detailDialog.refresh()")
    assert ui.find("compareCv").property("visible") is True
    assert ui.js("detailDialog.job.pay.low") == 50000
    ui.js("compareDialog.show(detailDialog.job.id)")
    ui.wait(500)
    dialog = ui.find("compareDialog")
    assert dialog.property("opened") is True
    assert "1 paragraph changed" in ui.find("compareSummary").property("text")
    assert ui.js("compareDialog.data.rows[0].adapted") == "I know Python<b> and SQL</b>."
    assert ui.js("compareDialog.data.rows[0].master") == "I know Python."
    ui.js("compareDialog.close()")
    ui.wait(300)
    ui.js("detailDialog.close()")
    ui.wait(350)
    assert ui.warnings == []


def test_insights_list_what_postings_pay(ui, monkeypatch):
    monkeypatch.setattr(ui_api.Api, "_cv_text", lambda self: "Python and Docker")
    ui.bridge._api._change("lead", lambda jobs, j: j.update(salary="50 000 - 60 000 CZK", desired_monthly_czk=55000), render=False)
    ui.js('openView("ins")')
    ui.wait(450)
    count = ui.find("listView").property("count")
    rows = [(ui.js(f"listView.model[{i}].primary"), ui.js(f"listView.model[{i}].secondary"), ui.js(f"listView.model[{i}].value")) for i in range(count)]
    pay = [r for r in rows if r[0] == "Junior Python Developer  ·  Acme"]
    assert pay and pay[0][2] == "60k" and "50–60k CZK" in pay[0][1] and "your target 55k: around your target" in pay[0][1]
    assert any("1 of 2 postings state it" in str(r[1]) for r in rows)
    ui.js("tableDialog.close()")
    ui.wait(350)
    assert ui.warnings == []


def test_scan_all_folders_switch_is_saved(ui):
    dialog = ui.find("optDialog")
    assert ui.bridge.options()["scan_all_folders"] is False
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    switch = ui.find("foldersSw")
    assert switch.property("checked") is False
    ui.bridge.saveOptions({**ui.bridge.options(), "scan_all_folders": True})
    ui.js("close()", dialog)
    ui.wait(350)
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    assert switch.property("checked") is True
    ui.js("close()", dialog)
    ui.wait(350)
    ui.bridge.saveOptions({**ui.bridge.options(), "scan_all_folders": False})
    assert ui.warnings == []


def test_application_only_cards_are_hidden_for_a_new_job_and_nothing_overflows(ui, monkeypatch):
    from jobhelper import planner
    for name in ("cv_lead.docx", "cv_lead.pdf"):
        (config.TAILORED_CV_DIR / name).write_bytes(b"x")
    for name in ("L_lead.docx", "L_lead.pdf"):
        (config.LETTERS_DIR / name).write_bytes(b"x")
    ui.bridge._api._change("lead", lambda jobs, j: j.update(tailored_cv_file="cv_lead.docx", cover_letter_file="L_lead.docx",
                                                             url="https://example.com/job"), render=False)
    monkeypatch.setattr(ui_api.Api, "_comparable", lambda self, j: True)
    monkeypatch.setattr(ui_api.Api, "_editable_letter", lambda self, j: True)
    _open_job(ui, "lead")
    for card in ("eventCard", "replyCard", "notesCard"):
        assert ui.find(card).property("visible") is False, card
    dialog_width = ui.js("detailDialog.width")
    footer = ui.find("detailFooter")
    assert footer.property("width") <= dialog_width                              # every button on, still inside the window
    assert ui.find("detailDocs").property("height") > 40                         # the document buttons wrapped to a second line
    assert ui.js("detailFlick.width") <= dialog_width
    planner.save_company_notes("Acme", {"contact": "Jana"})                      # notes kept from an earlier application
    ui.js("detailDialog.refresh()")
    assert ui.find("notesCard").property("visible") is True and ui.find("eventCard").property("visible") is False
    planner.save_company_notes("Acme", {})
    ui.js("detailDialog.close()")
    ui.wait(350)
    _open_job(ui, "app")
    for card in ("eventCard", "replyCard", "notesCard"):
        assert ui.find(card).property("visible") is True, card
    ui.js("detailDialog.close()")
    ui.wait(350)
    ui.bridge._api._change("lead", lambda jobs, j: j.update(tailored_cv_file="", cover_letter_file="", url=""), render=False)
    assert ui.warnings == []


def test_sidebar_counts_and_icons_share_one_column(ui):
    source = qml_source()
    assert 'objectName: "navTail_" + modelData.k; implicitWidth: 46' in source
    tail = source[source.index('objectName: "navTail_"'):][:1400]
    assert tail.count("anchors.centerIn: parent") >= 3                             # the pill, its number and the icon


def test_analytics_has_a_drawn_bar_chart_for_its_icon(ui):
    """A font's symbol for it could not be recognised at that size; three rising bars can."""
    icons = [i for i in ui.seen("navPanel") if i.objectName() == "navIcon_ana"]
    assert len(icons) == 5 and [i.property("visible") for i in icons].count(True) == 1     # one per row, shown on its own row only
    icon = next(i for i in icons if i.property("visible"))
    bars = sorted((b for b in icon.childItems() if b.property("radius") is not None), key=lambda b: b.property("x"))
    assert [b.property("height") for b in bars] == [6, 10, 14]                              # rising, left to right
    assert all(b.property("y") + b.property("height") == icon.property("height") for b in bars), "all stand on one line"
    assert all(b.property("x") + b.property("width") <= icon.property("width") for b in bars)
    assert "▦" not in qml_source()
    bulbs = [i for i in ui.seen("navPanel") if i.objectName() == "navIcon_ins" and i.property("visible")]
    assert len(bulbs) == 1 and len(bulbs[0].childItems()) == 3 and "✦" not in qml_source()    # drawn too: glass, neck, base
    assert ui.warnings == []


def test_filtered_out_list_opens_from_the_last_run_card(ui):
    from jobhelper import agent2_search
    agent2_search.save_dropped([
        {"title": "Software Engineering Intern", "company": "Acme", "source": "linkedin", "url": "https://example.com/1",
         "group": "skip word", "reason": 'The title has a word you skip ("intern").'},
        {"title": "Backend Developer", "company": "Beta", "source": "jobs.cz", "url": "",
         "group": "too many years", "reason": "Asks for 6+ years of experience."}])
    ui.bridge._api.note_run(filtered=2)
    ui.js("lastRunData = bridge.lastRun()")
    ui.wait(100)
    source = qml_source()
    assert 'lastRunData.filtered + " · see why ›"' in source and "win.openView(modelData.view)" in source
    ui.js('openView("skip")')
    ui.wait(450)
    assert ui.find("listView").property("count") == 2
    assert ui.js("listView.model[0].primary") == "Software Engineering Intern  ·  Acme"
    assert ui.js("listView.model[1].secondary") == "Asks for 6+ years of experience.   ·  jobs.cz"
    assert ui.js("listView.model[0].actions.length") == 1 and ui.js("listView.model[1].actions.length") == 0
    assert ui.js("tableDialog.groups.length") == 2                                # a filter chip per reason
    ui.js("tableDialog.close()")
    ui.wait(350)
    agent2_search.save_dropped([])
    assert ui.warnings == []


def test_polish_empty_filtered_row_older_master_note_and_old_letter_chip(ui, monkeypatch):
    source = qml_source()
    assert '" · see why ›" : "after next search"' in source                      # not a bare dash that looks broken
    assert 'letter_old_format === true ? "letter: older format, open in Word"' in source
    note = ui.find("olderMaster")
    monkeypatch.setattr(ui_api.Api, "cv_compare", lambda self, job_id: {
        "ok": True, "company": "Acme", "title": "Dev", "file": "x.docx", "older_master": True,
        "rows": [], "same": 3, "changed": 0, "added": 0, "removed": 0})
    ui.js('compareDialog.show("lead")')
    ui.wait(450)
    assert note.property("visible") is True
    ui.js("compareDialog.close()")
    ui.wait(300)
    monkeypatch.setattr(ui_api.Api, "cv_compare", lambda self, job_id: {
        "ok": True, "company": "Acme", "title": "Dev", "file": "x.docx", "older_master": False,
        "rows": [], "same": 3, "changed": 0, "added": 0, "removed": 0})
    ui.js('compareDialog.show("lead")')
    ui.wait(450)
    assert note.property("visible") is False
    ui.js("compareDialog.close()")
    ui.wait(300)
    assert ui.warnings == []


def test_level_words_and_years_are_edited_in_options(ui):
    dialog = ui.find("optDialog")
    start = ui.bridge.options()
    ui.bridge.saveOptions({**start, "level_words": ["senior", "lead"], "max_years": 6})
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    assert ui.js("levelTags.all().join(',')") == "senior,lead" and ui.find("yearsSpin").property("value") == 6
    ui.js('levelTags.add("Principal")')
    assert ui.js("levelTags.all().join(',')") == "senior,lead,principal"
    ui.js("close()", dialog)
    ui.wait(350)
    ui.bridge.saveOptions(start)
    assert ui.warnings == []


def test_help_has_a_what_changed_tab(ui):
    dialog = ui.find("helpDialog")
    ui.js("open()", dialog)
    ui.wait(450)
    assert ui.js("helpDialog.tab") == "how" and ui.js("helpDialog.changes.length") >= 2
    ui.js('helpDialog.tab = "new"')
    ui.wait(150)
    assert "2026" in ui.js("helpDialog.changes[0].date") and "<b>" in ui.js("helpDialog.changes[0].items[0]")
    ui.js('helpDialog.tab = "how"; helpDialog.close()')
    ui.wait(350)
    assert ui.warnings == []


def test_bring_back_button_in_the_filtered_out_list(ui):
    from jobhelper import agent2_search
    posting = store.new_job(id="back1", title="Senior Tester", company="Gamma", url="https://example.com/b")
    agent2_search.save_dropped([
        {"title": "Senior Tester", "company": "Gamma", "source": "linkedin", "url": "https://example.com/b",
         "group": "level", "reason": 'The title has a level you skip ("senior").', "job": posting},
        {"title": "Twin", "company": "Gamma", "source": "jobs.cz", "url": "", "group": "duplicate", "reason": "R.",
         "job": store.new_job(id="twin", title="Twin", company="Gamma")}])
    before = ui.js("bridge.counts().leads")
    ui.js('openView("skip")')
    ui.wait(450)
    assert [ui.js(f"listView.model[0].actions[{i}].label") for i in range(2)] == ["open ↗", "bring back"]
    assert ui.js("listView.model[1].actions.length") == 0                        # a duplicate would be removed again
    assert ui.js("bridge.bringBack(0).ok") is True
    ui.js("countsData = bridge.counts(); refreshOpenView()")
    ui.wait(300)
    assert ui.find("listView").property("count") == 1 and ui.js("countsData.leads") == before + 1
    ui.js("tableDialog.close()")
    ui.wait(350)
    jobs = [j for j in store.load_jobs() if j["id"] != "back1"]
    store.save_jobs(jobs)
    agent2_search.save_dropped([])
    ui.js("countsData = bridge.counts()")
    assert ui.warnings == []


# ── calendar file, thank-you note, answers, the last 7 days ─────────────────
def test_thank_you_card_appears_after_the_interview_day(ui, monkeypatch):
    opened = []
    monkeypatch.setattr(ui_api.Api, "_open", lambda self, target: opened.append(str(target)))
    today = date.today().isoformat()
    _open_job(ui)
    assert ui.find("thanksCard").property("visible") is False and ui.find("evCalendar").property("visible") is False
    ui.js(f'evDate.text = "{today}"; evTime.text = "00:00"; detailDialog.eventKind = "interview"; detailDialog.saveEvent()')
    assert ui.find("thanksCard").property("visible") is True and ui.find("evCalendar").property("visible") is True
    assert ui.js("bridge.exportEvent(detailDialog.job.id).ok") is True and opened[-1].endswith(".ics")
    ui.js("detailDialog.thanks = bridge.thanksDraft(detailDialog.job.id)")
    ui.wait(100)
    body = ui.find("thanksBody").property("text")
    assert "Thank you for the interview about the QA Engineer position" in body and "Beta" in body
    assert ui.js("bridge.markThanked(detailDialog.job.id)") is True
    ui.js("detailDialog.refresh()")
    assert ui.find("thanksCard").property("visible") is False                      # not offered twice for the same day
    ui.js("bridge.clearEvent(detailDialog.job.id); detailDialog.refresh(); upcomingData = bridge.upcoming(); detailDialog.close()")
    ui.wait(350)
    assert ui.warnings == []


def test_analytics_start_with_the_last_seven_days(ui):
    ui.js('openView("ana")')
    ui.wait(450)
    assert ui.js("listView.model[0].primary") == "The last 7 days" and ui.js("listView.model[0].heading") is True
    labels = [ui.js(f"listView.model[{i}].primary") for i in range(1, 6)]
    assert labels == ["New jobs found", "Ready to send", "Applications sent", "Replies received", "Follow-ups due"]
    assert ui.js("listView.model[6].primary") == "All time"
    ui.js("tableDialog.close()")
    ui.wait(350)
    assert ui.warnings == []


def test_start_up_check_shows_one_line_only_when_something_is_wrong(ui):
    banner = ui.find("healthBanner")
    assert ui.js("healthData.checks.length") > 5                                  # the check ran when the window started
    start = ui.js("healthData.problems")
    ui.js('healthData = { ok: true, checks: [], problems: 0, message: "" }')
    assert banner.property("visible") is False
    ui.js('healthData = { ok: false, problems: 2, message: "2 things need your attention. The first: No master CV yet.", '
          'checks: [{ name: "Master CV", ok: false, problem: "No master CV yet." }, { name: "Claude", ok: false, problem: "Not signed in." }, '
          '{ name: "Job list", ok: true, problem: "" }] }')
    assert banner.property("visible") is True and "No master CV yet" in ui.find("healthText").property("text")
    ui.js("showHealth()")
    ui.wait(450)
    assert ui.find("listView").property("count") == 3
    assert ui.js("listView.model[0].secondary") == "No master CV yet." and ui.js("listView.model[2].secondary") == "fine"
    ui.js("tableDialog.close()")
    ui.wait(350)
    ui.js("healthHidden = true")
    assert banner.property("visible") is False
    ui.js("healthHidden = false; healthData = bridge.selfCheck(true)")
    assert ui.js("healthData.problems") <= max(start, 3)
    ui.js('healthData = { ok: true, checks: [], problems: 0, message: "" }')
    assert ui.warnings == []


# ── the window is split into one file per window ────────────────────────────
def test_every_window_has_its_own_file_and_main_stays_small():
    folder = ROOT / "ui_qt"
    for name in ("TableDialog", "DetailDialog", "CompareDialog", "FormAnswersDialog", "HelpDialog", "PrepDialog", "LetterDialog",
                 "ResponseDialog", "LimitDialog", "WelcomeOverlay", "OptionsDialog", "ScheduleDialog", "SmallButton", "DetailCard"):
        assert (folder / f"{name}.qml").is_file(), name
    main = (folder / "Main.qml").read_text(encoding="utf-8")
    assert len(main.splitlines()) < 1400                                   # new windows go into their own file
    assert "\n    Dialog {" not in main and "\n    component " not in main


def test_a_split_file_never_hands_a_robot_its_own_empty_theme(ui):
    """Inside another file `theme: theme` means the robot's own (empty) property, not the app's colours."""
    for path in (ROOT / "ui_qt").glob("*.qml"):
        if path.name != "Main.qml":
            assert "theme: theme" not in path.read_text(encoding="utf-8"), path.name
    assert ui.js("logoBot.theme === win.appTheme && logoBot.theme.accent !== undefined") is True


def test_every_window_opens_from_the_main_file_without_warnings(ui):
    for name in Ui.WINDOWS:
        assert ui.win.findChild(QObject, name) is not None, name
    for code in ('openView("new")', "tableDialog.close()", "helpDialog.open()", "helpDialog.close()", "limitDialog.open()", "limitDialog.close()",
                 "schedDialog.open()", "schedDialog.close()", 'formDialog.show("app")', "formDialog.close()"):
        ui.js(code)
        ui.wait(380)
    assert ui.warnings == []


def test_start_up_check_still_runs_when_the_sign_in_check_is_slow(ui):
    """The splash gives up waiting after 6 seconds and marks the connection check as done;
    the self-check must still run when the answer arrives later."""
    ui.js('healthRan = false; healthData = { ok: true, checks: [], problems: 0, message: "" }; connChecked = true')
    ui.bridge.connectionsReady.emit(dict(ui.conn))
    ui.wait(150)
    assert ui.js("healthRan") is True and ui.js("healthData.checks.length") > 5
    ui.js('healthData = { ok: true, checks: [], problems: 0, message: "" }')
    ui.bridge.connectionsReady.emit(dict(ui.conn))                       # clean and already run: not repeated on every recheck
    ui.wait(150)
    assert ui.js("healthData.checks.length") == 0
    assert ui.warnings == []


# ── apply checklist, one PDF, search in Answers ─────────────────────────────
def test_apply_checklist_ticks_steps_and_marks_applied(ui, monkeypatch):
    opened = []
    monkeypatch.setattr(ui_api.Api, "_open", lambda self, target: opened.append(str(target)))
    monkeypatch.setattr(ui_api.Api, "combined_pdf", lambda self, job_id: {"ok": True, "file": "both.pdf", "pages": 3})
    monkeypatch.setattr(ui_api.Api, "_can_combine", lambda self, j: bool(j.get("tailored_cv_file")))
    for folder, name in ((config.TAILORED_CV_DIR, "cv_lead.docx"), (config.LETTERS_DIR, "L_lead.docx")):
        (folder / name).write_bytes(b"x")
    ui.bridge._api._change("lead", lambda jobs, j: j.update(tailored_cv_file="cv_lead.docx", cover_letter_file="L_lead.docx",
                                                             url="https://example.com/job", apply_steps=[]), render=False)
    _open_job(ui, "lead")
    assert ui.find("applyCard").property("visible") is True and ui.find("onePdf").property("visible") is True
    assert ui.js("detailDialog.applySteps.map(function (s) { return s.key; }).join(',')") == "posting,cv,letter,combined,answers"
    assert ui.find("applyCount").property("text").startswith("0 of 5 done")
    ui.js('detailDialog.doStep("posting"); detailDialog.doStep("cv"); detailDialog.doStep("combined")')
    assert opened[0] == "https://example.com/job" and "cv_lead." in opened[1]
    assert ui.js("detailDialog.stepsDone") == 3 and ui.js("detailDialog.job.apply_steps.join(',')") == "posting,cv,combined"
    ui.js('detailDialog.doStep("answers")')
    ui.wait(450)
    assert ui.find("formDialog").property("opened") is True and ui.js("formDialog.jobId") == "lead"
    assert ui.js("detailDialog.stepsDone") == 3                                   # opening is not doing: no tick yet
    ui.js('formDialog.fill({ items: [{ question: "Why us?", answer: "Because of your product.", check: "" }] }); formDialog.copy(0)')
    assert ui.js("detailDialog.stepsDone") == 4                                   # an answer was copied
    ui.js("formDialog.close()")
    ui.wait(350)
    ui.js("detailDialog.close()")
    ui.wait(350)
    _open_job(ui, "lead")                                                         # the ticks are still there
    assert ui.find("applyCount").property("text").startswith("4 of 5 done")
    ui.js("detailDialog.close()")
    ui.wait(350)
    _open_job(ui, "app")
    assert ui.find("applyCard").property("visible") is False                      # already applied: no checklist
    ui.js("detailDialog.close()")
    ui.wait(350)
    ui.bridge._api._change("lead", lambda jobs, j: j.update(tailored_cv_file="", cover_letter_file="", url="", apply_steps=[]), render=False)
    assert ui.warnings == []


def test_a_failed_step_is_not_ticked(ui, monkeypatch):
    monkeypatch.setattr(ui_api.Api, "combined_pdf", lambda self, job_id: {"ok": False, "msg": "The two PDF files could not be joined."})
    (config.TAILORED_CV_DIR / "cv_lead.docx").write_bytes(b"x")
    ui.bridge._api._change("lead", lambda jobs, j: j.update(tailored_cv_file="cv_lead.docx", url="", apply_steps=[]), render=False)
    _open_job(ui, "lead")
    ui.js('detailDialog.doStep("combined"); detailDialog.doStep("posting")')      # no single PDF, and this job has no link
    assert ui.js("detailDialog.stepsDone") == 0 and ui.js("(detailDialog.job.apply_steps || []).length") == 0
    ui.js("detailDialog.close()")
    ui.wait(350)
    ui.bridge._api._change("lead", lambda jobs, j: j.update(tailored_cv_file="", apply_steps=[]), render=False)
    assert ui.warnings == []


def test_options_show_your_own_letters_folder(ui, monkeypatch):
    opened = []
    monkeypatch.setattr(ui_api.Api, "_open", lambda self, target: opened.append(target))
    config.COVER_LETTERS_DIR.mkdir(parents=True, exist_ok=True)
    (config.COVER_LETTERS_DIR / "Alex_Novak_CoverLetter_Acme.docx").write_bytes(b"mine")
    dialog = ui.find("optDialog")
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    assert ui.js("optDialog.myLetters.count") == 1 and ui.js("optDialog.myLetters.name") == config.COVER_LETTERS_DIR.name
    assert ui.find("myLettersOpen").property("text") == "Open folder"
    assert ui.js("bridge.openMyLetters().ok") is True and opened == [config.COVER_LETTERS_DIR]
    ui.js("close()", dialog)
    ui.wait(350)
    (config.COVER_LETTERS_DIR / "Alex_Novak_CoverLetter_Acme.docx").unlink()
    assert ui.warnings == []


def test_without_a_city_the_search_does_not_start_and_options_open(ui):
    start = ui.bridge.options()
    ui.bridge.saveOptions({**start, "where": ""})
    ui.js("optsData = bridge.options()")
    ui.wait(100)
    assert ui.js('cityEN(optsData.where) || "No city yet"') == "No city yet"
    ui.js("startRun()")
    ui.wait(500)
    assert ui.win.property("appRunning") is False                                 # nothing was started
    assert ui.find("optDialog").property("opened") is True                        # straight to where the city is set
    assert "city" in ui.js("ticker.text").lower()
    assert ui.js("cityCombo.editText") == ""                                      # and no city is filled in for you
    ui.js("optDialog.close()")
    ui.wait(350)
    ui.bridge.saveOptions(start)
    ui.js("optsData = bridge.options()")
    assert ui.warnings == []


# ── hover that fades out cleanly ────────────────────────────────────────────
def test_hover_fades_to_the_same_colour_not_through_black(ui):
    """A colour animated to plain "transparent" passes through black and leaves a dark shadow behind."""
    assert ui.js('var c = clearOf(Qt.rgba(0.2, 0.4, 0.6, 1)); Math.round(c.r * 10) + "," + Math.round(c.g * 10) + "," + Math.round(c.b * 10) + "," + c.a') == "2,4,6,0"
    import re
    for path in sorted((ROOT / "ui_qt").glob("*.qml")):
        lines = path.read_text(encoding="utf-8").splitlines()
        for i, line in enumerate(lines):
            nearby = " ".join(lines[i:i + 3])
            if '"transparent"' in line and re.search(r"Behavior on (color|border\.color)", nearby) and "containsMouse" in line:
                raise AssertionError(f"{path.name}:{i + 1} fades a hover colour to plain transparent")


def _until(done, seconds=2.0):
    import time
    end = time.time() + seconds
    while time.time() < end and not done():
        time.sleep(0.01)
    return done()


def test_insights_are_warmed_up_in_the_background(ui, monkeypatch):
    calls = []
    monkeypatch.setattr(ui_api.Api, "read_ahead", lambda self: calls.append(1) or 0)
    ui.js("bridge.warmUp()")
    assert _until(lambda: calls and not ui.bridge._warming.locked()) and calls == [1]
    monkeypatch.setattr(ui_api.Api, "read_ahead", lambda self: (_ for _ in ()).throw(RuntimeError("no data yet")))
    ui.js("bridge.warmUp()")                                   # a failure in the head start is not shown to anyone
    assert _until(lambda: not ui.bridge._warming.locked())
    assert "bridge.warmUp()" in qml_source() and ui.warnings == []


def test_a_second_warm_up_does_not_start_while_the_first_is_reading(ui, monkeypatch):
    import threading
    started, go_on = [], threading.Event()
    monkeypatch.setattr(ui_api.Api, "read_ahead", lambda self: started.append(1) or go_on.wait(2))
    ui.js("bridge.warmUp()")
    assert _until(lambda: started)
    ui.js("bridge.warmUp()")
    ui.js("bridge.warmUp()")
    go_on.set()
    assert _until(lambda: not ui.bridge._warming.locked()) and started == [1]


def test_insights_are_not_prepared_while_the_welcome_screen_is_up(ui, monkeypatch):
    """Preparing them is heavy enough to make the welcome animation stutter, so it waits for the screen to close."""
    calls = []
    monkeypatch.setattr(ui_api.Api, "read_ahead", lambda self: calls.append(1) or 0)
    welcome = ui.find("welcome")
    was = welcome.property("shown")
    welcome.setProperty("shown", True)
    ui.js("reloadAll()")
    ui.wait(1100)
    assert calls == []
    welcome.setProperty("shown", False)
    ui.wait(300)
    assert calls == [], "not at once either: the screen is still fading out"
    ui.wait(900)
    assert _until(lambda: calls == [1])
    ui.js("reloadAll()")                                       # afterwards, every reload refreshes it straight away
    assert _until(lambda: calls == [1, 1])
    welcome.setProperty("shown", was)
    ui.wait(1000)
    assert ui.warnings == []


def test_the_start_up_check_comes_with_the_connection_check(ui, monkeypatch):
    """Both are done off the window's thread; the window only shows the result."""
    monkeypatch.setattr(ui_api.Api, "connections", lambda self: {"claude": True})
    monkeypatch.setattr(ui_api.Api, "self_check", lambda self, signed_in=None: {"ok": True, "problems": [], "asked": signed_in})
    got = []
    ui.bridge.connectionsReady.connect(got.append)
    ui.bridge.checkConnections()
    for _ in range(100):
        ui.wait(20)
        if got:
            break
    import json
    assert got and json.loads(got[-1]["health"]) == {"ok": True, "problems": [], "asked": True}
    assert "JSON.parse(info.health)" in qml_source()


def test_the_window_shows_the_start_up_check_it_was_sent(ui):
    was = ui.js("JSON.stringify(healthData)")
    ui.js("healthRan = false")
    ui.bridge.connectionsReady.emit({**ui.conn, "health": '{"ok": false, "problems": 1, "message": "No master CV yet.", "checks": []}'})
    ui.wait(80)
    assert ui.js("healthData.message") == "No master CV yet." and ui.js("healthRan") is True
    ui.js("healthRan = false")
    ui.bridge.connectionsReady.emit(dict(ui.conn))             # no check came with it: the window asks itself
    ui.wait(80)
    assert ui.js("healthData.checks.length") > 0
    ui.js("healthData = JSON.parse('%s')" % was.replace("\\", "\\\\").replace("'", "\\'"))
    assert ui.warnings == []



# ── answers for one job's form ──────────────────────────────────────────────
def test_form_answers_are_drafted_per_job_edited_and_kept(ui, monkeypatch):
    import time

    from jobhelper import llm
    monkeypatch.setattr(ui_api.Api, "_cv_text", lambda self: "Alex Novak\nPython, Docker.")
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"answers": [
        {"n": 1, "answer": "I like how Beta tests its products.", "needs_you": ""},
        {"n": 2, "answer": "I expect [your figure] CZK gross per month.", "needs_you": "Your salary figure"}]})
    source = qml_source()
    assert 'objectName: "answersButton"' not in source and "answersDialog" not in source      # the general Answers button is gone
    ui.js('formDialog.show("app")')
    ui.wait(450)
    dialog = ui.find("formDialog")
    assert dialog.property("opened") is True and ui.find("formFor").property("text") == "Beta  ·  QA Engineer"
    ui.js("formDialog.draft()")
    assert "Paste the questions" in ui.js("formDialog.err") and ui.js("formDialog.busy") is False     # red message, nothing asked
    ui.js('fmQuestions.text = (formDialog.info.common || []).join("\\n")')
    assert "Why do you want to work at Beta?" in ui.js("fmQuestions.text")
    ui.js('fmQuestions.text = "Why do you want to work at Beta?\\nWhat are your salary expectations?"; formDialog.draft()')
    assert ui.js("formDialog.busy") is True and ui.find("fmDraft").property("text") == "Asking Claude…"
    for _ in range(200):
        time.sleep(0.02)                                  # let the background worker run
        ui.wait(10)
        if ui.js("formDialog.busy") is False:
            break
    assert ui.js("formDialog.err") == "", ui.js("formDialog.err")
    assert ui.js("fmModel.count") == 2 and "1 need your attention" in ui.js("formDialog.note")
    assert ui.js("fmModel.get(1).check") == "Your salary figure" and ui.find("fmDraft").property("text") == "Draft again"
    ui.js('fmModel.setProperty(1, "answer", "I expect 60 000 CZK gross per month."); formDialog.dirty = true; formDialog.save()')
    assert ui.js("formDialog.note") == "Saved with this job." and ui.js("formDialog.dirty") is False
    ui.js("formDialog.close()")
    ui.wait(350)
    ui.js('formDialog.show("app")')                        # they are still there, and only for this job
    ui.wait(400)
    assert ui.js("fmModel.get(1).answer") == "I expect 60 000 CZK gross per month."
    ui.js("formDialog.close()")
    ui.wait(350)
    ui.js('formDialog.show("lead")')
    ui.wait(400)
    assert ui.js("fmModel.count") == 0 and ui.js("fmQuestions.text") == ""
    ui.js("formDialog.close()")
    ui.wait(350)
    ui.bridge.saveFormAnswers("app", [])
    assert ui.warnings == []


def test_form_answers_show_the_reason_when_claude_cannot_be_asked(ui, monkeypatch):
    import time

    from jobhelper import llm

    def limit(*a, **k):
        raise llm.LimitReached("Claude's 7-day limit is used up (99%). It resets at 18:00 on 08 Oct.")
    monkeypatch.setattr(ui_api.Api, "_cv_text", lambda self: "Alex Novak\nPython.")
    monkeypatch.setattr(llm, "ask_json", limit)
    ui.js('formDialog.show("app"); fmQuestions.text = "Why us?"; formDialog.draft()')
    for _ in range(200):
        time.sleep(0.02)
        ui.wait(10)
        if ui.js("formDialog.busy") is False:
            break
    assert "limit is used up" in ui.js("formDialog.err") and ui.find("formDialog").property("visible") is True    # stays open
    assert ui.js("fmQuestions.text") == "Why us?"                                 # what you pasted is not lost
    ui.js("formDialog.close()")
    ui.wait(350)
    assert ui.warnings == []


def test_form_answers_explain_every_wrong_step(ui, monkeypatch):
    """Whatever the user does wrong in this window gets a clear message, and nothing is lost silently."""
    monkeypatch.setattr(ui_api.Api, "make_form_answers", lambda self, job_id, text: pytest.fail("nothing may be drafted here"))
    ui.bridge.saveFormAnswers("app", [{"question": "Why us?", "answer": "Because of your product.", "check": ""},
                                      {"question": "Notice period?", "answer": "My notice period is [your notice period].", "check": "Fill it in"},
                                      {"question": "Anything else?", "answer": "", "check": ""}])
    ui.js('formDialog.show("app")')
    ui.wait(450)
    note = ui.find("fmNote")
    # pasted text that cannot be used: said while typing, in red, and Draft does nothing
    ui.js('fmQuestions.text = "' + "\\n".join(f"Question number {i}?" for i in range(18)) + '"')
    assert "That is 18 questions" in ui.find("fmHint").property("text")
    ui.js("formDialog.draft()")
    assert "one draft takes up to 15" in ui.find("fmHint").property("text") and ui.js("formDialog.busy") is False
    assert ui.js("formDialog.err") == ""                                           # said once, under the box, not twice
    ui.js('fmQuestions.text = "' + "x" * 320 + '"')
    assert "characters long" in ui.find("fmHint").property("text")
    ui.js('fmQuestions.text = "Why us?\\nwhy us?\\nSalary?"')
    assert ui.find("fmHint").property("text") == "2 questions found · 1 repeated line was left out · one Claude call"
    # drafting again would replace answers: asked first, never done on one click
    ui.js("formDialog.draft()")
    assert "replaces the 3 answers below" in note.property("text") and ui.js("formDialog.busy") is False
    ui.js('fmQuestions.text = "Why us?"')                                         # changing the text takes the confirmation back
    assert ui.js("formDialog.confirmDraft") is False
    # copying something unfinished
    ui.js("formDialog.copy(2)")
    assert "still empty" in note.property("text")
    ui.js("formDialog.copy(1)")
    assert "square brackets" in note.property("text")
    ui.js("formDialog.copy(0)")
    assert note.property("text") == "Copied."
    # saving says what is still open
    ui.js("formDialog.save()")
    assert note.property("text") == "Saved with this job. 1 answer is still empty. 1 answer still has a part in [square brackets] to fill in."
    # closing with an unsaved change: one warning, then it closes
    ui.js('fmModel.setProperty(0, "answer", "Changed and not saved."); formDialog.dirty = true; formDialog.tryClose()')
    ui.wait(250)
    dialog = ui.find("formDialog")
    assert dialog.property("visible") is True and "did not save" in note.property("text")
    assert ui.js("formDialog.closePolicy === 0") is True                           # Esc cannot throw the change away either
    ui.js("formDialog.tryClose()")
    ui.wait(450)
    assert dialog.property("visible") is False
    assert ui.bridge.formAnswers("app")["items"][0]["answer"] == "Because of your product."      # the discarded change was not saved
    # closing while Claude is working
    ui.js('formDialog.show("app"); formDialog.busy = true; formDialog.tryClose()')
    ui.wait(300)
    assert dialog.property("visible") is True and "still drafting" in note.property("text")
    ui.js("formDialog.busy = false; formDialog.tryClose()")
    ui.wait(450)
    ui.bridge.saveFormAnswers("app", [])
    assert ui.warnings == []


def test_a_message_behind_an_open_window_is_shown_on_top_of_it(ui):
    """The status line is hidden by any open window, so results and errors must also appear above it."""
    toast = ui.find("toast")
    ui.js('ticker.text = "said with no window open"')
    ui.wait(250)
    assert toast.property("visible") is False                                      # the status line itself is in view
    _open_job(ui, "app")
    ui.js('ticker.text = "No link for this job."')
    ui.wait(300)
    assert toast.property("visible") is True and ui.find("toastText").property("text") == "No link for this job."
    assert ui.js("toastBad") is True
    ui.js('ticker.text = "Opened the job posting in your browser."')
    ui.wait(250)
    assert ui.js("toastBad") is False and ui.find("toastText").property("text").startswith("Opened")
    ui.bridge.logChanged.emit("Scoring 3 of 14 jobs…")                             # a run's progress does not pop up
    ui.wait(250)
    assert ui.find("toastText").property("text").startswith("Opened")
    ui.js("toastTimer.interval = 150; toastTimer.restart()")
    ui.wait(600)
    assert toast.property("visible") is False                                      # and it goes away by itself
    ui.js("toastTimer.interval = 4200; detailDialog.close()")
    ui.wait(350)
    assert ui.warnings == []


@pytest.mark.parametrize("text, bad", [
    ("Couldn't make the single PDF.", True), ("No link for this job.", True), ("That job could not be opened.", True),
    ("These CVs cannot be compared.", True), ("The pasted reply could not be filed.", True), ("That answer is still empty.", True),
    ("Choose the city to search in first (Options).", True), ("Not prepared: the limit is used up", True),
    ("Date saved. You will be reminded the day before.", False), ("Moved to Applications (waiting).", False),
    ("Opened one PDF with the letter and the CV (2 pages).", False), ("Merged into one application.", False),
])
def test_which_messages_are_shown_as_errors(ui, text, bad):
    assert ui.js("looksBad(" + repr(text).replace("\\'", "'") + ")") is bad


def test_preparation_opens_as_a_document_from_both_windows(ui, monkeypatch):
    calls = []
    monkeypatch.setattr(ui_api.Api, "prep_document", lambda self, job_id, reveal=False: calls.append((job_id, reveal)) or {"ok": True, "file": "Alex_Beta_Prep.docx"})
    ui.bridge._api._change("app", lambda jobs, j: j.__setitem__("interview_prep", {"made": "2026-10-08", "summary": "s", "questions": [
        {"kind": "about you", "question": "Tell me about yourself.", "why": "", "answer": "I am a junior developer.", "check": ""}]}), render=False)
    _open_job(ui, "app")
    assert ui.find("prepDocument").property("visible") is True
    ui.js('prepDialog.show("app")')
    ui.wait(450)
    assert ui.find("prepDoc").property("text") == "Open as document" and ui.find("prepShow").property("text") == "Show in folder"
    assert ui.js('bridge.prepDocument("app", false).file') == "Alex_Beta_Prep.docx" and ui.js('bridge.prepDocument("app", true).ok') is True
    assert calls == [("app", False), ("app", True)]
    ui.js("prepDialog.close()")
    ui.wait(350)
    ui.js("detailDialog.close()")
    ui.wait(350)
    ui.bridge._api._change("app", lambda jobs, j: j.update(interview_prep={}), render=False)
    _open_job(ui, "app")
    assert ui.find("prepDocument").property("visible") is False                 # nothing prepared yet: no document button
    ui.js("detailDialog.close()")
    ui.wait(350)
    assert ui.warnings == []


def test_show_in_folder_buttons_and_the_laid_out_description(ui, monkeypatch):
    shown = []
    monkeypatch.setattr(ui_api.Api, "show_in_folder", lambda self, kind, job_id: shown.append((kind, job_id)) or {"ok": True, "file": "CV.pdf"})
    (config.TAILORED_CV_DIR / "cv_app.docx").write_bytes(b"x")
    ui.bridge._api._change("app", lambda jobs, j: j.update(tailored_cv_file="cv_app.docx",
                           description="We build tools. Requirements Strong Python Benefits Flexible hours"), render=False)
    _open_job(ui, "app")
    assert ui.find("showCv").property("visible") is True and ui.find("showLetter").property("visible") is False
    ui.js('detailDialog.showFile("cv")')
    assert shown == [("cv", "app")]
    assert ui.find("descText").property("text") == "We build tools.\n\nREQUIREMENTS\nStrong Python\n\nBENEFITS\nFlexible hours"
    ui.js("detailDialog.close()")
    ui.wait(350)
    ui.bridge._api._change("app", lambda jobs, j: j.update(tailored_cv_file="", description="We use Python, Docker and Kubernetes."), render=False)
    assert ui.warnings == []


def test_more_about_your_work_is_saved_and_checked(ui):
    dialog = ui.find("optDialog")
    start = ui.bridge.options()
    ui.bridge.saveOptions({**start, "where": "Prague", "about_experience": "At Acme I wrote the import service in Python."})
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    field = ui.find("expField")
    assert field.property("text") == "At Acme I wrote the import service in Python."
    assert ui.find("expCount").property("text") == "45 of 4000 characters"
    ui.js('expField.text = "' + "x" * 4010 + '"')
    assert ui.find("expCount").property("text") == "10 characters too long (the limit is 4000)"
    source = qml_source()
    assert "characters too long. Shorten it, then save." in source and "about_experience: expField.text.trim()" in source
    ui.js("close()", dialog)
    ui.wait(350)
    ui.bridge.saveOptions(start)
    assert ui.warnings == []


def test_preparation_window_groups_questions_by_who_asks(ui):
    assert ui.js('prepDialog.who({ audience: "hr", kind: "technical" })') == "hr"
    assert ui.js('prepDialog.who({ kind: "technical" })') == "technical" and ui.js('prepDialog.who({ kind: "practical" })') == "hr"   # older saved ones
    source = qml_source()
    assert "If you speak with HR or a recruiter: plain words, no jargon" in source
    assert "If you speak with someone who does this work or the hiring manager: precise answers" in source


def test_preparation_window_shows_what_you_can_ask_each_listener(ui):
    ui.bridge._api._change("app", lambda jobs, j: j.update(interview_prep={"made": "2026-10-08", "summary": "s", "questions": [
        {"audience": "hr", "kind": "about you", "question": "Tell me about yourself.", "why": "", "answer": "I am a junior developer.", "check": ""}],
        "ask_hr": ["What are the next steps?"], "ask_technical": ["What would I own first?"], "ask_them": []}), render=False)
    ui.js('prepDialog.show("app")')
    ui.wait(450)
    assert ui.find("askHr").property("visible") is True and ui.find("askTech").property("visible") is True
    assert ui.js("prepDialog.prep.ask_hr[0]") == "What are the next steps?"
    ui.js("prepDialog.close()")
    ui.wait(350)
    ui.bridge._api._change("app", lambda jobs, j: j.update(interview_prep={}), render=False)
    assert ui.warnings == []


def test_wanted_and_unwanted_roles_are_edited_in_options(ui):
    dialog = ui.find("optDialog")
    start = ui.bridge.options()
    ui.bridge.saveOptions({**start, "where": "Prague", "wanted_roles": ["backend developer"], "unwanted_roles": ["economist"]})
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    assert ui.js("wantedTags.all().join(',')") == "backend developer" and ui.js("unwantedTags.all().join(',')") == "economist"
    assert ui.js('unwantedTags.add("x")') is False and "job title" in ui.js("unwantedTags.err")       # wrong input is explained
    source = qml_source()
    assert "is under both the roles you want and the roles you do not want" in source
    note = ui.find("fieldsNote")                                                                      # your profession follows the roles as typed
    assert note.property("text").startswith("Your profession, as the app reads it from these roles: software and IT.")
    assert ui.js('wantedTags.add("registered nurse")') is True
    assert "these roles: software and IT, care and health." in note.property("text")
    ui.js('wantedTags.setAll(["beekeeper"])')
    assert note.property("text").startswith("The app cannot tell your profession from these roles or your CV") \
        or "from your CV" in note.property("text")
    assert "a title that names one is skipped before scoring" in source

    ui.js('wantedTags.setAll(["backend developer"])')                                                 # … or you pick it yourself
    names = ui.js("fieldsBox.read.all.join('|')").split("|")
    assert "software and IT" in names and "care and health" in names and ui.find("fieldsAuto").property("visible") is False
    assert note.property("text").endswith("Pick below to set it yourself.") and ui.js("optDialog.changed()") is False
    ui.js('optDialog.toggleField("law", fieldsBox.read.all); optDialog.toggleField("finance", fieldsBox.read.all)')
    assert ui.js("optDialog.pickedFields.join('|')") == "finance|law"                                 # kept in the app's order
    assert note.property("text").startswith("Your profession, as you picked it: finance, law.") and ui.find("fieldsAuto").property("visible") is True
    assert ui.js("optDialog.changed()") is True and ui.js("optDialog.collect().options.my_fields.join('|')") == "finance|law"
    ui.js('optDialog.toggleField("law", fieldsBox.read.all)')
    assert ui.js("optDialog.pickedFields.join('|')") == "finance"
    ui.js("optDialog.pickedFields = []")                                                              # "Let the app decide"
    assert "as the app reads it from these roles: software and IT" in note.property("text") and ui.js("optDialog.changed()") is False
    ui.js("close()", dialog)
    ui.wait(350)
    ui.bridge.saveOptions({**start, "where": "Prague", "wanted_roles": ["backend developer"], "my_fields": ["law", "no such field", "finance"]})
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    assert ui.js("optDialog.pickedFields.join('|')") == "finance|law" and ui.js("optDialog.changed()") is False   # an unknown name is dropped on load
    assert "Let the app decide" in source and "Pick it below if you want that." in source
    ui.js("close()", dialog)
    ui.wait(350)
    ui.bridge.saveOptions(start)
    assert ui.warnings == []


def test_the_preparation_can_be_edited_and_wrong_steps_are_explained(ui, monkeypatch):
    monkeypatch.setattr(ui_api.Api, "_cv_text", lambda self: "Alex Novak\nPython, Docker.")
    prep = {"made": "2026-10-08", "summary": "They will ask about Python.", "pitch": "I am a junior developer.", "fit": ["Python", "Docker"],
            "questions": [{"audience": "hr", "kind": "about you", "question": "Tell me about yourself.", "why": "", "answer": "I am a junior developer.", "check": ""}],
            "ask_hr": ["What are the next steps?"], "ask_technical": [], "ask_them": [], "checks": []}
    ui.bridge._api._change("app", lambda jobs, j: j.update(interview_prep=dict(prep)), render=False)
    ui.js('prepEditDialog.show("app")')
    ui.wait(450)
    dialog, note = ui.find("prepEditDialog"), ui.find("peNote")
    assert dialog.property("opened") is True and ui.js("peModel.count") == 5
    keys = [ui.js(f"peModel.get({i}).key") for i in range(5)]
    assert keys == ["summary", "fit", "pitch", "questions.0.answer", "ask_hr"]
    # an empty answer is refused, in red, and the window stays open
    ui.js('peModel.setProperty(3, "text", ""); prepEditDialog.dirty = true; prepEditDialog.save()')
    assert "is empty" in note.property("text") and dialog.property("visible") is True
    # too long is said before anything is sent
    ui.js('peModel.setProperty(3, "text", "I am a junior developer."); peModel.setProperty(2, "text", "' + "x" * 3010 + '"); prepEditDialog.save()')
    assert "longer than 3000 characters" in note.property("text")
    # closing with changes asks once
    ui.js('peModel.setProperty(2, "text", "My own pitch."); prepEditDialog.dirty = true; prepEditDialog.tryClose()')
    ui.wait(200)
    assert dialog.property("visible") is True and "did not save" in note.property("text") and ui.js("prepEditDialog.closePolicy === 0") is True
    ui.js("prepEditDialog.save()")
    assert note.property("text") == "Saved." and ui.js("prepEditDialog.dirty") is False
    assert ui.bridge.interviewPrep("app")["pitch"] == "My own pitch."
    ui.js("prepEditDialog.tryClose()")
    ui.wait(450)
    assert dialog.property("visible") is False
    # "Make again" would replace the edit: asked first
    _open_job(ui, "app")
    assert ui.js("detailDialog.job.prep_edited") is True
    source = qml_source()
    assert "Making it again replaces the changes you made in the sheet." in source and 'objectName: "prepEdit"' in source
    ui.js("detailDialog.close()")
    ui.wait(350)
    ui.bridge._api._change("app", lambda jobs, j: j.update(interview_prep={}), render=False)
    ui.js('prepEditDialog.show("app")')                                          # nothing to edit: said, no empty window
    ui.wait(250)
    assert dialog.property("visible") is False and "Prepare the interview first" in ui.js("ticker.text")
    assert ui.warnings == []


# ── every text field says when it is full or too long ───────────────────────
def test_text_fields_in_the_job_window_explain_their_limits(ui):
    _open_job(ui, "app")
    line = ui.find("notesLine")
    ui.js('cnContact.text = "' + "c" * 200 + '"; detailDialog.notesInfo = detailDialog.full(cnContact, "“Who you spoke to”"); detailDialog.notesBad = true')
    assert line.property("text") == "“Who you spoke to” takes up to 200 characters."
    ui.js('cnContact.text = "Jana"; cnNotes.text = "' + "n" * 2010 + '"')
    assert ui.find("notesOver").property("visible") is True and ui.find("notesOver").property("text") == "10 characters too long (the limit is 2000)"
    ui.js("detailDialog.saveNotes()")
    assert "10 characters too long" in line.property("text") and ui.js("detailDialog.notesBad") is True      # refused, in red
    from jobhelper import planner
    assert planner.company_notes("Beta")["updated"] == ""                                                    # and not saved cut
    ui.js('cnNotes.text = "Short note."; detailDialog.saveNotes()')
    assert line.property("text") == "Saved for Beta." and ui.js("detailDialog.notesBad") is False
    ui.js('evNote.text = "' + "e" * 200 + '"; detailDialog.eventErr = detailDialog.full(evNote, "The note")')
    assert ui.find("eventErr").property("text") == "The note takes up to 200 characters."
    ui.js('detailDialog.eventErr = ""; detailDialog.replyOpen = true; replyText.text = "' + "r" * 8005 + '"')
    assert ui.find("replyHint").property("text").startswith("5 characters too long: paste only the message itself")
    ui.js("detailDialog.fileReply()")
    assert "Paste only the message itself" in ui.js("detailDialog.replyErr") and ui.js("detailDialog.replyBusy") is False
    ui.js('replyText.text = ""; detailDialog.replyOpen = false; cnNotes.text = ""; cnContact.text = ""; detailDialog.saveNotes(); detailDialog.close()')
    ui.wait(350)
    assert ui.warnings == []


def test_form_answers_and_preparation_lists_explain_their_limits(ui, monkeypatch):
    ui.bridge.saveFormAnswers("app", [{"question": "Why us?", "answer": "Because.", "check": ""}])
    ui.js('formDialog.show("app")')
    ui.wait(450)
    ui.js('fmModel.setProperty(0, "answer", "' + "a" * 3004 + '"); formDialog.dirty = true; formDialog.save()')
    assert "is longer than 3000 characters" in ui.find("fmNote").property("text")
    assert ui.bridge.formAnswers("app")["items"][0]["answer"] == "Because."       # the long version was not saved
    ui.js("formDialog.dirty = false; formDialog.close()")
    ui.wait(350)
    assert ui.js("prepEditDialog.listProblem('one\\ntwo')") == ""
    many = "\\n".join(f"line {i}" for i in range(13))
    assert ui.js(f"prepEditDialog.listProblem('{many}')") == "13 lines; the most is 12. Remove some."
    assert ui.js("prepEditDialog.listProblem('" + "q" * 401 + "')") == "Line 1 is 401 characters long; the most is 400."
    ui.bridge.saveFormAnswers("app", [])
    source = qml_source()
    assert "“Right to work” takes up to" in source and "“You can start” takes up to" in source
    assert ui.warnings == []


# ── the older windows: every text field explains itself too ─────────────────
def test_number_fields_say_what_they_accept_instead_of_going_silent(ui):
    dialog = ui.find("optDialog")
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    spin = ui.find("scanSpin")                                                   # 1 to 365
    start = spin.property("value")
    ui.js('scanSpin.contentItem.text = "999"; scanSpin.contentItem.textEdited()')
    assert ui.js("scanSpin.note") == "Between 1 and 365" and ui.js("scanSpin.noteBad") is True
    assert spin.property("value") == start                                        # a wrong number is not taken
    ui.js("scanSpin.contentItem.editingFinished()")
    assert spin.property("value") == 365 and ui.js("scanSpin.note") == "Changed to 365, the highest allowed"
    assert ui.js("scanSpin.contentItem.text") == "365"
    ui.js('scanSpin.contentItem.text = "0"; scanSpin.contentItem.textEdited(); scanSpin.contentItem.editingFinished()')
    assert spin.property("value") == 1 and ui.js("scanSpin.note") == "Changed to 1, the lowest allowed"
    ui.js('scanSpin.contentItem.text = "45"; scanSpin.contentItem.textEdited()')
    assert spin.property("value") == 45 and ui.js("scanSpin.note") == ""           # a good number: taken at once, nothing said
    ui.js('scanSpin.contentItem.text = ""; scanSpin.contentItem.textEdited(); scanSpin.contentItem.editingFinished()')
    assert spin.property("value") == 45 and ui.js("scanSpin.note") == "Left at 45"
    ui.js(f"scanSpin.value = {start}; scanSpin.note = ''")
    ui.js("close()", dialog)
    ui.wait(350)
    assert ui.warnings == []


def test_career_site_name_is_checked(ui):
    dialog = ui.find("optDialog")
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    assert ui.js('optDialog.addCompany("' + "n" * 61 + '", "greenhouse", "acme")') is False
    assert ui.js("optDialog.coErr") == "The company name is 61 characters long; the most is 60."
    assert ui.js('optDialog.addCompany("Компания", "greenhouse", "acme")') is False and "Latin letters" in ui.js("optDialog.coErr")
    assert ui.js('optDialog.addCompany("Acme", "greenhouse", "")') is False and "token" in ui.js("optDialog.coErr")
    assert ui.find("coName").property("maximumLength") == 60
    ui.js("close()", dialog)
    ui.wait(350)
    assert ui.warnings == []


def test_letter_editor_explains_wrong_steps(ui, monkeypatch):
    from jobhelper import cover_letter
    paras = ["First paragraph of the letter.", "Second one.", "Third one.", "Fourth one.", "Fifth one."]
    monkeypatch.setattr(ui_api.Api, "letter_for_edit", lambda self, job_id: {"ok": True, "id": "lead", "company": "Acme", "title": "Dev",
                        "title_line": "Dev", "date_text": "", "paragraphs": list(paras), "min_words": cover_letter.MIN_WORDS, "max_words": cover_letter.MAX_WORDS})
    saved = []
    monkeypatch.setattr(ui_api.Api, "save_letter", lambda self, job_id, texts: saved.append(list(texts)) or {"ok": True, "warnings": []})
    ui.js('letterDialog.show("lead")')
    ui.wait(450)
    dialog, err = ui.find("letterDialog"), ui.find("letterErr")
    assert dialog.property("opened") is True
    ui.js('letterDialog.setText(1, "   "); letterDialog.setText(3, "")')
    ui.js("letterDialog.save()")
    assert err.property("text") == "Paragraph 2 and 4 are empty. A letter needs all 5 paragraphs." and saved == []
    ui.js("letterDialog.rewrite(1)")
    assert "Paragraph 2 is empty" in err.property("text") and ui.js("letterDialog.busyIndex") == -1      # nothing to rewrite: Claude is not asked
    ui.js('letterDialog.setText(1, "Second one, changed."); letterDialog.setText(3, "Fourth one."); letterDialog.tryClose()')
    ui.wait(200)
    assert dialog.property("visible") is True and "did not save" in err.property("text")                 # changes are not thrown away on one click
    ui.js("letterDialog.busyIndex = 2; letterDialog.tryClose()")
    assert "still rewriting paragraph 3" in err.property("text") and dialog.property("visible") is True
    ui.js("letterDialog.busyIndex = -1; letterDialog.save()")
    assert saved and saved[0][1] == "Second one, changed."
    assert ui.find("hintField").property("maximumLength") == 300
    ui.js('hintField.text = "' + "h" * 300 + '"; hintField.textEdited()')
    assert err.property("text") == "The note for a rewrite takes up to 300 characters."
    ui.js('hintField.text = ""; letterDialog.err = ""; letterDialog.tryClose()')
    ui.wait(450)
    assert dialog.property("visible") is False                                                           # saved: closes at once
    assert ui.warnings == []


# ── score again ─────────────────────────────────────────────────────────────
def test_score_again_from_the_job_window(ui, monkeypatch):
    import time

    from jobhelper import llm
    monkeypatch.setattr(ui_api.Api, "_cv_text", lambda self: "Alex Novak\nPython.")
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"match_score": 35, "match_reason": "Another field."})
    _open_job(ui, "lead")
    assert ui.find("scoreAgain").property("text") == "Score again" and ui.js("detailDialog.job.can_rescore") is True
    ui.js("detailDialog.scoreAgain()")
    assert ui.js("detailDialog.scoreBusy") is True and ui.find("scoreAgain").property("enabled") is False
    for _ in range(200):
        time.sleep(0.02)
        ui.wait(10)
        if ui.js("detailDialog.scoreBusy") is False:
            break
    assert ui.find("scoreNote").property("text") == "Scored again: 80 before, 35 now."
    assert ui.js("detailDialog.job.score") == 35 and ui.js("detailDialog.job.score_before") == 80
    ui.js("detailDialog.close()")
    ui.wait(350)
    _open_job(ui, "lead")                                                       # later it still says what it was
    assert ui.find("scoreNote").property("text") == "It was 80 before it was scored again."
    ui.js("detailDialog.close()")
    ui.wait(350)
    _open_job(ui, "app")
    assert ui.js("detailDialog.job.can_rescore") is False                         # an application is not scored again
    ui.js("detailDialog.close()")
    ui.wait(350)
    ui.bridge._api._change("lead", lambda jobs, j: j.update(match_score=80, match_reason="", score_before=None), render=False)
    assert ui.warnings == []


def test_score_again_says_why_it_failed_and_keeps_the_score(ui, monkeypatch):
    import time

    from jobhelper import llm

    def limit(*a, **k):
        raise llm.LimitReached("Claude's 7-day limit is used up (99%). It resets at 18:00.")
    monkeypatch.setattr(llm, "ask_json", limit)
    _open_job(ui, "lead")
    ui.js("detailDialog.scoreAgain()")
    for _ in range(200):
        time.sleep(0.02)
        ui.wait(10)
        if ui.js("detailDialog.scoreBusy") is False:
            break
    assert "Not scored again: Claude's 7-day limit is used up" in ui.find("scoreNote").property("text") and ui.js("detailDialog.scoreBad") is True
    assert ui.js("detailDialog.job.score") == 80
    ui.js("detailDialog.close()")
    ui.wait(350)
    assert ui.warnings == []


def test_score_all_again_asks_before_spending_the_limit(ui, monkeypatch):
    started = []
    monkeypatch.setattr(qt_bridge.Bridge, "rescoreAll", lambda self: started.append(1) or {"ok": False, "msg": "A run is already in progress."})
    ui.js('openView("new")')
    ui.wait(450)
    button = ui.find("rescoreAll")
    assert button.property("visible") is True and button.property("text") == "Score all again"
    source = qml_source()
    assert "Only scores change; no letter is written or removed. Press again to start." in source
    ui.js("tableDialog.confirmRescore = true; tableDialog.rescoreJobs = 1")
    assert button.property("text") == "Yes, score 1 again"
    ui.js("tableDialog.close()")
    ui.wait(350)
    assert ui.js("tableDialog.confirmRescore") is False and started == []         # closing takes the confirmation back; nothing started
    ui.js('openView("app")')
    ui.wait(400)
    assert button.property("visible") is False                                    # only in New Jobs
    ui.js("tableDialog.close()")
    ui.wait(350)
    assert ui.warnings == []


# ── Options: changes are not thrown away unasked ────────────────────────────
def test_options_ask_once_before_discarding_changes(ui):
    dialog = ui.find("optDialog")
    start = ui.bridge.options()
    ui.bridge.saveOptions({**start, "where": "Prague"})
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    assert ui.js("optDialog.changed()") is False
    ui.js("optDialog.tryClose()")                                               # nothing changed: closes at once
    ui.wait(450)
    assert dialog.property("visible") is False
    ui.js("open()", dialog)
    ui.wait(450)
    ui.js("thrSpin.value = thrSpin.value + 5")
    assert ui.js("optDialog.changed()") is True
    ui.js("optDialog.tryClose()")
    ui.wait(250)
    assert dialog.property("visible") is True and "did not save" in ui.js("optDialog.optErr")
    assert ui.find("optCancel").property("text") == "Discard changes"
    ui.js("thrSpin.value = thrSpin.value - 5")                                  # put back by hand: nothing to lose any more
    assert ui.js("optDialog.changed()") is False
    ui.js('skipTags.add("contractor"); optDialog.confirmClose = false; optDialog.tryClose()')
    ui.wait(200)
    assert dialog.property("visible") is True                                     # a tag counts as a change too
    ui.js("optDialog.tryClose()")                                               # the second time it goes
    ui.wait(450)
    assert dialog.property("visible") is False and "contractor" not in ui.bridge.options()["skip_words"]
    assert ui.js("optDialog.closePolicy === 0") is True                           # Esc cannot bypass the question
    ui.bridge.saveOptions(start)
    ui.js("optsData = bridge.options()")
    assert ui.warnings == []


def _wait_scored(ui):
    import time
    for _ in range(300):
        time.sleep(0.02)
        ui.wait(10)
        if ui.js("detailDialog.scoreBusy") is False:
            return
    raise AssertionError("the score never came back")


# ── what a score rests on, pasted text, documents left from a higher score ──
def test_a_title_only_score_is_marked_and_the_posting_can_be_pasted(ui, monkeypatch):
    from jobhelper import llm
    monkeypatch.setattr(ui_api.Api, "_cv_text", lambda self: "Alex Novak\nPython.")
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"match_score": 65, "match_reason": "Fits."})
    start = store.find(store.load_jobs(), "lead")
    ui.bridge._api._change("lead", lambda jobs, j: j.update(description="", match_score=40), render=False)
    ui.js('openView("new")')
    ui.wait(450)
    assert "scored from the title only" in ui.js("listView.model[0].secondary")
    ui.js("tableDialog.close()")
    ui.wait(350)
    _open_job(ui, "lead")
    assert ui.find("basisRow").property("visible") is True and ui.js("detailDialog.job.basis") == "title"
    ui.js('detailDialog.pasteOpen = true; pastedText.text = "too short"; detailDialog.savePasted()')
    assert ui.find("pasteHint").property("text") == "That is only 9 characters. Paste the whole posting (at least 200)."
    ui.js('pastedText.text = "' + "x" * 20010 + '"; detailDialog.pasteErr = ""')
    assert ui.find("pasteHint").property("text") == "10 characters too long (the limit is 20000)"
    ui.js("detailDialog.savePasted()")
    assert "the most is 20000" in ui.find("pasteHint").property("text") and ui.js("detailDialog.pasteOpen") is True      # refused, stays open
    ui.js('pastedText.text = "' + "We build internal tools in Python and test them well. " * 10 + '"; detailDialog.savePasted()')
    assert ui.js("detailDialog.pasteOpen") is False and ui.js("detailDialog.job.basis") == ""
    assert ui.find("basisRow").property("visible") is False
    assert ui.js("detailDialog.scoreBusy") is True                                   # the new text is scored at once, unasked
    assert ui.find("scoreNote").property("text") == "The posting's text is saved. Asking Claude…"
    _wait_scored(ui)
    assert ui.find("scoreNote").property("text") == "The posting's text is saved. Scored again: 40 before, 65 now."
    assert ui.js("detailDialog.job.score") == 65 and ui.js("detailDialog.scoreLead") == ""
    ui.js("detailDialog.close()")
    ui.wait(350)
    ui.bridge._api._change("lead", lambda jobs, j: j.update(description=start.get("description", ""), match_score=80, description_by_hand=False), render=False)
    assert ui.warnings == []


def test_documents_from_a_higher_score_are_marked_and_moved_only_after_asking(ui, monkeypatch):
    for folder, name in ((config.TAILORED_CV_DIR, "cv_lead.docx"), (config.LETTERS_DIR, "L_lead.docx")):
        (folder / name).write_bytes(b"x")
    monkeypatch.setattr(ui_api.Api, "_threshold", lambda self: 50)
    ui.bridge._api._change("lead", lambda jobs, j: j.update(tailored_cv_file="cv_lead.docx", cover_letter_file="L_lead.docx",
                                                             status="letter_ready", match_score=30), render=False)
    ui.js('openView("new")')
    ui.wait(450)
    assert "documents from a higher score" in ui.js("listView.model[0].secondary")
    ui.js("tableDialog.close()")
    ui.wait(350)
    _open_job(ui, "lead")
    assert ui.find("staleRow").property("visible") is True and "below your limit of 50" in ui.find("staleText").property("text")
    ui.js("detailDialog.shelve()")
    assert ui.find("shelveDocs").property("text") == "Yes, move them" and (config.LETTERS_DIR / "L_lead.docx").exists()      # asked first
    ui.js("detailDialog.shelve()")
    assert ui.find("staleText").property("text").endswith("files moved to the backup.") and not (config.LETTERS_DIR / "L_lead.docx").exists()
    assert ui.js("detailDialog.job.docs_stale") is False and ui.find("shelveDocs").property("visible") is False
    ui.js("detailDialog.close()")
    ui.wait(350)
    ui.bridge._api._change("lead", lambda jobs, j: j.update(match_score=80, status="letter_ready"), render=False)
    assert ui.warnings == []


def test_read_it_again_button_reports_both_outcomes(ui, monkeypatch):
    from jobhelper import agent2_search
    start = store.find(store.load_jobs(), "lead")
    ui.bridge._api._change("lead", lambda jobs, j: j.update(description="", match_score=40, url="https://example.com/job"), render=False)
    _open_job(ui, "lead")
    assert ui.find("rereadPosting").property("visible") is True
    got = ui.js("bridge.rereadPosting(detailDialog.job.id)")
    assert got["ok"] is False and "paste its text instead" in got["msg"]
    ui.js("detailDialog.readAgain()")                                                # the button: nothing new to read
    assert "paste its text instead" in ui.find("scoreNote").property("text") and ui.js("detailDialog.scoreBad") is True
    assert ui.js("detailDialog.scoreBusy") is False
    text = "We build internal tools in Python and test them well. " * 10
    monkeypatch.setattr(agent2_search, "refetch_description", lambda job: text)
    ui.js("win.appRunning = true; detailDialog.readAgain()")                         # read, but a search is running: no scoring now
    assert ui.find("scoreNote").property("text") == f"Read {len(text)} characters of the posting. A search or scan is running: press Score again when it has finished."
    assert ui.js("detailDialog.scoreBusy") is False and ui.js("detailDialog.job.basis") == "" and ui.find("basisRow").property("visible") is False
    ui.js("win.appRunning = false")

    from jobhelper import llm
    monkeypatch.setattr(ui_api.Api, "_cv_text", lambda self: "Alex Novak\nPython.")
    ui.bridge._api._change("lead", lambda jobs, j: j.update(description="", match_score=40), render=False)
    ui.js("detailDialog.refresh()")

    def limit(*a, **k):
        raise llm.LimitReached("Claude's 7-day limit is used up.")
    monkeypatch.setattr(llm, "ask_json", limit)
    ui.js("detailDialog.readAgain()")
    assert ui.find("rereadPosting").property("enabled") is False
    _wait_scored(ui)
    note = ui.find("scoreNote").property("text")                                     # the text is kept, and the reason it was not scored is said
    assert note.startswith(f"Read {len(text)} characters of the posting. Not scored again:") and ui.js("detailDialog.scoreBad") is True
    assert store.find(store.load_jobs(), "lead")["description"] == text and ui.js("detailDialog.job.score") == 40

    ui.bridge._api._change("lead", lambda jobs, j: j.update(description="", match_score=None), render=False)   # a job with no score yet
    ui.js("detailDialog.refresh(); detailDialog.readAgain()")
    assert ui.js("detailDialog.scoreBusy") is False and ui.js("ticker.text") == f"Read {len(text)} characters of the posting. The next search will score it."
    ui.js("detailDialog.close()")
    ui.wait(350)
    ui.bridge._api._change("lead", lambda jobs, j: j.update(description=start.get("description", ""), match_score=80, url=""), render=False)
    assert ui.warnings == []


# ── New Jobs: all documents below the limit at once, and the score a job had before ──
def test_documents_below_the_limit_are_moved_together_only_after_asking(ui, monkeypatch):
    for folder, name in ((config.TAILORED_CV_DIR, "cv_lead.docx"), (config.LETTERS_DIR, "L_lead.docx")):
        (folder / name).write_bytes(b"x")
    monkeypatch.setattr(ui_api.Api, "_threshold", lambda self: 50)
    ui.bridge._api._change("lead", lambda jobs, j: j.update(tailored_cv_file="cv_lead.docx", cover_letter_file="L_lead.docx",
                                                             status="letter_ready", match_score=30), render=False)
    ui.js('openView("new")')
    ui.wait(450)
    bar, button, text = ui.find("staleBar"), ui.find("shelveAll"), ui.find("staleBarText")
    assert bar.property("visible") is True and button.property("text") == "Move them to backup"
    assert text.property("text") == "1 job has documents from a score that is now below your limit."
    ui.js("tableDialog.shelveAll()")                                                # first press only asks, and names the jobs
    assert button.property("text") == "Yes, move documents of 1 job" and "A letter of your own is never moved." in text.property("text")
    title = store.find(store.load_jobs(), "lead")["title"]
    assert ui.find("staleNames").property("visible") is True and title in ui.find("staleNames").property("text")
    assert (config.LETTERS_DIR / "L_lead.docx").exists()
    ui.js('tableDialog.confirmStale = false')                                       # Cancel
    assert button.property("text") == "Move them to backup" and (config.LETTERS_DIR / "L_lead.docx").exists()
    ui.js("tableDialog.close()")
    ui.wait(350)
    ui.js('openView("new")')
    ui.wait(450)
    assert ui.js("tableDialog.confirmStale") is False                               # a closed window forgets the question

    ui.js("tableDialog.shelveAll()")
    ui.js("win.appRunning = true; tableDialog.shelveAll()")                         # a search started meanwhile
    assert "A search or scan is running" in ui.find("staleNote").property("text") and ui.js("tableDialog.staleBad") is True
    assert (config.LETTERS_DIR / "L_lead.docx").exists()
    ui.js("win.appRunning = false")

    ui.bridge._running = True                                                       # the bridge refuses too, whatever the window thinks
    try:
        assert "A search or scan is running" in ui.bridge.shelveAllDocs(1)["msg"] and (config.LETTERS_DIR / "L_lead.docx").exists()
    finally:
        ui.bridge._running = False

    ui.js("tableDialog.staleRows = [{id: 'x', title: 'T', company: 'C', score: 1}, {id: 'y', title: 'T', company: 'C', score: 1}]; tableDialog.confirmStale = true")
    ui.js("tableDialog.shelveAll()")                                                # the list is not what was asked about
    assert "The list changed: it is now 1 job" in ui.find("staleNote").property("text") and (config.LETTERS_DIR / "L_lead.docx").exists()
    assert ui.js("tableDialog.confirmStale") is True and ui.js("tableDialog.staleRows.length") == 1
    ui.js("tableDialog.shelveAll()")
    assert not (config.LETTERS_DIR / "L_lead.docx").exists() and not (config.TAILORED_CV_DIR / "cv_lead.docx").exists()
    assert "Documents of 1 job moved to the backup (2 files)." in ui.js("ticker.text")
    assert "documents from a higher score" not in ui.js("listView.model[0].secondary")
    assert ui.find("staleBarText").property("visible") is False                     # nothing left to move …

    undo = ui.find("undoRow")                                                       # … and what was moved can be put back
    assert bar.property("visible") is True and undo.property("visible") is True
    assert ui.find("undoText").property("text").startswith("Documents of 1 job were moved to the backup on 20")
    ui.js("win.appRunning = true; tableDialog.putBack()")
    assert "A search or scan is running. Put the documents back" in ui.find("staleNote").property("text") and not (config.LETTERS_DIR / "L_lead.docx").exists()
    ui.js("win.appRunning = false")
    ui.bridge._running = True
    try:
        assert "A search or scan is running" in ui.bridge.unshelveLast("")["msg"] and "A search or scan is running" in ui.bridge.unshelveDocs("lead")["msg"]
    finally:
        ui.bridge._running = False
    ui.js("tableDialog.lastShelved = {batch: 'gone', when: '2026-01-01T10:00:00', jobs: [{id: 'lead'}]}; tableDialog.putBack()")
    assert "The list changed. Look at it again" in ui.find("staleNote").property("text") and ui.js("tableDialog.staleBad") is True
    assert ui.js("tableDialog.lastShelved.batch") != "gone" and not (config.LETTERS_DIR / "L_lead.docx").exists()      # the real batch is shown again
    ui.js("tableDialog.putBack()")
    assert (config.LETTERS_DIR / "L_lead.docx").exists() and (config.TAILORED_CV_DIR / "cv_lead.docx").exists()
    assert "Documents of 1 job are back in place (2 files)." in ui.js("ticker.text")
    assert undo.property("visible") is False and ui.find("staleBarText").property("visible") is True                    # below the limit again, so offered again
    assert "documents from a higher score" in ui.js("listView.model[0].secondary")
    ui.js("tableDialog.close()")
    ui.wait(350)

    _open_job(ui, "lead")                                                           # the same for one job, in its window
    ui.js("detailDialog.shelve(); detailDialog.shelve()")
    assert not (config.LETTERS_DIR / "L_lead.docx").exists() and ui.find("unshelveDocs").property("visible") is True
    ui.js("detailDialog.close()")
    ui.wait(350)
    _open_job(ui, "lead")
    assert ui.find("staleRow").property("visible") is True and ui.find("staleText").property("text").startswith("The documents the app wrote for this job were moved to the backup on 20")
    ui.js("win.appRunning = true; detailDialog.unshelve()")
    assert "A search or scan is running" in ui.find("staleText").property("text") and ui.js("detailDialog.shelveBad") is True
    ui.js("win.appRunning = false")
    (config.LETTERS_DIR / "L_lead.docx").write_bytes(b"in the way")
    ui.js("detailDialog.unshelve()")
    assert "A file named L_lead.docx is already in its folder" in ui.find("staleText").property("text") and ui.js("detailDialog.shelveBad") is True
    (config.LETTERS_DIR / "L_lead.docx").unlink()
    ui.js("detailDialog.unshelve()")
    assert ui.find("staleText").property("text") == "2 files are back in place." and ui.js("detailDialog.shelveBad") is False
    assert (config.LETTERS_DIR / "L_lead.docx").read_bytes() == b"x" and ui.find("unshelveDocs").property("visible") is False
    assert ui.find("shelveDocs").property("visible") is True
    ui.js("detailDialog.close()")
    ui.wait(350)

    ui.js('openView("app")')
    ui.wait(400)
    assert bar.property("visible") is False                                         # only in New Jobs
    ui.js("tableDialog.close()")
    ui.wait(350)
    ui.bridge._api._change("lead", lambda jobs, j: j.update(match_score=80, status="letter_ready"), render=False)
    assert ui.warnings == []


def test_the_list_shows_the_score_before(ui):
    ui.bridge._api._change("lead", lambda jobs, j: j.update(score_before=55, match_score=30), render=False)
    ui.js('openView("new")')
    ui.wait(450)
    row = ui.js("(function(){ var m = listView.model; for (var i = 0; i < m.length; i++) if (m[i].id === 'lead') return m[i]; return {}; })()")
    assert row["was"] == "was 55" and row["value"] == "30" and ui.js("tableDialog.anyWas") is True
    assert not row["secondary"].startswith(" ") and "·   ·" not in row["secondary"]      # no separator before the first note
    ui.js("tableDialog.close()")
    ui.wait(350)
    ui.bridge._api._change("lead", lambda jobs, j: j.update(score_before=80, match_score=80), render=False)
    ui.js('openView("new")')
    ui.wait(450)
    assert ui.js("tableDialog.anyWas") is False and ui.js("listView.model.map(function (r) { return r.was || ''; }).join('')") == ""
    ui.js("tableDialog.close()")
    ui.wait(350)
    assert "the score it had before it was scored again" in qml_source()
    assert ui.warnings == []


# ── Options: old moved documents are removed only after asking, with the count ──
def test_old_moved_documents_are_removed_from_options_after_asking(ui):
    from datetime import datetime, timedelta

    from jobhelper import housekeeping, vault
    folder = vault.backups_dir() / housekeeping.SHELF
    folder.mkdir(parents=True, exist_ok=True)
    for old in folder.iterdir():
        old.unlink()

    def copy(name, days):
        path = folder / f"{name}.{(datetime.now() - timedelta(days=days, minutes=5)).strftime('%Y%m%d_%H%M%S')}.docx.enc"
        path.write_bytes(b"x" * 4096)
        return path
    fresh, old_one, old_two = copy("fresh", 1), copy("one", 40), copy("two", 90)
    dialog = ui.find("optDialog")
    ui.js("open()", dialog)
    ui.wait(450)
    note, button = ui.find("shelfNote"), ui.find("tidyShelf")
    assert note.property("text") == "3 moved documents are in the backup (12 KB). 2 are older than 30 days (8 KB)."
    assert button.property("text") == "Remove" and button.property("enabled") is True
    ui.js("shelfSpin.value = 200")                                                    # another age: counted again at once
    assert note.property("text").endswith("None is older than 200 days.") and button.property("enabled") is False
    ui.js("optDialog.tidyShelf()")
    assert ui.js("optDialog.confirmTidy") is False and old_two.exists()               # nothing to ask about
    ui.js("shelfSpin.value = 60; optDialog.tidyShelf()")                              # first press only asks
    assert button.property("text") == "Yes, move 1 to the Trash" and "The app can then no longer put them back" in note.property("text") and old_two.exists()
    assert ui.find("tidyCancel").property("visible") is True
    ui.js("shelfSpin.value = 30")                                                     # the age changed: the question is taken back
    assert ui.js("optDialog.confirmTidy") is False and button.property("text") == "Remove"
    ui.js("optDialog.tidyShelf(); win.appRunning = true; optDialog.tidyShelf()")      # a search started meanwhile
    assert "A search or scan is running" in note.property("text") and ui.js("optDialog.shelfBad") is True and old_one.exists()
    ui.js("win.appRunning = false")
    ui.bridge._running = True
    try:
        assert "A search or scan is running" in ui.bridge.tidyShelf(30, 2)["msg"] and old_one.exists()
    finally:
        ui.bridge._running = False
    ui.js("optDialog.tidyShelf()")
    copy("three", 50)                                                                 # the folder changed after the question
    ui.js("optDialog.tidyShelf()")
    assert "The number changed: it is now 3" in note.property("text") and button.property("text") == "Yes, move 3 to the Trash" and old_one.exists()
    ui.js("optDialog.tidyShelf()")
    assert note.property("text") == "Moved 3 old copies to the Trash (12 KB)." and ui.js("optDialog.shelfBad") is False
    assert (housekeeping.trash_dir(old_one) / old_one.name).exists()
    assert not old_one.exists() and not old_two.exists() and fresh.exists() and button.property("enabled") is False
    assert "wrapMode: Text.WordWrap" in qml_source().split('objectName: "shelfNote"')[1][:200]
    ui.js("close()", dialog)
    ui.wait(350)
    ui.js("open()", dialog)
    ui.wait(450)
    assert note.property("text") == "1 moved document is in the backup (4 KB). None is older than 30 days."
    ui.js("close()", dialog)
    ui.wait(350)
    fresh.unlink()
    assert ui.warnings == []


def test_jobs_that_could_not_be_put_back_are_listed_one_by_one(ui, monkeypatch):
    monkeypatch.setattr(ui_api.Api, "_threshold", lambda self: 50)
    jobs = store.load_jobs()
    for name in ("one", "two"):
        (config.LETTERS_DIR / f"L_{name}.docx").write_bytes(b"x")
        jobs.append(store.new_job(id=f"pb_{name}", title=f"Tester {name}", company="Acme <b>Labs</b>", match_score=20, status="letter_ready",
                                  cover_letter_file=f"L_{name}.docx", description="x" * 900))
    store.save_jobs(jobs)
    try:
        ui.js('openView("new")')
        ui.wait(450)
        ui.js("tableDialog.shelveAll(); tableDialog.shelveAll()")
        assert not (config.LETTERS_DIR / "L_one.docx").exists() and ui.js("tableDialog.staleFailed.length") == 0
        (config.LETTERS_DIR / "L_one.docx").write_bytes(b"in the way")
        ui.js("tableDialog.putBack()")
        failed = ui.find("staleFailed")
        assert ui.find("staleNote").property("text") == "Documents of 1 job are back in place (1 file). Not put back:" and ui.js("tableDialog.staleBad") is True
        assert failed.property("visible") is True
        assert failed.property("text") == "Tester one  ·  Acme <b>Labs</b>: A file named L_one.docx is already in its folder, so nothing was put back."
        assert (config.LETTERS_DIR / "L_two.docx").exists() and ui.find("undoRow").property("visible") is True          # the one left is still offered
        ui.js("tableDialog.putBack()")                                                    # nothing could be done this time: said per job too
        assert ui.find("staleNote").property("text") == "Couldn't put the documents back. The reason for each job:"
        assert "Tester one" in failed.property("text") and failed.property("text").count("\n") == 0
        (config.LETTERS_DIR / "L_one.docx").unlink()
        ui.js("tableDialog.putBack()")
        assert ui.js("tableDialog.staleFailed.length") == 0 and failed.property("visible") is False and (config.LETTERS_DIR / "L_one.docx").read_bytes() == b"x"
        source = qml_source().split('objectName: "staleFailed"')[1][:260]
        assert "textFormat: Text.PlainText" in qml_source().split('objectName: "staleFailed"')[0][-120:] + source and "wrapMode: Text.WordWrap" in source
        ui.js("tableDialog.close()")
        ui.wait(350)
        assert ui.js("tableDialog.staleFailed.length") == 0
    finally:
        store.save_jobs([j for j in store.load_jobs() if not j["id"].startswith("pb_")])
        for name in ("one", "two"):
            (config.LETTERS_DIR / f"L_{name}.docx").unlink(missing_ok=True)
    assert ui.warnings == []


# ── light theme: small text must read against the light surfaces ────────────
def _contrast(a, b):
    def lum(h):
        r, g, b_ = (int(h[i:i + 2], 16) / 255 for i in (1, 3, 5))
        f = lambda c: c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4   # noqa: E731
        return 0.2126 * f(r) + 0.7152 * f(g) + 0.0722 * f(b_)
    hi, lo = sorted((lum(a), lum(b)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


def test_light_theme_text_colours_read_on_light_surfaces():
    import re
    main = (ROOT / "ui_qt" / "Main.qml").read_text(encoding="utf-8")
    light = {name: value for name, value in re.findall(r'property color (\w+): win\.dark \? "#[0-9a-fA-F]{6}" : "(#[0-9a-fA-F]{6})"', main)}
    assert {"winbg", "surface", "surface2", "ink", "soft", "faint", "accent", "good", "warn", "crit"} <= set(light)
    for text in ("ink", "soft", "faint", "good", "warn", "crit"):
        for ground in ("winbg", "surface", "surface2"):
            assert _contrast(light[text], light[ground]) >= 4.5, (text, ground, round(_contrast(light[text], light[ground]), 2))
    assert _contrast(light["accentInk"], light["accent"]) >= 4.5 and _contrast(light["accentInk"], light["good"]) >= 4.5   # the label on a main button, the tick
    assert _contrast(light["warn"], light["warnbg"]) >= 4.5 and _contrast(light["good"], light["goodbg"]) >= 4.5 and _contrast(light["crit"], light["critbg"]) >= 4.5


def test_dark_theme_text_colours_read_on_dark_surfaces():
    import re
    main = (ROOT / "ui_qt" / "Main.qml").read_text(encoding="utf-8")
    dark = {name: value for name, value in re.findall(r'property color (\w+): win\.dark \? "(#[0-9a-fA-F]{6})" : "#[0-9a-fA-F]{6}"', main)}
    assert {"winbg", "surface", "surface2", "ink", "soft", "faint", "accent", "accentInk", "good", "warn", "crit"} <= set(dark)
    for text in ("ink", "soft", "faint", "good", "warn", "crit", "accent"):
        for ground in ("winbg", "surface", "surface2"):
            assert _contrast(dark[text], dark[ground]) >= 4.5, (text, ground, round(_contrast(dark[text], dark[ground]), 2))
    assert _contrast(dark["accentInk"], dark["accent"]) >= 4.5 and _contrast(dark["accentInk"], dark["good"]) >= 4.5   # the label on a main button, the tick
    assert _contrast(dark["warn"], dark["warnbg"]) >= 4.5 and _contrast(dark["good"], dark["goodbg"]) >= 4.5 and _contrast(dark["crit"], dark["critbg"]) >= 4.5


def test_the_resume_button_s_white_label_reads_on_its_blue(ui):
    assert _contrast("#ffffff", ui.js('"" + win.resumeBlue')) >= 4.5 and "win.resumeBlue" in qml_source()


def test_no_label_on_the_accent_is_fixed_white():
    """A main button, a picked chip and the tick take their label colour from the theme: white suits only the light theme."""
    for name in ("SmallButton", "PillButton", "TableDialog", "HelpDialog", "OptionsDialog", "DetailDialog", "FormAnswersDialog", "PrepDialog",
                 "PrepEditDialog", "LetterDialog", "CompareDialog"):
        text = (ROOT / "ui_qt" / f"{name}.qml").read_text(encoding="utf-8")
        assert '"#ffffff"' not in text and '"white"' not in text, name
    assert "theme.accentInk" in (ROOT / "ui_qt" / "SmallButton.qml").read_text(encoding="utf-8")


def test_the_label_on_a_main_button_follows_the_theme(ui):
    """Read from the running window: a colour that exists in the file but does not reach the label would pass the other tests."""
    start = ui.js("win.dark")
    try:
        ui.js("win.dark = true")
        assert ui.js('"" + theme.accentInk') == "#1b1305" and ui.js('"" + theme.accent') == "#f0a84a"
        ui.js('openView("new")')
        ui.wait(450)
        ui.js("tableDialog.confirmRescore = true; tableDialog.rescoreJobs = 1")              # a main-style button
        button = ui.find("rescoreAll")
        assert ui.js('"" + contentItem.color', button) == "#1b1305"
        ui.js("win.dark = false")
        ui.wait(50)
        assert ui.js('"" + theme.accentInk') == "#ffffff" and ui.js('"" + contentItem.color', button) == "#ffffff"
        ui.js("tableDialog.close()")
        ui.wait(350)
    finally:
        ui.js("win.dark = %s" % ("true" if start else "false"))
    assert ui.warnings == []


# ── text wraps, it is never cut off ─────────────────────────────────────────
def test_no_text_in_the_window_is_cut_with_an_ellipsis():
    for path in sorted((ROOT / "ui_qt").glob("*.qml")):
        text = path.read_text(encoding="utf-8")
        assert "elide:" not in text and "maximumLineCount" not in text, f"{path.name} cuts a text short; let it wrap"


def test_a_long_title_makes_its_row_taller_instead_of_being_cut(ui):
    long_title = "Senior Principal Associate Vice Coordinator of Cross-Functional Platform Reliability and Developer Experience Enablement (Remote, EMEA)"
    start = store.find(store.load_jobs(), "lead")["title"]
    ui.bridge._api._change("lead", lambda jobs, j: j.update(title=long_title), render=False)
    try:
        ui.js('openView("new")')
        ui.wait(450)
        heights = ui.js("(function(){ var out = []; for (var i = 0; i < listView.count; i++) { var it = listView.itemAtIndex(i); if (it) out.push(Math.round(it.height)); } return out.join(','); })()")
        heights = [int(h) for h in heights.split(",") if h]
        assert max(heights) > 54 and min(heights) >= 54                                 # the long one grew, the others kept their size
        ui.js("tableDialog.close()")
        ui.wait(350)
        _open_job(ui, "lead")
        title = ui.find("detailTitle")
        assert title.property("truncated") is False and title.property("lineCount") >= 2 and title.property("text") == long_title
        ui.js("detailDialog.close()")
        ui.wait(350)
        ui.js('ticker.text = "' + "A long message in the status line that has to be read to its last word. " * 4 + '"')
        ui.wait(100)
        assert ui.js("statusBar.height") > 46 and ui.js("ticker.truncated") is False     # the bar grows with its message
        ui.js('ticker.text = ""')
    finally:
        ui.bridge._api._change("lead", lambda jobs, j: j.update(title=start), render=False)
    assert ui.warnings == []


def test_help_describes_what_was_added_since_it_was_written():
    text = (ROOT / "ui_qt" / "HelpDialog.qml").read_text(encoding="utf-8")
    for words in ("Move documents to backup", "Put them back", "Trash", "Kinds of role you want", "Read it again", "was 55",
                  "no longer accepts applications", "app_trouble.log", "docs/GUIDE.md"):
        assert words in text, words


# ── the illustrated guide ───────────────────────────────────────────────────
def test_a_search_role_outside_software_can_be_added(ui):
    """The words the app knew for roles were all from software jobs: "registered nurse" was refused as not a role."""
    dialog = ui.find("optDialog")
    ui.js("open()", dialog)
    ui.wait(450)
    before = ui.js("rolesModel.count")
    assert ui.js('optDialog.addRole("registered nurse")') is True and ui.js("optDialog.roleErr") == ""       # a profession the app knows
    assert ui.js('optDialog.addRole("junior accountant")') is True
    assert ui.js('optDialog.addRole("beekeeper")') is False                                                   # unknown: asked once …
    assert "is not a role the app knows" in ui.js("optDialog.roleErr") and "press Add again" in ui.js("optDialog.roleErr")
    assert ui.js("rolesModel.count") == before + 2
    assert ui.js('optDialog.addRole("beekeeper")') is True and ui.js("optDialog.roleErr") == ""               # … then added
    assert ui.js('optDialog.addRole("asdfgh")') is False and ui.js('optDialog.addRole("zzzzqq")') is False     # a different word is asked about afresh
    assert ui.js("rolesModel.count") == before + 3
    assert ui.js('optDialog.addRole("план")') is False and "Latin letters" in ui.js("optDialog.roleErr")       # the other checks still hold
    assert ui.js('optDialog.addRole("beekeeper")') is False and "already in the list" in ui.js("optDialog.roleErr")
    assert ui.bridge.knownRole("staff nurse") is True and ui.bridge.knownRole("beekeeper") is False and ui.bridge.knownRole("") is False
    ui.js("optDialog.pendingRole = ''; optDialog.roleErr = ''; optDialog.load()")
    ui.js("close()", dialog)
    ui.wait(350)
    assert ui.warnings == []


def test_the_job_window_asks_before_closing_with_unsaved_text(ui):
    _open_job(ui, "app")
    dialog = ui.find("detailDialog")
    assert ui.js("detailDialog.pending.length") == 0 and ui.js("detailDialog.closePolicy") != 0           # nothing typed: Esc closes
    ui.js('cnNotes.text = "They asked about pytest."')
    assert ui.js("detailDialog.pending.join('|')") == "the notes about the company" and ui.js("detailDialog.closePolicy") == 0   # Esc alone no longer closes
    ui.js("detailDialog.tryClose()")                                                                    # the ✕ (and Esc, through the shortcut)
    note = ui.find("closeNote")
    assert dialog.property("opened") is True and note.property("visible") is True
    assert note.property("text") == "Not saved yet: the notes about the company. Save it, or close once more to discard it."
    ui.js("detailDialog.saveNotes()")                                                                   # saved: nothing pending any more
    assert ui.js("detailDialog.pending.length") == 0
    ui.js('evNote.text = "bring the laptop"; replyText.text = "Dear Alex, thank you for your time today."')
    assert ui.js("detailDialog.pending.join('|')") == "the pasted reply|the date to remember"
    ui.js("detailDialog.confirmClose = false; detailDialog.tryClose()")
    assert "the pasted reply, the date to remember" in note.property("text") and dialog.property("opened") is True
    ui.js("detailDialog.tryClose()")                                                                    # the second time closes and discards
    ui.wait(350)
    assert dialog.property("opened") is False
    _open_job(ui, "app")                                                                                # opened again: clean
    assert ui.js("detailDialog.pending.length") == 0 and ui.js("detailDialog.closeNote") == "" and ui.js("replyText.text") == ""
    ui.js("detailDialog.tryClose()")
    ui.wait(350)
    assert dialog.property("opened") is False
    ui.bridge._api.save_company_notes("Beta", {"contact": "", "salary": "", "asked": "", "notes": ""})
    _open_job(ui, "lead")                                                                               # a new job has none of those fields filled
    assert ui.js("detailDialog.pending.length") == 0
    ui.js("detailDialog.close()")
    ui.wait(350)
    assert ui.warnings == []


def test_a_reply_is_filed_under_another_application_from_its_window(ui, monkeypatch):
    from jobhelper import db
    db.upsert_response({"msg_id": "<move>", "received_date": "Tue, 18 Aug 2026 10:00:00 +0200", "from_addr": "hr@beta.example", "subject": "Interview",
                        "job_id": "", "company": "Beta", "role": "QA Engineer", "stage": "interview", "summary": "They invite you to a call."})
    dialog = ui.find("responseDialog")
    ui.js('responseDialog.show("<move>")')
    ui.wait(400)
    button, note = ui.find("fileUnder"), ui.find("fileNote")
    assert button.property("text") == "File under…" and ui.find("filePicker").property("visible") is False
    ui.js("win.appRunning = true; responseDialog.startPicking()")
    assert dialog.property("picking") is False and dialog.property("fileBad") is True and "running" in note.property("text")
    ui.js("win.appRunning = false; responseDialog.startPicking()")
    assert dialog.property("picking") is True and button.property("text") == "Cancel"
    assert ui.find("pickList").property("count") == 1                                # the one application of the test data
    ui.find("pickFilter").setProperty("text", "nothing like this")
    ui.js('responseDialog.targets = bridge.filingTargets("nothing like this")')
    ui.wait(60)
    assert ui.find("pickList").property("count") == 0 and "No application matches" in ui.find("pickEmpty").property("text")
    ui.js('responseDialog.fileUnder("lead")')                                        # a new job, not an application: refused, said in red
    assert dialog.property("fileBad") is True and "Mark that job as applied first" in note.property("text") and dialog.property("picking") is True
    ui.js('responseDialog.fileUnder("app")')
    ui.wait(80)
    assert dialog.property("fileBad") is False and "Filed under Beta · QA Engineer." in note.property("text") and dialog.property("picking") is False
    assert ui.js("responseDialog.r.job_id") == "app" and button.property("text") == "File under another…"
    assert ui.bridge.jobDetail("app")["status"] == "interview"
    ui.js("responseDialog.startPicking()")
    assert ui.find("fileNowhere").property("visible") is True
    ui.js('responseDialog.fileUnder("")')
    assert ui.js("responseDialog.r.job_id") == "" and "not filed under any application now" in note.property("text")
    ui.js("responseDialog.close()")
    ui.wait(350)
    ui.bridge._api._change("app", lambda jobs, j: j.update(status="waiting"), render=False)
    ui.bridge._set_running(True)
    assert "running" in ui.bridge.refileResponse("<move>", "app")["msg"]              # the backend refuses the same
    ui.bridge._set_running(False)
    ui.js('win.runMode = ""; win.resetRobots()')
    assert ui.warnings == []


def test_the_job_window_writes_documents_on_request(ui, monkeypatch):
    asked = []
    monkeypatch.setattr(ui_api.Api, "can_write_documents", lambda self, j: "" if j and j["id"] == "lead" else "no")
    monkeypatch.setattr(ui_api.Api, "write_documents", lambda self, job_id: asked.append(job_id) or {"ok": True, "cv": True, "letter": True,
                                                                                                      "msg": "Written: the CV and the letter."})
    ui.bridge._api._change("lead", lambda jobs, j: j.update(tailored_cv_file="", cover_letter_file=""), render=False)   # whatever earlier tests left
    ui.js('detailDialog.show("lead")')
    ui.wait(450)
    button, note = ui.find("writeDocs"), ui.find("docsNote")
    assert ui.find("writeDocsRow").property("visible") is True and button.property("text") == "Write CV and letter"
    assert "No CV or letter was written for this job" in note.property("text")
    ui.js("win.appRunning = true; detailDialog.writeDocs()")
    assert asked == [] and ui.find("detailDialog").property("docsBad") is True and "running" in note.property("text")
    ui.js("win.appRunning = false; detailDialog.writeDocs()")
    assert ui.find("detailDialog").property("docsBusy") is True and button.property("text") == "Writing…" and "about a minute" in note.property("text")
    ui.js("detailDialog.writeDocs()")                                # a second press while it is writing: said, not started twice
    for _ in range(100):
        ui.wait(20)
        if not ui.find("detailDialog").property("docsBusy"):
            break
    assert asked == ["lead"] and note.property("text") == "Written: the CV and the letter." and ui.find("detailDialog").property("docsBad") is False
    monkeypatch.setattr(ui_api.Api, "write_documents", lambda self, job_id: {"ok": False, "msg": "Claude's 5-hour limit is used up."})
    ui.js("detailDialog.writeDocs()")
    for _ in range(100):
        ui.wait(20)
        if not ui.find("detailDialog").property("docsBusy"):
            break
    assert note.property("text") == "Not written: Claude's 5-hour limit is used up." and ui.find("detailDialog").property("docsBad") is True
    ui.js("detailDialog.close()")
    ui.wait(350)
    ui.js('detailDialog.show("app")')                                # an application: the row is not offered
    ui.wait(450)
    assert ui.find("writeDocsRow").property("visible") is False
    ui.js("detailDialog.close()")
    ui.wait(350)
    assert ui.warnings == []


def test_help_shows_usage_and_the_log_of_problems(ui, monkeypatch):
    monkeypatch.setattr(ui_api.Api, "usage_overview", lambda self: {"calls": 12, "cost": 3.4, "kinds": [
        {"name": "Scoring jobs and writing letters", "calls": 9, "cost": 3.0}, {"name": "Reading replies", "calls": 3, "cost": 0.4}]})
    monkeypatch.setattr(ui_api.Api, "trouble_log", lambda self, most=25: {"lines": ["2026-10-09 12:15  a sound: it failed"], "total": 40, "file": "app_trouble.log"})
    shown = []
    monkeypatch.setattr(ui_api.Api, "show_trouble_log", lambda self: shown.append(1) or {"ok": True})
    ui.js("helpDialog.open()")
    ui.wait(400)
    assert ui.find("usageCard").property("visible") is False                         # on its own tab
    ui.js('helpDialog.tab = "use"')
    ui.wait(120)
    assert ui.find("usageCard").property("visible") is True and ui.find("logCard").property("visible") is True
    total = ui.find("usageTotal").property("text")
    assert "12 questions in all" in total and "$3.40" in total and "subscription only" in total
    texts = {t.property("text") for t in ui.seen("usageCard") if isinstance(t.property("text"), str)}
    assert {"Scoring jobs and writing letters", "9 questions", "≈ $3.00", "≈ $0.40"} <= texts
    log_texts = {t.property("text") for t in ui.seen("logCard") if isinstance(t.property("text"), str)}
    assert "2026-10-09 12:15  a sound: it failed" in log_texts and any("1 of 40 shown" in t for t in log_texts)
    ui.js("bridge.showTroubleLog()")
    assert shown == [1]
    monkeypatch.setattr(ui_api.Api, "usage_overview", lambda self: {"calls": 0, "cost": 0.0, "kinds": []})
    monkeypatch.setattr(ui_api.Api, "trouble_log", lambda self, most=25: {"lines": [], "total": 0, "file": "app_trouble.log"})
    ui.js("helpDialog.close()")
    ui.wait(300)
    ui.js('helpDialog.open(); helpDialog.tab = "use"')
    ui.wait(300)
    assert "Nothing yet" in ui.find("usageTotal").property("text") and ui.find("logShow").property("visible") is False
    ui.js('helpDialog.tab = "how"; helpDialog.close()')
    ui.wait(300)
    assert ui.warnings == []


def test_options_make_and_check_backups(ui, monkeypatch):
    asked = []
    monkeypatch.setattr(ui_api.Api, "backup_all", lambda self: asked.append("backup") or {"ok": True, "folder": "everything_x", "msg": "Backed up 12 files (1.0 MB), encrypted."})
    monkeypatch.setattr(ui_api.Api, "backup_check", lambda self: asked.append("check") or {"ok": False, "msg": "Something is wrong with the backups: jobs.json."})
    monkeypatch.setattr(ui_api.Api, "backup_export", lambda self, folder: asked.append(("export", folder)) or {"ok": True, "msg": "Saved 12 readable files."})
    dialog = ui.find("optDialog")
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    note = ui.find("storageNote")

    def done():
        for _ in range(150):
            ui.wait(20)
            if dialog.property("storageBusy") == "":
                return
        raise AssertionError("the backup action never answered")
    ui.js("win.appRunning = true; optDialog.storage('backup', '')", dialog)
    assert asked == [] and dialog.property("storageBad") is True and "running" in note.property("text")
    ui.js("win.appRunning = false; optDialog.storage('backup', '')", dialog)
    assert dialog.property("storageBusy") == "backup" and ui.find("backupNow").property("text") == "Backing up…"
    assert ui.find("backupCheck").property("enabled") is False                       # one at a time
    done()
    assert note.property("text") == "Backed up 12 files (1.0 MB), encrypted." and dialog.property("storageBad") is False
    ui.js("optDialog.storage('check', '')", dialog)
    done()
    assert dialog.property("storageBad") is True and "Something is wrong" in note.property("text")
    ui.js("optDialog.storage('export', 'file:///tmp/somewhere')", dialog)
    done()
    assert asked == ["backup", "check", ("export", "file:///tmp/somewhere")] and note.property("text") == "Saved 12 readable files."
    ui.bridge._set_running(True)
    ui.js("optDialog.storageBusy = ''; bridge.storage('backup', '')", dialog)         # the backend refuses the same during a run
    ui.wait(200)
    assert asked.count("backup") == 1
    ui.bridge._set_running(False)
    ui.js('win.runMode = ""; win.resetRobots()')
    ui.js("confirmClose = true; close()", dialog)
    ui.wait(350)
    assert ui.warnings == []


def test_the_account_card_shows_what_changed_in_a_watched_file(ui, monkeypatch):
    from jobhelper import integrity
    monkeypatch.setattr(integrity, "problems", lambda: [{"name": "cv_rules", "title": "CV rules", "state": "changed", "file": "cv_rules.md"},
                                                         {"name": "letter_rules", "title": "Cover letter rules", "state": "changed", "file": "l.md"}])
    monkeypatch.setattr(integrity, "what_changed", lambda name: {"title": "CV rules", "file": "cv_rules.md", "known": True, "more": 3,
                                                                 "added": ["Always add every skill."], "removed": ["Never add a skill."]} if name == "cv_rules"
                        else {"title": "Cover letter rules", "file": "l.md", "known": False, "added": [], "removed": [], "more": 0})
    ui.js("welcome.openCard()")
    ui.win.setProperty("connData", {**READY, "vault": "ok", "changed_files": ["CV rules", "Cover letter rules"]})
    ui.wait(80)
    button = ui.find("seeChanges")
    assert button.property("text") == "See what changed" and ui.js("welcome.changes.length") == 0
    ui.js("welcome.changes = bridge.whatChanged()")
    ui.wait(80)
    assert ui.js("welcome.changes.length") == 2 and button.property("text") == "Hide changes"
    assert "visible: welcome.changes.length > 0 && welcome.changedFiles.length > 0" in qml_source()
    texts = {t.property("text") for t in ui.seen("changesPanel") if isinstance(t.property("text"), str)}
    assert {"CV rules · cv_rules.md", "− Never add a skill.", "+ Always add every skill."} <= texts
    assert any("3 more changed lines" in t for t in texts) and any("No earlier copy of this file was kept" in t for t in texts)
    ui.js("welcome.changes = []")
    ui.win.setProperty("connData", dict(READY))
    ui.js("welcome.finish()")
    ui.wait(700)
    assert ui.js("welcome.changes.length") == 0, "put away when the card closes"
    assert ui.warnings == []


def test_options_take_a_salary_of_its_own_for_a_kind_of_role(ui):
    import json
    dialog = ui.find("optDialog")
    start = ui.bridge.options()
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    role, amount, err = ui.find("payRole"), ui.find("payAmount"), ui.find("payErr")
    for typed, value, part in (("", 40000, "kind of role in words"), ("7", 40000, "kind of role in words"), ("поддержка", 40000, "in English"),
                               ("support", 0, "a number above 0")):
        role.setProperty("text", typed); amount.setProperty("value", value)
        ui.js("optDialog.addPay()", dialog)
        assert part in err.property("text") and ui.js("payModel.count", dialog) == 0, typed
    role.setProperty("text", " support "); amount.setProperty("value", 42000)
    ui.js("optDialog.addPay()", dialog)
    assert ui.js("payModel.count", dialog) == 1 and err.property("visible") is False and role.property("text") == ""
    role.setProperty("text", "Support"); amount.setProperty("value", 50000)
    ui.js("optDialog.addPay()", dialog)
    assert "already in the list" in err.property("text") and ui.js("payModel.count", dialog) == 1
    got = json.loads(ui.js("JSON.stringify(optDialog.collect().options)", dialog))
    assert got["salary_by_role"] == [{"role": "support", "amount": 42000}]
    assert ui.bridge.saveOptions(got)["ok"] is True
    ui.js("confirmClose = true; close()", dialog)
    ui.wait(350)
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    assert ui.js("payModel.count", dialog) == 1 and ui.js("payModel.get(0).role", dialog) == "support"      # loaded again
    chips = {t.property("text") for t in ui.seen("payChips") if isinstance(t.property("text"), str)}
    assert any(t.startswith("support · 42 000 ") for t in chips)
    ui.js("payModel.remove(0); confirmClose = true; close()", dialog)
    ui.wait(350)
    assert ui.bridge.saveOptions({**ui.bridge.options(), "salary_by_role": start.get("salary_by_role", [])})["ok"] is True
    ui.js("optsData = bridge.options()")
    assert ui.warnings == []


def test_the_account_card_tries_the_saved_mailboxes(ui, monkeypatch):
    from jobhelper import email_client
    monkeypatch.setattr(config, "IMAP_USER", "me@gmail.com")
    monkeypatch.setattr(config, "IMAP_PASSWORD", "abcdabcdabcdabcd")
    monkeypatch.setattr(config, "IMAP_USER_2", "two@seznam.cz")
    monkeypatch.setattr(config, "IMAP_PASSWORD_2", "secret")
    monkeypatch.setattr(email_client, "try_login", lambda user, password, timeout=8.0, host="":
                        {"ok": True, "msg": ""} if "gmail" in user else {"ok": False, "msg": "The mail server refused this address and password."})
    welcome = ui.find("welcome")
    ui.js("welcome.openCard()")
    ui.win.setProperty("connData", {**READY, "email": True, "email2": True, "email2_addr": "two@seznam.cz"})
    ui.js("welcome.testMail()")
    assert welcome.property("testBusy") is True and ui.find("mailTest").property("text") == "Testing…"
    ui.js("welcome.testMail()")                                     # a second press while it is testing does nothing
    for _ in range(200):
        ui.wait(20)
        if not welcome.property("testBusy"):
            break
    note = welcome.property("testNote")
    assert "Signed in to me@gmail.com" in note and "refused" in note and welcome.property("testBad") is True
    assert "secret" not in note and "abcdabcd" not in note
    ui.win.setProperty("connData", dict(READY))
    ui.js("welcome.finish()")
    ui.wait(700)
    assert welcome.property("testNote") == "" and ui.warnings == []


def test_a_long_company_name_in_options_wraps_and_stays_clear_of_the_remove_mark(ui):
    dialog = ui.find("optDialog")
    ui.js("optsData = bridge.options()")
    ui.js("open()", dialog)
    ui.wait(450)
    ui.js('companiesModel.append({ name: "The Very Long Named International Engineering and Consulting Company of Central Europe Limited", ats: "greenhouse", token: "' + "x" * 60 + '" })', dialog)
    ui.wait(120)
    line = next(i for i in ui.seen("optScroll") if i.objectName() == "companyLine")
    name = next(i for i in line.childItems() if i.objectName() == "companyName")
    assert name.property("lineCount") >= 2 and name.property("truncated") is False
    assert line.mapToItem(None, 0, 0).x() + line.width() <= ui.find("optScroll").mapToItem(None, 0, 0).x() + ui.find("optScroll").width() - 30
    ui.js("companiesModel.remove(companiesModel.count - 1); confirmClose = true; close()", dialog)
    ui.wait(350)
    assert ui.warnings == []


def test_the_job_window_says_what_to_check_in_the_adapted_cv(ui):
    ui.bridge._api._change("lead", lambda jobs, j: j.update(tailored_cv_file="cv_lead.docx", cv_check=["cost center", "vendor master"]), render=False)
    ui.js('detailDialog.show("lead")')
    ui.wait(450)
    note = ui.find("cvCheck")
    assert ui.js("detailDialog.job.cv_check.length") == 2
    assert "It speaks of cost center, vendor master, which are terms of the posting that your own CV does not mention" in note.property("text")
    ui.js("detailDialog.close()")
    ui.wait(350)
    ui.bridge._api._change("lead", lambda jobs, j: j.update(cv_check=["cost center"]), render=False)
    ui.js('detailDialog.show("lead")')
    ui.wait(450)
    assert "which is a term of the posting" in note.property("text")
    ui.js("detailDialog.close()")
    ui.wait(350)
    ui.bridge._api._change("lead", lambda jobs, j: j.update(cover_letter_file="L_lead.docx", letter_check=["tax codes"]), render=False)
    ui.js('detailDialog.show("lead")')
    ui.wait(450)
    assert "It says you have done something with tax codes, which is a term of the posting" in ui.find("letterCheck").property("text")
    ui.js("detailDialog.close()")
    ui.wait(350)

    def clear(jobs, j):
        j.pop("cv_check", None)
        j.pop("letter_check", None)
        j["tailored_cv_file"] = j["cover_letter_file"] = ""
    ui.bridge._api._change("lead", clear, render=False)
    assert ui.bridge.jobDetail("lead")["letter_check"] == []
    assert ui.bridge.jobDetail("lead")["cv_check"] == []
    assert ui.warnings == []


def test_new_jobs_can_be_reviewed_one_by_one(ui, monkeypatch):
    """Going through thirty "stretch" jobs meant opening each from the list and coming back; now one window walks through them."""
    kept = store.load_jobs()
    text = "We build tools in Python. " * 30
    store.save_jobs(kept + [store.new_job(id=f"r{n}", title=f"Role {n}", company=f"Firm {n}", status="letter_ready", match_score=40 - n,
                                          match_tier="stretch", description=text, cover_letter_file=f"L{n}.docx", tailored_cv_file=f"C{n}.docx")
                            for n in range(4)])
    done = []
    monkeypatch.setattr(qt_bridge.Bridge, "applyJob", qt_bridge.Slot(str, result=bool)(lambda self, job_id: done.append(("apply", job_id)) or True))
    monkeypatch.setattr(ui_api.Api, "dismiss_job", lambda self, job_id: done.append(("dismiss", job_id)) or {"ok": True})
    try:
        ui.js('openView("new")')
        ui.wait(450)
        table, detail = ui.find("tableDialog"), ui.find("detailDialog")
        ui.js('tableDialog.groupFilter = "stretch"; tableDialog.apply()')
        ui.wait(120)
        start = ui.find("reviewStart")
        assert start.property("visible") is True and start.property("text") == "Review these 4 one by one"
        assert ui.js("tableDialog.reviewIds().join(',')") == "r0,r1,r2,r3"                  # the list as it is shown, in its order
        ui.js("detailDialog.startReview(tableDialog.reviewIds())")
        ui.wait(450)
        where = ui.find("reviewWhere")
        assert detail.property("opened") is True and ui.js("detailDialog.job.id") == "r0" and ui.find("reviewBar").property("visible") is True
        assert where.property("text").startswith("Reviewing 1 of 4 left.")
        assert ui.find("reviewPrev").property("enabled") is False and ui.find("reviewNext").property("enabled") is True
        ui.js("detailDialog.reviewStep(1)")                                                 # leave it as it is, on to the next
        ui.wait(200)
        assert ui.js("detailDialog.job.id") == "r1" and where.property("text").startswith("Reviewing 2 of 4 left.") and done == []
        ui.js("detailDialog.notInterested()")                                               # decided: the next one comes by itself
        ui.wait(250)
        assert done == [("dismiss", "r1")] and ui.js("detailDialog.job.id") == "r2" and detail.property("opened") is True
        assert where.property("text").startswith("Reviewing 2 of 3 left · 1 decided · Removed from New Jobs.")
        ui.js("detailDialog.reviewStep(-1)")
        ui.wait(200)
        assert ui.js("detailDialog.job.id") == "r0"
        # unsaved text is asked about before moving on, as everywhere in this window
        ui.js('cnNotes.text = "They pay monthly."')
        ui.js("detailDialog.reviewStep(1)")
        assert ui.js("detailDialog.job.id") == "r0" and "press Next once more" in ui.find("closeNote").property("text")
        ui.js("detailDialog.reviewStep(1)")
        ui.wait(200)
        assert ui.js("detailDialog.job.id") == "r2"
        ui.js("detailDialog.markApplied()")
        ui.wait(250)
        assert done[-1] == ("apply", "r2") and ui.js("detailDialog.job.id") == "r3" and table.property("opened") is True     # the list stays behind it
        ui.js("detailDialog.notInterested()")
        ui.wait(250)
        assert ui.js("detailDialog.job.id") == "r0" and where.property("text").startswith("Reviewing 1 of 1 left · 3 decided")
        assert ui.find("reviewNext").property("enabled") is False
        ui.js("detailDialog.notInterested()")                                               # the last one: the round ends
        ui.wait(450)
        assert detail.property("opened") is False and "That was the last of the 4 jobs." in ui.js("ticker.text")
        assert ui.js("detailDialog.reviewQueue.length") == 0
        # opened the ordinary way, the window has no review bar and closes after a decision as before
        ui.js('detailDialog.show("lead")')
        ui.wait(450)
        assert ui.find("reviewBar").property("visible") is False
        ui.js("detailDialog.startReview(['lead', 'r0']); detailDialog.close()")             # closing ends a round
        ui.wait(450)
        assert ui.js("detailDialog.reviewQueue.length") == 0
        ui.js('tableDialog.groupFilter = ""; tableDialog.apply(); tableDialog.close()')
        ui.wait(350)
    finally:
        store.save_jobs(kept)
        ui.js("win.reloadAll()")
    assert ui.warnings == []


# ── the assistant ───────────────────────────────────────────────────────────
def test_the_assistant_stands_on_the_line_blinks_and_waves(ui):
    bot = ui.find("assistantBot")
    assert bot.property("visible") and bot.property("x") + bot.property("width") <= 1000 and bot.property("y") + bot.property("height") <= 500
    for robot in ("a4", "a5"):                                  # it does not stand on a station of the line
        other = ui.find(robot)
        assert bot.property("x") > other.property("x") + 130
    lids = []
    for _ in range(90):                                         # the eyes close and open by themselves
        lids.append(bot.property("lid"))
        ui.wait(40)
        if max(lids) > 0.6 and lids[-1] < 0.2:
            break
    assert max(lids) > 0.6 and lids[-1] < 0.2, "it blinked"
    arm, bubble = ui.drawn("assistantArm", "stagePanel"), ui.drawn("assistantBubble", "stagePanel")
    assert arm.property("rotation") > -40 and bubble.property("opacity") == 0
    ui.js("stagePanel.assistant.greet()")
    ui.wait(500)
    assert bot.property("greeting") and arm.property("rotation") < -100, "the arm is up, waving"
    assert bubble.property("opacity") == 1 and bot.property("says").startswith("Hi!")
    assert bubble.mapToItem(bot.parentItem(), 0, 0).x() >= 0, "the bubble stays on the floor"
    wide = bubble.property("width")
    bot.setProperty("greeting", False)
    sizes = []
    for _ in range(12):                                         # on its way out the bubble keeps its words and its size
        sizes.append((bot.property("shown"), bubble.property("width")))
        ui.wait(20)
    assert set(sizes) == {("Hi! Tap me to fill in a form.", wide)}, "it used to switch to the other text and jump while fading"
    ui.wait(400)
    assert arm.property("rotation") > -40 and bubble.property("opacity") == 0
    bot.setProperty("hovered", True)
    ui.wait(400)
    assert bot.property("shown") == "Assistant · tap to ask" and bubble.property("opacity") == 1
    bot.setProperty("hovered", False)
    ui.wait(400)
    assert ui.warnings == []


def test_the_assistant_asks_for_your_details_first_and_checks_them(ui):
    from jobhelper import details
    dialog = ui.find("assistantDialog")
    ui.js('stagePanel.assistant.tapped()')
    ui.wait(500)
    assert dialog.property("opened") and dialog.property("page") == "details", "nothing is saved yet: the details come first"
    assert ui.js("fieldModel.count") == len(details.KEYS) and ui.js("talk.count") == 1
    ui.js("assistantDialog.saveDetails()")
    assert "Every field is empty" in ui.find("asDetailsNote").property("text") and not details.path().exists()
    ui.js('assistantDialog.setField("first_name", "Alex"); assistantDialog.setField("last_name", "Novak"); assistantDialog.setField("email", "alex@example")')
    assert dialog.property("dirty") and ui.js("assistantDialog.problem.field") == "email"
    assert "email address" in ui.find("asDetailsNote").property("text")
    ui.js("assistantDialog.saveDetails()")
    assert not details.path().exists(), "a wrong field is not saved"
    ui.js("assistantDialog.tryClose()")
    assert dialog.property("opened") and "did not save" in ui.find("asDetailsNote").property("text"), "closing with unsaved changes asks once"
    ui.js('assistantDialog.setField("email", "alex@example.com"); assistantDialog.setField("linkedin", "linkedin.com/in/alexnovak")')
    assert ui.js("!assistantDialog.problem.msg")
    ui.js("assistantDialog.saveDetails()")
    assert details.load()["email"] == "alex@example.com" and details.load()["linkedin"] == "https://linkedin.com/in/alexnovak"
    assert not dialog.property("dirty") and dialog.property("detailsSaved") and "Saved" in ui.find("asDetailsNote").property("text")
    ui.js("assistantDialog.tryClose()")
    ui.wait(400)
    assert dialog.property("visible") is False and ui.warnings == []


def test_the_assistant_takes_a_request_and_shows_what_it_filled(ui, monkeypatch):
    import threading
    dialog, gate, asked = ui.find("assistantDialog"), threading.Event(), []

    def answer(self, text, job_id=""):
        asked.append((text, job_id))
        gate.wait(5)
        return {"ok": True, "action": "fill", "host": "careers.northwind.example", "job_id": "lead", "job": "Junior Python Developer at Acme", "guessed": True,
                "filled": [{"label": "First name", "value": "Alex", "kind": "detail", "check": ""},
                           {"label": "When can you start?", "value": "I can start on [your start date].", "kind": "answer", "check": "Your start date"}],
                "left": [{"label": "Resume/CV", "why": "Attach the file yourself: a page cannot be handed a file from outside."}],
                "kept": 0, "boxes": 1, "say": "I filled in 2 fields on careers.northwind.example. I never press Submit."}
    monkeypatch.setattr(ui_api.Api, "assistant_ask", answer)
    ui.js("assistantDialog.show()")
    ui.wait(450)
    assert dialog.property("page") == "ask", "details are saved: it opens ready to be asked"
    assert [j["id"] for j in dialog.property("jobs").toVariant()] == ["", "lead", "app"]
    before = ui.js("talk.count")
    ui.js('assistantDialog.ask("   ")')
    assert "Write what you need" in ui.find("asErr").property("text") and ui.js("talk.count") == before and asked == []
    ui.js('assistantDialog.ask("' + "x" * 301 + '")')
    assert "longer than 300" in ui.find("asErr").property("text") and asked == []
    ui.js('assistantDialog.forJob("app")')
    ui.js('assistantDialog.ask("fill the form in chrome")')
    assert dialog.property("busy") and ui.find("assistantBot").property("busy") and ui.find("asFill").property("enabled") is False
    ui.js('assistantDialog.ask("again")')
    assert "still working" in ui.find("asErr").property("text")
    ui.js("assistantDialog.tryClose()")
    assert dialog.property("opened"), "it is not closed in the middle of a request"
    gate.set()
    for _ in range(100):
        if not dialog.property("busy"):
            break
        ui.wait(30)
    assert asked == [("fill the form in chrome", "app")]
    assert ui.js("talk.count") == before + 2 and ui.js("talk.get(talk.count - 1).text").startswith("I filled in 2 fields")
    assert ui.js("filledModel.count") == 2 and ui.js("filledModel.get(1).check") == "Your start date" and ui.js("leftModel.get(0).label") == "Resume/CV"
    monkeypatch.setattr(ui_api.Api, "assistant_ask", lambda self, text, job_id="": {"ok": False, "action": "fill", "msg": "Google Chrome is not open."})
    ui.js("assistantDialog.ask(\"fill it\")")
    for _ in range(100):
        if not dialog.property("busy"):
            break
        ui.wait(30)
    assert ui.js("talk.get(talk.count - 1).bad") is True and ui.js("talk.get(talk.count - 1).text") == "Google Chrome is not open."
    assert ui.js("filledModel.count") == 0, "an old report does not stay under a new answer"
    ui.js("assistantDialog.tryClose()")
    ui.wait(400)
    assert dialog.property("visible") is False and ui.warnings == []


def test_the_assistant_looks_at_the_pointer_nods_and_shakes_its_head(ui):
    bot = ui.find("assistantBot")
    eyes, head = ui.drawn("assistantEyes", "stagePanel"), ui.drawn("assistantHead", "stagePanel")
    rest_x, rest_y = eyes.property("x"), head.property("y")
    ui.js("stagePanel.assistant.lookAt(100, 100)")                # far to the left and above
    ui.wait(300)
    assert bot.property("lookX") == -1 and bot.property("lookY") == -1 and eyes.property("x") < rest_x - 3
    ui.js("stagePanel.assistant.lookAt(stagePanel.assistant.x + 60, stagePanel.assistant.y + 36)")
    ui.wait(300)
    assert 0 < bot.property("lookX") < 0.3, "a pointer close by moves the eyes a little"
    ui.js("stagePanel.assistant.lookAhead()")
    ui.wait(300)
    assert bot.property("lookX") == 0 and abs(eyes.property("x") - rest_x) < 0.01
    for move, part, name in (("nod", "y", "nodBy"), ("shake", "x", "shakeBy")):
        start = head.property(part)
        ui.js(f"stagePanel.assistant.{move}()")
        seen = []
        for _ in range(40):
            seen.append(head.property(part) - start)
            ui.wait(20)
        assert max(abs(s) for s in seen) > 3, f"the head moved ({move})"
        assert bot.property(name) == 0 and abs(head.property(part) - start) < 0.01, "and came back"
    assert head.property("y") == rest_y and ui.warnings == []


def _wait_idle(ui, dialog):
    for _ in range(100):
        if not dialog.property("busy"):
            return
        ui.wait(30)


def test_after_a_fill_the_assistant_offers_what_applies(ui, monkeypatch):
    dialog, did = ui.find("assistantDialog"), []
    fill = {"ok": True, "action": "fill", "host": "careers.acme.example", "url": "https://careers.acme.example/apply", "job_id": "lead",
            "job": "Junior Python Developer at Acme", "guessed": False, "kept": 0, "boxes": 0, "say": "I filled in 1 field.",
            "filled": [{"label": "Why Acme?", "value": "Because.", "kind": "answer", "check": ""}],
            "left": [{"label": "Resume/CV", "why": "Attach the file yourself: a page cannot be handed a file from outside."}],
            "wants_file": True, "has_files": True, "is_lead": True, "can_learn": True}
    monkeypatch.setattr(ui_api.Api, "assistant_ask", lambda self, text, job_id="": dict(fill))
    monkeypatch.setattr(ui_api.Api, "assistant_watch", lambda self: {"ok": False, "url": "", "sign": "", "empty": 0})
    monkeypatch.setattr(ui_api.Api, "assistant_files", lambda self, job_id: did.append(("files", job_id)) or {"ok": True, "say": "Shown in Finder: CV."})
    monkeypatch.setattr(ui_api.Api, "assistant_learn", lambda self: did.append(("learn",)) or {"ok": True, "learned": 1, "say": "I will use your wording for 1 question."})
    monkeypatch.setattr(ui_api.Api, "assistant_memory", lambda self: {"count": len([d for d in did if d[0] == "learn"])})
    monkeypatch.setattr(ui_api.Api, "assistant_forget", lambda self: did.append(("forget",)) or {"ok": True, "say": "Forgotten: 1 remembered answers."})
    ui.js("assistantDialog.show()")
    ui.wait(450)
    assert not ui.find("asNext").property("visible"), "nothing is offered before anything was done"
    ui.js('assistantDialog.ask("fill the form")')
    _wait_idle(ui, dialog)
    assert all(dialog.property(name) for name in ("canShowFiles", "canLearn", "canApply")) and not dialog.property("canOpenForm")
    assert ui.find("asNext").property("visible") and ui.find("asFiles").property("visible") and not ui.find("asOpenForm").property("visible")
    ui.js('assistantDialog.act("files", "", "Show my CV and letter in Finder")')
    _wait_idle(ui, dialog)
    ui.js('assistantDialog.act("learn", "", "Remember what I changed on the page")')
    _wait_idle(ui, dialog)
    assert did == [("files", "lead"), ("learn",)] and dialog.property("remembered") == 1
    assert ui.js("talk.get(talk.count - 1).text") == "I will use your wording for 1 question."
    # the form moves on to its next page
    assert dialog.property("watchSign") == "pending" and not dialog.property("nextPage")
    ui.bridge.formWatched.emit({"ok": True, "url": "https://careers.acme.example/apply", "sign": "4:a|b", "empty": 2})
    ui.bridge.formWatched.emit({"ok": True, "url": "https://careers.acme.example/apply", "sign": "4:a|b", "empty": 1})
    ui.wait(100)
    assert not dialog.property("nextPage"), "the same fields: the same page"
    ui.bridge.formWatched.emit({"ok": True, "url": "https://careers.acme.example/apply", "sign": "3:c|d", "empty": 0})
    assert not dialog.property("nextPage"), "other fields, but nothing empty to fill"
    ui.bridge.formWatched.emit({"ok": False, "url": "", "sign": "", "empty": 0})
    ui.bridge.formWatched.emit({"ok": True, "url": "https://careers.acme.example/apply", "sign": "3:c|d", "empty": 3})
    ui.wait(100)
    assert dialog.property("nextPage") and ui.find("asNextPage").property("visible")
    assert "other fields now" in ui.js("talk.get(talk.count - 1).text")
    # mark applied, from here
    assert store.find(store.load_jobs(), "lead")["status"] in store.LEAD_STATUSES
    ui.js("assistantDialog.markApplied()")
    ui.wait(200)
    assert store.find(store.load_jobs(), "lead")["status"] == "waiting" and not dialog.property("canApply")
    assert "Moved Junior Python Developer at Acme to Applications" in ui.js("talk.get(talk.count - 1).text")
    # the memory can be emptied, after asking once
    ui.js('assistantDialog.page = "details"')
    ui.wait(100)
    assert "1 answer of your own is remembered" in ui.find("asMemory").property("text")
    ui.js("assistantDialog.forget()")
    assert dialog.property("confirmForget") and "once more" in ui.find("asMemory").property("text") and ("forget",) not in did
    ui.js("assistantDialog.forget()")
    _wait_idle(ui, dialog)
    assert did[-1] == ("forget",) and ui.js("talk.get(talk.count - 1).text").startswith("Forgotten")
    # a form in a frame: the button opens it; a failure makes the robot shake its head
    opened = []
    monkeypatch.setattr(ui_api.Api, "assistant_ask", lambda self, text, job_id="": {
        "ok": False, "action": "fill", "msg": "The form on this page sits inside a frame.", "frame": "https://boards.example.net/embed?id=7"})
    monkeypatch.setattr(ui_api.Api, "assistant_open_form", lambda self, url: opened.append(url) or {"ok": True, "say": "I opened the form in a new tab."})
    ui.js('assistantDialog.ask("fill the form")')
    moved = []
    for _ in range(60):
        moved.append(abs(ui.find("assistantBot").property("shakeBy")))
        ui.wait(20)
    assert max(moved) > 3, "it shook its head"
    assert dialog.property("canOpenForm") and not dialog.property("canLearn") and ui.find("asOpenForm").property("visible")
    ui.js('assistantDialog.act("open_form", assistantDialog.last.frame, "Open the form in its own tab")')
    _wait_idle(ui, dialog)
    assert opened == ["https://boards.example.net/embed?id=7"]
    ui.js("assistantDialog.tryClose()")
    ui.wait(400)
    store.save_jobs([dict(j, status="letter_ready") if j["id"] == "lead" else j for j in store.load_jobs()])      # as the other tests expect it
    ui.js("win.reloadAll()")
    assert dialog.property("visible") is False and ui.warnings == []


def test_a_press_on_empty_space_ends_the_editing_of_a_field(ui):
    """The accent outline and the cursor of a text field used to stay when you clicked somewhere else in the window."""
    from PySide6.QtCore import QPoint, QPointF, Qt
    dialog = ui.find("assistantDialog")
    ui.js("assistantDialog.show()")
    ui.wait(500)
    ui.js('assistantDialog.page = "ask"')
    box = ui.find("asBox")
    ui.js("forceActiveFocus()", box)
    ui.wait(100)
    assert box.property("activeFocus") is True
    spot = box.mapToScene(QPointF(box.width() / 2, -170)).toPoint()             # in the conversation above the field: empty space
    QTest.mouseClick(ui.win, Qt.LeftButton, Qt.NoModifier, QPoint(spot.x(), spot.y()))
    ui.wait(150)
    assert box.property("activeFocus") is False, "the field let go"
    QTest.mouseClick(ui.win, Qt.LeftButton, Qt.NoModifier, box.mapToScene(QPointF(box.width() / 2, box.height() / 2)).toPoint())
    ui.wait(150)
    assert box.property("activeFocus") is True, "and a click in the field still starts editing"
    header = box.mapToScene(QPointF(box.width() / 2, 0)).toPoint()
    QTest.mouseClick(ui.win, Qt.LeftButton, Qt.NoModifier, QPoint(header.x(), int(dialog.property("y")) + 30))   # beside the title
    ui.wait(150)
    assert box.property("activeFocus") is False
    ui.js("assistantDialog.tryClose()")
    ui.wait(400)
    assert ui.warnings == []


def test_every_window_with_a_text_field_lets_go_of_it():
    """A window that has a field to type in must have a way to end the editing by pressing beside it."""
    fields = re.compile(r"\b(TextField|TextEdit|TextArea|AutoField|ThemedSpin|TimeField|TagEditor) \{")
    for path in sorted((ROOT / "ui_qt").glob("*Dialog.qml")) + [ROOT / "ui_qt" / "WelcomeOverlay.qml"]:
        text = path.read_text(encoding="utf-8")
        if fields.search(text):
            assert "FocusCatcher {}" in text or "Sink.forceActiveFocus()" in text, f"{path.name} keeps a field's outline when you press beside it"
        for scroll in re.findall(r"\n( +)Flickable \{\n(?:.*\n){0,6}?\1    (\w+) \{", text):
            if "FocusCatcher {}" in text:
                assert scroll[1] in ("FocusCatcher", "MouseArea"), f"{path.name}: a scroll area without a FocusCatcher as its first child"
