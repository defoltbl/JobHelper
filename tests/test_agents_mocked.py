"""Agent 1 & 3 logic tests with the LLM mocked — no subscription usage, no network."""
from pathlib import Path

import pytest

from jobhelper import agent1_cv, agent3_letters, config, llm, store

FIXTURE_DIR = Path(__file__).resolve().parent / "fixtures"


# ── Agent 1 ─────────────────────────────────────────────────────────────────
def test_read_word_cvs():
    """The CVs that come with the tests (never the ones in data/cv: a test does not read your own files)."""
    files = sorted(FIXTURE_DIR.glob("*cv*.docx")) + sorted(FIXTURE_DIR.glob("*CV*.docx"))
    assert files
    for p in files:
        assert len(agent1_cv.read_cv_text(p)) > 200


def test_read_txt_cv(isolate_data):
    cv = isolate_data / "cv" / "cv.txt"
    cv.parent.mkdir(parents=True, exist_ok=True)
    cv.write_text("Alex Novak\nJunior Software Developer\nPython, SQL, Git")
    assert "Junior Software Developer" in agent1_cv.read_cv_text(cv)


def test_analyze_saves_profile(isolate_data, monkeypatch):
    cv = isolate_data / "cv" / "cv.txt"
    cv.parent.mkdir(parents=True, exist_ok=True)
    cv.write_text("Alex Novak, Junior Software Developer, Python, SQL, Git, REST APIs.")

    fake_profile = {
        "summary": "Junior developer",
        "seniority": "junior",
        "years_experience": 1,
        "target_roles": ["Junior Python Developer"],
        "key_skills": ["Python", "SQL"],
        "search_queries": ["python junior", "backend developer"],
        "strengths": ["fast learner"],
        "gaps": ["limited production experience"],
    }
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: dict(fake_profile))

    profile = agent1_cv.analyze()
    assert profile["seniority"] == "junior"
    assert profile["_cv_file"] == "cv.txt"
    assert config.PROFILE_JSON.exists()
    assert agent1_cv.load_profile()["search_queries"] == ["python junior", "backend developer"]


def test_analyze_errors_without_cv(isolate_data):
    (isolate_data / "cv").mkdir(parents=True, exist_ok=True)
    with pytest.raises(FileNotFoundError):
        agent1_cv.analyze()


# ── Agent 3 ─────────────────────────────────────────────────────────────────
def test_process_writes_letter_for_strong_match(isolate_data, monkeypatch):
    jobs = [store.new_job(id="jobscz-1", title="Python Developer", company="Acme",
                          location="Praha", description="We need a Python dev.")]
    profile = {"summary": "Junior dev", "seniority": "junior", "years_experience": 1,
               "key_skills": ["Python"], "strengths": ["quick learner"]}

    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"match_score": 85, "match_reason": "Strong Python fit"})
    monkeypatch.setattr(llm, "ask", lambda *a, **k: "Dear team,\n\nI am excited to apply.\n\nRegards, Alex")

    stats = agent3_letters.process(jobs, profile, threshold=60)
    assert stats["letters"] == 1
    job = jobs[0]
    assert job["match_score"] == 85
    assert job["status"] == "letter_ready"
    assert job["cover_letter_file"].endswith(".md")
    assert (config.LETTERS_DIR / job["cover_letter_file"]).exists()


LETTER_BODY = [("I am writing to apply for the Python Developer role at Acme, because the work matches what I do every day "
                "and I enjoy building small, well tested services in Python and SQL. ") * 2] * 5          # about 340 words in all


def test_process_writes_tailored_cv(isolate_data, monkeypatch):
    jobs = [store.new_job(id="jobscz-9", title="Python Developer", company="Acme",
                          location="Praha", description="Python, SQL, REST")]
    # the score, then the letter: with no template of your own, one is made from the CV's head (hands-on F15)
    monkeypatch.setattr(llm, "ask_json", lambda prompt, *a, **k: {"title_line": "Junior Developer", "paragraphs": LETTER_BODY, "role_short": "Dev"}
                        if "cover letter" in prompt else {"match_score": 85, "match_reason": "fit"})

    def fake_ask(prompt, *a, **k):
        return "# Alex Novak\n\nPython Developer\n- Python, SQL" if "Master CV" in prompt else "Dear team, letter."

    monkeypatch.setattr(llm, "ask", fake_ask)
    stats = agent3_letters.process(jobs, {}, threshold=60,
                                   cv_text="Alex Novak\nJunior Developer\nalex@example.com · +420 600 000 000\nSkills: Python, SQL")
    assert stats["letters"] == 1
    assert stats["tailored_cvs"] == 1
    job = jobs[0]
    assert job["tailored_cv_file"].startswith("CV-")
    assert (config.TAILORED_CV_DIR / job["tailored_cv_file"]).exists()
    assert config.LETTER_TEMPLATE.exists() and job["cover_letter_file"].endswith(".docx"), "a template made from the CV's head"
    from jobhelper import ui_api
    assert ui_api.Api().editable_letter(job) is True, "so Edit letter works for a new person"


def test_process_no_tailored_cv_without_cv_text(isolate_data, monkeypatch):
    jobs = [store.new_job(id="jobscz-10", title="Dev", company="Acme", description="x")]
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"match_score": 85, "match_reason": "fit"})
    monkeypatch.setattr(llm, "ask", lambda *a, **k: "Dear team, letter.")
    stats = agent3_letters.process(jobs, {}, threshold=60)  # no cv_text
    assert stats["tailored_cvs"] == 0
    assert jobs[0]["tailored_cv_file"] == ""


def test_process_skips_weak_match(isolate_data, monkeypatch):
    jobs = [store.new_job(id="jobscz-2", title="ML Researcher", company="Lab",
                          description="PhD required.")]
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"match_score": 30, "match_reason": "Needs PhD"})
    # Letter step must NOT be called for a weak match.
    monkeypatch.setattr(llm, "ask", lambda *a, **k: (_ for _ in ()).throw(AssertionError("letter written for weak match")))
    stats = agent3_letters.process(jobs, {}, threshold=60)
    assert stats["letters"] == 0
    assert jobs[0]["match_score"] == 30
    assert jobs[0]["status"] == "found"  # unchanged
    assert jobs[0]["cover_letter_file"] == ""


def test_process_handles_missing_description(isolate_data, monkeypatch):
    jobs = [store.new_job(id="jobscz-3", title="Backend Engineer", company="Startup", description="")]
    captured = {}

    def fake_score(prompt, *a, **k):
        captured["prompt"] = prompt
        return {"match_score": 70, "match_reason": "ok"}

    monkeypatch.setattr(llm, "ask_json", fake_score)
    monkeypatch.setattr(llm, "ask", lambda *a, **k: "Hello letter")
    agent3_letters.process(jobs, {}, threshold=60)
    assert "No description was available" in captured["prompt"]


def test_saved_profile_is_dropped_when_the_cv_changes(isolate_data, monkeypatch):
    config.ensure_dirs()
    cv = config.CV_DIR / "cv.txt"
    cv.write_text("Alex Novak. Junior developer. Internship Feb - Mar 2024. " * 3, encoding="utf-8")
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"seniority": "junior", "years_experience": 0.2})
    agent1_cv.analyze()
    assert agent1_cv.load_profile(current_only=True)["years_experience"] == 0.2
    cv.write_text("Alex Novak. Junior developer. Internship Sep 2023 - Mar 2024. " * 3, encoding="utf-8")
    assert agent1_cv.load_profile(current_only=True) is None          # the pipeline will analyse the new CV
    assert agent1_cv.load_profile()["years_experience"] == 0.2         # the old analysis is still readable


def test_cv_analysis_asks_for_internships_in_the_experience_total():
    assert "INCLUDING internships" in agent1_cv.PROMPT_TEMPLATE and '"internships"' in agent1_cv.PROMPT_TEMPLATE


def test_letters_below_fifty_are_written_when_the_threshold_allows_and_marked_stretch(isolate_data, monkeypatch):
    jobs = [store.new_job(id="jobscz-1", title="IT Support Specialist", company="Acme", location="Praha",
                          description="We need IT support.")]
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"match_score": 35, "match_reason": "Different function."})
    monkeypatch.setattr(llm, "ask", lambda *a, **k: "Dear team,\n\nI would like to apply.\n\nRegards, Alex")
    assert agent3_letters.process([dict(jobs[0])], {"summary": "x"}, threshold=50)["letters"] == 0
    stats = agent3_letters.process(jobs, {"summary": "x"}, threshold=30)
    assert stats["letters"] == 1
    assert jobs[0]["status"] == "letter_ready" and jobs[0]["match_tier"] == "stretch"
    assert "Stretch match (35%)" in jobs[0]["notes"]


def test_tiers():
    assert [agent3_letters._tier(s) for s in (95, 60, 59, 50, 49, 30)] == \
        ["strong", "strong", "moderate", "moderate", "stretch", "stretch"]


# ── no company name: no letter addressed to nobody ──────────────────────────
def test_no_documents_are_written_for_a_posting_without_a_company(isolate_data, monkeypatch, capsys):
    jobs = [store.new_job(id="p1", title="Junior AI Engineer", company="", location="Praha",
                          description="We are looking for a junior engineer who builds AI features in Python."),
            store.new_job(id="p2", title="UX Designer", company=" ", location="Praha",
                          description="Detail nabídky UX Designer Firma: Thales DIS s.r.o. ( zaměstnavatel ) Místo pracoviště: Praha")]
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"scores": [{"id": "J1", "match_score": 75, "match_reason": "from the title"},
                                                                   {"id": "J2", "match_score": 70, "match_reason": "fits"}]})
    wrote = []
    monkeypatch.setattr(llm, "ask", lambda *a, **k: wrote.append(1) or "Dear team,\n\nI am excited to apply.\n\nRegards, Alex")
    stats = agent3_letters.process(jobs, {"seniority": "junior"}, threshold=60)
    nameless, named = jobs
    assert nameless["cover_letter_file"] == "" and nameless["status"] == "found" and nameless["match_score"] == 75
    assert nameless["hold"] == agent3_letters.NO_COMPANY and nameless.get("doc_attempts", 0) == 0
    assert named["company"] == "Thales DIS s.r.o" and named["status"] == "letter_ready" and "hold" not in named   # name found in the text
    assert stats["letters"] == 1 and len(wrote) == 1
    assert "no company name in the posting" in capsys.readouterr().out


def test_the_hold_is_lifted_once_the_company_is_known(isolate_data, monkeypatch):
    job = store.new_job(id="p1", title="Junior AI Engineer", company="", match_score=75, description="We build AI features in Python.")
    monkeypatch.setattr(llm, "ask", lambda *a, **k: "Dear team,\n\nI am excited to apply.\n\nRegards, Alex")
    agent3_letters.process([job], {}, threshold=60)
    assert job["hold"] and job["cover_letter_file"] == ""
    job["company"] = "Acme"                                           # you filled it in, or a later search found it
    agent3_letters.process([job], {}, threshold=60)
    assert "hold" not in job and job["status"] == "letter_ready"


# ── the app's own checks on a score ─────────────────────────────────────────
CV_UA = "Alex Novak\nNative Polish speaker, B2 English.\nLANGUAGES\nPolish – Native · English – B2 · Czech – Beginner\nPython, SQL."


def test_a_required_language_the_cv_lacks_caps_the_score(isolate_data, monkeypatch):
    jobs = [store.new_job(id="q", title="Junior Quality Engineer", company="Acme",
                          description="Excellent communication skills in Czech or Slovak and English (B2+). Testing experience welcome."),
            store.new_job(id="ok", title="Python Developer", company="Beta", description="Fluent English is required. We use Python.")]
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"scores": [{"id": "J1", "match_score": 55, "match_reason": "Meets the experience bar."},
                                                                   {"id": "J2", "match_score": 70, "match_reason": "Good fit."}]})
    monkeypatch.setattr(llm, "ask", lambda *a, **k: "Dear team,\n\nI am excited to apply.\n\nRegards, Alex")
    stats = agent3_letters.process(jobs, {}, threshold=30, cv_text=CV_UA)
    capped, fine = jobs
    assert capped["match_score"] == agent3_letters.LANGUAGE_CAP and capped["cover_letter_file"] == "" and capped["status"] == "found"
    assert capped["score_capped"] == "Requires Czech or Slovak, which your CV does not show."
    assert capped["match_reason"].startswith("Requires Czech or Slovak, which your CV does not show. The app lowered the score from 55 to 25.")
    assert fine["match_score"] == 70 and "score_capped" not in fine and stats["capped"] == 1


def test_scoring_is_told_what_you_are_looking_for(isolate_data, monkeypatch):
    from jobhelper import ui_api
    config.ensure_dirs()
    prompts = []
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or {"match_score": 10, "match_reason": "r"})
    job = lambda: [store.new_job(id="e", title="Associate Economist", company="Acme", description="Economics.")]   # noqa: E731
    agent3_letters.process(job(), {"target_roles": ["backend developer", "ML engineer"]}, threshold=60)
    assert "looking for these kinds of role: backend developer, ML engineer" in prompts[-1] and "scores at most 40" in prompts[-1]
    ui_api.Api().save_options({"wanted_roles": ["data engineer"], "unwanted_roles": ["economist", "sales"]})
    agent3_letters.process(job(), {"target_roles": ["backend developer"]}, threshold=60)
    assert "kinds of role: data engineer." in prompts[-1] and "backend developer" not in prompts[-1]            # your list replaces the CV's
    assert "does NOT want these kinds of role: economist, sales" in prompts[-1] and "at most 20" in prompts[-1]
    ui_api.Api().save_options({"wanted_roles": [], "unwanted_roles": []})
    agent3_letters.process(job(), {"target_roles": ["backend developer"]}, threshold=60)
    assert "looking for these kinds of role" not in prompts[-1]                                               # emptied on purpose: eligibility only
    assert agent3_letters.direction_rules(None, None) == "" and agent3_letters.direction_rules(["  "], []) == ""


# ── score again ─────────────────────────────────────────────────────────────
def _scored(id, score, **kw):
    return store.new_job(id=id, title=kw.pop("title", f"Role {id}"), company=f"Co {id}", match_score=score, match_reason="old reason",
                         description=kw.pop("description", "We use Python."), **kw)


def test_rescore_changes_only_the_score_and_remembers_the_old_one(isolate_data, monkeypatch):
    jobs = [_scored("e", 80, title="Associate Economist", status="letter_ready", cover_letter_file="L_e.docx", tailored_cv_file="CV_e.docx", match_tier="strong"),
            _scored("d", 68, title="Python Developer"), _scored("same", 50), _scored("skip", 40), store.new_job(id="none", title="T", company="C")]
    prompts = []
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or {"scores": [
        {"id": "J1", "match_score": 30, "match_reason": "Another field."}, {"id": "J2", "match_score": 75, "match_reason": "Good fit."},
        {"id": "J3", "match_score": 50, "match_reason": "Same as before."}]})
    monkeypatch.setattr(llm, "ask", lambda *a, **k: pytest.fail("no letter or CV may be written here"))
    seen = []
    stats = agent3_letters.rescore(jobs, ["e", "d", "same", "missing-id"], {"target_roles": ["Python developer"]}, CV_UA,
                                   progress=lambda d, t: seen.append((d, t)))
    e, d, same, skip, none = jobs
    assert (e["match_score"], e["score_before"], e["match_reason"]) == (30, 80, "Another field.")
    assert e["cover_letter_file"] == "L_e.docx" and e["status"] == "letter_ready" and e["match_tier"] == "stretch"   # documents stay, the tier follows
    assert (d["match_score"], d["score_before"]) == (75, 68) and "score_before" not in same and same["match_reason"] == "Same as before."
    assert skip["match_score"] == 40 and skip["match_reason"] == "old reason" and none["match_score"] is None      # not asked for: untouched
    assert stats["scored"] == 3 and stats["failed"] == 0
    assert [(c["id"], c["before"], c["after"]) for c in stats["changed"]] == [("e", 80, 30), ("d", 68, 75)]
    assert len(prompts) == 1 and "looking for these kinds of role: Python developer" in prompts[0] and seen[-1] == (3, 3)


def test_rescore_asks_alone_for_what_a_batch_left_out_and_applies_the_language_check(isolate_data, monkeypatch):
    jobs = [_scored("a", 60), _scored("b", 60, description="Excellent communication skills in Czech or Slovak and English.")]
    replies = [{"scores": [{"id": "J1", "match_score": 62, "match_reason": "ok"}, {"id": "J9", "match_score": 99, "match_reason": "not ours"}]},
               {"match_score": 58, "match_reason": "Meets the bar."}]
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: replies.pop(0))
    stats = agent3_letters.rescore(jobs, ["a", "b"], {}, CV_UA)
    assert jobs[0]["match_score"] == 62 and jobs[1]["match_score"] == agent3_letters.LANGUAGE_CAP and jobs[1]["score_before"] == 60
    assert jobs[1]["match_reason"].startswith("Requires Czech or Slovak") and stats["scored"] == 2 and replies == []


def test_rescore_counts_what_could_not_be_scored_and_stops_on_the_limit(isolate_data, monkeypatch):
    jobs = [_scored("a", 60)]
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"match_score": "high", "match_reason": "?"})
    stats = agent3_letters.rescore(jobs, ["a"], {}, "")
    assert stats == {"scored": 0, "changed": [], "failed": 1} and jobs[0]["match_score"] == 60 and jobs[0]["match_reason"] == "old reason"

    def limit(*a, **k):
        raise llm.LimitReached("Claude's 5-hour limit is used up (99%).")
    monkeypatch.setattr(llm, "ask_json", limit)
    with pytest.raises(llm.LimitReached):
        agent3_letters.rescore(jobs, ["a"], {}, "")
    waited = []

    def on_limit(error):
        waited.append(str(error))
        monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"match_score": 44, "match_reason": "after the pause"})
    agent3_letters.rescore(jobs, ["a"], {}, "", on_limit=on_limit)
    assert len(waited) == 1 and jobs[0]["match_score"] == 44                    # paused, then the same job was asked again


# ── what a score rests on ───────────────────────────────────────────────────
def test_text_basis():
    assert agent3_letters.text_basis({"description": ""}) == "title" and agent3_letters.text_basis({}) == "title"
    assert agent3_letters.text_basis({"description": "   \n "}) == "title"
    assert agent3_letters.text_basis({"description": "A short summary of the ad."}) == "summary"
    assert agent3_letters.text_basis({"description": "x" * agent3_letters.SHORT_TEXT}) == ""


def test_a_title_alone_cannot_score_high_and_gets_no_documents(isolate_data, monkeypatch, capsys):
    jobs = [store.new_job(id="t", title="Junior AI Engineer", company="Acme", description=""),
            store.new_job(id="s", title="Python Developer", company="Beta", description="A short summary: Python, tests.")]
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"scores": [{"id": "J1", "match_score": 75, "match_reason": "The title fits."},
                                                                   {"id": "J2", "match_score": 70, "match_reason": "Fits."}]})
    wrote = []
    monkeypatch.setattr(llm, "ask", lambda *a, **k: wrote.append(1) or "Dear team,\n\nI am excited to apply.\n\nRegards, Alex")
    stats = agent3_letters.process(jobs, {}, threshold=30)
    blind, short = jobs
    assert blind["match_score"] == agent3_letters.NO_TEXT_CAP and blind["score_capped"].startswith("Scored from the title alone")
    assert "lowered the score from 75 to 40" in blind["match_reason"]
    assert blind["hold"] == agent3_letters.NO_TEXT and blind["cover_letter_file"] == "" and blind["status"] == "found"
    assert short["match_score"] == 70 and short["status"] == "letter_ready"            # a short summary is marked, not capped
    assert stats["letters"] == 1 and len(wrote) == 1 and "no posting text" in capsys.readouterr().out


def test_a_low_title_only_score_is_left_as_it_is():
    job = store.new_job(id="t", title="Economist", company="Acme")
    assert agent3_letters.apply_score(job, {"match_score": 20, "match_reason": "Another field."}) is True
    assert job["match_score"] == 20 and "score_capped" not in job and job["match_reason"] == "Another field."


# ── found by reading agent1_cv line by line (10 October 2026) ───────────────
def test_a_cv_that_keeps_its_jobs_in_a_table_is_read_whole(tmp_path, isolate_data):
    """Only the paragraphs of a Word CV were read. Many CVs keep their jobs or skills in a table: those lines were missing
    from everything the app knows about the person, so true experience looked like a claim the CV does not back."""
    import docx as word
    from jobhelper import agent1_cv, config, vault
    cv = word.Document()
    cv.add_paragraph("Alex Novak")
    cv.add_paragraph("Junior Software Developer")
    table = cv.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "2023 - 2025", "Backend developer at Northwind: Python, PostgreSQL"
    table.cell(1, 0).merge(table.cell(1, 1)).text = "Skills: Docker, Kafka"
    cv.add_paragraph("Education")
    path = tmp_path / "cv.docx"
    cv.save(str(path))
    lines = agent1_cv.read_cv_text(path).splitlines()
    assert lines == ["Alex Novak", "Junior Software Developer", "2023 - 2025 | Backend developer at Northwind: Python, PostgreSQL", "Skills: Docker, Kafka",
                     "Education"], "in the order of the page, a merged cell once"
    plain = word.Document()
    for text in ("Alex Novak", "", "Profile"):
        plain.add_paragraph(text)
    plain.save(str(path))
    assert agent1_cv.read_cv_text(path) == "Alex Novak\n\nProfile", "a CV without tables reads exactly as before"
    config.ensure_dirs()
    for damaged in ("not json at all", "[1, 2]", '"text"'):
        vault.write_text(config.PROFILE_JSON, damaged)
        assert agent1_cv.load_profile() is None and agent1_cv.load_profile(current_only=True) is None, "a damaged analysis is made again, not crashed on"



def test_the_letter_template_is_kept_yours_and_follows_the_cv_when_the_app_made_it(isolate_data, monkeypatch):
    """Review of the template from the CV's head: a template of your own is never touched, one that went missing is not
    replaced, and one the app made follows a change of your contacts."""
    from jobhelper import cover_letter, integrity
    cv = "Alex Novak\nJunior Developer\nalex@example.com · +420 600 000 000\nPROFILE\nText."
    state = {"now": "absent"}
    monkeypatch.setattr(integrity, "check", lambda: [{"name": "letter_template", "state": state["now"]}])
    agent3_letters._keep_template(cv)
    assert cover_letter.app_made_head(config.LETTER_TEMPLATE)[:3] == ["Alex Novak", "Junior Developer", "alex@example.com"]
    state["now"] = "ok"                                   # still as the app made it
    agent3_letters._keep_template(cv.replace("alex@example.com", "novak@example.com"))
    assert cover_letter.app_made_head(config.LETTER_TEMPLATE)[2] == "novak@example.com", "your new contacts"
    from jobhelper import vault
    assert list((vault.backups_dir() / "letter_versions").glob("*")), "the template as it was is kept before it is rebuilt"
    state["now"] = "changed"                              # you edited it in Word: it is yours now
    edited = config.LETTER_TEMPLATE.read_bytes()
    agent3_letters._keep_template(cv)
    assert config.LETTER_TEMPLATE.read_bytes() == edited, "an app-made template you edited is never rebuilt"
    state["now"] = "ok"
    import docx
    own = docx.Document()
    for text in ("Mine", "x", "y", "z", "", "d", "Dear", "1", "2", "3", "4", "5", "Bye", "Mine"):
        own.add_paragraph(text)
    own.save(str(config.LETTER_TEMPLATE))
    before = config.LETTER_TEMPLATE.read_bytes()
    agent3_letters._keep_template(cv)
    assert config.LETTER_TEMPLATE.read_bytes() == before, "a template of your own is never touched"
    config.LETTER_TEMPLATE.unlink()
    state["now"] = "missing"
    agent3_letters._keep_template(cv)
    assert not config.LETTER_TEMPLATE.exists(), "a missing template of yours is not replaced: the start-up check says it is gone"
    monkeypatch.setattr(cover_letter, "make_template", lambda *a, **k: (_ for _ in ()).throw(ValueError("bad text")))
    state["now"] = "absent"
    agent3_letters._keep_template(cv)                     # a template that cannot be made: the plain letter, no crash
    assert not config.LETTER_TEMPLATE.exists()
