"""Offline tests for the cover-letter writer (no LLM, no network)."""
from datetime import date
from pathlib import Path

import docx
import pytest

from jobhelper import config, cover_letter, llm, store

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "tests" / "fixtures" / "letter_example_deutscheborse.docx"      # never your own template in data/
# A frozen copy of the master CV the approved examples were made from, so these
# tests do not change when the real CV in data/cv/ is edited.
MASTER = Path(__file__).parent / "fixtures" / "Candidate_CV_EN.docx"
EXAMPLE = Path(__file__).parent / "fixtures" / "letter_example_deutscheborse.docx"

pytestmark = pytest.mark.skipif(not (TEMPLATE.exists() and MASTER.exists()), reason="template/master not present")


def _example():
    paras = [p.text for p in docx.Document(str(EXAMPLE)).paragraphs]
    return paras[1], paras[7:12], paras


def _master_text():
    return "\n".join(p.text for p in docx.Document(str(MASTER)).paragraphs)


def test_rebuilds_the_approved_example_exactly(tmp_path):
    title, body, example_paras = _example()
    out = tmp_path / "letter.docx"
    cover_letter.build_docx(TEMPLATE, out, title, body, when=date(2026, 10, 6))
    assert [p.text for p in docx.Document(str(out)).paragraphs] == example_paras
    assert cover_letter.verify(TEMPLATE, out) == []


def test_approved_example_passes_the_content_checks(tmp_path, monkeypatch):
    rules = tmp_path / "cover_letter_instruction.md"            # its own rules, never your real ones
    rules.write_text("Do not claim: Anything else (Copilot, Docker, and so on).\n", encoding="utf-8")
    monkeypatch.setattr(config, "LETTER_INSTRUCTION", rules)
    _, body, _ = _example()
    unverified = cover_letter.unverified_terms(_master_text())
    assert "Kubernetes" in unverified and "Copilot" in unverified and "Docker" not in unverified
    assert cover_letter.content_problems(body, unverified) == []


def test_gap_admissions_are_rejected_and_removed():
    _, body, _ = _example()
    bad = list(body)
    bad[3] += " Although I have not worked with Kubernetes yet, I am a fast learner."
    problems = cover_letter.content_problems(bad, ["Kubernetes"])
    assert any("admits a gap" in p for p in problems)
    assert "Although" not in cover_letter.enforce(bad, ["Kubernetes"])[3]


def test_unbacked_experience_claim_is_rejected_but_a_wish_is_fine():
    _, body, _ = _example()
    claim = list(body)
    claim[1] += " I have worked with Kubernetes clusters in production."
    assert any("as experience" in p for p in cover_letter.content_problems(claim, ["Kubernetes"]))
    assert "Kubernetes clusters in production" not in cover_letter.enforce(claim, ["Kubernetes"])[1]
    wish = list(body)
    wish[3] += " I would like to work with Terraform next."
    assert cover_letter.content_problems(wish, ["Terraform", "Kubernetes"]) == []


def test_date_format():
    assert cover_letter.letter_date(date(2026, 10, 6)) == "6 October 2026"


def test_write_end_to_end_with_mocked_model(isolate_data, monkeypatch):
    title, body, example_paras = _example()
    monkeypatch.setattr(config, "LETTER_TEMPLATE", TEMPLATE)
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {
        "title_line": title, "paragraphs": body, "role_short": "Middleware Admin"})
    job = store.new_job(id="jobscz-1", title="Junior Middleware Administrator (f/m/d)",
                        company="Deutsche Börse Group", description="Linux, Kubernetes, Kafka")
    name = cover_letter.write(job, _master_text(), MASTER)
    assert name == "Candidate_CoverLetter_DeutscheBorse_MiddlewareAdmin.docx"
    got = [p.text for p in docx.Document(str(config.LETTERS_DIR / name)).paragraphs]
    assert got[1] == "Junior Software Developer" != title, "under the name stands what the CV says you are, not the role applied for"
    assert got[7:12] == body
    assert got[5] == cover_letter.letter_date()            # today's date
    assert got[0] == example_paras[0] and got[2:5] == example_paras[2:5] and got[12:] == example_paras[12:]   # head + closing untouched
    job2 = store.new_job(id="jobscz-2", title="Junior Middleware Administrator (f/m/d)", company="Northwind", description="Linux")
    bare = "\n".join(line for i, line in enumerate(_master_text().splitlines()) if i != 1)      # a CV with no line under the name
    assert cover_letter.own_title(bare) == ""
    name2 = cover_letter.write(job2, bare, MASTER)
    assert docx.Document(str(config.LETTERS_DIR / name2)).paragraphs[1].text == title, "then the role, as before"


@pytest.mark.parametrize("second, title", [
    ("Junior Software Developer", "Junior Software Developer"), ("  Data   Analyst  ", "Data Analyst"), ("QA Engineer | Python, SQL", "QA Engineer | Python, SQL"),
    ("alex@example.com · +420 123 456 789", ""), ("Brno, Czechia · linkedin.com/in/alexnovak", ""), ("+420 123 456 789", ""), ("PROFILE", ""),
    ("SUMMARY", ""), ("I am a developer who likes clean code and good tests.", ""), ("https://alexnovak.example", ""), ("2019 - 2024", ""),
    ("A very long line that goes on about many things and is clearly not a title of a person at all", "")])
def test_your_own_title_is_the_line_under_your_name(second, title):
    assert cover_letter.own_title(f"Alex Novak\n\n{second}\nPROFILE\nBuilds services.") == title
    assert cover_letter.own_title("Alex Novak") == "" and cover_letter.own_title("") == "" and cover_letter.own_title(None) == ""


def test_the_title_of_a_written_letter_can_be_corrected_and_nothing_else_changes(tmp_path):
    title, body, example_paras = _example()
    out = tmp_path / "letter.docx"
    cover_letter.build_docx(TEMPLATE, out, "Group Full-Stack Developer – Payments", body, when=date(2026, 10, 6))
    assert cover_letter.set_title(out, "Junior Software Developer") is True
    got = [p.text for p in docx.Document(str(out)).paragraphs]
    assert got[1] == "Junior Software Developer" and got[:1] + got[2:] == example_paras[:1] + example_paras[2:]
    assert cover_letter.verify(TEMPLATE, out, []) == [] and list(tmp_path.iterdir()) == [out], "a sound letter, and no draft left behind"
    assert cover_letter.set_title(out, "Junior Software Developer") is False and cover_letter.set_title(out, "  ") is False
    short = tmp_path / "short.docx"
    docx.Document().save(str(short))
    with pytest.raises(ValueError):
        cover_letter.set_title(short, "Anything")           # not built from the letter template: refused, not rewritten


# ── first person only (the first real run produced a letter about "Alex ... he") ──
def test_third_person_letter_is_detected():
    about_him = ["The role matches the way Alex already works.", "He would like to bring the same drive to SAP."]
    as_him = ["The role matches the way I already work.", "My thesis was presented at two conferences."]
    assert cover_letter.third_person(about_him, "Alex") is True
    assert cover_letter.third_person(["PostgreSQL work gives the clearest evidence.", "He can start soon."], "Alex") is True
    assert cover_letter.third_person(as_him, "Alex") is False
    assert "third person" in cover_letter.content_problems(about_him, [], "Alex")[0]
    assert cover_letter.candidate_first_name("Alex Novak\nJunior Software Developer") == "Alex"
    assert cover_letter.candidate_first_name("") == ""


def test_third_person_letter_is_retried_and_never_saved(isolate_data, monkeypatch):
    import shutil
    shutil.copy2(TEMPLATE, config.LETTER_TEMPLATE)
    calls = []
    para = ("The role matches the way Alex already works across backend systems and data pipelines. " * 6).strip()

    def ask_json(prompt, **kw):
        calls.append(prompt)
        return {"title_line": "Junior Developer", "paragraphs": [para] * 5, "role_short": "Dev"}

    monkeypatch.setattr(llm, "ask_json", ask_json)
    job = store.new_job(id="linkedin-9", title="Junior Developer", company="Acme", description="Python")
    with pytest.raises(ValueError, match="third person"):
        cover_letter.write(job, _master_text(), MASTER)
    assert len(calls) == 2 and "first person" in calls[1]          # told what was wrong on the retry
    assert list(config.LETTERS_DIR.glob("*.docx")) == []


# ── the posting's words stated as the candidate's own experience ────────────
ACCOUNTANT = ("Junior accountant with one year of practice in accounts payable. Processed about 300 supplier invoices a month in SAP, "
              "prepared weekly payment runs, reconciled supplier statements and helped with the month-end closing.")
POSTING = ("You will post invoices with the correct tax codes and cost centers, keep the vendor master data up to date and confirm "
           "intercompany balances with our entities. Analyst I level. We work in SAP.")


def test_a_letter_that_states_the_posting_s_terms_as_experience_is_noticed():
    claims = cover_letter.borrowed_claims
    said = ["In my current role I process a high volume of invoices in SAP, assigning the correct tax codes and cost centers, "
            "which matches the core requirement of posting invoices for your mother entity."]
    assert {"tax codes", "cost centers"} <= {c.lower() for c in claims(said, POSTING, ACCOUNTANT)}
    # speaking of the role, of the reader or of the future is what a letter is for
    for fine in ("The AP Accountant role caught my attention because it covers tax codes, cost centers and vendor master data.",
                 "Your team keeps the vendor master data current and confirms intercompany balances, which is work I respect.",
                 "I would like to grow into vendor master data and intercompany balances.",
                 "I perform the month-end closing, skills that should translate well to confirming intercompany balances.",
                 "I reconciled supplier statements and prepared payment runs in SAP.",
                 "The Analyst I opening stood out to me because it covers cost centers and tax codes."):
        assert claims([fine], POSTING, ACCOUNTANT) == [], fine
    assert claims([], POSTING, ACCOUNTANT) == [] and claims(said, "", ACCOUNTANT) == []
    # audit: a wish word later in the sentence freed the claim before it
    wished = ["I assign tax codes and cost centers to every invoice and would love to keep doing so."]
    assert {"tax codes", "cost centers"} <= {c.lower() for c in claims(wished, POSTING, ACCOUNTANT)}
    # what the person wrote in their own notes is theirs
    assert claims(said, POSTING, ACCOUNTANT + " I assigned tax codes and cost centers to every invoice.") == []


def test_a_letter_that_borrows_is_sent_back_and_what_is_left_is_noted(isolate_data, monkeypatch):
    title, body, _ = _example()
    monkeypatch.setattr(config, "LETTER_TEMPLATE", TEMPLATE)
    posting = "You will run quarterly zebra migrations and tune walrus pipelines for our platform."
    borrowed = list(body)
    borrowed[2] = body[2].rstrip() + " I ran zebra migrations every quarter."
    prompts, answers = [], [{"title_line": title, "paragraphs": borrowed, "role_short": "Admin"}, {"title_line": title, "paragraphs": body, "role_short": "Admin"}]
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or answers[min(len(prompts) - 1, 1)])
    job = store.new_job(id="j1", title="Administrator", company="Acme", description=posting)
    name = cover_letter.write(job, _master_text(), MASTER)
    assert len(prompts) == 2 and "The letter presents 'zebra migrations' as the candidate's own experience" in prompts[1]
    assert "letter_check" not in job
    assert "zebra" not in "\n".join(p.text for p in docx.Document(str(config.LETTERS_DIR / name)).paragraphs)
    # the model insists: nothing is cut out silently, the person is told what to look at
    prompts.clear()
    answers[1] = answers[0]
    job = store.new_job(id="j2", title="Administrator", company="Acme", description=posting)
    cover_letter.write(job, _master_text(), MASTER)
    assert job["letter_check"] == ["zebra migrations"]


def test_a_letter_figure_nobody_gave_is_sent_back_and_what_is_left_is_noted(isolate_data, monkeypatch):
    title, body, _ = _example()
    monkeypatch.setattr(config, "LETTER_TEMPLATE", TEMPLATE)
    invented = list(body)
    invented[2] = body[2].rstrip() + " My work saved the team 37 hours a week."
    prompts, answers = [], [{"title_line": title, "paragraphs": invented, "role_short": "Admin"},
                            {"title_line": title, "paragraphs": body, "role_short": "Admin"}]
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or answers[min(len(prompts) - 1, 1)])
    job = store.new_job(id="f1", title="Administrator", company="Acme", description="We run Linux servers.")
    cover_letter.write(job, _master_text(), MASTER)
    assert len(prompts) == 2 and "The letter states the figure 37" in prompts[1] and "letter_figures" not in job
    prompts.clear()
    answers[1] = answers[0]                                   # the model insists: the person is told which figure to check
    job = store.new_job(id="f2", title="Administrator", company="Acme", description="We run Linux servers.")
    cover_letter.write(job, _master_text(), MASTER)
    assert job["letter_figures"] == ["37"]


def test_a_letter_that_moves_work_to_another_employer_is_sent_back_and_what_is_left_is_noted(isolate_data, monkeypatch):
    """Seen in the hands-on check: the letter put the REST API work at the internship, not at the job where it was done."""
    title, body, _ = _example()
    monkeypatch.setattr(config, "LETTER_TEMPLATE", TEMPLATE)
    cv = _master_text() + ("\nEXPERIENCE\nJunior Developer – Northwind Parcel s.r.o., Brno (2024 – present): built an internal REST API in Django.\n"
                           "IT Intern – Fabrikam Labs, Brno (2023 – 2024): maintained Python scripts for data imports.\n")
    moved = list(body)
    moved[2] = body[2].rstrip() + " At Fabrikam Labs I built internal REST APIs."
    prompts, answers = [], [{"title_line": title, "paragraphs": moved, "role_short": "Admin"},
                            {"title_line": title, "paragraphs": body, "role_short": "Admin"}]
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or answers[min(len(prompts) - 1, 1)])
    job = store.new_job(id="m1", title="Administrator", company="Acme", description="We run Linux servers.")
    cover_letter.write(job, cv, MASTER)          # (the fixture CV has a "Northwind Labs" too: "Labs" names neither)
    assert len(prompts) == 2 and "The letter puts 'REST', 'APIs' at Fabrikam" in prompts[1] and "letter_misplaced" not in job
    prompts.clear()
    answers[1] = answers[0]                                   # the model insists: the person is told what to look at
    job = store.new_job(id="m2", title="Administrator", company="Acme", description="We run Linux servers.")
    cover_letter.write(job, cv, MASTER)
    assert job["letter_misplaced"] == ["'REST', 'APIs' at Fabrikam"] and "letter_check" not in job, "a note of its own"
    assert any("at Fabrikam" in w for w in cover_letter.letter_warnings(moved, cv)), "a hand-edited letter is warned too"
    assert not any("at Fabrikam" in w for w in cover_letter.letter_warnings(moved, cv, "Fabrikam")), "unless it is written to them"


def test_a_rewrite_that_moves_work_is_sent_back_once_then_given(isolate_data, monkeypatch):
    title, body, _ = _example()
    cv = _master_text() + "\nEXPERIENCE\nJunior Developer – Northwind Parcel s.r.o., Brno (2024 – present): built an internal REST API in Django.\n" \
                          "IT Intern – Fabrikam Labs, Brno (2023 – 2024): maintained Python scripts for data imports.\n"
    moved = body[2].rstrip() + " At Fabrikam I built internal REST APIs."
    prompts = []
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or {"paragraph": moved})
    job = store.new_job(id="r1", title="Administrator", company="Acme", description="We run Linux servers.")
    assert cover_letter.rewrite_paragraph(job, cv, list(body), 2) == moved, "the second answer is given: you see it, then a warning"
    assert len(prompts) == 2 and "It puts 'REST', 'APIs' at Fabrikam" in prompts[1]


def test_a_rewrite_with_a_real_problem_and_moved_work_is_refused(isolate_data, monkeypatch):
    title, body, _ = _example()
    cv = _master_text() + "\nEXPERIENCE\nIT Intern – Fabrikam Labs, Brno (2023 – 2024): maintained Python scripts.\n" \
                          "Junior Developer – Northwind Parcel s.r.o., Brno (2024 – present): built an internal REST API in Django.\n"
    bad = body[2].rstrip() + " At Fabrikam I built internal REST APIs and deployed Kubernetes daily."
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: {"paragraph": bad})
    monkeypatch.setattr(cover_letter, "unverified_terms", lambda cv_text: ["Kubernetes"])
    job = store.new_job(id="r2", title="Administrator", company="Acme", description="We run Linux servers.")
    with pytest.raises(ValueError, match="Kubernetes"):
        cover_letter.rewrite_paragraph(job, cv, list(body), 2)
