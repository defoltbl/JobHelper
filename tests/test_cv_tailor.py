"""Offline tests for the CV adaptation (no LLM, no network)."""
import zipfile
from pathlib import Path

import docx
import pytest

from jobhelper import word_document, config, cv_tailor, llm, store
from jobhelper.word_document import text_of

ROOT = Path(__file__).resolve().parent.parent
# A frozen copy of the master CV the approved examples were made from, so these
# tests do not change when the real CV in data/cv/ is edited.
MASTER = Path(__file__).parent / "fixtures" / "Candidate_CV_EN.docx"
EXAMPLE = Path(__file__).parent / "fixtures" / "cv_example_deutscheborse.docx"

pytestmark = pytest.mark.skipif(not MASTER.exists(), reason="master CV not present")


def _example_content():
    paras = [p.text for p in docx.Document(str(EXAMPLE)).paragraphs]
    profile = paras[4]
    rows = []
    for text in paras[6:13]:
        label, items = text.split(": ", 1)
        rows.append({"label": label, "items": items})
    return profile, rows, paras


def test_rebuilds_the_approved_example_exactly(tmp_path):
    profile, rows, example_paras = _example_content()
    out = tmp_path / "out.docx"
    cv_tailor.build_docx(MASTER, out, profile, rows)
    assert [p.text for p in docx.Document(str(out)).paragraphs] == example_paras
    assert cv_tailor.verify(MASTER, out) == []


def test_only_document_xml_changes(tmp_path):
    profile, rows, _ = _example_content()
    out = tmp_path / "out.docx"
    cv_tailor.build_docx(MASTER, out, profile, rows)
    with zipfile.ZipFile(MASTER) as zm, zipfile.ZipFile(out) as zo:
        assert zm.namelist() == zo.namelist()
        for name in zm.namelist():
            if name != "word/document.xml":
                assert zm.read(name) == zo.read(name)


def test_approved_example_passes_the_content_checks():
    profile, rows, _ = _example_content()
    master_text = "\n".join(p.text for p in docx.Document(str(MASTER)).paragraphs)
    unverified = cv_tailor.unverified_terms(master_text)
    assert "Kubernetes" in unverified and "Docker" not in unverified
    m = docx.Document(str(MASTER)).paragraphs
    assert cv_tailor.content_problems(profile, rows, m[4].text, cv_tailor.master_rows(MASTER),
                                      unverified, master_text) == []


def test_unverified_tech_is_rejected_and_stripped():
    unverified = ["Kubernetes", "Kafka"]
    profile = "MSc graduate. Experienced with Kubernetes in production. Keen to grow into Kafka."
    rows = [{"label": "Containers", "items": "Docker, Kubernetes, Docker Compose (multi-service stacks)"}]
    problems = cv_tailor.content_problems(profile, rows * 4, "x" * 500, ["y" * 100] * 6, unverified)
    assert any("KEY SKILLS" in p for p in problems) and any("only as an interest" in p for p in problems)
    new_profile, new_rows = cv_tailor.enforce(profile, rows, unverified)
    assert "Experienced with Kubernetes" not in new_profile
    assert "Keen to grow into Kafka." in new_profile          # interest sentence is allowed
    assert new_rows[0]["items"] == "Docker, Docker Compose (multi-service stacks)"


def test_clean_removes_em_dash_and_emoji():
    assert cv_tailor.clean("Azure — App Service 🚀") == "Azure – App Service"


def test_file_name_follows_the_instruction():
    job = {"company": "Deutsche Börse Group", "title": "Junior Middleware Administrator (f/m/d)"}
    assert cv_tailor.file_stem(MASTER, job, "Middleware Admin") == "Candidate_CV_DeutscheBorse_MiddlewareAdmin"


def test_tailor_end_to_end_with_mocked_model(isolate_data, monkeypatch):
    profile, rows, example_paras = _example_content()
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {
        "profile": profile, "skills": rows, "role_short": "Middleware Admin"})
    job = store.new_job(id="linkedin-1", title="Junior Middleware Administrator",
                        company="Deutsche Börse Group", description="Linux, Kubernetes, Kafka")
    cv_text = "\n".join(p.text for p in docx.Document(str(MASTER)).paragraphs)
    name = cv_tailor.tailor(job, cv_text, MASTER)
    assert name == "Candidate_CV_DeutscheBorse_MiddlewareAdmin.docx"
    out = config.TAILORED_CV_DIR / name
    assert [p.text for p in docx.Document(str(out)).paragraphs] == example_paras


def test_skill_terms_must_exist_in_the_master():
    master_text = "\n".join(p.text for p in docx.Document(str(MASTER)).paragraphs)
    assert cv_tailor.unknown_skill_terms("Docker, PostgreSQL, schema design, GitHub Actions", master_text) == []
    assert cv_tailor.unknown_skill_terms("Docker, Snowflake, AWS S3", master_text) == ["Snowflake", "AWS", "S3"]
    _, rows = cv_tailor.enforce("Fine.", [{"label": "Data", "items": "PostgreSQL, Snowflake, pgvector"}], [], master_text)
    assert rows[0]["items"] == "PostgreSQL, pgvector"


@pytest.mark.parametrize("master, prefix", [
    ("JaneDoe_CV_EN.docx", "JaneDoe"), ("Jane Doe Resume.docx", "JaneDoe"), ("jane-doe-cv.docx", "janedoe"),
    ("JaneDoe.docx", "JaneDoe"), ("CV.docx", "Candidate"), ("Candidate_CV_EN.docx", "Candidate"),
])
def test_file_names_work_for_any_persons_cv(master, prefix):
    job = store.new_job(id="x", title="Data Analyst (m/f/d)", company="Acme s.r.o.")
    stem = cv_tailor.file_stem(master, job, "Data Analyst")
    assert stem == f"{prefix}_CV_Acme_DataAnalyst"
    from jobhelper import cover_letter
    assert cover_letter.file_stem(master, job, "Data Analyst") == f"{prefix}_CoverLetter_Acme_DataAnalyst"


def test_without_a_rules_file_the_general_rules_are_used(isolate_data):
    assert not config.CV_INSTRUCTION.exists()
    assert cv_tailor.load_rules("Jane Doe") == cv_tailor.GENERIC_RULES
    config.CV_INSTRUCTION.write_text("## 1. What changes\nAdapt the profile of {{first_name}}.\n## 2. Skipped\nx\n", encoding="utf-8")
    assert cv_tailor.load_rules("Jane Doe\nAnalyst") == "## 1. What changes\nAdapt the profile of Jane."


# ── any CV layout: the skills rows end at the next heading, whatever it is called ──
def _cv_docx(path, lines):
    import docx
    doc = docx.Document()
    for line in lines:
        doc.add_paragraph(line)
    doc.save(str(path))
    return path


OTHER_LAYOUT = ["Alex Novak", "Prague · alex@example.com", "SUMMARY", "Accountant with three years in a shared service centre.",
                "SKILLS", "Accounting: IFRS, VAT, accounts payable", "Tools: SAP, Excel", "Languages: English C1",
                "EXPERIENCE", "Junior Accountant, Example s.r.o. (2023 - 2026)", "Month-end close and VAT returns for two entities.",
                "Accounts payable for 40 suppliers.", "EDUCATION", "BSc Economics, 2023", "INTERESTS", "Running"]


def _body(path):
    import zipfile

    from lxml import etree
    with zipfile.ZipFile(path) as z:
        return etree.fromstring(z.read("word/document.xml")).find(f"{{{word_document.W}}}body")


def test_another_cv_layout_keeps_its_experience_section(tmp_path):
    master = _cv_docx(tmp_path / "AlexNovak_CV.docx", OTHER_LAYOUT)
    sec = cv_tailor.sections(_body(master))
    assert text_of(sec["profile"]).startswith("Accountant with three years")
    assert [text_of(p) for p in sec["rows"]] == OTHER_LAYOUT[5:8]                 # the three skill rows, and not one line more
    assert cv_tailor.adaptable(master) is True and cv_tailor.master_rows(master) == OTHER_LAYOUT[5:8]
    out = tmp_path / "adapted.docx"
    cv_tailor.build_docx(master, out, "Accountant who closes the books on time.",
                         [{"label": "Accounting", "items": "IFRS, VAT"}, {"label": "Tools", "items": "SAP, Excel"}])
    text = [text_of(p) for p in cv_tailor.sections(_body(out))["paras"]]
    for line in OTHER_LAYOUT[8:]:                                                           # everything after the skills is still there
        assert line in text, line
    assert "Accountant who closes the books on time." in text and "Accounting: IFRS, VAT" in text and OTHER_LAYOUT[7] not in text
    assert cv_tailor.verify(master, out) == []


def test_a_section_that_went_missing_fails_the_check(tmp_path):
    master = _cv_docx(tmp_path / "master.docx", OTHER_LAYOUT)
    broken = _cv_docx(tmp_path / "broken.docx", [x for x in OTHER_LAYOUT if x not in OTHER_LAYOUT[8:12]])     # no EXPERIENCE at all
    problems = cv_tailor.verify(master, broken)
    assert any("section of the master CV is missing" in p or "locked paragraph" in p for p in problems)


@pytest.mark.parametrize("line, heading", [
    ("EXPERIENCE", True), ("W O R K   E X P E R I E N C E", True), ("WORK EXPERIENCE & PRACTICAL TRAINING", True), ("ZKUŠENOSTI", True),
    ("Accounting: IFRS, VAT", False), ("AWS: EC2, S3", False), ("Python, SQL, Docker", False), ("", False), ("2023", False),
    ("A" * 61, False), ("Junior Accountant, Example s.r.o.", False),
])
def test_what_counts_as_a_heading(line, heading):
    assert cv_tailor._is_heading(line) is heading


def test_a_cv_without_a_profile_or_with_endless_rows_is_left_alone(tmp_path):
    no_profile = _cv_docx(tmp_path / "a.docx", ["Alex Novak", "EXPERIENCE", "Junior Accountant", "EDUCATION", "BSc"])
    assert cv_tailor.adaptable(no_profile) is False and cv_tailor.adaptable(tmp_path / "missing.docx") is False
    (tmp_path / "not_word.docx").write_bytes(b"plain text")
    assert cv_tailor.adaptable(tmp_path / "not_word.docx") is False
    endless = _cv_docx(tmp_path / "b.docx", ["Alex Novak", "PROFILE", "A developer.", "KEY SKILLS"] + [f"Row {n}: thing" for n in range(20)])
    sec = cv_tailor.sections(_body(endless))
    assert sec["rows"] == [] and sec["rows_unclear"] is True and sec["profile"] is not None     # the profile is still adapted
    heading_after = _cv_docx(tmp_path / "c.docx", ["Alex Novak", "PROFILE", "KEY SKILLS", "Tools: Excel"])
    assert cv_tailor.sections(_body(heading_after))["profile"] is None                         # an empty profile section is not a profile


def test_a_word_cv_without_a_profile_section_still_gets_a_profile_note(isolate_data, tmp_path, monkeypatch):
    from jobhelper import agent3_letters, config, llm
    config.ensure_dirs()
    master = _cv_docx(config.CV_DIR / "AlexNovak_CV.docx", ["Alex Novak", "EXPERIENCE", "Junior Accountant", "EDUCATION", "BSc"])
    monkeypatch.setattr(cv_tailor, "tailor", lambda *a, **k: (_ for _ in ()).throw(AssertionError("no profile section: not adaptable")))
    monkeypatch.setattr(llm, "ask", lambda *a, **k: "Accountant with three years of month-end closing.")
    job = store.new_job(id="j1", title="Accountant", company="Example", description="IFRS")
    name = agent3_letters._make_tailored_cv(job, "Alex Novak\nAccountant", master, dict(title="Accountant", company="Example", location="", description="IFRS"))
    assert name.endswith(".md") and "month-end closing" in (config.TAILORED_CV_DIR / name).read_text(encoding="utf-8")


def test_a_replaced_document_is_kept_in_the_backup(isolate_data):
    from jobhelper import agent3_letters, config, vault
    config.ensure_dirs()
    old = config.TAILORED_CV_DIR / "Old_CV.docx"
    old.write_bytes(b"the earlier adapted CV")
    agent3_letters._retire(old)
    kept = list((vault.backups_dir() / "letter_versions").iterdir())
    assert not old.exists() and len(kept) == 1 and vault.read_bytes(kept[0]) == b"the earlier adapted CV"
    agent3_letters._retire(old)                                        # nothing there: nothing happens
    assert len(list((vault.backups_dir() / "letter_versions").iterdir())) == 1


# ── the posting's words stated as the person's experience, in any profession ──
ACCOUNTANT_CV = ("Junior accountant with a bachelor's degree in finance and one year of practice in accounts payable. "
                 "Processed about 300 supplier invoices a month in SAP, prepared weekly payment runs, reconciled supplier statements "
                 "and helped with the month-end closing. Bank reconciliation, VAT returns, SAP FI, MS Excel.")
AP_POSTING = ("You will post invoices with the correct tax codes and cost centers, keep the vendor master data up to date within our "
              "Procure to Pay process, and confirm intercompany balances. Experience with Coupa is a plus. We work in SAP.")


def test_terms_of_the_posting_that_the_cv_knows_nothing_of_are_found():
    """Seen with an accountant's CV: the adapted profile said "coding tax and cost center data" and "vendor master
    data upkeep". The fixed never-claim list knows technologies only, so nothing objected."""
    draft = ("Junior accountant with one year of practice in accounts payable. Experienced in processing supplier invoices, "
             "coding tax codes and cost center data, and vendor master data upkeep in SAP. Keen to grow into intercompany balances and Coupa.")
    found = [p.lower() for p in cv_tailor.borrowed_phrases(draft, AP_POSTING, ACCOUNTANT_CV)]
    assert {"tax codes", "cost center", "vendor master", "master data", "intercompany balances"} <= set(found)
    claimed = [p.lower() for p in cv_tailor.borrowed_claims(draft, AP_POSTING, ACCOUNTANT_CV)]
    assert {"tax codes", "cost center", "vendor master", "master data"} <= set(claimed)
    assert "intercompany balances" not in claimed, "named as an interest: that is the honest way"


def test_wording_the_cv_backs_is_left_alone():
    fine = ("Junior accountant with one year of practice in accounts payable. Experienced in processing supplier invoices, "
            "payment runs and reconciling supplier statements in SAP, and in supporting the month-end closing.")
    assert cv_tailor.borrowed_phrases(fine, AP_POSTING, ACCOUNTANT_CV) == []
    # one shared word is not enough; endings do not matter; what the person wrote in their notes counts as theirs
    assert cv_tailor.borrowed_phrases("I reconciled invoices and posted invoice batches.", AP_POSTING, ACCOUNTANT_CV) == []
    notes = ACCOUNTANT_CV + "\nI also kept the vendor master data up to date."
    assert "vendor master" not in [p.lower() for p in cv_tailor.borrowed_phrases("Kept vendor master data current.", AP_POSTING, notes)]
    assert cv_tailor.borrowed_phrases("", AP_POSTING, ACCOUNTANT_CV) == [] and cv_tailor.borrowed_phrases("Cost centers.", "", ACCOUNTANT_CV) == []
    assert cv_tailor.borrowed_claims("", AP_POSTING, ACCOUNTANT_CV) == []


def test_a_draft_that_borrows_from_the_posting_is_sent_back_and_what_is_left_is_noted(isolate_data, monkeypatch):
    profile, rows, _ = _example_content()
    cv_text = "\n".join(p.text for p in docx.Document(str(MASTER)).paragraphs)
    posting = "You will run quarterly zebra migrations and tune walrus pipelines for our platform."
    assert "zebra" not in cv_text.lower() and "walrus" not in cv_text.lower()
    borrowed = profile.rstrip() + " Experienced in zebra migrations."
    prompts, answers = [], [{"profile": borrowed, "skills": rows, "role_short": "Admin"}, {"profile": profile, "skills": rows, "role_short": "Admin"}]
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or answers[min(len(prompts) - 1, 1)])
    job = store.new_job(id="j1", title="Administrator", company="Acme", description=posting)
    name = cv_tailor.tailor(job, cv_text, MASTER)
    assert len(prompts) == 2 and "PROFILE presents 'zebra migrations' as the candidate's experience" in prompts[1]
    assert "cv_check" not in job                                            # the second draft is clean
    text = "\n".join(p.text for p in docx.Document(str(config.TAILORED_CV_DIR / name)).paragraphs)
    assert "zebra" not in text
    # the model insists: the app does not cut the sentence out, it tells the person what to look at
    prompts.clear()
    answers[1] = answers[0]
    job = store.new_job(id="j2", title="Administrator", company="Acme", description=posting)
    cv_tailor.tailor(job, cv_text, MASTER)
    assert job["cv_check"] == ["zebra migrations"]
    # named as an interest, it is neither sent back nor noted
    prompts.clear()
    keen = {"profile": profile.rstrip() + " Keen to grow into zebra migrations.", "skills": rows, "role_short": "Admin"}
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or keen)
    job = store.new_job(id="j3", title="Administrator", company="Acme", description=posting)
    cv_tailor.tailor(job, cv_text, MASTER)
    assert not any("zebra migrations' as the candidate" in prompt for prompt in prompts) and "cv_check" not in job


def test_a_skill_inside_another_technologys_name_is_not_in_the_cv():
    """"Java" was found inside "JavaScript"; a part of a word otherwise still counts ("SQL" in "PostgreSQL", "API" in "APIs")."""
    assert cv_tailor.unknown_skill_terms("Java, Spring", "JavaScript, TypeScript") == ["Java", "Spring"]
    assert cv_tailor.unknown_skill_terms("REST APIs, SQL, Docker, API", "RESTful API design, PostgreSQL, Dockerized apps") == ["APIs"]
    assert cv_tailor.unknown_skill_terms("Java", "Java and JavaScript") == []


def test_an_invented_figure_or_name_in_the_profile_is_sent_back_then_cut_or_noted(isolate_data, monkeypatch):
    """The CV is the document employers see; its PROFILE had no figure check and no name check."""
    profile, rows, _ = _example_content()
    cv_text = "\n".join(p.text for p in docx.Document(str(MASTER)).paragraphs)
    assert "Initech" not in cv_text and "8 years" not in cv_text
    invented = profile.rstrip() + " Ran Docker clusters for 8 years at Initech."
    prompts = []
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or {"profile": invented, "skills": rows, "role_short": "Admin"})
    job = store.new_job(id="j1", title="Administrator", company="Acme", description="We run containers.")
    name = cv_tailor.tailor(job, cv_text, MASTER)
    assert len(prompts) == 2 and "PROFILE states the figure 8" in prompts[1] and "Initech" not in prompts[1], "a name is noted, not sent back"
    text = "\n".join(p.text for p in docx.Document(str(config.TAILORED_CV_DIR / name)).paragraphs)
    assert "8 years" not in text and "Initech" not in text, "the sentence with the invented figure is cut out"
    noted = profile.rstrip() + " Led the platform team at Initech."
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: {"profile": noted, "skills": rows, "role_short": "Admin"})
    job = store.new_job(id="j2", title="Administrator", company="Acme", description="We run containers.")
    cv_tailor.tailor(job, cv_text, MASTER)
    assert "Initech" in job["cv_check"], "a name the CV does not give is put in front of the person"
    assert cv_tailor.profile_problems("Keen to grow into Kafka and Kubernetes.", job, cv_text) == ([], []), "an interest claims nothing"
    rust = store.new_job(id="j3", title="Rust Developer", company="Acme", description="")
    assert "Rust" in cv_tailor.profile_problems("I built many Rust services.", rust, cv_text)[1], "the job title backs no name"
    # a profile that is only the invented sentence is kept, and its figure noted, rather than failing the CV
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: {"profile": "Developer with 8 years of Python.", "skills": rows, "role_short": "Dev"})
    job = store.new_job(id="j4", title="Administrator", company="Acme", description="")
    cv_tailor.tailor(job, cv_text, MASTER)
    assert job["cv_check"][0] == "the figure 8", "named as a figure, first, so the cap of eight never hides it"



def test_the_cv_note_lists_a_figure_first_and_a_cut_never_breaks_a_sentence(isolate_data, monkeypatch):
    profile, rows, _ = _example_content()
    cv_text = "\n".join(p.text for p in docx.Document(str(MASTER)).paragraphs)
    posting = "You will run quarterly zebra migrations, tune walrus pipelines and build otter dashboards."
    # an abbreviation: "e.g." is not a sentence end, so nothing is cut, and the figure is noted first
    crowded = ("Backend engineer, e.g. 47 years of Python. Experienced in zebra migrations, walrus pipelines and otter dashboards. "
               "Worked with Initech, Globex, Hooli, Umbrella, Soylent, Vandelay, Wonka and Stark on 47 releases.")
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: {"profile": crowded, "skills": rows, "role_short": "Dev"})
    job = store.new_job(id="j1", title="Developer", company="Acme", description=posting, cv_check=["old"])
    name = cv_tailor.tailor(job, cv_text, MASTER)
    text = "\n".join(p.text for p in docx.Document(str(config.TAILORED_CV_DIR / name)).paragraphs)
    assert "e.g. 47 years of Python" in text, "no half sentence is left behind"
    assert job["cv_check"][0] == "the figure 47" and len(job["cv_check"]) <= 8 and job["cv_check"].count("the figure 47") == 1
    # a clean draft clears an old note
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: {"profile": profile, "skills": rows, "role_short": "Dev"})
    cv_tailor.tailor(job, cv_text, MASTER)
    assert "cv_check" not in job


# ── a plain skills list: reordered, never cut to one item a row (hands-on check, 29 Oct) ──
PLAIN_SKILLS = "Python, Django, Flask, PostgreSQL, SQL, Git, Docker, pytest, REST APIs, Linux"


def _plain_master(path: Path) -> Path:
    """A one-page CV of an invented person whose KEY SKILLS is one line without a label."""
    d = docx.Document()
    for text in ("Mara Example", "Junior Software Developer | Brno | mara@example.com", "PROFILE",
                 "Junior Python developer with two years of practice building small web services and test automation.",
                 "KEY SKILLS", PLAIN_SKILLS, "EXPERIENCE",
                 "Junior Developer – Northwind Parcel s.r.o., Brno (2024 – present): built an internal REST API in Django.",
                 "EDUCATION", "Bc. Applied Informatics – Example University, Brno (2020 – 2023)"):
        d.add_paragraph(text)
    d.save(str(path))
    return path


def test_a_plain_skills_list_keeps_its_shape_and_every_skill(isolate_data, monkeypatch, tmp_path):
    """Seen in the hands-on check: KEY SKILLS came out as "Django: Python / PostgreSQL: PostgreSQL / …", and Flask,
    SQL, Git and Linux were gone, with no note."""
    from jobhelper import pdf
    master = _plain_master(tmp_path / "Mara_CV.docx")
    cv_text = "\n".join(p.text for p in docx.Document(str(master)).paragraphs)
    shape = cv_tailor.MasterSkills(cv_tailor.master_rows(master))
    assert shape.labelled is False and shape.limits == (1, 1) and "Flask" in shape.items()
    labelled = [{"label": "Django", "items": "Python, Django"}, {"label": "PostgreSQL", "items": "PostgreSQL, SQL"},
                {"label": "Docker", "items": "Docker, Git"}, {"label": "pytest", "items": "pytest"}, {"label": "Other", "items": "REST APIs"}]
    one_row_short = [{"label": "", "items": "Django, Python, REST APIs, pytest, PostgreSQL"}]
    prompts, answers = [], [{"profile": "Junior Python developer.", "skills": labelled, "role_short": "Python Dev"},
                            {"profile": "Junior Python developer.", "skills": one_row_short, "role_short": "Python Dev"}]
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or answers[min(len(prompts) - 1, 1)])
    monkeypatch.setattr(pdf, "to_pdf", lambda path: None)
    job = store.new_job(id="p1", title="Python Developer", company="Contoso", description="Django, REST APIs, pytest.")
    name = cv_tailor.tailor(job, cv_text, master)
    assert "These rows have no labels. Return exactly 1 row" in prompts[0]
    assert "KEY SKILLS has 5 rows; use exactly 1, as the master does." in prompts[1]
    texts = [p.text for p in docx.Document(str(config.TAILORED_CV_DIR / name)).paragraphs]
    row = texts[texts.index("KEY SKILLS") + 1]
    assert row.startswith("Django, Python, REST APIs") and ":" not in row, "the posting's skills first, no label"
    items = [i.strip() for i in row.split(",")]
    for skill in ("Flask", "SQL", "Git", "Linux", "Docker"):
        assert skill in items, f"{skill} from the master is kept"
    assert texts[texts.index("KEY SKILLS") + 2] == "EXPERIENCE" and cv_tailor.verify(master, config.TAILORED_CV_DIR / name) == []


def test_a_skill_left_out_to_fit_the_page_is_named_for_the_person(isolate_data, monkeypatch, tmp_path):
    from jobhelper import pdf
    master = _plain_master(tmp_path / "Mara_CV.docx")
    cv_text = "\n".join(p.text for p in docx.Document(str(master)).paragraphs)
    long_row = [{"label": "", "items": PLAIN_SKILLS + ", " + ", ".join(["Python"] * 60)}]
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: {"profile": "Junior Python developer.", "skills": long_row, "role_short": "Dev"})
    monkeypatch.setattr(pdf, "to_pdf", lambda path: None)
    job = store.new_job(id="p2", title="Python Developer", company="Contoso", description="Python.")
    cv_tailor.tailor(job, cv_text, master)
    assert job["cv_left_out"][0] == "Python" and "cv_check" not in job, "never cut silently, in a note of its own"
    assert cv_tailor._drop_last_item([{"label": "", "items": "Git"}]) == "", "a last item is never taken"


def test_the_general_rules_ask_for_working_independently_only_when_the_cv_says_it():
    """Seen in the hands-on check: the profile gained "Works independently on backend tasks…", which the CV never said."""
    assert "keep a sentence about working independently only if the master's PROFILE has one" in cv_tailor.GENERIC_RULES
    assert "keep one sentence about working independently" not in cv_tailor.GENERIC_RULES


def test_the_shape_of_the_master_rows_decides_what_is_kept():
    labelled = cv_tailor.MasterSkills(["Languages: Python, SQL", "Tools: Git, Docker"])
    assert labelled.labelled and not labelled.plain and labelled.limits == (2, 8)
    assert labelled.missing([{"label": "Code", "items": "Python"}]) == [], "labelled rows may leave out what the posting ignores"
    assert labelled.shape_rule() == "" and labelled.row({"label": "Code", "items": "Python."}) == {"label": "Code", "items": "Python"}
    assert cv_tailor.MasterSkills(["a: b"] * 6).limits == (4, 8)
    for neither in ([], ["Python, Django", "Languages: English"]):        # no rows found, or a mix: the general rules alone
        shape = cv_tailor.MasterSkills(neither)
        assert not shape.plain and shape.limits == (cv_tailor.MIN_ROWS, cv_tailor.MAX_ROWS) and shape.shape_rule() == ""
        assert shape.missing([{"label": "X", "items": "Python"}]) == []
    plain = cv_tailor.MasterSkills([PLAIN_SKILLS])
    assert plain.row({"label": "Backend", "items": "Python"}) == {"label": "", "items": "Python"}, "a plain list has no labels"
    assert plain.missing([{"label": "", "items": "Python, Django, Flask, PostgreSQL, SQL, Git, Docker, pytest, REST API, Linux"}]) == []


def test_a_plain_list_sent_back_twice_in_rows_is_put_together_again(isolate_data, monkeypatch, tmp_path):
    from jobhelper import pdf
    master = _plain_master(tmp_path / "Mara_CV.docx")
    cv_text = "\n".join(p.text for p in docx.Document(str(master)).paragraphs)
    rows = [{"label": "Web", "items": "Django, Flask"}, {"label": "Data", "items": "PostgreSQL"}, {"label": "Tools", "items": "Git"}]
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: {"profile": "Junior Python developer.", "skills": rows, "role_short": "Dev"})
    monkeypatch.setattr(pdf, "to_pdf", lambda path: None)
    job = store.new_job(id="p3", title="Python Developer", company="Contoso", description="Django.")
    name = cv_tailor.tailor(job, cv_text, master)
    texts = [p.text for p in docx.Document(str(config.TAILORED_CV_DIR / name)).paragraphs]
    row = texts[texts.index("KEY SKILLS") + 1]
    assert texts[texts.index("KEY SKILLS") + 2] == "EXPERIENCE", "one row, as in the master"
    assert row.startswith("Django, Flask, PostgreSQL, Git") and set(PLAIN_SKILLS.split(", ")) <= set(row.split(", "))
