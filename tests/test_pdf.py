"""PDF copies through LibreOffice. LibreOffice itself is stubbed, except for one
real conversion that only runs where it is installed."""
import shutil
from pathlib import Path
from types import SimpleNamespace

import docx
import pytest

from jobhelper import agent3_letters, config, cv_tailor, llm, pdf, store

ROOT = Path(__file__).resolve().parent.parent
# A frozen copy of the master CV the approved examples were made from, so these
# tests do not change when the real CV in data/cv/ is edited.
MASTER = Path(__file__).parent / "fixtures" / "Candidate_CV_EN.docx"
EXAMPLE = Path(__file__).parent / "fixtures" / "cv_example_deutscheborse.docx"
REAL_SOFFICE = next((p for p in ("/Applications/LibreOffice.app/Contents/MacOS/soffice",)
                     if Path(p).is_file()), "") or shutil.which("soffice") or ""


@pytest.fixture
def fake_office(monkeypatch):
    """A LibreOffice that 'converts' by writing a marker PDF; records its calls."""
    calls, state = [], {"fail": False}

    def run(cmd, **kwargs):
        calls.append(cmd)
        if state["fail"]:
            return SimpleNamespace(returncode=1, stdout="", stderr="boom")
        outdir, src = Path(cmd[cmd.index("--outdir") + 1]), Path(cmd[-1])
        (outdir / (src.stem + ".pdf")).write_bytes(b"%PDF-1.4 fake")
        return SimpleNamespace(returncode=0, stdout="", stderr="")

    monkeypatch.setattr(pdf, "find_soffice", lambda: "/fake/soffice")
    monkeypatch.setattr(pdf.subprocess, "run", run)
    return SimpleNamespace(calls=calls, state=state)


def test_nothing_happens_without_libreoffice(tmp_path):
    doc = tmp_path / "cv.docx"
    doc.write_bytes(b"x")
    assert pdf.available() is False
    assert pdf.to_pdf(doc) is None and not (tmp_path / "cv.pdf").exists()


def test_converts_next_to_the_document_with_a_private_profile(fake_office, tmp_path):
    doc = tmp_path / "My CV.docx"
    doc.write_bytes(b"x")
    out = pdf.to_pdf(doc)
    assert out == tmp_path / "My CV.pdf" and out.read_bytes().startswith(b"%PDF")
    cmd = fake_office.calls[0]
    assert cmd[0] == "/fake/soffice" and "--headless" in cmd and cmd[-1] == str(doc)
    assert any(arg.startswith("-env:UserInstallation=file://") for arg in cmd)
    assert pdf.pdf_for(doc) == out


def test_failed_conversion_returns_none(fake_office, tmp_path):
    doc = tmp_path / "cv.docx"
    doc.write_bytes(b"x")
    fake_office.state["fail"] = True
    assert pdf.to_pdf(doc) is None and pdf.pdf_for(doc) is None


def test_only_docx_files_are_converted(fake_office, tmp_path):
    note = tmp_path / "letter.md"
    note.write_text("x", encoding="utf-8")
    assert pdf.to_pdf(note) is None and pdf.to_pdf(tmp_path / "missing.docx") is None
    assert fake_office.calls == []


def test_page_count_of_an_unreadable_file_is_zero(tmp_path):
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"not a pdf")
    assert pdf.page_count(bad) == 0 and pdf.page_count(tmp_path / "none.pdf") == 0


# ── the one-page rule on the rendered CV ────────────────────────────────────
def _example():
    paras = [p.text for p in docx.Document(str(EXAMPLE)).paragraphs]
    rows = [dict(zip(("label", "items"), t.split(": ", 1))) for t in paras[6:13]]
    return paras[4], rows


def _tailor(monkeypatch, pages, master=1):
    profile, rows = _example()
    seen = []
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"profile": profile, "skills": rows, "role_short": "Admin"})
    monkeypatch.setattr(pdf, "to_pdf", lambda path: Path(path).with_suffix(".pdf"))
    monkeypatch.setattr(pdf, "page_count", lambda path: (seen.append(1), pages[min(len(seen), len(pages)) - 1])[1])
    monkeypatch.setattr(cv_tailor, "master_pages", lambda path: master)
    job = store.new_job(id="linkedin-1", title="Junior Middleware Administrator", company="Deutsche Börse Group",
                        description="Linux")
    cv_text = "\n".join(p.text for p in docx.Document(str(MASTER)).paragraphs)
    name = cv_tailor.tailor(job, cv_text, MASTER)
    return sum(len(r["items"]) for r in rows), [p.text for p in docx.Document(str(config.TAILORED_CV_DIR / name)).paragraphs]


@pytest.mark.skipif(not MASTER.exists(), reason="master CV not present")
def test_cv_over_one_page_is_trimmed_until_it_fits(isolate_data, monkeypatch):
    profile, rows = _example()
    original = [f"{r['label']}: {r['items']}" for r in rows]
    _, paras = _tailor(monkeypatch, pages=[2, 2, 1])
    now = paras[6:13]
    assert sum(len(a) for a in now) < sum(len(a) for a in original)        # two items were dropped
    assert sum(1 for a, b in zip(original, now) if a != b) in (1, 2)
    assert paras[4] == profile                                              # the profile is never cut


@pytest.mark.skipif(not MASTER.exists(), reason="master CV not present")
def test_one_page_cv_is_left_alone(isolate_data, monkeypatch):
    _, rows = _example()
    _, paras = _tailor(monkeypatch, pages=[1])
    assert paras[6:13] == [f"{r['label']}: {r['items']}" for r in rows]


@pytest.mark.skipif(not MASTER.exists(), reason="master CV not present")
def test_no_trimming_when_the_master_itself_is_not_one_page(isolate_data, monkeypatch):
    _, rows = _example()
    _, paras = _tailor(monkeypatch, pages=[2, 2, 2], master=2)
    assert paras[6:13] == [f"{r['label']}: {r['items']}" for r in rows]


@pytest.mark.skipif(not MASTER.exists(), reason="master CV not present")
def test_trimming_gives_up_after_a_few_tries(isolate_data, monkeypatch):
    _, paras = _tailor(monkeypatch, pages=[2] * 50)
    assert len(paras[6:13]) == 7                       # rows are shortened, never removed


# ── agent 3 makes the PDF copies ────────────────────────────────────────────
def test_letter_and_cv_get_pdf_copies_and_old_ones_are_removed(isolate_data, monkeypatch, fake_office):
    config.ensure_dirs()
    for folder, old in ((config.TAILORED_CV_DIR, "old_cv"), (config.LETTERS_DIR, "old_letter")):
        (folder / f"{old}.docx").write_bytes(b"x")
        (folder / f"{old}.pdf").write_bytes(b"x")
    config.LETTER_TEMPLATE.write_bytes(b"template")

    def tailor(job, cv_text, cv_path):
        (config.TAILORED_CV_DIR / "new_cv.docx").write_bytes(b"x")
        return "new_cv.docx"

    def write(job, cv_text, cv_path):
        (config.LETTERS_DIR / "new_letter.docx").write_bytes(b"x")
        return "new_letter.docx"

    monkeypatch.setattr(cv_tailor, "tailor", tailor)
    monkeypatch.setattr(cv_tailor, "adaptable", lambda path: True)          # the stand-in master is a Word CV with a profile
    monkeypatch.setattr(agent3_letters.cover_letter, "write", write)
    job = store.new_job(id="j", title="Dev", company="Acme", source="linkedin",
                        tailored_cv_file="old_cv.docx", cover_letter_file="old_letter.docx")
    fields = dict(title="Dev", company="Acme", location="", description="", profile="")
    assert agent3_letters._make_tailored_cv(job, "cv text", "master.docx", fields) == "new_cv.docx"
    assert agent3_letters._make_letter(job, "cv text", "master.docx", fields) == "new_letter.docx"
    assert sorted(p.name for p in config.TAILORED_CV_DIR.iterdir()) == ["new_cv.docx", "new_cv.pdf"]
    assert sorted(p.name for p in config.LETTERS_DIR.iterdir()) == ["new_letter.docx", "new_letter.pdf"]


# ── the real thing, where LibreOffice is installed ──────────────────────────
@pytest.mark.skipif(not REAL_SOFFICE, reason="LibreOffice is not installed")
def test_real_conversion_gives_a_one_page_pdf(tmp_path, monkeypatch):
    monkeypatch.setattr(pdf, "find_soffice", lambda: REAL_SOFFICE)
    doc = tmp_path / "example.docx"
    shutil.copy2(EXAMPLE, doc)
    out = pdf.to_pdf(doc)
    assert out is not None and out.read_bytes().startswith(b"%PDF")
    assert pdf.page_count(out) == 1
