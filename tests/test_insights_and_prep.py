"""Insights worked out locally, and interview preparation built from the CV."""
from pathlib import Path
import pytest

from jobhelper import agent2_search, config, db, insights, interview_prep, llm, store, ui_api

CV = ("Alex Novak\nJunior Software Developer\nPython, Kotlin, SQL, PostgreSQL, FastAPI, Docker, Azure, Git, Linux, React.\n"
      "Built a document search service with FastAPI and pgvector. English B2.")


def _app(id, title, status, source="linkedin", **kw):
    return store.new_job(id=id, title=title, company=f"Co {id}", status=status, source=source, **kw)


# ── what is working ─────────────────────────────────────────────────────────
@pytest.mark.parametrize("title, kind", [
    ("Junior AI Engineer", "AI / machine learning"), ("Machine Learning Engineer", "AI / machine learning"),
    ("Junior Data Scientist", "AI / machine learning"), ("Data Analyst", "Data"), ("Junior Data Engineer", "Data"),
    ("QA Tester", "QA / testing"), ("Test Automation Engineer", "QA / testing"),
    ("IT Support Specialist", "IT support"), ("Help Desk Technician", "IT support"), ("Technical Support Engineer", "IT support"),
    ("Android Developer", "Mobile"), ("Junior Python Developer", "Software development"), ("Backend Engineer", "Software development"),
    ("Marketing Consultant", "Sales and marketing"), ("Staff Nurse", "Care and health"), ("Junior Accountant", "Finance"),
    ("Beekeeper", "Other"), ("(existing application)", insights.ROLE_UNKNOWN), ("", insights.ROLE_UNKNOWN),
])
def test_role_types(title, kind):
    assert insights.role_type(title) == kind


def test_what_works_counts_outcomes_per_role_and_source(isolate_data):
    jobs = [_app("1", "Junior Python Developer", "interview"), _app("2", "Backend Engineer", "rejected"),
            _app("3", "Software Engineer", "waiting"),
            _app("4", "IT Support", "rejected", "jobs.cz"), _app("5", "Help Desk", "rejected", "jobs.cz"),
            _app("6", "Service Desk Agent", "rejected", "jobs.cz"),
            _app("7", "QA Tester", "rejected"),                                   # too few
            _app("8", "Data Analyst", "rejected"),                                # reached an interview earlier
            _app("9", "(existing application)", "waiting", "manual"),
            store.new_job(id="lead", title="Junior Python Developer", company="X", status="found")]
    responses = [{"job_id": "8", "stage": "interview"}, {"job_id": "8", "stage": "rejection"},
                 {"job_id": "2", "stage": "acknowledged"}]
    w = insights.what_works(jobs, responses)
    by = {r["name"]: r for r in w["by_role"]}
    assert w["applications"] == 9 and w["by_role"][-1]["name"] == insights.ROLE_UNKNOWN
    dev = by["Software development"]
    assert (dev["applications"], dev["interviews"], dev["rejected"], dev["waiting"], dev["replied"]) == (3, 1, 1, 1, 1)
    assert dev["verdict"] == "working: 1 of 3 reached an interview"
    assert by["IT support"]["verdict"] == "not working: all 3 rejected"
    assert by["QA / testing"]["verdict"] == "too few to say"
    assert by["Data"]["interviews"] == 1                                          # an earlier interview still counts
    src = {r["name"]: r for r in w["by_source"]}
    assert src["jobs.cz"]["applications"] == 3 and src["your own applications"]["applications"] == 1
    assert sum(r["applications"] for r in w["by_source"]) == 9 == sum(r["applications"] for r in w["by_role"])


# ── search keywords ─────────────────────────────────────────────────────────
def test_keyword_stats_show_which_keywords_waste_scoring():
    def lead(i, title, score, query=None, **kw):
        j = store.new_job(id=str(i), title=title, company="C", match_score=score, **kw)
        if query:
            j["query"] = query
        return j

    jobs = [lead(1, "IT Support Specialist", 10, "it support"), lead(2, "IT Support Engineer", 15, "it support"),
            lead(3, "Senior IT Support", 20, "it support"),
            lead(4, "Junior Python Developer", 70, "junior python developer"), lead(5, "Python Developer", 35, "junior python developer"),
            lead(6, "Python Developer II", 25, "junior python developer"),
            lead(7, "QA Tester", 40),                                             # no keyword recorded: matched by title
            lead(8, "Marketing Manager", 5),                                      # belongs to no keyword
            lead(9, "Old application", 90, "it support", source="manual", status="waiting")]
    rows = {r["keyword"]: r for r in insights.keyword_stats(jobs, ["it support", "junior python developer", "qa tester", "llm engineer"], 30)}
    assert rows["it support"]["verdict"] == "consider dropping: 3 scored, none reached 30" and rows["it support"]["average"] == 15
    assert rows["junior python developer"]["verdict"] == "working: 1 of 3 scored 50 or more"
    assert (rows["qa tester"]["leads"], rows["qa tester"]["guessed"], rows["qa tester"]["verdict"]) == (1, 1, "too few to say")
    assert rows["llm engineer"]["verdict"] == "no jobs found with it yet" and rows["llm engineer"]["average"] is None
    assert sum(r["leads"] for r in rows.values()) == 7                            # the manual one and the unmatched one are left out


def test_search_remembers_which_keyword_found_a_job(monkeypatch):
    ad = ("We are looking for a junior Python developer to join our team in Prague. "
          "You will work with the backend team and write clean, tested code every day.")
    monkeypatch.setattr(agent2_search, "active_providers", lambda source, max_age_days=0: [
        ("one", lambda q, w, n: [store.new_job(id="x-" + q, title="Dev " + q, company="Acme " + q, location="Prague", description=ad)])])
    found = agent2_search.search(["python developer", "qa tester"], where="Prague")
    assert {j["query"] for j in found} == {"python developer", "qa tester"}


# ── skills to learn ─────────────────────────────────────────────────────────
def test_skill_gaps_lists_what_postings_ask_and_the_cv_lacks():
    jobs = [store.new_job(id="1", title="A", company="A", match_score=60, description="Python, AWS, Kubernetes and Docker. Fluent German."),
            store.new_job(id="2", title="B", company="B", match_score=10, description="AWS and Kafka. Agile team, Jira."),
            store.new_job(id="3", title="C", company="C", status="waiting", description="Kubernetes, PostgreSQL."),
            store.new_job(id="4", title="D", company="D", description=""),
            store.new_job(id="5", title="E", company="E", status="not_interested", description="Rust and Scala")]
    g = insights.skill_gaps(jobs, CV)
    assert g["postings"] == 3 and g["relevant"] == 2
    skills = {r["skill"]: r for r in g["skills"]}
    assert set(skills) == {"AWS", "Kubernetes", "Kafka"}                          # not Python/Docker (in the CV), not Agile/Jira
    assert (skills["Kubernetes"]["postings"], skills["Kubernetes"]["relevant"]) == (2, 2)
    assert (skills["AWS"]["postings"], skills["AWS"]["relevant"], skills["AWS"]["share"]) == (2, 1, 67)
    assert g["skills"][0]["skill"] == "Kubernetes"                                # asked most in postings close to you
    assert g["languages"] == [{"language": "German", "postings": 1}]
    assert "Python" in g["in_cv"] and "English" not in g["in_cv"]


def test_skill_gaps_never_change_the_cv_or_a_letter(isolate_data, monkeypatch):
    import inspect
    assert "tailored" not in inspect.getsource(insights) and "LETTERS_DIR" not in inspect.getsource(insights)


def test_api_insights_puts_it_together(isolate_data, monkeypatch):
    config.ensure_dirs()
    monkeypatch.setattr(ui_api.Api, "_cv_text", lambda self: CV)
    job = store.new_job(id="1", title="Junior Python Developer", company="A", match_score=60, description="Python and AWS.")
    store.save_jobs([job, _app("2", "QA Tester", "rejected")])
    api = ui_api.Api()
    api.save_options({**api.options(), "queries": ["junior python developer"], "threshold": 30})
    data = api.insights()
    assert data["threshold"] == 30 and data["keywords"][0]["keyword"] == "junior python developer"
    assert data["gaps"]["skills"][0]["skill"] == "AWS" and data["works"]["applications"] == 1


# ── interview preparation ───────────────────────────────────────────────────
GOOD = {"summary": "Expect questions about Python and your projects.",
        "questions": [
            {"kind": "technical", "question": "How did you build the RAG assistant?", "why": "It is your main project.",
             "answer": "I built it with FastAPI and pgvector, and I ran the model locally."},
            {"kind": "about you", "question": "Tell me about yourself.", "why": "Opening.",
             "answer": "I am a junior developer who works with Python and Kotlin."},
            {"kind": "technical", "question": "Have you used Kubernetes?", "why": "It is in the posting.",
             "answer": "I have not used Kubernetes in a project; the closest I have done is Docker Compose."},
            {"kind": "nonsense kind", "question": "Why us?", "why": "", "answer": "The work matches what I build."},
            {"kind": "practical", "question": "", "answer": "empty question is dropped"}],
        "ask_them": ["What does the first month look like?"], "revise": ["Kubernetes basics"]}


def _job():
    return store.new_job(id="i1", title="Junior Python Developer", company="Acme", status="interview",
                         description="Python, FastAPI, Kubernetes.")


def test_prep_is_cleaned_grouped_and_honest_bridges_are_accepted(isolate_data, monkeypatch):
    prompts = []
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or GOOD)
    prep = interview_prep.generate(_job(), CV, [{"stage": "interview", "summary": "Invitation to a 45-minute call."}])
    assert len(prompts) == 1                                                       # nothing to fix: one call
    assert [q["kind"] for q in prep["questions"]] == ["about you", "about you", "technical", "technical"]
    assert all(q["check"] == "" for q in prep["questions"])                        # "have not used Kubernetes" is not a claim
    assert prep["ask_them"] == ["What does the first month look like?"] and prep["made"]
    assert "Invitation to a 45-minute call." in prompts[0] and "Alex Novak" in prompts[0]
    # nothing about the person is written into the code: unknown facts are named as unknown
    assert "Right to work: not stated" in prompts[0] and "Can start: not stated" in prompts[0]
    assert "Czech Republic" not in prompts[0].split("Master CV")[0]


def test_an_answer_claiming_what_the_cv_lacks_is_retried_then_marked(isolate_data, monkeypatch):
    bad = {"summary": "s", "questions": [
        {"kind": "technical", "question": "Kubernetes?", "why": "", "answer": "I have two years of production experience with Kubernetes."},
        {"kind": "about you", "question": "You?", "why": "", "answer": "Alex is a junior developer."}]}
    calls = []
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: calls.append(prompt) or bad)
    prep = interview_prep.generate(_job(), CV)
    assert len(calls) == 2 and "Kubernetes" in calls[1] and "first person" in calls[1]
    marked = {q["question"]: q["check"] for q in prep["questions"]}
    assert marked["Kubernetes?"] == "Kubernetes"                                   # still wrong after the retry: shown, with a warning


def test_unusable_reply_is_an_error(isolate_data, monkeypatch):
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"questions": []})
    with pytest.raises(ValueError, match="no questions"):
        interview_prep.generate(_job(), CV)


def test_prep_text_for_copying():
    prep = interview_prep._clean(GOOD)
    text = interview_prep.as_text(_job(), prep)
    assert text.startswith("Interview preparation: Junior Python Developer at Acme")
    assert "IF YOU SPEAK WITH HR OR A RECRUITER" in text and "IF YOU SPEAK WITH SOMEONE WHO DOES THIS WORK OR THE HIRING MANAGER" in text
    assert text.index("IF YOU SPEAK WITH HR") < text.index("IF YOU SPEAK WITH SOMEONE WHO DOES THIS WORK")
    assert "QUESTIONS TO ASK THEM" in text and "REVISE BEFORE THE INTERVIEW" in text
    assert "—" not in text


def test_api_makes_keeps_and_returns_the_preparation(isolate_data, monkeypatch):
    config.ensure_dirs()
    monkeypatch.setattr(ui_api.Api, "_cv_text", lambda self: CV)
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: GOOD)
    store.save_jobs([_job(), store.new_job(id="lead", title="Dev", company="B")])
    db.upsert_response({"msg_id": "<1@x>", "job_id": "i1", "company": "Acme", "stage": "interview", "summary": "Call on Friday."})
    api = ui_api.Api()
    assert api.interview_prep("i1") == {} and api.job_detail("i1")["has_prep"] is False
    assert api.make_interview_prep("i1") == {"ok": True}
    saved = api.interview_prep("i1")
    assert saved["company"] == "Acme" and len(saved["questions"]) == 4 and api.job_detail("i1")["has_prep"] is True
    assert "Tell me about yourself." in api.prep_text("i1")
    assert api.make_interview_prep("lead")["ok"] is False                          # only for applications
    assert api.make_interview_prep("nope")["ok"] is False and api.prep_text("lead") == ""

    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: (_ for _ in ()).throw(llm.LimitReached("Claude's 5-hour limit is used up (99%).")))
    failed = api.make_interview_prep("i1")
    assert failed["ok"] is False and "limit is used up" in failed["msg"]
    assert len(api.interview_prep("i1")["questions"]) == 4                         # the earlier one is kept


# ── trends, time to a decision, why jobs score low ──────────────────────────
from datetime import date, datetime, timedelta  # noqa: E402
from email.utils import format_datetime  # noqa: E402

WED = date(2026, 10, 7)                       # a Wednesday; its week starts on Monday 5 October


def _mail(job_id, day, stage):
    return {"job_id": job_id, "stage": stage, "received_date": format_datetime(datetime.combine(day, datetime.min.time()))}


def test_trends_count_by_calendar_week(isolate_data):
    jobs = [_app("1", "Dev", "waiting", date_applied="2026-10-05"), _app("2", "Dev", "rejected", date_applied="2026-10-07"),
            _app("3", "Dev", "waiting", date_applied="2026-10-04"),                 # Sunday: the week before
            _app("4", "Dev", "waiting", date_applied=""),
            store.new_job(id="lead", title="Dev", company="X", status="found", date_applied="2026-10-06")]
    mails = [_mail("2", date(2026, 10, 6), "rejection"), _mail("1", date(2026, 10, 5), "acknowledged"),
             _mail("3", date(2026, 9, 30), "interview"), _mail("9", date(2026, 9, 1), "assessment"), {"job_id": "x", "stage": "other", "received_date": "?"}]
    got = insights.trends(jobs, mails, weeks=8, today=WED)
    assert got[0] == {"week": "2026-10-05", "current": True, "applications": 2, "replies": 2, "invitations": 0, "rejections": 1}
    assert got[1] == {"week": "2026-09-28", "current": False, "applications": 1, "replies": 1, "invitations": 1, "rejections": 0}
    assert got[-1]["week"] == "2026-08-31" and got[-1]["invitations"] == 1          # empty older weeks are cut off
    assert [t["week"] for t in insights.trends([], [], today=WED)] == ["2026-10-05"]


def test_reply_times_and_who_has_waited_too_long(isolate_data):
    def decided(i, applied, after, title="Junior Python Developer", stage="rejection"):
        return _app(str(i), title, "rejected", date_applied=applied), _mail(str(i), date.fromisoformat(applied) + timedelta(days=after), stage)
    pairs = [decided(1, "2026-09-01", 2), decided(2, "2026-09-01", 4), decided(3, "2026-09-01", 6),
             decided(4, "2026-09-01", 8, "IT Support"), decided(5, "2026-09-01", 20, "IT Support", "interview")]
    jobs = [p[0] for p in pairs] + [_app("w1", "Dev", "waiting", date_applied="2026-10-06"),
                                    _app("w2", "Dev", "waiting", date_applied="2026-08-01"),
                                    _app("w3", "Dev", "waiting", date_applied=""),
                                    _app("early", "Dev", "rejected", date_applied="2026-09-10")]
    mails = [p[1] for p in pairs] + [_mail("1", date(2026, 9, 2), "acknowledged"),           # a confirmation is not a decision
                                     _mail("1", date(2026, 9, 9), "rejection"),              # only the first decision counts
                                     _mail("early", date(2026, 9, 1), "rejection")]          # before the apply date: ignored
    got = insights.reply_times(jobs, mails, today=WED)
    assert (got["decisions"], got["middle"], got["most_within"], got["slowest"]) == (5, 6, 20, 20)
    assert got["by_role"] == [{"name": "Software development", "decisions": 3, "middle": 4, "fastest": 2, "slowest": 6}]
    assert (got["waiting"], got["waiting_late"]) == (2, 1) and got["late"][0]["id"] == "w2" and got["late"][0]["days"] == 67
    empty = insights.reply_times([_app("w", "Dev", "waiting", date_applied="2026-01-01")], [], today=WED)
    assert (empty["decisions"], empty["waiting_late"], empty["late"]) == (0, 0, [])           # nothing to compare with: nobody is "late"


def test_low_score_reasons_are_read_from_the_explanations():
    def scored(i, score, reason):
        return store.new_job(id=str(i), title="T", company="C", match_score=score, match_reason=reason)
    jobs = [scored(1, 10, "This is a sales role, completely unrelated to the candidate's software background."),
            scored(2, 15, "Requires 8-10 years of commercial Java experience; the candidate has no Java."),
            scored(3, 20, "German at an advanced level is a must-have requirement."),
            scored(4, 25, "The profile lacks Kubernetes and Terraform, and there is no mention of Java."),
            scored(5, 5, "A poor match overall."),
            scored(6, 80, "Unrelated words in a high score do not count: lacks nothing."),
            scored(7, 10, ""), store.new_job(id="8", title="T", company="C")]
    got = insights.low_score_reasons(jobs, "Python, Docker", threshold=30)
    assert got["jobs"] == 5 and got["threshold"] == 30
    by = {r["reason"]: (r["jobs"], r["share"]) for r in got["rows"]}
    assert by == {"Technology you have not used": (2, 40), "More experience than you have": (1, 20), "A different field": (1, 20),
                  "Another language required": (1, 20), insights.LOW_OTHER: (1, 20)}
    assert got["rows"][0]["reason"] == "Technology you have not used" and got["rows"][-1]["reason"] == insights.LOW_OTHER
    assert got["technologies"] == [{"skill": "Java", "jobs": 2}]                    # named twice, and not in the CV
    assert insights.low_score_reasons([], "")["rows"] == []


# ── the preparation as a sheet to read before the call ──────────────────────
FULL = {
    "summary": "They will test your Python and how you think about robots.",
    "company": ["BRIGHTLINE builds humanoid robots.", "Engineering is in Prague."], "company_say": "From what I have read, BRIGHTLINE trains robots from human demonstrations.",
    "role": ["Train and evaluate models.", "Build training pipelines."], "role_plain": "You train the models that run on the robot.",
    "fit": ["Python and PyTorch in my reinforcement learning research.", "A FastAPI service deployed with Docker."],
    "gap": "The posting asks for three years; you have internships. Say so and show how fast you learn.",
    "key_question": {"question": "This role asks for 3+ years. Why you?", "answer": "I do not have three years, but I have built and shipped real projects."},
    "pitch": "I recently finished my MSc and built a RAG assistant with FastAPI.",
    "concepts": [{"term": "Imitation learning", "meaning": "Learning a policy from demonstrations."}, {"term": "", "meaning": "dropped"}],
    "reminders": ["Lead with the RL project."],
    "questions": [{"kind": "technical", "question": "How does PPO work?", "why": "Core RL.", "answer": "I used it in my research project."},
                  {"kind": "about you", "question": "Tell me about yourself.", "why": "Opening.", "answer": "I am a junior developer."}],
    "ask_them": ["What does the first month look like?"], "revise": ["Policy gradients"],
}


def test_generate_keeps_the_new_parts_and_checks_what_you_would_say(monkeypatch):
    prompts = []
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or dict(FULL))
    prep = interview_prep.generate({"title": "Machine Learning Engineer", "company": "BRIGHTLINE", "description": "Robots."}, CV)
    assert len(prompts) == 1 and "Do NOT add facts from" in prompts[0]             # company facts only from the posting
    assert prep["company"] == FULL["company"] and prep["pitch"] == FULL["pitch"] and prep["key_question"]["question"].startswith("This role asks")
    assert prep["concepts"] == [{"term": "Imitation learning", "meaning": "Learning a policy from demonstrations."}]
    assert prep["checks"] == [] and prep["reminders"] == ["Lead with the RL project."]


def test_a_pitch_that_claims_too_much_is_sent_back_and_then_flagged(monkeypatch):
    prompts = []
    bad = dict(FULL, pitch="I have three years of production experience with Kubernetes.")
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or dict(bad))
    prep = interview_prep.generate({"title": "ML Engineer", "company": "BRIGHTLINE"}, CV)
    assert len(prompts) == 2 and "the opening pitch" in prompts[1] and "Kubernetes" in prompts[1]
    assert prep["checks"] == ["The opening pitch mentions Kubernetes as experience, which is not in your CV."]


def test_an_older_saved_preparation_without_the_new_parts_still_works(monkeypatch):
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"summary": "s", "questions": FULL["questions"]})
    prep = interview_prep.generate({"title": "T", "company": "C"}, CV)
    assert prep["company"] == [] and prep["pitch"] == "" and prep["key_question"] == {"question": "", "answer": ""}


def test_two_documents_one_for_hr_one_for_a_specialist_in_the_order_asked(tmp_path):
    """Company, you, projects and skills, their questions with answers, your questions: in that order, in both files."""
    import docx

    from jobhelper import prep_doc
    job = {"title": "Machine Learning Engineer", "company": "BRIGHTLINE", "event": {"date": "2026-10-13", "time": "11:30", "kind": "interview", "note": "Google Meet"}}
    prep = dict(FULL, made="2026-10-08", checks=["The opening pitch mentions X."])
    prep["concepts"] = prep["concepts"][:1]
    prep["questions"] = [dict(q, check="Kubernetes" if q["kind"] == "technical" else "") for q in FULL["questions"]]
    prep["projects"] = [{"name": "Game-AI agent", "what": "I built an agent that learns from play. I used PyTorch.", "plain": "I built a program that learns to play a game.",
                         "tools": ["Python", "PyTorch"]}]
    prep["skills"] = [{"skill": "Python", "proof": "I used it in every project in my CV."}]
    read = lambda path: [p.text for p in docx.Document(str(path)).paragraphs if p.text.strip()]   # noqa: E731
    hr = read(prep_doc.build(job, prep, tmp_path / "out" / "hr.docx", contact="Petra", audience="hr"))
    tech = read(prep_doc.build(job, prep, tmp_path / "out" / "tech.docx", contact="Petra", audience="technical"))
    assert hr[0] == "BRIGHTLINE – Interview preparation: HR" and tech[0] == "BRIGHTLINE – Interview preparation: Technical"
    assert hr[1] == "Machine Learning Engineer" and hr[2] == "Tuesday 13 October 2026, 11:30 · Google Meet · with Petra"
    order_hr = ["1. About the company", "2. About you", "3. Your projects and skills", "4. Questions HR may ask you", "5. Questions you can ask HR",
                "Before the call"]
    order_tech = ["1. About the company", "2. About you", "3. Your projects and skills", "4. Questions a specialist may ask you",
                  "Concepts to explain", "5. Questions you can ask them", "Before the call"]
    assert [hr.index(h) for h in order_hr] == sorted(hr.index(h) for h in order_hr)
    assert [tech.index(h) for h in order_tech] == sorted(tech.index(h) for h in order_tech)
    # each file has its own questions, and its own way of telling the projects
    hr_qs = {q["question"] for q in prep["questions"] if prep_doc._audience(q) == "hr"}
    tech_qs = {q["question"] for q in prep["questions"] if prep_doc._audience(q) == "technical"}
    assert hr_qs <= set(hr) and not (tech_qs - hr_qs) & set(hr)
    assert tech_qs <= set(tech) and not (hr_qs - tech_qs) & set(tech)
    assert "“I built a program that learns to play a game.”" in hr and "Tools: Python, PyTorch" not in hr
    assert "“I built an agent that learns from play. I used PyTorch.”" in tech and "Tools: Python, PyTorch" in tech
    prep["projects"][0]["tech"] = [{"tool": "PyTorch", "about": "A library for neural networks.", "use": "It trained the policy.", "why": "It is common in research."}]
    prep["projects"][0]["result"], prep["projects"][0]["ask"] = "Shown at two conferences.", ["Why imitation learning first?"]
    prep["skills"][0]["detail"] = "I can explain how I tested it."
    tech = read(prep_doc.build(job, prep, tmp_path / "out" / "tech2.docx", audience="technical"))
    hr = read(prep_doc.build(job, prep, tmp_path / "out" / "hr2.docx", audience="hr"))
    line = "PyTorch. What it is: A library for neural networks. In this project: It trained the policy. Why: It is common in research."
    assert line in tech and line not in hr
    assert "Result: Shown at two conferences." in tech and "Why imitation learning first?" in tech and "Why imitation learning first?" not in hr
    assert "Python: I used it in every project in my CV. I can explain how I tested it." in tech
    assert "Python: I used it in every project in my CV." in hr
    assert "Python: I used it in every project in my CV." in hr
    assert "Imitation learning: Learning a policy from demonstrations." in tech and "Concepts to explain" not in hr
    assert prep_doc.COMPANY_NOTE in hr and prep_doc.COMPANY_NOTE in tech
    assert "Check this answer before you say it. It mentions Kubernetes. Your CV does not show this." in tech
    assert "Check: The opening pitch mentions X." in hr and all("—" not in "\n".join(doc) for doc in (hr, tech))
    # HR always gets the question about the next steps, whatever was drafted
    assert hr.index(interview_prep.NEXT_STEPS_QUESTION) > hr.index("5. Questions you can ask HR")
    # an older preparation (no projects, no skills, one mixed list of questions to ask) still makes both files
    old = {"questions": FULL["questions"], "summary": "s", "ask_them": ["What does a normal day look like?"]}
    for audience in prep_doc.AUDIENCES:
        lines = read(prep_doc.build({"title": "T", "company": "C"}, old, tmp_path / f"old_{audience}.docx", audience=audience))
        assert prep_doc.NEW_PART in lines and "What does a normal day look like?" in lines and "Before the call" in lines


def test_the_text_the_app_writes_itself_follows_simplified_technical_english():
    """ASD-STE100: at most 20 words in a sentence of instruction, 25 in a description; no em dash."""
    import re

    from jobhelper import prep_doc
    texts = [prep_doc.COMPANY_NOTE, prep_doc.NEW_PART, *prep_doc.NOTES.values(), *prep_doc.FIXED_REMINDERS, interview_prep.NEXT_STEPS_QUESTION]
    for text in texts:
        for sentence in re.split(r"(?<=[.!?])\s+", text):
            assert len(sentence.split()) <= 20, sentence
        assert "—" not in text
    assert "ASD-STE100" in interview_prep.SYSTEM and "20 words" in interview_prep.SYSTEM


def test_sheet_file_name_and_date_line():
    from jobhelper import prep_doc
    assert prep_doc.file_name("AlexNovak", "Actum Digital, a.s.") == "AlexNovak_ActumDigital_Prep_HR.docx"
    assert prep_doc.file_name("AlexNovak", "Actum Digital, a.s.", "technical") == "AlexNovak_ActumDigital_Prep_Technical.docx"
    assert prep_doc.file_name("", "") == "Candidate_Company_Prep_HR.docx"
    assert prep_doc.when_text({"date": "2026-10-13", "time": ""}) == "Tuesday 13 October 2026" and prep_doc.when_text({}) == "" and prep_doc.when_text(None) == ""
    assert prep_doc.when_text({"date": None}) == ""


def test_api_writes_and_opens_both_documents(isolate_data, monkeypatch):
    config.ensure_dirs()
    opened, shown = [], []
    monkeypatch.setattr(ui_api.Api, "_open", lambda self, target: opened.append(target))
    monkeypatch.setattr(ui_api.Api, "_reveal", lambda self, target: shown.append(target))
    store.save_jobs([store.new_job(id="a", title="ML Engineer", company="BRIGHTLINE", status="interview"), store.new_job(id="b", title="T", company="C", status="waiting")])
    api = ui_api.Api()
    assert "Prepare the interview first" in api.prep_document("b")["msg"] and api.prep_document("zzz")["ok"] is False
    api._change("a", lambda jobs, j: j.__setitem__("interview_prep", dict(FULL, made="2026-10-08", checks=[])), render=False)
    got = api.prep_document("a")
    folder = config.DATA_DIR / "interview_prep"
    hr, tech = folder / "Candidate_Brightline_Prep_HR.docx", folder / "Candidate_Brightline_Prep_Technical.docx"
    assert got == {"ok": True, "file": hr.name, "files": [hr.name, tech.name]}
    assert opened == [hr, tech] and hr.is_file() and tech.is_file()
    assert api.prep_document("a", reveal=True)["ok"] and shown == [hr] and len(opened) == 2


def test_hr_questions_always_ask_about_the_next_steps():
    assert interview_prep.with_next_steps(["Who would I work with?"])[0] == interview_prep.NEXT_STEPS_QUESTION
    asked = ["What does the hiring process look like from here?", "Who would I work with?"]
    assert interview_prep.with_next_steps(asked) == asked                          # already asked: not twice
    assert interview_prep.with_next_steps([]) == [interview_prep.NEXT_STEPS_QUESTION]


# ── work experience, told for the role ──────────────────────────────────────
WORK = [{"job": "Software Developer Intern, Acme, Feb-Aug 2026", "story": "I worked on the import service in Python. I added its tests.",
         "gained": ["Working in a team's codebase"], "if_true": ["Did you review other people's code?"], "thin": "The CV gives no results for this job."},
        {"job": "", "story": "dropped: no job named"}, {"job": "Tutor", "story": ""}, "junk"]


def test_experience_is_kept_per_job_with_reminders_not_claims(monkeypatch):
    prompts = []
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or dict(FULL, experience=WORK))
    prep = interview_prep.generate({"title": "ML Engineer", "company": "BRIGHTLINE"}, CV)
    assert prep["experience"] == [dict(WORK[0], plain="", typical=[])]
    assert "one entry for EVERY job and internship in the CV" in prompts[0] and "do not add duties, tools, team sizes or results" in prompts[0]
    assert "(nothing written yet)" in prompts[0]                                  # no notes: said so, not left blank


def test_what_you_wrote_about_your_work_counts_like_the_cv(monkeypatch):
    story = dict(WORK[0], story="I ran our services on Kubernetes and wrote the deployment scripts.")
    prompts = []
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or dict(FULL, experience=[story]))
    without = interview_prep.generate({"title": "T", "company": "C"}, CV)
    assert len(prompts) == 2 and "the story about Software Developer Intern" in prompts[1]      # not in the CV: sent back, then flagged
    assert without["checks"] == ["The story about Software Developer Intern, Acme, Feb-Aug 2026 mentions Kubernetes as experience, which is not in your CV."]
    prompts.clear()
    notes = "At Acme I deployed our services to Kubernetes and wrote the deployment scripts."
    with_notes = interview_prep.generate({"title": "T", "company": "C"}, CV, notes=notes)
    assert len(prompts) == 1 and notes in prompts[0] and with_notes["checks"] == []            # you said it is true: accepted


def test_the_sheet_has_the_work_experience_section(tmp_path):
    import docx

    from jobhelper import prep_doc
    prep = dict(FULL, experience=[WORK[0]], made="2026-10-08", checks=[])
    read = lambda audience: [p.text for p in docx.Document(str(prep_doc.build({"title": "ML Engineer", "company": "BRIGHTLINE"}, prep,      # noqa: E731
                                                                                    tmp_path / f"{audience}.docx", audience=audience))).paragraphs if p.text.strip()]
    lines = read("technical")
    start = lines.index("Your work experience, told for this role")
    assert lines.index("2. About you") < lines.index("What to lead with") < start < lines.index("The honest gap. Say it openly; do not hide it.")
    block = lines[start:start + 12]
    assert "Software Developer Intern, Acme, Feb-Aug 2026" in block and "“I worked on the import service in Python. I added its tests.”" in block
    assert "Ask yourself these questions. Add a point only if it is true:" in block and "Did you review other people's code?" in block
    assert any(x.startswith("Your CV says little here: The CV gives no results for this job.") and "More about your work" in x for x in block)
    hr = read("hr")                                       # HR gets the story, without the specialist's working notes
    assert "Your work experience, in plain words" in hr and "Did you review other people's code?" not in hr


def test_experience_notes_come_from_options_and_reach_both_drafts(isolate_data, monkeypatch):
    from jobhelper import form_answers
    config.ensure_dirs()
    api = ui_api.Api()
    assert config.experience_notes() == "" and api.options()["about_experience"] == ""
    refused = api.save_options({"about_experience": "  At Acme I wrote the import service.  " + "x" * 5000})
    assert refused["ok"] is False and "Shorten it" in refused["msg"] and config.experience_notes() == ""       # too long: refused, never cut silently
    assert api.save_options({"about_experience": "  At Acme I wrote the import service.  " + "x" * 3000})["ok"] is True
    assert config.experience_notes().startswith("At Acme I wrote the import service.") and len(config.experience_notes()) <= config.MAX_EXPERIENCE_NOTES
    api.save_options({"about_experience": "At Acme I wrote the import service."})
    monkeypatch.setattr(ui_api.Api, "_cv_text", lambda self: CV)
    store.save_jobs([store.new_job(id="a", title="Dev", company="Acme", status="waiting")])
    seen = {}
    monkeypatch.setattr(interview_prep, "generate", lambda j, cv, emails, notes="": seen.update(prep=notes) or dict(FULL, made="x", checks=[], experience=[]))
    monkeypatch.setattr(form_answers, "generate", lambda j, cv, qs, notes="": seen.update(form=notes) or [{"question": "Q?", "answer": "A.", "check": ""}])
    assert api.make_interview_prep("a")["ok"] and api.make_form_answers("a", "Why us?")["ok"]
    assert seen == {"prep": "At Acme I wrote the import service.", "form": "At Acme I wrote the import service."}


# ── two audiences: HR in plain words, the specialist in detail ────────────────
TWO = dict(FULL, questions=[
    {"audience": "technical", "kind": "technical", "question": "How does PPO work?", "why": "", "answer": "It clips the policy update."},
    {"audience": "hr", "kind": "about you", "question": "Tell me about your last project.", "why": "", "answer": "I taught a program to play a game by copying good players."},
    {"audience": "technical", "kind": "your projects", "question": "Tell me about your last project.", "why": "", "answer": "A route planner for delivery vans in Python."},
    {"audience": "nobody", "kind": "practical", "question": "When can you start?", "why": "", "answer": "In two weeks."},
    {"kind": "technical", "question": "What is a policy?", "why": "", "answer": "A mapping from states to actions."}],
    experience=[dict(WORK[0], plain="I helped build a tool that loads data, and I checked that it works.",
                     typical=["An intern in this kind of role usually fixes small bugs.", "Usually writes tests for their own changes."])])


def test_questions_are_split_by_who_asks(monkeypatch):
    prompts = []
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or dict(TWO))
    prep = interview_prep.generate({"title": "ML Engineer", "company": "BRIGHTLINE"}, CV)
    assert [(q["audience"], q["question"]) for q in prep["questions"]] == [
        ("hr", "Tell me about your last project."), ("hr", "When can you start?"),                 # an unknown audience: by its kind
        ("technical", "How does PPO work?"), ("technical", "What is a policy?"), ("technical", "Tell me about your last project.")]
    assert "They usually do not know the specialist side of the work" in prompts[0] and "about 7 questions a recruiter" in prompts[0]
    assert "it is never written as something this candidate did" in prompts[0]
    assert prep["experience"][0]["typical"][0].startswith("An intern in this kind of role usually")


def test_the_sheet_has_both_sets_and_marks_usual_tasks_as_general(tmp_path):
    import docx

    from jobhelper import prep_doc
    prep = dict(interview_prep._clean(TWO), made="2026-10-08", checks=[])
    read = lambda audience: [p.text for p in docx.Document(str(prep_doc.build({"title": "ML Engineer", "company": "BRIGHTLINE"}, prep,      # noqa: E731
                                                                                    tmp_path / f"{audience}.docx", audience=audience))).paragraphs if p.text.strip()]
    hr, tech = read("hr"), read("technical")
    assert "I taught a program to play a game by copying good players." in hr                     # the plain answer is in the HR file
    assert "A route planner for delivery vans in Python." in tech and "It clips the policy update." in tech
    assert "A route planner for delivery vans in Python." not in hr
    assert hr[hr.index("4. Questions HR may ask you") + 1] == prep_doc.NOTES["hr"]
    assert "“I helped build a tool that loads data, and I checked that it works.”" in hr               # the plain story, for HR
    general = tech.index("What people in this type of job usually do. This is general information about the role, not your record:")
    assert tech[general + 1] == "An intern in this kind of role usually fixes small bugs."
    assert general < tech.index("Ask yourself these questions. Add a point only if it is true:")


def test_a_plain_story_is_held_to_the_cv_too(monkeypatch):
    bad = dict(TWO, experience=[dict(TWO["experience"][0], plain="I ran the company's Kubernetes clusters.")])
    prompts = []
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or dict(bad))
    interview_prep.generate({"title": "T", "company": "C"}, CV)
    assert len(prompts) == 2 and "the plain story about Software Developer Intern" in prompts[1]


# ── what you can ask them, per listener ─────────────────────────────────────
ASK = dict(TWO, ask_hr=["What are the next steps after this call?", "How big is the team I would join?"],
           ask_technical=["What would I own in the first three months?"], ask_them=[])


def test_questions_to_ask_are_kept_for_each_listener(monkeypatch):
    prompts = []
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or dict(ASK))
    prep = interview_prep.generate({"title": "ML Engineer", "company": "BRIGHTLINE"}, CV)
    assert prep["ask_hr"] == ASK["ask_hr"] and prep["ask_technical"] == ASK["ask_technical"]
    assert '"ask_hr"' in prompts[0] and '"ask_technical"' in prompts[0]
    text = interview_prep.as_text({"title": "ML Engineer", "company": "BRIGHTLINE"}, prep)
    assert "QUESTIONS TO ASK HR OR THE RECRUITER" in text and "QUESTIONS TO ASK THE PEOPLE YOU WOULD WORK WITH OR THE MANAGER" in text


def test_the_sheet_puts_your_questions_under_each_set(tmp_path):
    import docx

    from jobhelper import prep_doc
    prep = dict(interview_prep._clean(ASK), made="2026-10-08", checks=[])
    read = lambda p, name, audience: [x.text for x in docx.Document(str(prep_doc.build({"title": "T", "company": "BRIGHTLINE"}, p,      # noqa: E731
                                                                                         tmp_path / name, audience=audience))).paragraphs if x.text.strip()]
    hr, tech = read(prep, "hr.docx", "hr"), read(prep, "tech.docx", "technical")
    assert "What are the next steps after this call?" in hr[hr.index("5. Questions you can ask HR"):]
    assert "What would I own in the first three months?" in tech[tech.index("5. Questions you can ask them"):]
    assert "What would I own in the first three months?" not in hr and "What are the next steps after this call?" not in tech
    assert interview_prep.NEXT_STEPS_QUESTION not in hr                         # the drafted list already asks about the next steps
    old = dict(interview_prep._clean(dict(TWO, ask_them=["An older mixed question?"])), made="x", checks=[])
    old["ask_hr"], old["ask_technical"] = [], []
    for audience in ("hr", "technical"):                                        # an older saved preparation still shows its list
        assert "An older mixed question?" in read(old, f"old_{audience}.docx", audience)


# ── questions to ask must fit the person asking ─────────────────────────────
@pytest.mark.parametrize("question, kept", [
    ("Who would this role report to, and how large is the team in Prague?", True),
    ("What does the first project usually look like for a new engineer?", True),
    ("How many years has the team been working on this product?", True),                       # years, but not about the candidate
    ("What does the relocation and visa sponsorship process look like in practice?", False),   # the right to work is already stated
    ("Do you help with the work permit?", False),
    ("Given the 3+ years requirement, what level of autonomy would you expect?", False),       # points at the candidate's weak spot
    ("Considering my limited experience with robots, how would you onboard me?", False),
    ("Since I haven't used Isaac Sim, is training provided?", False),
    ("How flexible is the 3+ years requirement?", False),
])
def test_questions_the_candidate_should_not_ask_are_removed(question, kept):
    facts = {"right_to_work": "full right to work in the Czech Republic", "location": "Prague"}
    assert (interview_prep.fit_questions([question], facts) == [question]) is kept


def test_a_visa_question_stays_when_nothing_is_known_about_the_right_to_work():
    q = "Do you sponsor a work permit for this role?"
    assert interview_prep.fit_questions([q], {"right_to_work": "", "location": ""}) == [q]


def test_generate_applies_the_check_to_both_lists(isolate_data, monkeypatch):
    config.ensure_dirs()
    ui_api.Api().save_options({"where": "Prague", "about_right_to_work": "full right to work in the Czech Republic"})
    reply = dict(ASK, ask_hr=["What are the next steps?", "What does visa sponsorship look like?"],
                 ask_technical=["Given the 3+ years requirement, how much autonomy would I have?", "What would I own first?"])
    prompts = []
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or dict(reply))
    prep = interview_prep.generate({"title": "ML Engineer", "company": "BRIGHTLINE"}, CV)
    assert prep["ask_hr"] == ["What are the next steps?"] and prep["ask_technical"] == ["What would I own first?"]
    assert "never one that points at the candidate's own weak spot" in prompts[0]


# ── editing the saved preparation by hand ───────────────────────────────────
def _saved():
    return dict(interview_prep._clean(dict(ASK, experience=TWO["experience"])), made="2026-10-08", checks=[])


def test_parts_follow_the_order_of_the_sheet():
    got = interview_prep.parts(_saved())
    keys = [p["key"] for p in got]
    assert keys[:6] == ["summary", "company", "company_say", "role", "role_plain", "fit"]
    assert keys.index("experience.0.story") < keys.index("gap") < keys.index("key_question.question") < keys.index("pitch") \
        < keys.index("questions.0.answer") < keys.index("ask_hr") < keys.index("reminders")
    by = {p["key"]: p for p in got}
    assert by["fit"]["kind"] == "list" and by["fit"]["text"].count("\n") == 1
    assert by["questions.0.answer"]["label"] == "HR asks: Tell me about your last project."
    assert by["experience.0.plain"]["label"].startswith("The same for HR, in plain words: Software Developer Intern")
    assert interview_prep.parts({"questions": FULL["questions"]})[0]["key"] == "questions.0.answer"      # an older, smaller preparation


def test_edits_are_put_in_and_everything_else_is_kept():
    prep = _saved()
    new, problem = interview_prep.apply_parts(prep, [
        {"key": "pitch", "text": "  I finished my MSc and I build things that work — carefully.  "},
        {"key": "fit", "text": "• Python in my research project\n\n- A FastAPI service\n"},
        {"key": "questions.0.answer", "text": "I taught a program to play a game."},
        {"key": "ask_hr", "text": "What are the next steps?"},
        {"key": "made_up.key", "text": "ignored"}, "junk"], CV)
    assert problem == "" and new["pitch"] == "I finished my MSc and I build things that work – carefully."      # em dash cleaned
    assert new["fit"] == ["Python in my research project", "A FastAPI service"] and new["ask_hr"] == ["What are the next steps?"]
    assert new["questions"][0]["answer"] == "I taught a program to play a game." and new["questions"][1]["answer"] == prep["questions"][1]["answer"]
    assert new["gap"] == prep["gap"] and new["edited"] and "edited" not in prep                                 # the old one is untouched


@pytest.mark.parametrize("edit, problem", [
    ({"key": "questions.0.answer", "text": "   "}, "is empty. Write it, or leave the draft as it was."),
    ({"key": "experience.0.story", "text": ""}, "is empty"),
    ({"key": "key_question.answer", "text": ""}, "“Your answer to the key question” is empty"),
    ({"key": "pitch", "text": "x" * 3001}, "is 3001 characters long; the most is 3000. Shorten it."),
    ({"key": "fit", "text": "\n".join(f"point {i}" for i in range(13))}, "has 13 lines; the most is 12. Remove some."),
    ({"key": "ask_hr", "text": "q" * 401}, "is 401 characters long; the most is 400. Shorten it."),
])
def test_a_wrong_edit_is_refused_with_the_reason_and_nothing_changes(edit, problem):
    prep = _saved()
    new, said = interview_prep.apply_parts(prep, [{"key": "gap", "text": "A good change that must not be half-saved."}, edit], CV)
    assert problem in said and new is prep and prep["gap"] == FULL["gap"]


def test_an_optional_part_may_be_emptied_and_a_claim_you_add_is_flagged():
    prep = _saved()
    new, problem = interview_prep.apply_parts(prep, [{"key": "company_say", "text": ""}, {"key": "revise", "text": ""},
                                                    {"key": "pitch", "text": "I ran production Kubernetes clusters for two years."}], CV)
    assert problem == "" and new["company_say"] == "" and new["revise"] == []
    assert new["checks"] == ["The opening pitch mentions Kubernetes as experience, which is not in your CV."]   # saved, but marked for you


def test_api_edits_the_saved_preparation(isolate_data, monkeypatch):
    config.ensure_dirs()
    monkeypatch.setattr(ui_api.Api, "_cv_text", lambda self: CV)
    store.save_jobs([store.new_job(id="a", title="ML Engineer", company="BRIGHTLINE", status="interview"), store.new_job(id="b", title="T", company="C", status="waiting")])
    api = ui_api.Api()
    assert "Prepare the interview first" in api.prep_parts("b")["msg"] and api.save_prep_parts("b", [])["ok"] is False
    api._change("a", lambda jobs, j: j.update(interview_prep=_saved()), render=False)
    got = api.prep_parts("a")
    assert got["ok"] and got["company"] == "BRIGHTLINE" and got["edited"] == "" and api.job_detail("a")["prep_edited"] is False
    refused = api.save_prep_parts("a", [{"key": "questions.0.answer", "text": ""}])
    assert refused["ok"] is False and "is empty" in refused["msg"]
    assert api.save_prep_parts("a", [{"key": "pitch", "text": "My own pitch."}]) == {"ok": True, "to_check": 0}
    assert ui_api.Api().interview_prep("a")["pitch"] == "My own pitch." and api.job_detail("a")["prep_edited"] is True
    assert api.prep_parts("a")["edited"] != ""


def test_the_preparation_is_not_worded_for_one_profession():
    """It spoke of "an engineer" and "the stack", whoever the candidate was."""
    prompt = interview_prep.PROMPT.lower()
    for word in ("engineer", "stack", "the technology", "ml research"):
        assert word not in prompt, word
    assert "whatever the profession" in prompt and "someone who does this work" in prompt
    assert "engineer" not in " ".join(interview_prep.AUDIENCE_TITLES.values()).lower()
    source = Path(interview_prep.__file__).read_text(encoding="utf-8").lower()
    assert "engineer" not in source


# ── found by reading the document modules line by line (10 October 2026) ────
def test_one_sentence_where_a_list_was_asked_is_one_item_not_one_per_letter():
    from jobhelper import interview_prep
    prep = interview_prep._clean({
        "questions": [{"question": "Why us?", "answer": "Because of your SQL reports.", "kind": "about you", "audience": "hr"}],
        "company": "They build planning tools.", "fit": ["SQL reports", 3, None, {"x": 1}], "role": {"not": "a list"}, "revise": None,
        "experience": [{"job": "Analyst, Northwind", "story": "I built reports.", "gained": "Care with numbers.", "typical": 7, "if_true": ["Did you test it?"]}],
        "concepts": "none", "ask_hr": "How does the team work?"})
    assert prep["company"] == ["They build planning tools."] and prep["ask_hr"] == ["How does the team work?"]
    assert prep["fit"] == ["SQL reports", "3"] and prep["role"] == [] and prep["revise"] == [] and prep["concepts"] == []
    assert prep["experience"][0]["gained"] == ["Care with numbers."] and prep["experience"][0]["typical"] == [] and prep["experience"][0]["if_true"] == ["Did you test it?"]
    assert interview_prep._clean({"questions": "Why us?", "experience": "long ago"})["questions"] == []
    old = {"questions": [{"audience": "panel", "kind": "about you", "question": "Q?", "answer": "A."}]}
    assert "PANEL" in interview_prep.as_text({"title": "Dev", "company": "Acme"}, old), "a saved preparation with an audience nobody knows still prints"


def test_a_cv_or_template_paragraph_that_cannot_be_rewritten_is_said_in_words(tmp_path):
    """Text inside a link or a form control has no plain run to write into: the builders used to stop with "list index out of range"."""
    from lxml import etree
    from jobhelper import cover_letter, cv_tailor
    W = cv_tailor.W
    para = etree.fromstring(f'<w:p xmlns:w="{W}"><w:hyperlink><w:r><w:t>text in a link</w:t></w:r></w:hyperlink></w:p>')
    with pytest.raises(ValueError, match="Retype it as plain text"):
        cover_letter._set_para(para, "new")
    assert cv_tailor._is_heading("WORK EXPERIENCE") and not cv_tailor._is_heading("Skills: Python") and not cv_tailor._is_heading("ab")


def test_projects_and_skills_are_asked_for_kept_held_to_the_cv_and_editable(monkeypatch):
    prompts = []
    reply = dict(FULL, ask_hr=["Who would I work with?"],
                 projects=[{"name": "Game-AI agent", "what": "I built an agent that learns from play.", "plain": "I built a game program.", "tools": ["PyTorch"],
                            "tech": [{"tool": "PyTorch", "about": "A library for neural networks.", "use": "It trained the policy.", "why": "It is the usual choice in research."},
                                     {"tool": "", "use": "no tool: dropped"}],
                            "result": "Shown at two conferences.", "ask": ["Why imitation learning first?"]},
                           {"name": "", "what": "no name: dropped"}, "not a project"],
                 skills=[{"skill": "Python", "proof": "Used in every project.", "detail": "I wrote the training loop."}, {"skill": "Go", "proof": ""}])
    monkeypatch.setattr(llm, "ask_json", lambda prompt, **k: prompts.append(prompt) or reply)
    prep = interview_prep.generate({"title": "ML Engineer", "company": "BRIGHTLINE"}, CV)
    assert '"projects"' in prompts[0] and '"skills"' in prompts[0] and "next steps" in prompts[0] and "ASD-STE100" in interview_prep.SYSTEM
    assert prep["projects"] == [{"name": "Game-AI agent", "what": "I built an agent that learns from play.", "plain": "I built a game program.", "tools": ["PyTorch"],
                                 "tech": [{"tool": "PyTorch", "about": "A library for neural networks.", "use": "It trained the policy.",
                                           "why": "It is the usual choice in research."}],
                                 "result": "Shown at two conferences.", "ask": ["Why imitation learning first?"]}]
    assert prep["skills"] == [{"skill": "Python", "proof": "Used in every project.", "detail": "I wrote the training loop."}]
    assert '"about"' in prompts[0] and '"why"' in prompts[0] and "never as a detail or a" in prompts[0]
    assert "technical reason it suits that task" in prompts[0] and "not about this application" in prompts[0]
    assert prep["ask_hr"][0] == interview_prep.NEXT_STEPS_QUESTION and "Who would I work with?" in prep["ask_hr"]
    labels = {p["key"]: p["label"] for p in interview_prep.parts(prep)}
    assert labels["projects.0.what"] == "Your project: Game-AI agent" and labels["skills.0.proof"] == "Where your CV shows it: Python"
    new, problem = interview_prep.apply_parts(prep, [{"key": "projects.0.what", "text": "I built and tested an agent."}], CV)
    assert problem == "" and new["projects"][0]["what"] == "I built and tested an agent."
    _, problem = interview_prep.apply_parts(prep, [{"key": "skills.0.proof", "text": " "}], CV)
    assert "is empty" in problem
    text = interview_prep.as_text({"title": "ML Engineer", "company": "BRIGHTLINE"}, prep)
    assert "YOUR PROJECTS" in text and "Tools: PyTorch" in text and "- Python: Used in every project." in text
    assert "- PyTorch: A library for neural networks. It trained the policy. It is the usual choice in research." in text
