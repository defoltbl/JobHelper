"""Data access behind the desktop app (jobhelper.ui_api.Api)."""
import json
from types import SimpleNamespace

import pytest

from jobhelper import api_core, api_documents, api_insights, api_system, config, store, ui_api, vault


@pytest.fixture
def api(isolate_data):
    return ui_api.Api()


def _seed(*jobs):
    config.ensure_dirs()
    store.save_jobs(list(jobs))


def test_options_defaults_then_saved_values_win(api):
    opts = api.options()
    assert opts["where"] == "" and opts["source"] == "czech" and opts["queries"] == []       # no city is assumed
    api.save_options({"where": "Brno", "threshold": 70, "queries": ["qa engineer"]})
    opts = api.options()
    assert (opts["where"], opts["threshold"], opts["queries"]) == ("Brno", 70, ["qa engineer"])
    assert opts["per_query"] == 15                      # untouched keys keep their defaults


def test_options_seed_keywords_from_profile(api):
    config.ensure_dirs()
    config.PROFILE_JSON.write_text(json.dumps({"search_queries": ["junior python developer"]}), encoding="utf-8")
    assert api.options()["queries"] == ["junior python developer"]


def test_companies_roundtrip_drops_incomplete_rows(api):
    result = api.save_companies([
        {"name": "Proton", "ats": "greenhouse", "token": "proton"},
        {"name": "", "ats": "lever", "token": "acme"},           # name falls back to the token
        {"name": "Broken", "ats": "", "token": "x"},
        {"name": "Broken2", "ats": "lever", "token": ""},
    ])
    assert result["ok"] is False and result["count"] == 2 and result["msg"].startswith("Not saved: Broken, Broken2."), "each one named"
    assert api.companies() == [{"name": "Proton", "ats": "greenhouse", "token": "proton"},
                               {"name": "acme", "ats": "lever", "token": "acme"}]


def test_companies_survive_a_corrupt_file(api):
    config.ensure_dirs()
    (config.DATA_DIR / "companies.json").write_text("{oops", encoding="utf-8")
    assert api.companies() == []


def test_counts_split_leads_and_applications(api):
    _seed(store.new_job(id="a", status="found"), store.new_job(id="b", status="letter_ready"),
          store.new_job(id="c", status="waiting"), store.new_job(id="d", status="rejected"))
    counts = api.counts()
    assert (counts["leads"], counts["applications"]) == (2, 2)


def test_leads_sorted_by_score_and_respect_the_floor(api, monkeypatch):
    _seed(store.new_job(id="low", title="Low", company="A", match_score=30),
          store.new_job(id="high", title="High", company="B", match_score=90),
          store.new_job(id="new", title="Unscored", company="C"),
          store.new_job(id="app", title="Applied", company="D", match_score=95, status="waiting"))
    assert [j["id"] for j in api.leads()] == ["high", "low", "new"]
    monkeypatch.setattr(config, "MIN_LEAD_SCORE", 50)
    assert [j["id"] for j in api.leads()] == ["high", "new"]     # unscored leads are never hidden


def test_apply_job_moves_lead_to_applications(api):
    _seed(store.new_job(id="x", title="Dev", company="Acme", status="letter_ready"))
    assert api.apply_job("x") == {"ok": True}
    job = store.find(store.load_jobs(), "x")
    assert job["status"] == "waiting" and job["date_applied"]
    assert api.counts()["leads"] == 0 and api.counts()["applications"] == 1
    assert config.EXCEL_PATH.exists() and config.NEW_JOBS_EXCEL.exists()


def test_apply_job_unknown_id(api):
    _seed()
    assert api.apply_job("nope") == {"ok": False}


def test_open_doc_prefers_docx_and_reports_missing(api, monkeypatch):
    opened = []
    monkeypatch.setattr(api_core.Desktop, "open", lambda self, path: opened.append(path))
    config.ensure_dirs()
    (config.LETTERS_DIR / "letter.md").write_text("x", encoding="utf-8")
    (config.LETTERS_DIR / "letter.docx").write_bytes(b"x")
    assert api.open_doc("letter", "letter.md") == {"ok": True}
    assert opened == [config.LETTERS_DIR / "letter.docx"]
    assert api.open_doc("letter", "missing.md") == {"ok": False}
    assert api.open_doc("unknown-kind", "letter.md") == {"ok": False}


def test_open_url_needs_a_link(api, monkeypatch):
    opened = []
    monkeypatch.setattr(api_core.Desktop, "open", lambda self, path: opened.append(path))
    _seed(store.new_job(id="with", url="https://example.com/job"), store.new_job(id="without"))
    assert api.open_url("with")["ok"] and opened == ["https://example.com/job"]
    assert api.open_url("without") == {"ok": False}


def _fake_cli(monkeypatch, stdout="", error=None):
    def run():
        if error:
            raise error
        return SimpleNamespace(stdout=stdout, stderr="", returncode=0)
    monkeypatch.setattr(api_system, "_claude_auth_status", run)


def test_connections_reads_account_and_plan(api, monkeypatch):
    _fake_cli(monkeypatch, json.dumps({"loggedIn": True, "email": "me@example.com", "subscriptionType": "pro"}))
    conn = api.connections()
    assert (conn["claude"], conn["plan"], conn["claude_email"]) == (True, "pro", "me@example.com")
    assert conn["cv"] == ""


def test_connections_logged_out_hides_account(api, monkeypatch):
    _fake_cli(monkeypatch, json.dumps({"loggedIn": False, "email": "old@example.com", "subscriptionType": "pro"}))
    conn = api.connections()
    assert (conn["claude"], conn["plan"], conn["claude_email"]) == (False, "", "")


def test_connections_when_cli_is_missing(api, monkeypatch):
    _fake_cli(monkeypatch, error=FileNotFoundError("claude"))
    assert api.connections()["claude"] is False


def test_connections_reports_master_cv(api, monkeypatch):
    _fake_cli(monkeypatch, "not json")
    config.ensure_dirs()
    (config.CV_DIR / "my_cv.txt").write_text("Alex, junior developer", encoding="utf-8")
    assert api.connections()["cv"] == "my_cv.txt"


# ── job detail, follow-ups, PDF ─────────────────────────────────────────────
from datetime import date as _date, timedelta as _timedelta  # noqa: E402

_OLD_DAY = _date.today() - _timedelta(days=20)          # quiet long enough to be due, not long enough to be stale
OLD = _OLD_DAY.isoformat()


def test_job_detail_compares_keywords_with_the_master_cv(api, monkeypatch):
    monkeypatch.setattr(api_core.CvSource, "text", lambda self: "Python, Docker, FastAPI. English B2.")
    _seed(store.new_job(id="x", title="Backend Developer", company="Acme", location="Prague", url="https://e.x/j",
                        description="Python, FastAPI, Kubernetes and Kafka. English.", match_score=72,
                        match_tier="strong", match_reason="Good Python fit.", status="letter_ready"))
    d = api.job_detail("x")
    assert (d["title"], d["company"], d["score"], d["tier"], d["reason"]) == \
        ("Backend Developer", "Acme", 72, "strong", "Good Python fit.")
    assert d["matched"] == ["Python", "FastAPI", "English"] and d["missing"] == ["Kubernetes", "Kafka"]
    assert d["is_lead"] is True and d["followup"]["due"] is False
    assert d["cv"] == "" and d["cv_pdf"] is False


def test_job_detail_without_description_or_unknown_id(api):
    _seed(store.new_job(id="x", title="Dev", company="Acme"))
    d = api.job_detail("x")
    assert d["matched"] == [] and d["missing"] == [] and d["description"] == "" and d["score"] is None
    assert api.job_detail("missing") == {}


def test_job_detail_reports_existing_pdf_copies(api):
    config.ensure_dirs()
    (config.TAILORED_CV_DIR / "cv.docx").write_bytes(b"x")
    (config.TAILORED_CV_DIR / "cv.pdf").write_bytes(b"x")
    (config.LETTERS_DIR / "letter.docx").write_bytes(b"x")
    _seed(store.new_job(id="x", title="Dev", company="Acme", tailored_cv_file="cv.docx", cover_letter_file="letter.docx"))
    d = api.job_detail("x")
    assert (d["cv"], d["cv_pdf"], d["letter"], d["letter_pdf"]) == ("cv.docx", True, "letter.docx", False)


def test_open_doc_pdf_only_opens_an_existing_pdf(api, monkeypatch):
    opened = []
    monkeypatch.setattr(api_core.Desktop, "open", lambda self, path: opened.append(path.name))
    config.ensure_dirs()
    (config.TAILORED_CV_DIR / "cv.docx").write_bytes(b"x")
    assert api.open_doc("cv_pdf", "cv.docx") == {"ok": False}
    (config.TAILORED_CV_DIR / "cv.pdf").write_bytes(b"x")
    assert api.open_doc("cv_pdf", "cv.docx") == {"ok": True}
    assert api.open_doc("cv", "cv.docx") == {"ok": True}
    assert opened == ["cv.pdf", "cv.docx"]


def test_followups_are_counted_and_listed_first(api):
    from jobhelper import db
    _seed(store.new_job(id="fresh", title="A", company="A", status="waiting", date_applied="2099-01-01",
                        last_update="2099-01-01"),
          store.new_job(id="quiet", title="B", company="B", status="waiting", date_applied=OLD, last_update=OLD),
          store.new_job(id="answered", title="C", company="C", status="waiting", date_applied=OLD, last_update=OLD),
          store.new_job(id="rejected", title="D", company="D", status="rejected", date_applied=OLD, last_update=OLD))
    db.upsert_response({"msg_id": "<1@x>", "job_id": "answered", "company": "C", "stage": "acknowledged",
                        "received_date": "Wed, 01 Jan 2098 10:00:00 +0000"})
    assert api.counts()["followups"] == 1
    rows = api.applications()
    assert rows[0]["id"] == "quiet" and rows[0]["followup"] is True and rows[0]["days_waiting"] > 10
    assert [r["followup"] for r in rows[1:]] == [False, False, False]
    assert rows[-1]["id"] == "rejected"


def test_applications_are_grouped_interview_waiting_rejected(api):
    _seed(store.new_job(id="r-new", title="R", company="R", status="rejected", last_update="2026-10-06"),
          store.new_job(id="w-old", title="W", company="W", status="waiting", last_update="2026-08-01"),
          store.new_job(id="i-old", title="I", company="I", status="interview", last_update="2026-07-01"),
          store.new_job(id="w-new", title="W2", company="W2", status="waiting", last_update="2026-10-01"),
          store.new_job(id="offer", title="O", company="O", status="offer", last_update="2026-06-01"),
          store.new_job(id="hold", title="H", company="H", status="on_hold", last_update="2026-10-05"),
          store.new_job(id="r-old", title="R2", company="R2", status="rejected", last_update="2026-07-01"),
          store.new_job(id="lead", title="L", company="L", status="found"))
    assert [r["id"] for r in api.applications()] == ["offer", "i-old", "w-new", "w-old", "hold", "r-new", "r-old"]


def test_followup_days_come_from_options_and_zero_turns_them_off(api):
    _seed(store.new_job(id="quiet", title="B", company="B", status="waiting", date_applied=OLD, last_update=OLD))
    assert api.options()["followup_days"] == 10
    api.save_options({"followup_days": 0})
    assert api.counts()["followups"] == 0 and api.applications()[0]["followup"] is False


def test_mark_followed_up_is_saved(api):
    _seed(store.new_job(id="quiet", title="B", company="B", status="waiting", date_applied=OLD, last_update=OLD))
    assert api.mark_followed_up("quiet") == {"ok": True}
    job = store.find(store.load_jobs(), "quiet")
    assert job["followup_date"] and "Followed up" in job["notes"]
    assert api.counts()["followups"] == 0
    assert api.mark_followed_up("nope") == {"ok": False}


def test_followup_draft_signs_with_the_name_from_the_cv_file(api):
    config.ensure_dirs()
    (config.CV_DIR / "JaneDoe_CV.txt").write_text("Jane", encoding="utf-8")
    _seed(store.new_job(id="quiet", title="QA Engineer", company="Acme", status="waiting", date_applied=OLD))
    draft = api.followup_draft("quiet")
    assert draft["body"].endswith("Jane Doe")
    assert f"{_OLD_DAY.day} {_OLD_DAY.strftime('%B %Y')}" in draft["body"]
    assert api.followup_draft("nope") == {}


def test_responses_are_grouped_by_stage_and_sorted_by_real_date(api):
    from jobhelper import db
    rows = [("<1@x>", "Wed, 2 Sep 2026 09:00:00 +0200", "rejection", "Old rejection"),
            ("<2@x>", "Mon, 5 Oct 2026 10:00:00 +0200", "rejection", "New rejection"),
            ("<3@x>", "Thu, 1 Oct 2026 08:00:00 +0000", "interview", "Interview"),
            ("<4@x>", "Fri, 25 Sep 2026 08:00:00 +0000", "acknowledged", "Ack older"),
            ("<5@x>", "Tue, 6 Oct 2026 08:00:00 +0000", "acknowledged", "Ack newer"),
            ("<6@x>", "Sat, 3 Oct 2026 08:00:00 +0000", "assessment", "Test task"),
            ("<7@x>", "not a date", "other", "Unreadable date"),
            ("<8@x>", "2026-10-04T12:00:00", "other", "ISO date")]
    for msg_id, when, stage, company in rows:
        db.upsert_response({"msg_id": msg_id, "received_date": when, "stage": stage, "company": company})
    # the database itself: newest first, whatever weekday the header starts with
    assert [r["company"] for r in db.all_responses()][:3] == ["Ack newer", "New rejection", "ISO date"]
    assert db.all_responses()[-1]["company"] == "Unreadable date"
    listed = api.responses()
    assert [r["company"] for r in listed] == ["Interview", "Test task", "Ack newer", "Ack older",
                                              "ISO date", "Unreadable date", "New rejection", "Old rejection"]
    assert listed[0]["date"] == "2026-10-01" and listed[5]["date"] == ""


def test_reply_detail_and_analytics_that_add_up(api):
    from jobhelper import db
    _seed(store.new_job(id="a", title="QA Engineer", company="Acme", status="rejected", notes="[2026-10-07] Marked rejected by you."),
          store.new_job(id="b", title="Dev", company="Beta", status="waiting"),
          store.new_job(id="c", title="Dev", company="Gamma", status="interview"),
          store.new_job(id="lead", title="Dev", company="Delta", status="found"))
    db.upsert_response({"msg_id": "<1@x>", "job_id": "b", "company": "Beta", "role": "Dev", "stage": "acknowledged",
                        "received_date": "Mon, 5 Oct 2026 10:15:00 +0200", "subject": "We got it", "from_addr": "hr@beta.com",
                        "summary": "Received.", "next_step": "Wait"})
    db.upsert_response({"msg_id": "<2@x>", "job_id": "gone", "company": "Old", "stage": "rejection",
                        "rejection_category": "position_filled", "received_date": "not a date"})
    d = api.response_detail("<1@x>")
    assert (d["company"], d["stage"], d["date"], d["next_step"]) == ("Beta", "acknowledged", "2026-10-05 10:15", "Wait")
    assert (d["job_id"], d["job_title"], d["job_status"], d["category"]) == ("b", "Dev", "waiting", "")
    orphan = api.response_detail("<2@x>")
    assert orphan["job_id"] == "" and orphan["category"] == "position filled" and orphan["date"] == ""
    assert api.response_detail("<nope>") == {}
    assert [r["id"] for r in api.responses()] == ["<1@x>", "<2@x>"]

    an = api.analytics()
    assert an["applications_tracked"] == 3 == sum(an["applications_by_status"].values())
    assert an["rejected_by_hand"] == 1
    assert an["response_emails"] == 2 == sum(an["by_stage"].values())
    assert sum(an["rejection_reasons"].values()) == an["by_stage"]["rejection"]


def test_options_default_skip_words(api):
    opts = api.options()
    assert opts["skip_words"] == ["intern", "internship"] and opts["keep_words"] == []
    api.save_options({**opts, "skip_words": [], "keep_words": ["x"]})
    assert api.options()["skip_words"] == [] and api.options()["keep_words"] == ["x"]


# ── what changed ────────────────────────────────────────────────────────────
def test_changelog_is_read_from_the_file_newest_first(api, monkeypatch, tmp_path):
    (tmp_path / "CHANGELOG.md").write_text(
        "# What changed\n\nIntro text.\n- a bullet before any date is ignored\n\n"
        "## 7 October 2026\n\n- **Compare CV**: shows <only> what differs & nothing else.\n- Second line.\n\n"
        "## Empty section\n\n## 6 October 2026\n- Older change.\n", encoding="utf-8")
    monkeypatch.setattr(config, "ROOT", tmp_path)
    got = api.changelog()
    assert [e["date"] for e in got] == ["7 October 2026", "6 October 2026"]
    assert got[0]["items"] == ["<b>Compare CV</b>: shows &lt;only&gt; what differs &amp; nothing else.", "Second line."]
    (tmp_path / "CHANGELOG.md").unlink()
    assert api.changelog() == []                                   # a missing file is an empty list, not a crash


def test_the_real_changelog_is_well_formed():
    import re
    got = ui_api.Api().changelog()
    assert len(got) >= 2 and all(e["items"] for e in got)
    # each heading is "Unreleased", a version with its date, a date (before versions), or the oldest "Earlier"
    for entry in got:
        assert re.fullmatch(r"Unreleased|Earlier|(\d+\.\d+\.\d+ – )?\d{1,2} \w+ \d{4}", entry["date"]), entry["date"]
    assert all(entry["date"] != "Unreleased" for entry in got[1:])          # only the newest entry may be unreleased
    for entry in got:
        for item in entry["items"]:
            assert item.count("<b>") == item.count("</b>") and "**" not in item and "—" not in item


def test_job_window_says_why_a_good_score_has_no_documents(api):
    from jobhelper import agent3_letters
    store.save_jobs([store.new_job(id="h", title="Junior AI Engineer", company="", match_score=75, hold=agent3_letters.NO_COMPANY)])
    jobs = store.load_jobs()
    jobs[0]["hold"] = agent3_letters.NO_COMPANY
    store.save_jobs(jobs)
    assert "no company name" in api.job_detail("h")["hold"]
    jobs[0]["cover_letter_file"] = "L.docx"                          # once a letter exists the note is gone
    store.save_jobs(jobs)
    assert api.job_detail("h")["hold"] == ""


# ── show a document in its folder ───────────────────────────────────────────
def test_show_in_folder_reveals_the_pdf_when_there_is_one(api, monkeypatch):
    shown = []
    monkeypatch.setattr(api_core.Desktop, "reveal", lambda self, target: shown.append(target))
    config.ensure_dirs()
    (config.TAILORED_CV_DIR / "CV_a.docx").write_bytes(b"cv")
    (config.TAILORED_CV_DIR / "CV_a.pdf").write_bytes(b"%PDF")
    (config.LETTERS_DIR / "Letter_a.docx").write_bytes(b"letter")
    config.COVER_LETTERS_DIR.mkdir(parents=True, exist_ok=True)
    (config.COVER_LETTERS_DIR / "Mine_CoverLetter_Beta.docx").write_bytes(b"mine")
    store.save_jobs([store.new_job(id="a", title="Dev", company="Acme", tailored_cv_file="CV_a.docx", cover_letter_file="Letter_a.docx"),
                     store.new_job(id="own", title="Dev", company="Beta", status="waiting", source="manual", cover_letter_file="Mine_CoverLetter_Beta.docx"),
                     store.new_job(id="bare", title="Dev", company="C")])
    assert api.show_in_folder("cv", "a") == {"ok": True, "file": "CV_a.pdf"} and shown[-1] == config.TAILORED_CV_DIR / "CV_a.pdf"
    assert api.show_in_folder("letter", "a") == {"ok": True, "file": "Letter_a.docx"}          # no PDF copy: the Word file
    assert api.show_in_folder("letter", "own")["file"] == "Mine_CoverLetter_Beta.docx"         # a letter of your own
    assert "no such document" in api.show_in_folder("cv", "bare")["msg"]
    assert "Make the single PDF first" in api.show_in_folder("combined", "a")["msg"]
    assert "not in the list" in api.show_in_folder("cv", "zzz")["msg"] and api.show_in_folder("wrong", "a")["ok"] is False
    (config.TAILORED_CV_DIR / "CV_a.docx").unlink(); (config.TAILORED_CV_DIR / "CV_a.pdf").unlink()
    assert "not on the disk" in api.show_in_folder("cv", "a")["msg"] and len(shown) == 3


def test_job_window_gets_the_description_laid_out(api):
    store.save_jobs([store.new_job(id="a", title="Dev", company="Acme",
                                   description="We build tools. Requirements Strong Python Benefits Flexible hours Show more Show less")])
    detail = api.job_detail("a")
    assert detail["description_text"] == "We build tools.\n\nREQUIREMENTS\nStrong Python\n\nBENEFITS\nFlexible hours"
    assert detail["description"].endswith("Show more Show less")           # the stored text itself is untouched


# ── score again ─────────────────────────────────────────────────────────────
def test_score_again_for_one_job(api, monkeypatch):
    from jobhelper import llm
    store.save_jobs([store.new_job(id="e", title="Associate Economist", company="Acme", match_score=80, match_reason="old", description="Economics."),
                     store.new_job(id="app", title="Dev", company="B", status="waiting", match_score=70),
                     store.new_job(id="none", title="Dev", company="C")])
    assert api.rescore_ids() == ["e"] and api.rescore_plan() == {"jobs": 1, "calls": 1}
    assert api.job_detail("e")["can_rescore"] is True and api.job_detail("app")["can_rescore"] is False and api.job_detail("none")["can_rescore"] is False
    for job_id in ("app", "none", "zzz"):
        assert "Only a new job that already has a score" in api.rescore_one(job_id)["msg"]
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"match_score": 30, "match_reason": "Another field."})
    assert api.rescore_one("e") == {"ok": True, "score": 30, "before": 80}
    detail = ui_api.Api().job_detail("e")                                      # saved
    assert (detail["score"], detail["score_before"], detail["reason"]) == (30, 80, "Another field.")
    assert api.rescore_one("e") == {"ok": True, "score": 30, "before": 30}

    def limit(*a, **k):
        raise llm.LimitReached("Claude's 5-hour limit is used up (99%). It resets at 22:50.")
    monkeypatch.setattr(llm, "ask_json", limit)
    got = api.rescore_one("e")
    assert got["ok"] is False and "limit is used up" in got["msg"] and api.job_detail("e")["score"] == 30      # nothing lost
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"match_score": None})
    assert "no usable score" in api.rescore_one("e")["msg"]


# ── what a score rests on, pasted posting text, documents left from a higher score ──
LONG_AD = "We are hiring a junior Python developer to build internal tools. " * 8


def _with_docs(job_id, score, **kw):
    (config.LETTERS_DIR / f"L_{job_id}.docx").write_bytes(b"letter")
    (config.TAILORED_CV_DIR / f"CV_{job_id}.docx").write_bytes(b"cv")
    return store.new_job(id=job_id, title="Dev", company=f"Co {job_id}", match_score=score, status="letter_ready",
                         cover_letter_file=f"L_{job_id}.docx", tailored_cv_file=f"CV_{job_id}.docx", description=LONG_AD, **kw)


def test_list_and_window_say_what_a_score_rests_on(api):
    config.ensure_dirs()
    store.save_jobs([store.new_job(id="t", title="A", company="C", match_score=40), store.new_job(id="s", title="B", company="C", match_score=60, description="Short."),
                     store.new_job(id="f", title="C", company="C", match_score=60, description=LONG_AD), store.new_job(id="u", title="D", company="C")])
    assert {r["id"]: r["basis"] for r in api.leads()} == {"t": "title", "s": "summary", "f": "", "u": ""}       # unscored: nothing to qualify
    assert api.job_detail("t")["basis"] == "title" and api.job_detail("t")["can_paste_text"] is True
    assert api.job_detail("f")["can_paste_text"] is False


def test_pasted_posting_text_is_checked_and_kept(api):
    from jobhelper import agent3_letters
    config.ensure_dirs()
    job = store.new_job(id="t", title="A", company="C", match_score=40)
    job["hold"], job["doc_attempts"] = agent3_letters.NO_TEXT, 2
    store.save_jobs([job, store.new_job(id="app", title="B", company="C", status="waiting")])
    assert "only 5 characters. Paste the whole posting (at least 200)" in api.set_description("t", "short")["msg"]
    assert "20001 characters; the most is 20000" in api.set_description("t", "x" * 20001)["msg"]
    assert "only be added to a new job" in api.set_description("app", LONG_AD)["msg"]
    assert api.set_description("zzz", LONG_AD)["ok"] is False
    assert api.set_description("t", "  " + LONG_AD + "\r\n\r\n\r\n\r\nSecond part.  ")["ok"] is True
    saved = store.find(store.load_jobs(), "t")
    assert saved["description"].startswith("We are hiring") and saved["description"].endswith("\n\nSecond part.") and saved["description_by_hand"] is True
    assert "hold" not in saved and saved["doc_attempts"] == 0 and api.job_detail("t")["basis"] == ""


def test_documents_left_from_a_higher_score_are_marked_and_can_be_put_away(api):
    config.ensure_dirs()
    api.save_options({"threshold": 50})
    (config.LETTERS_DIR / "Mine.docx").write_bytes(b"my own letter")
    store.save_jobs([_with_docs("low", 30), _with_docs("ok", 70), _with_docs("app", 20, ),
                     store.new_job(id="own", title="Dev", company="Z", match_score=10, source="manual", status="letter_ready", cover_letter_file="Mine.docx"),
                     store.new_job(id="bare", title="Dev", company="Y", match_score=10)])
    jobs = store.load_jobs(); store.find(jobs, "app")["status"] = "waiting"; store.save_jobs(jobs)
    assert {r["id"]: r["stale_docs"] for r in api.leads()} == {"ok": False, "low": True, "own": False, "bare": False}
    assert api.job_detail("low")["docs_stale"] is True and api.job_detail("low")["threshold"] == 50 and api.job_detail("app")["docs_stale"] is False
    assert api.shelve_docs("low") == {"ok": True, "moved": 2}
    low = store.find(store.load_jobs(), "low")
    assert (low["cover_letter_file"], low["tailored_cv_file"], low["status"], low["match_score"]) == ("", "", "found", 30)
    assert not (config.LETTERS_DIR / "L_low.docx").exists() and len(list((vault.backups_dir() / "below_your_limit").iterdir())) == 2
    assert api.job_detail("low")["docs_stale"] is False
    assert "not for an application you sent" in api.shelve_docs("app")["msg"] and (config.LETTERS_DIR / "L_app.docx").exists()
    assert "no documents written by the app" in api.shelve_docs("own")["msg"] and (config.LETTERS_DIR / "Mine.docx").read_bytes() == b"my own letter"
    assert api.shelve_docs("bare")["ok"] is False and api.shelve_docs("zzz")["ok"] is False


def test_read_a_posting_again_from_the_job_window(api, monkeypatch):
    from jobhelper import agent2_search, agent3_letters
    config.ensure_dirs()
    job = store.new_job(id="t", title="Tester", company="C", match_score=40, url="https://x/t")
    job["hold"] = agent3_letters.NO_TEXT
    store.save_jobs([job, store.new_job(id="nolink", title="T", company="C", match_score=40),
                     store.new_job(id="app", title="T", company="C", status="waiting", url="https://x/a")])
    assert "still cannot be read" in api.reread_posting("t")["msg"]                       # the board gave nothing again
    assert "no link" in api.reread_posting("nolink")["msg"] and "Only a new job" in api.reread_posting("app")["msg"]
    monkeypatch.setattr(agent2_search, "refetch_description", lambda j: LONG_AD)
    assert api.reread_posting("t") == {"ok": True, "chars": len(LONG_AD)}
    saved = store.find(store.load_jobs(), "t")
    assert saved["description"] == LONG_AD and "hold" not in saved and api.job_detail("t")["basis"] == ""


def test_all_documents_below_the_limit_are_moved_in_one_go(api):
    config.ensure_dirs()
    api.save_options({"threshold": 50})
    (config.LETTERS_DIR / "Mine.docx").write_bytes(b"my own letter")
    assert api.stale_docs_list() == [] and "No new job has documents" in api.shelve_all_docs()["msg"]
    store.save_jobs([_with_docs("low", 30), _with_docs("lower", 10), _with_docs("ok", 70), _with_docs("app", 20),
                     store.new_job(id="own", title="Dev", company="Z", match_score=10, source="manual", status="letter_ready", cover_letter_file="Mine.docx")])
    jobs = store.load_jobs(); store.find(jobs, "app")["status"] = "waiting"; store.save_jobs(jobs)
    rows = api.stale_docs_list()
    assert [r["id"] for r in rows] == ["low", "lower"] and set(rows[0]) == {"id", "title", "company", "score"} and rows[0]["score"] == 30
    refused = api.shelve_all_docs(5)                                    # the confirmation named another number of jobs
    assert refused["ok"] is False and refused["changed"] is True and "it is now 2 jobs" in refused["msg"]
    assert (config.LETTERS_DIR / "L_low.docx").exists()
    assert api.shelve_all_docs(2) == {"ok": True, "jobs": 2, "files": 4, "failed": []}
    for name in ("low", "lower"):
        job = store.find(store.load_jobs(), name)
        assert (job["cover_letter_file"], job["tailored_cv_file"], job["status"]) == ("", "", "found")
        assert not (config.LETTERS_DIR / f"L_{name}.docx").exists()
    assert len(list((vault.backups_dir() / "below_your_limit").iterdir())) == 4
    for name in ("ok", "app"):                                          # above the limit, and an application: untouched
        assert (config.LETTERS_DIR / f"L_{name}.docx").exists() and store.find(store.load_jobs(), name)["cover_letter_file"]
    assert (config.LETTERS_DIR / "Mine.docx").read_bytes() == b"my own letter" and store.find(store.load_jobs(), "own")["cover_letter_file"] == "Mine.docx"
    assert api.stale_docs_list() == [] and api.shelve_all_docs()["ok"] is False


def test_one_job_that_cannot_be_moved_does_not_stop_the_others(api, monkeypatch):
    config.ensure_dirs()
    api.save_options({"threshold": 50})
    store.save_jobs([_with_docs("a", 30), _with_docs("b", 20)])
    real = api_documents.DocumentsService.shelve_docs
    monkeypatch.setattr(api_documents.DocumentsService, "shelve_docs",
                        lambda self, job_id, batch="": {"ok": False, "msg": "no"} if job_id == "a" else real(self, job_id, batch))
    got = api.shelve_all_docs(2)
    assert got["ok"] is True and got["jobs"] == 1 and got["files"] == 2
    assert got["failed"] == [{"title": "Dev", "company": "Co a", "why": "no"}]             # each one that failed, with its own reason
    monkeypatch.setattr(api_documents.DocumentsService, "shelve_docs", lambda self, job_id, batch="": {"ok": False})
    nothing = api.shelve_all_docs()
    assert nothing["ok"] is False and nothing["msg"] == "Couldn't move the documents."
    assert nothing["failed"] == [{"title": "Dev", "company": "Co a", "why": "Couldn't move the documents."}]


def test_the_list_shows_what_a_job_scored_before(api):
    config.ensure_dirs()
    changed, same, never, odd = (store.new_job(id=i, title="Dev", company=i, match_score=s) for i, s in (("changed", 25), ("same", 60), ("never", 50), ("odd", 40)))
    changed["score_before"], same["score_before"], odd["score_before"] = 55, 60, "55"
    waiting = store.new_job(id="waiting", title="Dev", company="w")
    waiting["score_before"] = 40                                        # its text arrived; the new score is still to come
    store.save_jobs([changed, same, never, odd, waiting])
    assert {r["id"]: r["before"] for r in api.leads()} == {"changed": 55, "same": None, "never": None, "odd": None, "waiting": 40}
    assert api.job_detail("changed")["score_before"] == 55 and api.job_detail("same")["score_before"] is None


# ── documents moved to the backup can be put back ───────────────────────────
def test_moved_documents_can_be_put_back_for_one_job(api):
    config.ensure_dirs()
    api.save_options({"threshold": 50})
    job = _with_docs("low", 30)
    job["match_tier"], job["apply_steps"] = "possible", ["cv"]
    store.save_jobs([job, _with_docs("ok", 70)])
    assert api.job_detail("low")["can_unshelve"] is False and "No documents of this job" in api.unshelve_docs("low")["msg"]
    assert api.shelve_docs("low") == {"ok": True, "moved": 2}
    shelf = vault.backups_dir() / "below_your_limit"
    detail = api.job_detail("low")
    assert detail["can_unshelve"] is True and detail["shelved_when"][:4].isdigit() and len(list(shelf.iterdir())) == 2
    assert api.unshelve_docs("low") == {"ok": True, "restored": 2}
    back = store.find(store.load_jobs(), "low")
    assert (back["cover_letter_file"], back["tailored_cv_file"], back["status"], back["match_tier"], back["apply_steps"]) \
        == ("L_low.docx", "CV_low.docx", "letter_ready", "possible", ["cv"]) and "shelved" not in back
    assert (config.LETTERS_DIR / "L_low.docx").read_bytes() == b"letter" and (config.TAILORED_CV_DIR / "CV_low.docx").read_bytes() == b"cv"
    assert list(shelf.iterdir()) == [] and api.job_detail("low")["docs_stale"] is True and api.job_detail("low")["can_unshelve"] is False
    assert api.unshelve_docs("low")["ok"] is False and api.unshelve_docs("zzz")["ok"] is False


def test_putting_back_changes_nothing_unless_all_of_it_can_be(api):
    config.ensure_dirs()
    api.save_options({"threshold": 50})
    shelf = vault.backups_dir() / "below_your_limit"

    def moved(job_id):
        store.save_jobs([_with_docs(job_id, 30)])
        for old in list(shelf.iterdir()) if shelf.exists() else []:
            old.unlink()
        assert api.shelve_docs(job_id)["ok"] is True
        return sorted(shelf.iterdir())

    copies = moved("a")                                             # a file with that name is there again
    (config.LETTERS_DIR / "L_a.docx").write_bytes(b"newer")
    assert "A file named L_a.docx is already in its folder" in api.unshelve_docs("a")["msg"]
    assert (config.LETTERS_DIR / "L_a.docx").read_bytes() == b"newer" and not (config.TAILORED_CV_DIR / "CV_a.docx").exists() and sorted(shelf.iterdir()) == copies
    (config.LETTERS_DIR / "L_a.docx").unlink()

    copies = moved("b")                                             # a backup copy is gone
    copies[0].unlink()
    assert "is missing, so nothing was put back" in api.unshelve_docs("b")["msg"]
    assert not (config.LETTERS_DIR / "L_b.docx").exists() and not (config.TAILORED_CV_DIR / "CV_b.docx").exists() and len(list(shelf.iterdir())) == 1

    moved("c")                                                      # the app wrote other documents since
    jobs = store.load_jobs(); store.find(jobs, "c").update(tailored_cv_file="CV_new.docx", cover_letter_file="L_new.docx", status="letter_ready"); store.save_jobs(jobs)
    assert api.job_detail("c")["can_unshelve"] is False and "Other documents were written for this job since" in api.unshelve_docs("c")["msg"]
    assert store.find(store.load_jobs(), "c")["cover_letter_file"] == "L_new.docx" and len(list(shelf.iterdir())) == 2

    moved("d")                                                      # you said you are not interested: nothing to put back for
    jobs = store.load_jobs(); store.find(jobs, "d")["status"] = store.NOT_INTERESTED; store.save_jobs(jobs)
    assert api.unshelve_docs("d")["ok"] is False and api.job_detail("d")["can_unshelve"] is False
    # it turned out to be an application (a reply showed you had applied): those are the documents you sent
    jobs = store.load_jobs(); store.find(jobs, "d")["status"] = "waiting"; store.save_jobs(jobs)
    assert api.job_detail("d")["can_unshelve"] is True and api.unshelve_docs("d") == {"ok": True, "restored": 2}
    back = store.find(store.load_jobs(), "d")
    assert back["cover_letter_file"] == "L_d.docx" and back["status"] == "waiting", "it stays an application"

    moved("e")                                                      # a note that points outside the folders is not followed
    jobs = store.load_jobs(); record = store.find(jobs, "e")["shelved"]
    record["files"][0]["name"] = "../../escaped.docx"; store.save_jobs(jobs)
    assert api.unshelve_docs("e") == {"ok": True, "restored": 2}
    assert not (config.DATA_DIR / "escaped.docx").exists() and not (config.DATA_DIR.parent / "escaped.docx").exists()

    moved("f")
    jobs = store.load_jobs(); store.find(jobs, "f")["shelved"]["files"] = ["broken"]; store.save_jobs(jobs)
    assert "note of what was moved is damaged" in api.unshelve_docs("f")["msg"]


def test_a_letter_of_your_own_is_untouched_by_moving_and_putting_back(api):
    config.ensure_dirs()
    api.save_options({"threshold": 50})
    (config.LETTERS_DIR / "Mine.docx").write_bytes(b"my own letter")
    (config.TAILORED_CV_DIR / "CV_own.docx").write_bytes(b"cv")
    store.save_jobs([store.new_job(id="own", title="Dev", company="Z", match_score=10, source="manual", status="letter_ready",
                                   cover_letter_file="Mine.docx", tailored_cv_file="CV_own.docx")])
    assert api.shelve_docs("own") == {"ok": True, "moved": 1}
    job = store.find(store.load_jobs(), "own")
    assert job["cover_letter_file"] == "Mine.docx" and job["shelved"]["cover_letter_file"] == "" and (config.LETTERS_DIR / "Mine.docx").read_bytes() == b"my own letter"
    assert api.job_detail("own")["can_unshelve"] is True and api.unshelve_docs("own") == {"ok": True, "restored": 1}
    job = store.find(store.load_jobs(), "own")
    assert (job["cover_letter_file"], job["tailored_cv_file"]) == ("Mine.docx", "CV_own.docx") and (config.LETTERS_DIR / "Mine.docx").read_bytes() == b"my own letter"


def test_the_last_batch_is_put_back_together(api, monkeypatch):
    from datetime import datetime, timedelta
    config.ensure_dirs()
    api.save_options({"threshold": 50})
    assert api.last_shelved() == {} and "No moved documents are left" in api.unshelve_last()["msg"]
    store.save_jobs([_with_docs("a", 30), _with_docs("b", 20), _with_docs("c", 10), _with_docs("ok", 70)])
    assert api.shelve_docs("c")["ok"] is True                        # moved alone, earlier
    jobs = store.load_jobs(); store.find(jobs, "c")["shelved"]["when"] = (datetime.now() - timedelta(hours=2)).isoformat(timespec="seconds"); store.save_jobs(jobs)
    assert api.shelve_all_docs(2)["ok"] is True
    last = api.last_shelved()
    assert sorted(r["id"] for r in last["jobs"]) == ["a", "b"] and set(last["jobs"][0]) == {"id", "title", "company"} and last["batch"]
    refused = api.unshelve_last("another-batch")
    assert refused["ok"] is False and refused["changed"] is True and not (config.LETTERS_DIR / "L_a.docx").exists()
    assert api.unshelve_last(last["batch"]) == {"ok": True, "jobs": 2, "files": 4, "failed": []}
    assert (config.LETTERS_DIR / "L_a.docx").exists() and (config.TAILORED_CV_DIR / "CV_b.docx").exists() and not (config.LETTERS_DIR / "L_c.docx").exists()
    assert [r["id"] for r in api.last_shelved()["jobs"]] == ["c"]     # the earlier one is now the last
    assert sorted(r["id"] for r in api.stale_docs_list()) == ["a", "b"]

    jobs = store.load_jobs(); store.find(jobs, "c")["shelved"]["when"] = (datetime.now() - timedelta(days=api.UNDO_DAYS + 1)).isoformat(timespec="seconds"); store.save_jobs(jobs)
    assert api.last_shelved() == {} and api.job_detail("c")["can_unshelve"] is True      # no longer offered in the list, still possible in the job window
    assert api.unshelve_docs("c")["ok"] is True

    assert api.shelve_all_docs(3)["ok"] is True                      # one of the three cannot come back: the others still do
    last = api.last_shelved()
    (config.LETTERS_DIR / "L_a.docx").write_bytes(b"in the way")
    got = api.unshelve_last(last["batch"])
    assert got["ok"] is True and got["jobs"] == 2 and len(got["failed"]) == 1
    assert got["failed"][0]["title"] == "Dev" and got["failed"][0]["company"] == "Co a" and "A file named L_a.docx is already in its folder" in got["failed"][0]["why"]
    stuck = api.last_shelved()                                       # the one left is now the last batch, and still cannot come back
    again = api.unshelve_last(stuck["batch"])
    assert again["ok"] is False and again["msg"] == "Couldn't put the documents back." and "already in its folder" in again["failed"][0]["why"]


def test_options_can_say_which_professions_are_taken_as_yours(api, monkeypatch):
    from jobhelper import agent1_cv, prefilter
    real = api.profession_fields
    assert real([])["all"] == list(prefilter.FIELDS) and len(real([])["all"]) > 5            # every profession the app knows, for the pick list
    monkeypatch.setattr(api, "profession_fields", lambda wanted=None: {k: v for k, v in real(wanted).items() if k != "all"})
    monkeypatch.setattr(agent1_cv, "load_profile", lambda: {"target_roles": ["Backend Developer"], "adjacent_roles": ["Data Analyst"]})
    monkeypatch.setattr(api_core.CvSource, "text", lambda self: "Nurse on a ward, then a nurse in a clinic.")
    assert api.profession_fields() == {"fields": ["software and IT"], "from_cv": False}                    # nothing typed: the CV analysis' roles
    assert api.profession_fields(["registered nurse"]) == {"fields": ["software and IT", "care and health"], "from_cv": False}   # adjacent roles count too
    monkeypatch.setattr(agent1_cv, "load_profile", lambda: {})
    assert api.profession_fields(["registered nurse", "", "  "]) == {"fields": ["care and health"], "from_cv": False}
    assert api.profession_fields(["beekeeper"]) == {"fields": ["care and health"], "from_cv": True}       # roles say nothing: the CV does
    monkeypatch.setattr(api_core.CvSource, "text", lambda self: "")
    assert api.profession_fields(["beekeeper"]) == {"fields": [], "from_cv": True} and api.profession_fields(None)["fields"] == []

    def broken():
        raise OSError("no profile")
    monkeypatch.setattr(agent1_cv, "load_profile", broken)
    assert api.profession_fields(["sales manager"]) == {"fields": ["sales and marketing"], "from_cv": False}


def test_what_a_score_rests_on_is_said_only_for_a_new_job_with_a_score(api):
    config.ensure_dirs()
    store.save_jobs([store.new_job(id="lead", title="Dev", company="A", match_score=40, description="Short summary."),
                     store.new_job(id="unscored", title="Dev", company="B", description="Short summary."),
                     store.new_job(id="app", title="Dev", company="C", match_score=64, status="waiting", description="Short summary."),
                     store.new_job(id="rejected", title="Dev", company="D", match_score=64, status="rejected")])
    assert [api.job_detail(i)["basis"] for i in ("lead", "unscored", "app", "rejected")] == ["summary", "", "", ""]
    assert api.job_detail("unscored")["can_paste_text"] is True and api.job_detail("app")["can_paste_text"] is False


def test_moving_and_putting_back_reach_a_running_search_s_own_list(api):
    """A search holds its own copy of the jobs and saves it when it ends: an edit made meanwhile must be in that copy too."""
    config.ensure_dirs()
    api.save_options({"threshold": 50})
    store.save_jobs([_with_docs("low", 30)])
    api.live = store.load_jobs()                                   # what a running search holds
    try:
        assert api.shelve_docs("low") == {"ok": True, "moved": 2}
        twin = store.find(api.live, "low")
        assert (twin["cover_letter_file"], twin["tailored_cv_file"], twin["status"]) == ("", "", "found") and len(twin["shelved"]["files"]) == 2
        store.save_jobs(api.live)                                  # the search ends and writes its copy
        assert api.job_detail("low")["can_unshelve"] is True and not (config.LETTERS_DIR / "L_low.docx").exists()
        api.live = store.load_jobs()
        assert api.unshelve_docs("low") == {"ok": True, "restored": 2}
        twin = store.find(api.live, "low")
        assert (twin["cover_letter_file"], twin["tailored_cv_file"], twin["status"]) == ("L_low.docx", "CV_low.docx", "letter_ready") and "shelved" not in twin
        store.save_jobs(api.live)
        assert (config.LETTERS_DIR / "L_low.docx").exists() and api.job_detail("low")["docs_stale"] is True
    finally:
        api.live = None


def test_a_closed_posting_is_marked_in_the_list_and_the_window(api):
    config.ensure_dirs()
    lead, app = _with_docs("lead", 70), _with_docs("app", 70)
    lead["closed"], app["closed"], app["status"] = True, True, "waiting"
    store.save_jobs([lead, app, _with_docs("open", 60)])
    assert {r["id"]: r["closed"] for r in api.leads()} == {"lead": True, "open": False}
    assert api.job_detail("lead")["closed"] is True and api.job_detail("open")["closed"] is False and api.job_detail("app")["closed"] is False


def test_the_back_end_is_a_facade_over_services_that_do_not_know_each_other():
    """`Api` only hands calls on; each public name has one owner; a service asks another one through `Api`."""
    import ast
    import inspect
    from pathlib import Path

    api = ui_api.Api()
    owners = {}
    for name, value in vars(ui_api.Api).items():
        if isinstance(value, ui_api._Via):
            owners.setdefault(name, value._service)
    services = {attr: getattr(api, attr) for attr in set(owners.values())}
    assert len(services) == 10 and len(owners) > 150
    for attr, service in services.items():                      # every public method of a service is on the facade, with that owner
        public = {n for n in vars(type(service)) if not n.startswith("_")}
        assert {n for n, o in owners.items() if o == attr} == public, attr
    for path in Path(ui_api.__file__).parent.glob("api_*.py"):   # no service imports another service
        source = path.read_text(encoding="utf-8")
        for node in ast.walk(ast.parse(source)):
            if isinstance(node, ast.ImportFrom) and (node.module or "").startswith("api_") and node.module not in ("api_base", "api_core"):
                raise AssertionError(f"{path.name} imports {node.module}")
    assert callable(api.shelve_docs) and callable(api.save_letter) and callable(api.set_event)
    assert "self._app" in inspect.getsource(type(services["_talk"]))      # the talk reaches the rest of the app only through the facade


# ── filing a reply under the right application ──────────────────────────────
def _replies_case():
    from jobhelper import db
    store.save_jobs([store.new_job(id="right", title="Marketing Consultant", company="Actum", status="waiting", date_applied="2026-08-01"),
                     store.new_job(id="wrong", title="Software Engineer", company="Bigsoft", status="interview", date_applied="2026-07-01"),
                     store.new_job(id="lead", title="Tester", company="Cyan", status="found")])
    db.upsert_response({"msg_id": "<invite>", "received_date": "Tue, 18 Aug 2026 10:00:00 +0200", "from_addr": "hr@actum.example",
                        "subject": "Interview", "job_id": "wrong", "company": "Actum", "role": "Marketing Consultant", "stage": "interview",
                        "summary": "They invite you to a call."})
    db.upsert_response({"msg_id": "<loose>", "received_date": "Wed, 12 Aug 2026 10:00:00 +0200", "from_addr": "hr@good.example",
                        "subject": "No", "job_id": "", "company": "Good", "role": "", "stage": "rejection", "summary": "A rejection."})
    return db


def test_a_reply_can_be_filed_under_the_application_it_belongs_to(isolate_data):
    db = _replies_case()
    api = ui_api.Api()
    assert [t["id"] for t in api.filing_targets()] == ["right", "wrong"]            # applications only, newest first
    assert [t["id"] for t in api.filing_targets("actum consult")] == ["right"] and api.filing_targets("zzz") == []
    result = api.refile_response("<invite>", "right")
    assert result["ok"] is True and "Filed under Actum · Marketing Consultant." in result["msg"]
    assert "status changed from waiting to interview" in result["msg"]
    assert "Bigsoft keeps its status (interview)" in result["msg"]                  # the one it left is yours to judge
    jobs = {j["id"]: j for j in store.load_jobs()}
    assert jobs["right"]["status"] == "interview" and jobs["wrong"]["status"] == "interview"
    assert api.response_detail("<invite>")["job_id"] == "right"
    assert [r["job_id"] for r in db.all_responses() if r["msg_id"] == "<invite>"] == ["right"]
    # the reader had taken the company from the email; filed by hand, the reply is listed under the application's company
    db.upsert_response({"msg_id": "<teams>", "received_date": "Tue, 18 Aug 2026 10:00:00 +0200", "from_addr": "noreply@video.example", "subject": "Int",
                        "job_id": "wrong", "company": "Bigsoft", "role": "", "stage": "acknowledged", "summary": "A call link."})
    assert api.refile_response("<teams>", "right")["ok"] is True
    assert [r["company"] for r in db.all_responses() if r["msg_id"] == "<teams>"] == ["Actum"]
    assert [r["company"] for r in db.all_responses() if r["msg_id"] == "<invite>"] == ["Actum"]         # already right: left as it was
    # a rejection closes the application it is filed under, from any status
    result = api.refile_response("<loose>", "right")
    assert result["ok"] is True and store.find(store.load_jobs(), "right")["status"] == "rejected"


def test_filing_a_reply_refuses_what_makes_no_sense(isolate_data):
    _replies_case()
    api = ui_api.Api()
    assert "not on file" in api.refile_response("<nope>", "right")["msg"]
    assert "already filed there" in api.refile_response("<invite>", "wrong")["msg"]
    assert "not filed under any application" in api.refile_response("<loose>", "")["msg"]
    assert "not in the list any more" in api.refile_response("<invite>", "gone")["msg"]
    assert "Mark that job as applied first" in api.refile_response("<invite>", "lead")["msg"]
    assert api.response_detail("<invite>")["job_id"] == "wrong"                     # nothing changed by any of these
    # an invitation does not pull a finished application backwards
    jobs = store.load_jobs(); store.find(jobs, "right")["status"] = "offer"; store.save_jobs(jobs)
    result = api.refile_response("<invite>", "right")
    assert result["ok"] is True and "status changed" not in result["msg"] and store.find(store.load_jobs(), "right")["status"] == "offer"


def test_a_reply_can_be_filed_under_no_application(isolate_data):
    _replies_case()
    api = ui_api.Api()
    result = api.refile_response("<invite>", "")
    assert result["ok"] is True and "not filed under any application now" in result["msg"] and "Bigsoft keeps its status" in result["msg"]
    assert api.response_detail("<invite>")["job_id"] == "" and store.find(store.load_jobs(), "wrong")["status"] == "interview"


# ── documents for one job, on request ───────────────────────────────────────
def test_documents_are_written_for_one_job_on_request(isolate_data, monkeypatch):
    from jobhelper import agent1_cv, agent3_letters
    text = "We build planning tools in Python and SQL. " * 30
    store.save_jobs([store.new_job(id="low", title="Python Developer", company="Acme", status="found", match_score=25, description=text),
                     store.new_job(id="half", title="Tester", company="Borg", status="letter_ready", match_score=55, description=text,
                                   cover_letter_file="My_own_letter.docx", source="manual"),
                     store.new_job(id="done", title="Dev", company="Cyan", status="letter_ready", match_score=70, description=text,
                                   cover_letter_file="L.docx", tailored_cv_file="CV.docx"),
                     store.new_job(id="bare", title="Dev", company="Dyne", status="found", match_score=35, description=""),
                     store.new_job(id="app", title="Dev", company="Eon", status="waiting", description=text)])
    monkeypatch.setattr(api_core.CvSource, "text", lambda self: "Alex Novak\nPython, SQL.")
    monkeypatch.setattr(agent1_cv, "find_cv", lambda: None)
    calls = []
    monkeypatch.setattr(agent3_letters, "_make_tailored_cv", lambda job, cv_text, cv_path, fields: calls.append(("cv", job["id"])) or f"CV_{job['id']}.docx")
    monkeypatch.setattr(agent3_letters, "_make_letter", lambda job, cv_text, cv_path, fields: calls.append(("letter", job["id"])) or f"L_{job['id']}.docx")
    api = ui_api.Api()
    assert api.job_detail("low")["can_write_docs"] is True and api.job_detail("done")["can_write_docs"] is False
    result = api.write_documents("low")                                # below the limit: written because you asked
    assert result["ok"] is True and result["msg"] == "Written: the CV and the letter."
    low = store.find(store.load_jobs(), "low")
    assert (low["tailored_cv_file"], low["cover_letter_file"], low["status"]) == ("CV_low.docx", "L_low.docx", "letter_ready")
    assert api.job_detail("low")["can_write_docs"] is False
    # your own letter is never replaced: only the CV is written
    result = api.write_documents("half")
    half = store.find(store.load_jobs(), "half")
    assert result["msg"] == "Written: the CV." and half["cover_letter_file"] == "My_own_letter.docx" and half["tailored_cv_file"] == "CV_half.docx"
    assert calls == [("cv", "low"), ("letter", "low"), ("cv", "half")]
    # refused, each with its reason, and nothing is asked of Claude
    assert "already has its CV and its letter" in api.write_documents("done")["msg"]
    assert "a letter written from a title alone would be generic" in api.write_documents("bare")["msg"]
    assert "for a new job only" in api.write_documents("app")["msg"] and "for a new job only" in api.write_documents("nope")["msg"]
    monkeypatch.setattr(api_core.CvSource, "text", lambda self: "")
    store.save_jobs(store.load_jobs() + [store.new_job(id="x", title="Dev", company="Fox", status="found", description=text)])
    assert "master CV could not be read" in api.write_documents("x")["msg"]
    assert len(calls) == 3
    # Claude stops answering after the CV was written: the CV is kept on the job, not paid for twice
    from jobhelper import llm
    monkeypatch.setattr(api_core.CvSource, "text", lambda self: "Alex Novak\nPython, SQL.")
    store.save_jobs(store.load_jobs() + [store.new_job(id="stuck", title="Python Developer", company="Fern", status="found",
                                                       match_score=40, description=text)])
    monkeypatch.setattr(agent3_letters, "_make_letter", lambda *a, **k: (_ for _ in ()).throw(llm.CallTimedOut()))
    result = api.write_documents("stuck")
    stuck = store.find(store.load_jobs(), "stuck")
    assert result["ok"] is True and result["cv"] is True and result["letter"] is False      # the window refreshes and shows the CV
    assert result["msg"].startswith("Written: the CV. Not written: the letter.") and "did not answer" in result["msg"]
    assert stuck["tailored_cv_file"] == "CV_stuck.docx" and stuck["status"] == "letter_ready"
    # it stops at the CV: nothing was made, so nothing is said to be written and the job stays as it was
    monkeypatch.setattr(agent3_letters, "_make_tailored_cv", lambda *a, **k: (_ for _ in ()).throw(llm.CallTimedOut()))
    store.save_jobs(store.load_jobs() + [store.new_job(id="early", title="Python Developer", company="Fern", status="found",
                                                       match_score=40, description=text)])
    before = store.find(store.load_jobs(), "early")
    result = api.write_documents("early")
    assert result["ok"] is False and "did not answer" in result["msg"] and store.find(store.load_jobs(), "early") == before



def test_a_failure_while_writing_documents_is_said_and_changes_nothing(isolate_data, monkeypatch):
    from jobhelper import agent1_cv, agent3_letters, llm
    text = "We build planning tools in Python and SQL. " * 30
    store.save_jobs([store.new_job(id="a", title="Dev", company="Acme", status="found", match_score=25, description=text)])
    monkeypatch.setattr(api_core.CvSource, "text", lambda self: "Alex Novak\nPython.")
    monkeypatch.setattr(agent1_cv, "find_cv", lambda: None)
    monkeypatch.setattr(agent3_letters, "_make_tailored_cv", lambda *a: (_ for _ in ()).throw(RuntimeError("no connection")))
    monkeypatch.setattr(agent3_letters, "_make_letter", lambda *a: (_ for _ in ()).throw(RuntimeError("no connection")))
    api = ui_api.Api()
    result = api.write_documents("a")
    assert result["ok"] is False and "Nothing could be written" in result["msg"]
    job = store.find(store.load_jobs(), "a")
    assert job["status"] == "found" and not job["tailored_cv_file"] and not job["cover_letter_file"]
    monkeypatch.setattr(agent3_letters, "generate_for", lambda *a, **k: (_ for _ in ()).throw(llm.LimitReached("Claude's 5-hour limit is used up.")))
    assert "limit is used up" in api.write_documents("a")["msg"]


# ── usage, the log of problems and backups, all from the window ─────────────
def test_usage_overview_adds_up_by_kind_of_work(isolate_data):
    from jobhelper import usage
    api = ui_api.Api()
    assert api.usage_overview() == {"calls": 0, "cost": 0.0, "kinds": []}
    for tag, cost in (("agent3_letters", 1.2), ("agent3_letters", 0.3), ("agent5_responses", 0.5), ("usage_ping", 0.01),
                      ("limit_check", 0.02), ("some_new_kind", 0.04), ("direction_trial", 0.1)):
        usage.record({"cost_usd": cost}, tag=tag)
    seen = api.usage_overview()
    assert seen["calls"] == 7 and seen["cost"] == 2.17
    assert [(k["name"], k["calls"], k["cost"]) for k in seen["kinds"]] == [
        ("Scoring jobs and writing letters", 2, 1.5), ("Reading replies", 1, 0.5), ("Trials and other work", 1, 0.1),
        ("Some new kind", 1, 0.04), ("Asking how much is left", 2, 0.03)]


def test_the_log_of_problems_is_read_newest_first(isolate_data):
    from jobhelper import trouble
    api = ui_api.Api()
    assert api.trouble_log() == {"lines": [], "total": 0, "file": "app_trouble.log"}
    assert "nothing has gone wrong" in api.show_trouble_log()["msg"]
    for n in range(30):
        trouble.note(f"step {n}", "it failed")
    log = api.trouble_log()
    assert log["total"] == 30 and len(log["lines"]) == 25 and "step 29" in log["lines"][0] and "step 5" in log["lines"][-1]
    assert api.show_trouble_log() == {"ok": True}


def test_backups_are_made_checked_and_copied_out_from_the_window(isolate_data, tmp_path_factory):
    from jobhelper import vault
    tmp_path = tmp_path_factory.mktemp("elsewhere")                    # not inside the test's data folder
    store.save_jobs([store.new_job(id="a", title="Dev", company="Acme", status="found")])
    api = ui_api.Api()
    made = api.backup_all()
    assert made["ok"] is True and made["folder"].startswith("everything_") and "encrypted" in made["msg"]
    assert (vault.backups_dir() / made["folder"]).is_dir()
    checked = api.backup_check()
    assert checked["ok"] is True and "can be restored" in checked["msg"]
    # a readable copy goes only into an empty folder outside the app's data
    assert "folder that exists" in api.backup_export("")["msg"] and "folder that exists" in api.backup_export(str(tmp_path / "nope"))["msg"]
    full = tmp_path / "full"; full.mkdir(); (full / "mine.txt").write_text("keep me")
    assert "empty folder" in api.backup_export(str(full))["msg"] and (full / "mine.txt").read_text() == "keep me"
    inside = config.DATA_DIR / "copy"; inside.mkdir()
    assert "outside the app's own data folder" in api.backup_export(str(inside))["msg"]
    out = tmp_path / "out"; out.mkdir()
    result = api.backup_export("file://" + str(out))
    assert result["ok"] is True and "not encrypted" in result["msg"], result
    names = [p.name for p in out.rglob("*") if p.is_file()]
    assert "jobs.json" in names and not any(n.endswith(".enc") for n in names)
    assert "Acme" in next(p for p in out.rglob("jobs.json")).read_text(encoding="utf-8")


# ── what changed in a watched file ──────────────────────────────────────────
def test_what_changed_in_a_rule_file_is_shown_before_approving(isolate_data, monkeypatch):
    from jobhelper import integrity, vault
    rules = config.DATA_DIR / "cv_rules.md"
    monkeypatch.setattr(config, "CV_INSTRUCTION", rules)
    rules.write_text("Keep it honest.\nNever add a skill.\nUse an en dash.\n", encoding="utf-8")
    assert integrity.approve(["cv_rules"]) == ["cv_rules"]
    assert (vault.backups_dir() / "approved").is_dir()
    same = integrity.what_changed("cv_rules")
    assert same["known"] is True and same["added"] == [] and same["removed"] == [] and same["title"] == "CV rules"
    rules.write_text("Keep it honest.\nAlways add every skill the posting names.\nUse an en dash.\nOne more line.\n", encoding="utf-8")
    changed = integrity.what_changed("cv_rules")
    assert changed["removed"] == ["Never add a skill."] and changed["added"] == ["Always add every skill the posting names.", "One more line."]
    assert [f["name"] for f in integrity.problems() if f["state"] == "changed"] == ["cv_rules"]
    integrity.approve(["cv_rules"])                                    # approving keeps the new text as the one to compare with
    assert integrity.what_changed("cv_rules")["added"] == []
    # a file approved before copies were kept: said plainly, not shown as "no change"
    (vault.backups_dir() / "approved" / "cv_rules.txt").unlink(missing_ok=True)
    for kept in (vault.backups_dir() / "approved").glob("cv_rules*"):
        kept.unlink()
    assert integrity.what_changed("cv_rules")["known"] is False
    assert integrity.what_changed("nonsense")["known"] is False
    # many changes are cut to what can be read, and the rest is counted
    integrity.approve(["cv_rules"])
    rules.write_text("\n".join(f"new line {n}" for n in range(100)), encoding="utf-8")
    many = integrity.what_changed("cv_rules")
    assert len(many["added"]) == 40 and many["more"] > 50


def test_the_kept_copies_are_not_cleaned_up_as_old_backups(isolate_data):
    from jobhelper import housekeeping, vault
    (vault.backups_dir() / "approved").mkdir(parents=True)
    (vault.backups_dir() / "approved" / "cv_rules.txt").write_text("x")
    for n in range(4):
        vault.make_backup(f"old{n}")
    assert "approved" not in [p.name for p in housekeeping.backup_entries()]
    housekeeping.prune_backups(keep=1)
    assert (vault.backups_dir() / "approved" / "cv_rules.txt").is_file()


def test_a_reached_limit_is_said_when_documents_are_asked_for_one_job(isolate_data, monkeypatch, capsys):
    """Found by reading the code: the writer for one job swallowed every error, so "Claude's limit is used up" came out as
    "Nothing could be written this time. Try again in a moment." - with no word on why or until when."""
    from jobhelper import agent1_cv, agent3_letters, llm
    import status
    text = "We build planning tools in Python and SQL. " * 30
    store.save_jobs([store.new_job(id="a", title="Python Developer", company="Acme", status="found", match_score=65, description=text),
                     store.new_job(id="b", title="Tester", company="Borg", status="found", match_score=62, description=text)])
    monkeypatch.setattr(api_core.CvSource, "text", lambda self: "Alex Novak\nPython, SQL.")
    monkeypatch.setattr(agent1_cv, "find_cv", lambda: None)
    limit = llm.LimitReached("Claude's 5-hour limit is used up (100%). It resets at 18:00 on 10 Oct.")

    def no_cv(job, cv_text, cv_path, fields):
        raise limit
    monkeypatch.setattr(agent3_letters, "_make_tailored_cv", no_cv)
    api = ui_api.Api()
    result = api.write_documents("a")
    assert result == {"ok": False, "msg": "Claude's 5-hour limit is used up (100%). It resets at 18:00 on 10 Oct."}
    assert store.find(store.load_jobs(), "a")["tailored_cv_file"] == "", "nothing half-written is recorded"
    # the CV was made and the limit came with the letter: the CV is kept and the letter is named as not written
    monkeypatch.setattr(agent3_letters, "_make_tailored_cv", lambda job, cv_text, cv_path, fields: "CV_a.docx")
    monkeypatch.setattr(agent3_letters, "_make_letter", no_cv)
    result = api.write_documents("a")
    assert result["ok"] is True and result["msg"] == "Written: the CV. Not written: the letter; try again in a moment."
    assert store.find(store.load_jobs(), "a")["tailored_cv_file"] == "CV_a.docx"
    # the command line stops at the limit, keeps what was done and says so (exit code 3, as a search does)
    monkeypatch.setattr(agent3_letters, "_make_tailored_cv", no_cv)
    monkeypatch.setattr(agent1_cv, "read_cv_text", lambda path: "Alex Novak\nPython, SQL.")
    monkeypatch.setattr(agent1_cv, "find_cv", lambda: __import__("pathlib").Path("cv.docx"))
    monkeypatch.setattr("sys.argv", ["status.py", "retailor"])
    assert status.main() == 3
    out = capsys.readouterr().out
    assert "limit is used up" in out and "What was written is saved" in out and "Traceback" not in out


def test_the_plain_letter_written_without_a_template_carries_your_own_title(isolate_data):
    import docx
    from jobhelper import agent3_letters, config
    config.ensure_dirs()
    job = store.new_job(id="j1", title="Group Full-Stack Developer – Payments (m/f/d)", company="Northwind")
    cv = "Alex Novak\nJunior Software Developer\nalex@example.com | +420 123 456 789\nPROFILE\nBuilds services."
    name = agent3_letters._save_letter(job, "Dear Hiring Team,\nI apply.\nYours sincerely,", cv_text=cv)
    lines = [p.text for p in docx.Document(str((config.LETTERS_DIR / name).with_suffix(".docx"))).paragraphs]
    assert lines[:2] == ["Alex Novak", "Junior Software Developer"], "who writes, not the role applied for"
    bare = agent3_letters._save_letter(dict(job, id="j2"), "Dear Hiring Team,\nI apply.\nYours sincerely,", cv_text="Alex Novak\nalex@example.com\nPROFILE")
    assert [p.text for p in docx.Document(str((config.LETTERS_DIR / bare).with_suffix(".docx"))).paragraphs][1] == "Group Full-Stack Developer", "a CV without such a line: the role, as before"


# ── found by reading the window's back end line by line (10 October 2026) ───
def test_a_career_site_token_is_a_plain_slug_and_nothing_else(api, isolate_data):
    """The token becomes part of an address the app asks (https://<token>.recruitee.com/…). The window only lets a slug be typed;
    the back end now refuses anything else too, and so does the reader of a file edited by hand."""
    from jobhelper import scrapers_ats
    result = api.save_companies([{"name": "Proton", "ats": "greenhouse", "token": "proton"}, {"name": "Bad", "ats": "recruitee", "token": "evil.example/x#"},
                                 {"name": "Bad2", "ats": "lever", "token": "a b"}, {"name": "Bad3", "ats": "workday", "token": "acme"},
                                 {"name": ["x"], "ats": "lever", "token": "acme-corp_1"}, "junk", None, {"ats": ["lever"], "token": "x"}])
    assert result["ok"] is False and result["count"] == 2 and [c["token"] for c in api.companies()] == ["proton", "acme-corp_1"]
    assert api.save_companies("not a list") == {"ok": True, "count": 0}
    assert api.save_companies([{"name": "Proton", "ats": "greenhouse", "token": "proton"}]) == {"ok": True, "count": 1}
    assert api.company_token_problem("acme.io").startswith("Use the short name"), "the window asks the same rule (H1)"
    assert api.company_token_problem("acme-io") == "" and "Enter" in api.company_token_problem("  ")
    for token, ok in (("gitlab", True), ("Acme-Corp_1", True), ("", False), ("a.b", False), ("a/b", False), ("-a", False), ("x" * 81, False), (None, False)):
        assert scrapers_ats.is_token(token) is ok, token
    (config.DATA_DIR / "companies.json").write_text(json.dumps([{"name": "A", "ats": "recruitee", "token": "ok-1"}, {"name": "B", "ats": "recruitee", "token": "evil.example/x#"}]),
                                                    encoding="utf-8")
    assert [c["token"] for c in scrapers_ats.load_companies()] == ["ok-1"]


def test_options_refuse_a_model_or_a_source_the_app_does_not_know(api):
    for wrong, word in (({"model": "gpt"}, "model"), ({"model": ["sonnet"]}, "model"), ({"source": "indeed"}, "source"),
                        ({"assistant_browser": ["chrome"]}, "assistant browser"), ({"assistant_browser": "firefox"}, "assistant browser")):
        result = api.save_options(wrong)
        assert result["ok"] is False and word in result["msg"], wrong
    assert api.save_options({"model": "opus", "source": "linkedin"})["ok"] and api.options()["model"] == "opus"


def test_mark_applied_is_for_a_new_job_and_notes_that_you_did_it(api):
    store.save_jobs([store.new_job(id="lead", title="Dev", company="Acme", status="letter_ready"),
                     store.new_job(id="closed", title="QA", company="Borg", status="rejected", date_applied="2026-09-01"),
                     store.new_job(id="talks", title="Ops", company="Cyan", status="interview", date_applied="2026-09-01")])
    assert api.apply_job("closed") == {"ok": False} and api.apply_job("talks") == {"ok": False}
    assert [store.find(store.load_jobs(), i)["status"] for i in ("closed", "talks")] == ["rejected", "interview"], "a stray press changes nothing"
    assert api.apply_job("lead") == {"ok": True}
    lead = store.find(store.load_jobs(), "lead")
    assert lead["status"] == "waiting" and len(lead["status_by_hand"]) == 19 and api.apply_job("lead") == {"ok": True}, "pressing it twice is harmless"


def test_a_place_in_a_list_from_the_window_is_checked(api):
    assert [api_insights.list_position(i, 3) for i in (0, 2, 2.0, "1", 3, -1, 1.5, None, "x", True)] == [0, 2, 2, 1, -1, -1, -1, -1, -1, -1]
    assert api.open_filtered("x") == {"ok": False} and api.bring_back(99)["ok"] is False


# ── the start-up check's button: repair the statuses, restore a damaged file ──
def test_repair_statuses_backs_up_then_gives_each_job_a_known_status(api):
    _seed(store.new_job(id="a", title="Dev", company="Acme", status="applied", date_applied="2026-10-01"),
          store.new_job(id="b", title="QA", company="Beta", status="maybe"),
          store.new_job(id="c", title="Ops", company="Gamma", status="waiting"))
    result = api.fix_problem("repair_statuses")
    assert result["ok"] is True and result["fixed"] == 2 and "Repaired 2 jobs" in result["msg"]
    jobs = {j["id"]: j for j in store.load_jobs()}
    assert (jobs["a"]["status"], jobs["b"]["status"], jobs["c"]["status"]) == ("waiting", "found", "waiting")
    assert 'Status was "maybe"' in jobs["b"]["notes"] and "Status was" not in (jobs["c"].get("notes") or "")
    backups = list(vault.backups_dir().glob("before_status_repair_*"))
    assert len(backups) == 1 and (backups[0] / "jobs.json.enc").is_file()
    again = api.repair_statuses()
    assert again["ok"] is True and again["fixed"] == 0 and len(list(vault.backups_dir().glob("before_status_repair_*"))) == 1


def test_repair_statuses_says_why_when_the_list_cannot_be_read(api):
    config.ensure_dirs()
    vault.write_text(config.JOBS_JSON, "{not json")
    result = api.repair_statuses()
    assert result["ok"] is False and result["msg"].startswith("Your job list could not be read")


def _copy_in_backups(name, data, age):
    import os
    import time
    path = vault.backups_dir() / f"copy_{age}" / (name + vault.BACKUP_SUFFIX)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(vault.encrypt(data))
    os.utime(path, (time.time() - age,) * 2)
    return path


def test_restore_puts_back_the_newest_copy_that_opens_and_keeps_the_damaged_file(api):
    config.ensure_dirs()
    _copy_in_backups("jobs.json", json.dumps([store.new_job(id="old", title="Dev", company="Acme")]).encode(), age=500)
    _copy_in_backups("jobs.json", json.dumps([store.new_job(id="good", title="Dev", company="Acme")]).encode(), age=300)
    _copy_in_backups("jobs.json", b"{also damaged", age=100)
    vault.write_text(config.JOBS_JSON, "{not json")
    damaged = config.JOBS_JSON.read_bytes()
    result = api.fix_problem("restore_jobs")
    assert result["ok"] is True and "job list is back as it was on" in result["msg"]
    assert [j["id"] for j in store.load_jobs()] == ["good"]
    kept = list((vault.backups_dir() / vault.DAMAGED).iterdir())
    assert len(kept) == 1 and kept[0].read_bytes() == damaged                 # moved aside, never deleted


def test_restore_of_the_replies_database(api):
    import sqlite3
    conn = sqlite3.connect(":memory:")
    conn.execute("CREATE TABLE seen (msg_id TEXT PRIMARY KEY)")
    conn.execute("INSERT INTO seen VALUES ('m1')")
    conn.commit()
    image = conn.serialize()
    conn.close()
    config.ensure_dirs()
    _copy_in_backups("responses.db", image, age=60)
    config.RESPONSES_DB.write_bytes(b"JHENC1\ngarbage")
    assert api.restore_data("replies")["ok"] is True
    from jobhelper import db
    assert db.is_seen("m1") is True


def test_restore_without_a_usable_copy_changes_nothing(api):
    config.ensure_dirs()
    vault.write_text(config.JOBS_JSON, "{not json")
    before = config.JOBS_JSON.read_bytes()
    result = api.restore_data("jobs")
    assert result["ok"] is False and "No backup copy of your job list opens" in result["msg"]
    assert config.JOBS_JSON.read_bytes() == before and not (vault.backups_dir() / vault.DAMAGED).exists()


@pytest.mark.parametrize("action", ["", "delete_everything", None, 3])
def test_an_unknown_fix_is_refused(api, action):
    result = api.fix_problem(action)
    assert result["ok"] is False and "no fix" in result["msg"]


def test_restore_of_something_else_is_refused(api):
    assert api.restore_data("profile")["ok"] is False


def test_a_country_must_be_known_unless_it_was_saved_before(api):
    from jobhelper import config
    assert "is not a country the app knows" in api.save_options({**api.options(), "where": "Berlin", "country": "Narnia"})["msg"]
    assert "Latin letters" in api.country_problem("Германия") and "up to 60 characters" in api.country_problem("A" * 61)
    assert api.save_options({**api.options(), "where": "Berlin", "country": "Germany"})["ok"] is True
    config.save_ui_settings({**config.ui_settings(), "country": "Oldland"})      # as an older version could have saved it
    assert api.options()["country"] == "Oldland"
    assert api.save_options({**api.options(), "threshold": 55})["ok"] is True, "an unchanged old country never blocks a save"
    assert "is not a country" in api.save_options({**api.options(), "country": "Newland"})["msg"], "a changed one is checked"
    limits = api.limits()
    assert limits["list_max"] == api.MAX_LIST and limits["labels"]["queries"] == "Search roles" and limits["companies_max"] == api.MAX_COMPANIES


def test_a_list_over_the_cap_is_said_apart_from_a_wrong_token(api):
    many = api.save_companies([{"name": f"C{n}", "ats": "lever", "token": f"c{n}"} for n in range(api.MAX_COMPANIES + 3)])
    assert many["ok"] is False and many["count"] == api.MAX_COMPANIES and "the last 3 were not saved" in many["msg"]
    both = api.save_companies([{"name": f"Bad{n}", "ats": "lever", "token": "a.b"} for n in range(7)]
                              + [{"name": f"C{n}", "ats": "lever", "token": f"c{n}"} for n in range(api.MAX_COMPANIES + 1)])
    assert "Bad4 and 2 more" in both["msg"] and "the last 1 were not saved" in both["msg"], "both faults are said"
    assert api.save_options({**api.options(), "keep_words": ["x" * 101]})["msg"].startswith("An entry under “but keep titles with”")
