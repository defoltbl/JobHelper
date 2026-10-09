from openpyxl import load_workbook

from jobhelper import agent4_tracker, email_client, store


def test_new_jobs_hides_low_scores(isolate_data, monkeypatch):
    from jobhelper import config
    monkeypatch.setattr(config, "MIN_LEAD_SCORE", 40)
    jobs = [
        store.new_job(id="a", title="Good", company="A", match_score=70, status="found"),
        store.new_job(id="b", title="Weak", company="B", match_score=12, status="found"),
        store.new_job(id="c", title="Unscored", company="C", match_score=None, status="found"),
    ]
    agent4_tracker.render_new_jobs(jobs)
    nw = load_workbook(isolate_data / "new_jobs.xlsx").active
    titles = {nw.cell(row=r, column=2).value for r in range(2, nw.max_row + 1)}
    assert "Good" in titles and "Unscored" in titles
    assert "Weak" not in titles  # scored below floor → hidden


def test_render_splits_leads_and_applications(isolate_data):
    jobs = [
        store.new_job(id="a", title="Dev", company="Acme", match_score=80, status="letter_ready"),
        store.new_job(id="b", title="Eng", company="DataCorp", match_score=50, status="found"),
        store.new_job(id="c", title="QA", company="Applied Co", status="waiting"),
        store.new_job(id="d", title="Ops", company="Rejected Co", status="rejected"),
    ]
    agent4_tracker.render_all(jobs)
    assert isolate_data.joinpath("applications.xlsx").exists()
    assert isolate_data.joinpath("new_jobs.xlsx").exists()

    # New-jobs file has only the leads (found / letter_ready).
    nw = load_workbook(isolate_data / "new_jobs.xlsx").active
    nw_titles = {nw.cell(row=r, column=2).value for r in range(2, nw.max_row + 1)}
    assert nw_titles == {"Dev", "Eng"}

    # Applications file has only the applied/rejected ones, with progress columns.
    ap = load_workbook(isolate_data / "applications.xlsx").active
    ap_headers = [c.value for c in ap[1]]
    assert "Stages" in ap_headers and "Result" in ap_headers
    ap_titles = {ap.cell(row=r, column=2).value for r in range(2, ap.max_row + 1)}
    assert ap_titles == {"QA", "Ops"}


def test_company_tokens_strips_noise():
    toks = email_client.company_tokens("Acme Group s.r.o.")
    assert "acme" in toks
    assert "sro" not in toks  # stopword-ish removed


def test_match_job_by_company():
    jobs = [store.new_job(id="a", company="ESET Research"),
            store.new_job(id="b", company="Siemens Mobility")]
    m = email_client.match_job("careers@eset.cz", "Interview invitation", "About your ESET application", jobs)
    assert m["id"] == "a"


def test_match_job_none_when_no_overlap():
    jobs = [store.new_job(id="a", company="ESET Research")]
    assert email_client.match_job("noreply@newsletter.com", "Weekly digest", "jobs for you", jobs) is None


def test_match_job_is_word_level_not_substring():
    # "der" (from "DER Touristik") must NOT match inside "under review".
    jobs = [
        store.new_job(id="der", company="DER Touristik CZ a.s"),
        store.new_job(id="ms", company="Microsoft"),
    ]
    m = email_client.match_job("Microsoft Careers <no-reply@microsoft.com>",
                               "Thank you for your application!",
                               "Your application is under review.", jobs)
    assert m["id"] == "ms"


def test_stage_to_status_mapping():
    assert email_client.STAGE_TO_STATUS["interview"] == "interview"
    assert email_client.STAGE_TO_STATUS["rejection"] == "rejected"
    assert email_client.STAGE_TO_STATUS["offer"] == "offer"


def test_the_excel_table_shows_what_you_ask_in_your_currency(isolate_data):
    from openpyxl import load_workbook

    from jobhelper import agent4_tracker, config, store, ui_api
    jobs = [store.new_job(id="a", title="Python Developer", company="Acme", status="found", desired_monthly_czk=58000, desired_annual_czk=696000)]
    api = ui_api.Api()
    assert api.save_options({**api.options(), "where": "Berlin", "salary_target": 3200, "salary_currency": "EUR"})["ok"] is True
    agent4_tracker.render_all(jobs)
    sheet = load_workbook(config.NEW_JOBS_EXCEL).active
    head = [c.value for c in sheet[1]]
    row = dict(zip(head, [c.value for c in sheet[2]]))
    assert row["Desired Monthly (EUR)"] == 3200 and row["Desired Annual (EUR)"] == 38400 and "Desired Monthly (CZK)" not in head
    assert api.save_options({**api.options(), "where": "Prague", "salary_target": 0})["ok"] is True
    agent4_tracker.render_all(jobs)
    sheet = load_workbook(config.NEW_JOBS_EXCEL).active
    row = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert row["Desired Monthly (CZK)"] == 58000 and row["Desired Annual (CZK)"] == 696000       # the app's estimate, in crowns
