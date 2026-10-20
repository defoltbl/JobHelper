"""The window, loaded offscreen: dates, company notes, replies, answers for a form and the text fields.

Its tests share one window and follow on from each other, so they run on one worker in this order (see
tests/window_harness.py). The other window test files run beside it, each with a window of its own."""
from datetime import date, timedelta
import re

import pytest

pytest.importorskip("PySide6")
pytestmark = pytest.mark.xdist_group("test_window_jobs")

from PySide6.QtCore import QMetaObject, QObject  # noqa: E402

from jobhelper import api_core, config, store, ui_api  # noqa: E402
from tests.conftest import ROOT  # noqa: E402
from tests.window_harness import (Ui, _border_with_focus, _open, _open_job, _until, _until_closed, _until_open, _until_true, qml_source)  # noqa: E402


# ── dates, company notes, pasted replies, CV comparison, pay ────────────────


def test_a_date_on_a_job_shows_on_the_main_screen_and_is_announced(ui):
    card = ui.find("upcomingCard")
    assert card.property("visible") is False                                       # nothing planned: no empty card
    _open_job(ui)
    from jobhelper import planner
    assert ui.find("eventKinds").property("count") == len(planner.KINDS)   # the kinds the check accepts, from the back end
    ui.js('evDate.text = "next week"; detailDialog.saveEvent()')
    assert "YYYY-MM-DD" in ui.find("eventErr").property("text") and ui.find("eventErr").property("visible") is True
    assert ui.find("detailDialog").property("opened") is True and ui.js("upcomingData.length") == 0
    tomorrow = (date.today() + timedelta(days=1)).isoformat()
    ui.js(f'evDate.text = "{tomorrow}"; evTime.text = "14:30"; evNote.text = "Teams call"; detailDialog.eventKind = "call"; detailDialog.saveEvent()')
    assert ui.js("detailDialog.eventErr") == "" and ui.js("detailDialog.job.event.kind") == "call"
    assert ui.js("upcomingData.length") == 1 and ui.js("upcomingData[0].label") == "tomorrow"
    assert card.property("visible") is True
    ui.sent.clear()
    assert ui.js("bridge.remind()") == 1 and ui.js("bridge.remind()") == 0         # once a day
    assert ui.sent == [("Tomorrow at 14:30: call · Beta", "Reminder")]
    ui.js('detailDialog.show("app")')                                              # reopening shows what was saved
    _until_open(ui, "detailDialog")
    assert ui.find("evDate").property("text") == tomorrow and ui.find("evNote").property("text") == "Teams call"
    ui.js("bridge.clearEvent(detailDialog.job.id); detailDialog.refresh(); upcomingData = bridge.upcoming()")
    assert card.property("visible") is False and ui.find("evDate").property("text") == ""
    ui.js("detailDialog.close()")
    _until_closed(ui, "detailDialog")
    assert ui.warnings == []


def test_company_notes_are_kept_for_the_company(ui):
    from jobhelper import planner
    _open_job(ui)
    ui.js('cnContact.text = "Jana, HR"; cnSalary.text = "55k"; cnAsked.text = "SQL joins"; cnNotes.text = "Nice team"; detailDialog.saveNotes()')
    assert ui.js("detailDialog.notesInfo") == "Saved for Beta."
    assert planner.company_notes("beta")["asked"] == "SQL joins"
    ui.js("detailDialog.close()")
    _until_closed(ui, "detailDialog")
    _open_job(ui)
    assert ui.find("cnContact").property("text") == "Jana, HR" and ui.find("cnNotes").property("text") == "Nice team"
    ui.js('cnContact.text = ""; cnSalary.text = ""; cnAsked.text = ""; cnNotes.text = ""; detailDialog.saveNotes()')
    assert planner.company_notes("Beta")["updated"] == ""
    ui.js("detailDialog.close()")
    _until_closed(ui, "detailDialog")
    assert ui.warnings == []


def test_a_pasted_reply_is_filed_from_the_job_window(ui, monkeypatch):
    import time

    from jobhelper import db, llm
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"stage": "acknowledged", "summary": "They got it."})
    _open_job(ui)
    assert ui.js("detailDialog.replyOpen") is False
    ui.js('detailDialog.replyOpen = true; replyText.text = "ok"; detailDialog.fileReply()')
    assert "too short" in ui.js("detailDialog.replyErr") and ui.js("detailDialog.replyBusy") is False
    ui.js('replyText.text = "Hello, thank you for applying. We received your application and will be in touch."; detailDialog.fileReply()')
    assert ui.js("detailDialog.replyBusy") is True
    for _ in range(200):
        time.sleep(0.02)                                  # let the background worker run
        ui.wait(10)
        if ui.js("detailDialog.replyBusy") is False:
            break
    assert ui.js("detailDialog.replyErr") == "", ui.js("detailDialog.replyErr")
    assert ui.js("detailDialog.replyInfo") == "Filed as: acknowledged." and ui.js("detailDialog.replyOpen") is False
    assert ui.find("replyText").property("text") == ""
    rows = [r for r in db.all_responses() if r["msg_id"].startswith("pasted-")]
    assert len(rows) == 1 and rows[0]["job_id"] == "app"
    ui.js("detailDialog.close()")
    _until_closed(ui, "detailDialog")
    assert ui.warnings == []


def test_cv_comparison_window(ui, monkeypatch):
    import docx

    from jobhelper import agent1_cv

    def make(path, lines):
        d = docx.Document()
        for line in lines:
            d.add_paragraph(line)
        d.save(str(path))
        return path

    master = make(config.DATA_DIR / "master.docx", ["Alex Novak", "PROFILE", "I know Python."])
    make(config.TAILORED_CV_DIR / "cv_lead.docx", ["Alex Novak", "PROFILE", "I know Python and SQL."])
    monkeypatch.setattr(agent1_cv, "find_cv", lambda: master)
    _open_job(ui, "lead")
    assert ui.find("compareCv").property("visible") is False                       # no adapted CV yet
    ui.js("compareDialog.show(detailDialog.job.id)")
    ui.wait(150)
    assert ui.find("compareDialog").property("visible") is False                   # and nothing opens
    ui.bridge._api.editor.change("lead", lambda jobs, j: j.update(tailored_cv_file="cv_lead.docx", salary="50 000 - 60 000 CZK"), render=False)
    ui.js("detailDialog.refresh()")
    assert ui.find("compareCv").property("visible") is True
    assert ui.js("detailDialog.job.pay.low") == 50000
    ui.js("compareDialog.show(detailDialog.job.id)")
    _until_open(ui, "compareDialog")
    dialog = ui.find("compareDialog")
    assert dialog.property("opened") is True
    assert "1 paragraph changed" in ui.find("compareSummary").property("text")
    assert ui.js("compareDialog.data.rows[0].adapted") == "I know Python<b> and SQL</b>."
    assert ui.js("compareDialog.data.rows[0].master") == "I know Python."
    ui.js("compareDialog.close()")
    _until_closed(ui, "compareDialog")
    ui.js("detailDialog.close()")
    _until_closed(ui, "detailDialog")
    assert ui.warnings == []


def test_insights_list_what_postings_pay(ui, monkeypatch):
    monkeypatch.setattr(api_core.CvSource, "text", lambda self: "Python and Docker")
    ui.bridge._api.editor.change("lead", lambda jobs, j: j.update(salary="50 000 - 60 000 CZK", desired_monthly_czk=55000), render=False)
    ui.js('openView("ins")')
    _until_open(ui, "tableDialog")
    count = ui.find("listView").property("count")
    rows = [(ui.js(f"listView.model[{i}].primary"), ui.js(f"listView.model[{i}].secondary"), ui.js(f"listView.model[{i}].value")) for i in range(count)]
    pay = [r for r in rows if r[0] == "Junior Python Developer  ·  Acme"]
    assert pay and pay[0][2] == "60k" and "50–60k CZK" in pay[0][1] and "your target 55k: around your target" in pay[0][1]
    assert any("1 of 2 postings state it" in str(r[1]) for r in rows)
    ui.js("tableDialog.close()")
    _until_closed(ui, "tableDialog")
    assert ui.warnings == []


def test_scan_all_folders_switch_is_saved(ui):
    dialog = ui.find("optDialog")
    assert ui.bridge.options()["scan_all_folders"] is False
    ui.js("optsData = bridge.options()")
    _open(ui, dialog, 450)
    switch = ui.find("foldersSw")
    assert switch.property("checked") is False
    ui.bridge.saveOptions({**ui.bridge.options(), "scan_all_folders": True})
    ui.js("close()", dialog)
    _until_closed(ui, dialog)
    ui.js("optsData = bridge.options()")
    _open(ui, dialog, 450)
    assert switch.property("checked") is True
    ui.js("close()", dialog)
    _until_closed(ui, dialog)
    ui.bridge.saveOptions({**ui.bridge.options(), "scan_all_folders": False})
    assert ui.warnings == []


def test_application_only_cards_are_hidden_for_a_new_job_and_nothing_overflows(ui, monkeypatch):
    from jobhelper import planner
    for name in ("cv_lead.docx", "cv_lead.pdf"):
        (config.TAILORED_CV_DIR / name).write_bytes(b"x")
    for name in ("L_lead.docx", "L_lead.pdf"):
        (config.LETTERS_DIR / name).write_bytes(b"x")
    ui.bridge._api.editor.change("lead", lambda jobs, j: j.update(tailored_cv_file="cv_lead.docx", cover_letter_file="L_lead.docx",
                                                             url="https://example.com/job"), render=False)
    monkeypatch.setattr(ui_api.Api, "comparable", lambda self, j: True)
    monkeypatch.setattr(ui_api.Api, "editable_letter", lambda self, j: True)
    _open_job(ui, "lead")
    for card in ("eventCard", "replyCard", "notesCard"):
        assert ui.find(card).property("visible") is False, card
    dialog_width = ui.js("detailDialog.width")
    footer = ui.find("detailFooter")
    assert footer.property("width") <= dialog_width                              # every button on, still inside the window
    assert ui.find("detailDocs").property("height") > 40                         # the document buttons wrapped to a second line
    assert ui.js("detailFlick.width") <= dialog_width
    planner.save_company_notes("Acme", {"contact": "Jana"})                      # notes kept from an earlier application
    ui.js("detailDialog.refresh()")
    assert ui.find("notesCard").property("visible") is True and ui.find("eventCard").property("visible") is False
    planner.save_company_notes("Acme", {})
    ui.js("detailDialog.close()")
    _until_closed(ui, "detailDialog")
    _open_job(ui, "app")
    for card in ("eventCard", "replyCard", "notesCard"):
        assert ui.find(card).property("visible") is True, card
    ui.js("detailDialog.close()")
    _until_closed(ui, "detailDialog")
    ui.bridge._api.editor.change("lead", lambda jobs, j: j.update(tailored_cv_file="", cover_letter_file="", url=""), render=False)
    assert ui.warnings == []


def test_sidebar_counts_and_icons_share_one_column(ui):
    source = qml_source()
    assert 'objectName: "navTail_" + modelData.k; implicitWidth: 46' in source
    tail = source[source.index('objectName: "navTail_"'):][:1400]
    assert tail.count("anchors.centerIn: parent") >= 3                             # the pill, its number and the icon


def test_analytics_has_a_drawn_bar_chart_for_its_icon(ui):
    """A font's symbol for it could not be recognised at that size; three rising bars can."""
    icons = [i for i in ui.seen("navPanel") if i.objectName() == "navIcon_ana"]
    assert len(icons) == 5 and [i.property("visible") for i in icons].count(True) == 1     # one per row, shown on its own row only
    icon = next(i for i in icons if i.property("visible"))
    bars = sorted((b for b in icon.childItems() if b.property("radius") is not None), key=lambda b: b.property("x"))
    assert [b.property("height") for b in bars] == [6, 10, 14]                              # rising, left to right
    assert all(b.property("y") + b.property("height") == icon.property("height") for b in bars), "all stand on one line"
    assert all(b.property("x") + b.property("width") <= icon.property("width") for b in bars)
    assert "▦" not in qml_source()
    bulbs = [i for i in ui.seen("navPanel") if i.objectName() == "navIcon_ins" and i.property("visible")]
    assert len(bulbs) == 1 and len(bulbs[0].childItems()) == 3 and "✦" not in qml_source()    # drawn too: glass, neck, base
    assert ui.warnings == []


def test_filtered_out_list_opens_from_the_last_run_card(ui):
    from jobhelper import agent2_search
    agent2_search.save_dropped([
        {"title": "Software Engineering Intern", "company": "Acme", "source": "linkedin", "url": "https://example.com/1",
         "group": "skip word", "reason": 'The title has a word you skip ("intern").'},
        {"title": "Backend Developer", "company": "Beta", "source": "jobs.cz", "url": "",
         "group": "too many years", "reason": "Asks for 6+ years of experience."}])
    ui.bridge._api.note_run(filtered=2)
    ui.js("lastRunData = bridge.lastRun()")
    ui.wait(100)
    source = qml_source()
    assert 'lastRunData.filtered + " · see why ›"' in source and "win.openView(modelData.view)" in source
    ui.js('openView("skip")')
    _until_open(ui, "tableDialog")
    assert ui.find("listView").property("count") == 2
    assert ui.js("listView.model[0].primary") == "Software Engineering Intern  ·  Acme"
    assert ui.js("listView.model[1].secondary") == "Asks for 6+ years of experience.   ·  jobs.cz"
    assert ui.js("listView.model[0].actions.length") == 1 and ui.js("listView.model[1].actions.length") == 0
    assert ui.js("tableDialog.groups.length") == 2                                # a filter chip per reason
    ui.js("tableDialog.close()")
    _until_closed(ui, "tableDialog")
    agent2_search.save_dropped([])
    assert ui.warnings == []


def test_polish_empty_filtered_row_older_master_note_and_old_letter_chip(ui, monkeypatch):
    source = qml_source()
    assert '" · see why ›" : "after next search"' in source                      # not a bare dash that looks broken
    assert 'letter_old_format === true ? "letter: older format, open in Word"' in source
    note = ui.find("olderMaster")
    monkeypatch.setattr(ui_api.Api, "cv_compare", lambda self, job_id: {
        "ok": True, "company": "Acme", "title": "Dev", "file": "x.docx", "older_master": True,
        "rows": [], "same": 3, "changed": 0, "added": 0, "removed": 0})
    ui.js('compareDialog.show("lead")')
    _until_open(ui, "compareDialog")
    assert note.property("visible") is True
    ui.js("compareDialog.close()")
    _until_closed(ui, "compareDialog")
    monkeypatch.setattr(ui_api.Api, "cv_compare", lambda self, job_id: {
        "ok": True, "company": "Acme", "title": "Dev", "file": "x.docx", "older_master": False,
        "rows": [], "same": 3, "changed": 0, "added": 0, "removed": 0})
    ui.js('compareDialog.show("lead")')
    _until_open(ui, "compareDialog")
    assert note.property("visible") is False
    ui.js("compareDialog.close()")
    _until_closed(ui, "compareDialog")
    assert ui.warnings == []


def test_level_words_and_years_are_edited_in_options(ui):
    dialog = ui.find("optDialog")
    start = ui.bridge.options()
    ui.bridge.saveOptions({**start, "level_words": ["senior", "lead"], "max_years": 6})
    ui.js("optsData = bridge.options()")
    _open(ui, dialog, 450)
    assert ui.js("levelTags.all().join(',')") == "senior,lead" and ui.find("yearsSpin").property("value") == 6
    ui.js('levelTags.add("Principal")')
    assert ui.js("levelTags.all().join(',')") == "senior,lead,principal"
    ui.js("close()", dialog)
    _until_closed(ui, dialog)
    ui.bridge.saveOptions(start)
    assert ui.warnings == []


def test_help_has_a_what_changed_tab(ui):
    dialog = ui.find("helpDialog")
    _open(ui, dialog, 450)
    assert ui.js("helpDialog.tab") == "how" and ui.js("helpDialog.changes.length") >= 2
    ui.js('helpDialog.tab = "new"')
    ui.wait(150)
    first = ui.js("helpDialog.changes[0].date")
    assert (first == "Unreleased" or "2026" in first) and "<b>" in ui.js("helpDialog.changes.map(c => c.items.join(' ')).join(' ')")
    ui.js('helpDialog.tab = "how"; helpDialog.close()')
    _until_closed(ui, "helpDialog")
    assert ui.warnings == []


def test_bring_back_button_in_the_filtered_out_list(ui):
    from jobhelper import agent2_search
    posting = store.new_job(id="back1", title="Senior Tester", company="Gamma", url="https://example.com/b")
    agent2_search.save_dropped([
        {"title": "Senior Tester", "company": "Gamma", "source": "linkedin", "url": "https://example.com/b",
         "group": "level", "reason": 'The title has a level you skip ("senior").', "job": posting},
        {"title": "Twin", "company": "Gamma", "source": "jobs.cz", "url": "", "group": "duplicate", "reason": "R.",
         "job": store.new_job(id="twin", title="Twin", company="Gamma")}])
    before = ui.js("bridge.counts().leads")
    ui.js('openView("skip")')
    _until_open(ui, "tableDialog")
    assert [ui.js(f"listView.model[0].actions[{i}].label") for i in range(2)] == ["open ↗", "bring back"]
    assert ui.js("listView.model[1].actions.length") == 0                        # a duplicate would be removed again
    assert ui.js("bridge.bringBack(0).ok") is True
    ui.js("countsData = bridge.counts(); refreshOpenView()")
    ui.wait(300)
    assert ui.find("listView").property("count") == 1 and ui.js("countsData.leads") == before + 1
    ui.js("tableDialog.close()")
    _until_closed(ui, "tableDialog")
    jobs = [j for j in store.load_jobs() if j["id"] != "back1"]
    store.save_jobs(jobs)
    agent2_search.save_dropped([])
    ui.js("countsData = bridge.counts()")
    assert ui.warnings == []


# ── calendar file, thank-you note, answers, the last 7 days ─────────────────
def test_thank_you_card_appears_after_the_interview_day(ui, monkeypatch):
    opened = []
    monkeypatch.setattr(api_core.Desktop, "open", lambda self, target: opened.append(str(target)))
    today = date.today().isoformat()
    _open_job(ui)
    assert ui.find("thanksCard").property("visible") is False and ui.find("evCalendar").property("visible") is False
    ui.js(f'evDate.text = "{today}"; evTime.text = "00:00"; detailDialog.eventKind = "interview"; detailDialog.saveEvent()')
    assert ui.find("thanksCard").property("visible") is True and ui.find("evCalendar").property("visible") is True
    assert ui.js("bridge.exportEvent(detailDialog.job.id).ok") is True and opened[-1].endswith(".ics")
    ui.js("detailDialog.thanks = bridge.thanksDraft(detailDialog.job.id)")
    ui.wait(100)
    body = ui.find("thanksBody").property("text")
    assert "Thank you for the interview about the QA Engineer position" in body and "Beta" in body
    assert ui.js("bridge.markThanked(detailDialog.job.id)") is True
    ui.js("detailDialog.refresh()")
    assert ui.find("thanksCard").property("visible") is False                      # not offered twice for the same day
    ui.js("bridge.clearEvent(detailDialog.job.id); detailDialog.refresh(); upcomingData = bridge.upcoming(); detailDialog.close()")
    _until_closed(ui, "detailDialog")
    assert ui.warnings == []


def test_analytics_start_with_the_last_seven_days(ui):
    ui.js('openView("ana")')
    _until_open(ui, "tableDialog")
    assert ui.js("listView.model[0].primary") == "The last 7 days" and ui.js("listView.model[0].heading") is True
    labels = [ui.js(f"listView.model[{i}].primary") for i in range(1, 6)]
    assert labels == ["New jobs found", "Ready to send", "Applications sent", "Replies received", "Follow-ups due"]
    assert ui.js("listView.model[6].primary") == "All time"
    ui.js("tableDialog.close()")
    _until_closed(ui, "tableDialog")
    assert ui.warnings == []


def test_start_up_check_shows_one_line_only_when_something_is_wrong(ui):
    banner = ui.find("healthBanner")
    assert ui.js("healthData.checks.length") > 5                                  # the check ran when the window started
    start = ui.js("healthData.problems")
    ui.js('healthData = { ok: true, checks: [], problems: 0, message: "" }')
    assert banner.property("visible") is False
    ui.js('healthData = { ok: false, problems: 2, message: "2 things need your attention. The first: No master CV yet.", '
          'checks: [{ name: "Master CV", ok: false, problem: "No master CV yet." }, { name: "Claude", ok: false, problem: "Not signed in." }, '
          '{ name: "Job list", ok: true, problem: "" }] }')
    assert banner.property("visible") is True and "No master CV yet" in ui.find("healthText").property("text")
    ui.js("showHealth()")
    ui.wait(450)
    assert ui.find("listView").property("count") == 3
    assert ui.js("listView.model[0].secondary") == "No master CV yet." and ui.js("listView.model[2].secondary") == "fine"
    ui.js("tableDialog.close()")
    _until_closed(ui, "tableDialog")
    ui.js("healthHidden = true")
    assert banner.property("visible") is False
    ui.js("healthHidden = false; healthData = bridge.selfCheck(true)")
    assert ui.js("healthData.problems") <= max(start, 3)
    ui.js('healthData = { ok: true, checks: [], problems: 0, message: "" }')
    assert ui.warnings == []


def test_start_up_check_offers_the_fix_and_says_what_it_did(ui):
    banner, fix = ui.find("healthBanner"), ui.find("healthFix")
    ui.js('healthNote = ""; healthNoteOk = false; healthHidden = false')
    ui.js('healthData = { ok: false, problems: 1, message: "Your job list is damaged.", action: "restore_jobs", '
          'action_label: "Restore from backup", checks: [{ name: "Job list", ok: false, problem: "Your job list is damaged.", action: "restore_jobs" }] }')
    assert banner.property("visible") is True and fix.property("visible") is True and fix.property("text") == "Restore from backup"
    ui.js('healthData = { ok: false, problems: 1, message: "Paste your key.", action: "", action_label: "", checks: [] }')
    assert fix.property("visible") is False                                      # nothing the app can do: no button
    ui.js('healthData = { ok: false, problems: 1, message: "Odd.", action: "no_such_fix", action_label: "Repair", checks: [] }')
    ui.js("fixHealth()")                                                          # the back end refuses: a red line under it
    note = ui.find("healthNote")
    assert note.property("visible") is True and "no fix" in note.property("text") and ui.js("healthNoteOk") is False
    assert banner.property("visible") is True
    ui.js('healthNoteOk = true; healthNote = "Repaired 2 jobs."; healthData = { ok: true, checks: [], problems: 0, message: "" }')
    assert banner.property("visible") is True and ui.find("healthText").property("text") == "Repaired 2 jobs."   # green: put right
    assert fix.property("visible") is False
    ui.js('healthHidden = true; healthNote = ""; healthNoteOk = false')
    assert banner.property("visible") is False
    ui.js('healthHidden = false')
    assert ui.warnings == []


# ── the window is split into one file per window ────────────────────────────
def test_every_window_has_its_own_file_and_main_stays_small():
    folder = ROOT / "ui_qt"
    for name in ("TableDialog", "DetailDialog", "CompareDialog", "FormAnswersDialog", "HelpDialog", "PrepDialog", "LetterDialog",
                 "ResponseDialog", "LimitDialog", "WelcomeOverlay", "OptionsDialog", "ScheduleDialog", "SmallButton", "DetailCard"):
        assert (folder / f"{name}.qml").is_file(), name
    main = (folder / "Main.qml").read_text(encoding="utf-8")
    assert len(main.splitlines()) < 1400                                   # new windows go into their own file
    assert "\n    Dialog {" not in main and "\n    component " not in main


def test_a_split_file_never_hands_a_robot_its_own_empty_theme(ui):
    """Inside another file `theme: theme` means the robot's own (empty) property, not the app's colours."""
    for path in (ROOT / "ui_qt").glob("*.qml"):
        if path.name != "Main.qml":
            assert "theme: theme" not in path.read_text(encoding="utf-8"), path.name
    assert ui.js("logoBot.theme === win.appTheme && logoBot.theme.accent !== undefined") is True


def test_every_window_opens_from_the_main_file_without_warnings(ui):
    for name in Ui.WINDOWS:
        assert ui.win.findChild(QObject, name) is not None, name
    for code in ('openView("new")', "tableDialog.close()", "helpDialog.open()", "helpDialog.close()", "limitDialog.open()", "limitDialog.close()",
                 "schedDialog.open()", "schedDialog.close()", 'formDialog.show("app")', "formDialog.close()"):
        ui.js(code)
        ui.wait(380)
    assert ui.warnings == []


def test_start_up_check_still_runs_when_the_sign_in_check_is_slow(ui):
    """The splash gives up waiting after 6 seconds and marks the connection check as done;
    the self-check must still run when the answer arrives later."""
    ui.js('healthRan = false; healthData = { ok: true, checks: [], problems: 0, message: "" }; connChecked = true')
    ui.bridge.connectionsReady.emit(dict(ui.conn))
    ui.wait(150)
    assert ui.js("healthRan") is True and ui.js("healthData.checks.length") > 5
    ui.js('healthData = { ok: true, checks: [], problems: 0, message: "" }')
    ui.bridge.connectionsReady.emit(dict(ui.conn))                       # clean and already run: not repeated on every recheck
    ui.wait(150)
    assert ui.js("healthData.checks.length") == 0
    assert ui.warnings == []


# ── apply checklist, one PDF, search in Answers ─────────────────────────────
def test_apply_checklist_ticks_steps_and_marks_applied(ui, monkeypatch):
    opened = []
    monkeypatch.setattr(api_core.Desktop, "open", lambda self, target: opened.append(str(target)))
    monkeypatch.setattr(ui_api.Api, "combined_pdf", lambda self, job_id: {"ok": True, "file": "both.pdf", "pages": 3})
    monkeypatch.setattr(ui_api.Api, "can_combine", lambda self, j: bool(j.get("tailored_cv_file")))
    for folder, name in ((config.TAILORED_CV_DIR, "cv_lead.docx"), (config.LETTERS_DIR, "L_lead.docx")):
        (folder / name).write_bytes(b"x")
    ui.bridge._api.editor.change("lead", lambda jobs, j: j.update(tailored_cv_file="cv_lead.docx", cover_letter_file="L_lead.docx",
                                                             url="https://example.com/job", apply_steps=[]), render=False)
    _open_job(ui, "lead")
    assert ui.find("applyCard").property("visible") is True and ui.find("onePdf").property("visible") is True
    assert ui.js("detailDialog.applySteps.map(function (s) { return s.key; }).join(',')") == "posting,cv,letter,combined,answers"
    assert ui.find("applyCount").property("text").startswith("0 of 5 done")
    ui.js('detailDialog.doStep("posting"); detailDialog.doStep("cv"); detailDialog.doStep("combined")')
    assert opened[0] == "https://example.com/job" and "cv_lead." in opened[1]
    assert ui.js("detailDialog.stepsDone") == 3 and ui.js("detailDialog.job.apply_steps.join(',')") == "posting,cv,combined"
    ui.js('detailDialog.doStep("answers")')
    _until_open(ui, "formDialog")
    assert ui.find("formDialog").property("opened") is True and ui.js("formDialog.jobId") == "lead"
    assert ui.js("detailDialog.stepsDone") == 3                                   # opening is not doing: no tick yet
    ui.js('formDialog.fill({ items: [{ question: "Why us?", answer: "Because of your product.", check: "" }] }); formDialog.copy(0)')
    assert ui.js("detailDialog.stepsDone") == 4                                   # an answer was copied
    ui.js("formDialog.close()")
    _until_closed(ui, "formDialog")
    ui.js("detailDialog.close()")
    _until_closed(ui, "detailDialog")
    _open_job(ui, "lead")                                                         # the ticks are still there
    assert ui.find("applyCount").property("text").startswith("4 of 5 done")
    ui.js("detailDialog.close()")
    _until_closed(ui, "detailDialog")
    _open_job(ui, "app")
    assert ui.find("applyCard").property("visible") is False                      # already applied: no checklist
    ui.js("detailDialog.close()")
    _until_closed(ui, "detailDialog")
    ui.bridge._api.editor.change("lead", lambda jobs, j: j.update(tailored_cv_file="", cover_letter_file="", url="", apply_steps=[]), render=False)
    assert ui.warnings == []


def test_a_failed_step_is_not_ticked(ui, monkeypatch):
    monkeypatch.setattr(ui_api.Api, "combined_pdf", lambda self, job_id: {"ok": False, "msg": "The two PDF files could not be joined."})
    (config.TAILORED_CV_DIR / "cv_lead.docx").write_bytes(b"x")
    ui.bridge._api.editor.change("lead", lambda jobs, j: j.update(tailored_cv_file="cv_lead.docx", url="", apply_steps=[]), render=False)
    _open_job(ui, "lead")
    ui.js('detailDialog.doStep("combined"); detailDialog.doStep("posting")')      # no single PDF, and this job has no link
    assert ui.js("detailDialog.stepsDone") == 0 and ui.js("(detailDialog.job.apply_steps || []).length") == 0
    ui.js("detailDialog.close()")
    _until_closed(ui, "detailDialog")
    ui.bridge._api.editor.change("lead", lambda jobs, j: j.update(tailored_cv_file="", apply_steps=[]), render=False)
    assert ui.warnings == []


def test_about_you_shows_your_own_letters_folder(ui, monkeypatch):
    opened = []
    monkeypatch.setattr(api_core.Desktop, "open", lambda self, target: opened.append(target))
    config.COVER_LETTERS_DIR.mkdir(parents=True, exist_ok=True)
    (config.COVER_LETTERS_DIR / "Alex_Novak_CoverLetter_Acme.docx").write_bytes(b"mine")
    ui.js("aboutDialog.show()")
    _until_open(ui, "aboutDialog")
    assert ui.js("aboutDialog.myLetters.count") == 1 and ui.js("aboutDialog.myLetters.name") == config.COVER_LETTERS_DIR.name
    assert ui.find("myLettersOpen").property("text") == "Open folder"
    assert ui.js("bridge.openMyLetters().ok") is True and opened == [config.COVER_LETTERS_DIR]
    ui.js("close()", ui.find("aboutDialog"))
    _until_closed(ui, "aboutDialog")
    (config.COVER_LETTERS_DIR / "Alex_Novak_CoverLetter_Acme.docx").unlink()
    assert ui.warnings == []


def test_without_a_city_the_search_does_not_start_and_options_open(ui):
    start = ui.bridge.options()
    ui.bridge.saveOptions({**start, "where": ""})
    ui.js("optsData = bridge.options()")
    ui.wait(100)
    assert ui.js('cityEN(optsData.where) || "No city yet"') == "No city yet"
    ui.js("startRun()")
    _until_open(ui, "optDialog")
    assert ui.win.property("appRunning") is False                                 # nothing was started
    assert ui.find("optDialog").property("opened") is True                        # straight to where the city is set
    assert "city" in ui.js("ticker.text").lower()
    assert ui.js("cityCombo.editText") == ""                                      # and no city is filled in for you
    ui.js("optDialog.close()")
    _until_closed(ui, "optDialog")
    ui.bridge.saveOptions(start)
    ui.js("optsData = bridge.options()")
    assert ui.warnings == []


# ── hover that fades out cleanly ────────────────────────────────────────────
def test_hover_fades_to_the_same_colour_not_through_black(ui):
    """A colour animated to plain "transparent" passes through black and leaves a dark shadow behind."""
    assert ui.js('var c = clearOf(Qt.rgba(0.2, 0.4, 0.6, 1)); Math.round(c.r * 10) + "," + Math.round(c.g * 10) + "," + Math.round(c.b * 10) + "," + c.a') == "2,4,6,0"
    import re
    for path in sorted((ROOT / "ui_qt").glob("*.qml")):
        lines = path.read_text(encoding="utf-8").splitlines()
        for i, line in enumerate(lines):
            nearby = " ".join(lines[i:i + 3])
            if '"transparent"' in line and re.search(r"Behavior on (color|border\.color)", nearby) and "containsMouse" in line:
                raise AssertionError(f"{path.name}:{i + 1} fades a hover colour to plain transparent")


def test_insights_are_warmed_up_in_the_background(ui, monkeypatch):
    calls = []
    monkeypatch.setattr(ui_api.Api, "read_ahead", lambda self: calls.append(1) or 0)
    ui.js("bridge.warmUp()")
    assert _until(lambda: calls and not ui.bridge._warming.locked()) and calls == [1]
    monkeypatch.setattr(ui_api.Api, "read_ahead", lambda self: (_ for _ in ()).throw(RuntimeError("no data yet")))
    ui.js("bridge.warmUp()")                                   # a failure in the head start is not shown to anyone
    assert _until(lambda: not ui.bridge._warming.locked())
    assert "bridge.warmUp()" in qml_source() and ui.warnings == []


def test_a_second_warm_up_does_not_start_while_the_first_is_reading(ui, monkeypatch):
    import threading
    started, go_on = [], threading.Event()
    monkeypatch.setattr(ui_api.Api, "read_ahead", lambda self: started.append(1) or go_on.wait(2))
    ui.js("bridge.warmUp()")
    assert _until(lambda: started)
    ui.js("bridge.warmUp()")
    ui.js("bridge.warmUp()")
    go_on.set()
    assert _until(lambda: not ui.bridge._warming.locked()) and started == [1]


def test_insights_are_not_prepared_while_the_welcome_screen_is_up(ui, monkeypatch):
    """Preparing them is heavy enough to make the welcome animation stutter, so it waits for the screen to close."""
    calls = []
    monkeypatch.setattr(ui_api.Api, "read_ahead", lambda self: calls.append(1) or 0)
    welcome = ui.find("welcome")
    was = welcome.property("shown")
    welcome.setProperty("shown", True)
    ui.js("reloadAll()")
    ui.wait(1100)
    assert calls == []
    welcome.setProperty("shown", False)
    ui.wait(300)
    assert calls == [], "not at once either: the screen is still fading out"
    ui.wait(900)
    assert _until(lambda: calls == [1])
    ui.js("reloadAll()")                                       # afterwards, every reload refreshes it straight away
    assert _until(lambda: calls == [1, 1])
    welcome.setProperty("shown", was)
    ui.wait(1000)
    assert ui.warnings == []


def test_the_start_up_check_comes_with_the_connection_check(ui, monkeypatch):
    """Both are done off the window's thread; the window only shows the result."""
    monkeypatch.setattr(ui_api.Api, "connections", lambda self: {"claude": True})
    monkeypatch.setattr(ui_api.Api, "self_check", lambda self, signed_in=None: {"ok": True, "problems": [], "asked": signed_in})
    got = []
    ui.bridge.connectionsReady.connect(got.append)
    ui.bridge.checkConnections()
    for _ in range(100):
        ui.wait(20)
        if got:
            break
    import json
    assert got and json.loads(got[-1]["health"]) == {"ok": True, "problems": [], "asked": True}
    assert "JSON.parse(info.health)" in qml_source()


def test_the_window_shows_the_start_up_check_it_was_sent(ui):
    was = ui.js("JSON.stringify(healthData)")
    ui.js("healthRan = false")
    ui.bridge.connectionsReady.emit({**ui.conn, "health": '{"ok": false, "problems": 1, "message": "No master CV yet.", "checks": []}'})
    ui.wait(80)
    assert ui.js("healthData.message") == "No master CV yet." and ui.js("healthRan") is True
    ui.js("healthRan = false")
    ui.bridge.connectionsReady.emit(dict(ui.conn))             # no check came with it: the window asks itself
    ui.wait(80)
    assert ui.js("healthData.checks.length") > 0
    ui.js("healthData = JSON.parse('%s')" % was.replace("\\", "\\\\").replace("'", "\\'"))
    assert ui.warnings == []


# ── answers for one job's form ──────────────────────────────────────────────
def test_form_answers_are_drafted_per_job_edited_and_kept(ui, monkeypatch):
    import time

    from jobhelper import llm
    monkeypatch.setattr(api_core.CvSource, "text", lambda self: "Alex Novak\nPython, Docker.")
    monkeypatch.setattr(llm, "ask_json", lambda *a, **k: {"answers": [
        {"n": 1, "answer": "I like how Beta tests its products.", "needs_you": ""},
        {"n": 2, "answer": "I expect [your figure] CZK gross per month.", "needs_you": "Your salary figure"}]})
    source = qml_source()
    assert 'objectName: "answersButton"' not in source and "answersDialog" not in source      # the general Answers button is gone
    ui.js('formDialog.show("app")')
    _until_open(ui, "formDialog")
    dialog = ui.find("formDialog")
    assert dialog.property("opened") is True and ui.find("formFor").property("text") == "Beta  ·  QA Engineer"
    ui.js("formDialog.draft()")
    assert "Paste the questions" in ui.js("formDialog.err") and ui.js("formDialog.busy") is False     # red message, nothing asked
    ui.js('fmQuestions.text = (formDialog.info.common || []).join("\\n")')
    assert "Why do you want to work at Beta?" in ui.js("fmQuestions.text")
    ui.js('fmQuestions.text = "Why do you want to work at Beta?\\nWhat are your salary expectations?"; formDialog.draft()')
    assert ui.js("formDialog.busy") is True and ui.find("fmDraft").property("text") == "Asking Claude…"
    for _ in range(200):
        time.sleep(0.02)                                  # let the background worker run
        ui.wait(10)
        if ui.js("formDialog.busy") is False:
            break
    assert ui.js("formDialog.err") == "", ui.js("formDialog.err")
    assert ui.js("fmModel.count") == 2 and "1 need your attention" in ui.js("formDialog.note")
    assert ui.js("fmModel.get(1).check") == "Your salary figure" and ui.find("fmDraft").property("text") == "Draft again"
    ui.js('fmModel.setProperty(1, "answer", "I expect 60 000 CZK gross per month."); formDialog.dirty = true; formDialog.save()')
    assert ui.js("formDialog.note") == "Saved with this job." and ui.js("formDialog.dirty") is False
    ui.js("formDialog.close()")
    _until_closed(ui, "formDialog")
    ui.js('formDialog.show("app")')                        # they are still there, and only for this job
    _until_open(ui, "formDialog")
    assert ui.js("fmModel.get(1).answer") == "I expect 60 000 CZK gross per month."
    ui.js("formDialog.close()")
    _until_closed(ui, "formDialog")
    ui.js('formDialog.show("lead")')
    _until_open(ui, "formDialog")
    assert ui.js("fmModel.count") == 0 and ui.js("fmQuestions.text") == ""
    ui.js("formDialog.close()")
    _until_closed(ui, "formDialog")
    ui.bridge.saveFormAnswers("app", [])
    assert ui.warnings == []


def test_form_answers_show_the_reason_when_claude_cannot_be_asked(ui, monkeypatch):
    import time

    from jobhelper import llm

    def limit(*a, **k):
        raise llm.LimitReached("Claude's 7-day limit is used up (99%). It resets at 18:00 on 08 Oct.")
    monkeypatch.setattr(api_core.CvSource, "text", lambda self: "Alex Novak\nPython.")
    monkeypatch.setattr(llm, "ask_json", limit)
    ui.js('formDialog.show("app"); fmQuestions.text = "Why us?"; formDialog.draft()')
    for _ in range(200):
        time.sleep(0.02)
        ui.wait(10)
        if ui.js("formDialog.busy") is False:
            break
    assert "limit is used up" in ui.js("formDialog.err") and ui.find("formDialog").property("visible") is True    # stays open
    assert ui.js("fmQuestions.text") == "Why us?"                                 # what you pasted is not lost
    ui.js("formDialog.close()")
    _until_closed(ui, "formDialog")
    assert ui.warnings == []


def test_form_answers_explain_every_wrong_step(ui, monkeypatch):
    """Whatever the user does wrong in this window gets a clear message, and nothing is lost silently."""
    monkeypatch.setattr(ui_api.Api, "make_form_answers", lambda self, job_id, text: pytest.fail("nothing may be drafted here"))
    ui.bridge.saveFormAnswers("app", [{"question": "Why us?", "answer": "Because of your product.", "check": ""},
                                      {"question": "Notice period?", "answer": "My notice period is [your notice period].", "check": "Fill it in"},
                                      {"question": "Anything else?", "answer": "", "check": ""}])
    ui.js('formDialog.show("app")')
    _until_open(ui, "formDialog")
    note = ui.find("fmNote")
    # pasted text that cannot be used: said while typing, in red, and Draft does nothing
    ui.js('fmQuestions.text = "' + "\\n".join(f"Question number {i}?" for i in range(18)) + '"')
    assert "That is 18 questions" in ui.find("fmHint").property("text")
    ui.js("formDialog.draft()")
    assert "one draft takes up to 15" in ui.find("fmHint").property("text") and ui.js("formDialog.busy") is False
    assert ui.js("formDialog.err") == ""                                           # said once, under the box, not twice
    ui.js('fmQuestions.text = "' + "x" * 320 + '"')
    assert "characters long" in ui.find("fmHint").property("text")
    ui.js('fmQuestions.text = "Why us?\\nwhy us?\\nSalary?"')
    assert ui.find("fmHint").property("text") == "2 questions found · 1 repeated line was left out · one Claude call"
    # drafting again would replace answers: asked first, never done on one click
    ui.js("formDialog.draft()")
    assert "replaces the 3 answers below" in note.property("text") and ui.js("formDialog.busy") is False
    ui.js('fmQuestions.text = "Why us?"')                                         # changing the text takes the confirmation back
    assert ui.js("formDialog.confirmDraft") is False
    # copying something unfinished
    ui.js("formDialog.copy(2)")
    assert "still empty" in note.property("text")
    ui.js("formDialog.copy(1)")
    assert "square brackets" in note.property("text")
    ui.js("formDialog.copy(0)")
    assert note.property("text") == "Copied."
    # saving says what is still open
    ui.js("formDialog.save()")
    assert note.property("text") == "Saved with this job. 1 answer is still empty. 1 answer still has a part in [square brackets] to fill in."
    # closing with an unsaved change: one warning, then it closes
    ui.js('fmModel.setProperty(0, "answer", "Changed and not saved."); formDialog.dirty = true; formDialog.tryClose()')
    ui.wait(250)
    dialog = ui.find("formDialog")
    assert dialog.property("visible") is True and "did not save" in note.property("text")
    assert ui.js("formDialog.closePolicy === 0") is True                           # Esc cannot throw the change away either
    ui.js("formDialog.tryClose()")
    ui.wait(450)
    assert dialog.property("visible") is False
    assert ui.bridge.formAnswers("app")["items"][0]["answer"] == "Because of your product."      # the discarded change was not saved
    # closing while Claude is working
    ui.js('formDialog.show("app"); formDialog.busy = true; formDialog.tryClose()')
    ui.wait(300)
    assert dialog.property("visible") is True and "still drafting" in note.property("text")
    ui.js("formDialog.busy = false; formDialog.tryClose()")
    ui.wait(450)
    ui.bridge.saveFormAnswers("app", [])
    assert ui.warnings == []


def test_a_message_behind_an_open_window_is_shown_on_top_of_it(ui):
    """The status line is hidden by any open window, so results and errors must also appear above it."""
    toast = ui.find("toast")
    ui.js('ticker.text = "said with no window open"')
    ui.wait(250)
    assert toast.property("visible") is False                                      # the status line itself is in view
    _open_job(ui, "app")
    ui.js('ticker.text = "No link for this job."')
    ui.wait(300)
    assert toast.property("visible") is True and ui.find("toastText").property("text") == "No link for this job."
    assert ui.js("toastBad") is True
    ui.js('ticker.text = "Opened the job posting in your browser."')
    ui.wait(250)
    assert ui.js("toastBad") is False and ui.find("toastText").property("text").startswith("Opened")
    ui.bridge.logChanged.emit("Scoring 3 of 14 jobs…")                             # a run's progress does not pop up
    ui.wait(250)
    assert ui.find("toastText").property("text").startswith("Opened")
    ui.js("toastTimer.interval = 150; toastTimer.restart()")
    ui.wait(600)
    assert toast.property("visible") is False                                      # and it goes away by itself
    ui.js("toastTimer.interval = 4200; detailDialog.close()")
    _until_closed(ui, "detailDialog")
    assert ui.warnings == []


@pytest.mark.parametrize("text, bad", [
    ("Couldn't make the single PDF.", True), ("No link for this job.", True), ("That job could not be opened.", True),
    ("These CVs cannot be compared.", True), ("The pasted reply could not be filed.", True), ("That answer is still empty.", True),
    ("Choose the city to search in first (Options).", True), ("Not prepared: the limit is used up", True),
    ("Date saved. You will be reminded the day before.", False), ("Moved to Applications (waiting).", False),
    ("Opened one PDF with the letter and the CV (2 pages).", False), ("Merged into one application.", False),
])
def test_which_messages_are_shown_as_errors(ui, text, bad):
    assert ui.js("looksBad(" + repr(text).replace("\\'", "'") + ")") is bad


def test_preparation_opens_as_a_document_from_both_windows(ui, monkeypatch):
    calls = []
    monkeypatch.setattr(ui_api.Api, "prep_document", lambda self, job_id, reveal=False: calls.append((job_id, reveal)) or {"ok": True, "file": "Alex_Beta_Prep_HR.docx",
                                                                                                                  "files": ["Alex_Beta_Prep_HR.docx", "Alex_Beta_Prep_Technical.docx"]})
    ui.bridge._api.editor.change("app", lambda jobs, j: j.__setitem__("interview_prep", {"made": "2026-10-08", "summary": "s", "questions": [
        {"kind": "about you", "question": "Tell me about yourself.", "why": "", "answer": "I am a junior developer.", "check": ""}]}), render=False)
    _open_job(ui, "app")
    assert ui.find("prepDocument").property("visible") is True
    ui.js('prepDialog.show("app")')
    _until_open(ui, "prepDialog")
    assert ui.find("prepDoc").property("text") == "Open as documents" and ui.find("prepShow").property("text") == "Show in folder"
    assert ui.js('bridge.prepDocument("app", false).files.length') == 2 and ui.js('bridge.prepDocument("app", true).ok') is True
    assert calls == [("app", False), ("app", True)]
    ui.js("prepDialog.note = ''"); QMetaObject.invokeMethod(ui.find("prepDoc"), "clicked"); ui.wait(30)
    assert ui.js("prepDialog.note") == "Opened in Word: Alex_Beta_Prep_HR.docx and Alex_Beta_Prep_Technical.docx."
    ui.js("prepDialog.close()")
    _until_closed(ui, "prepDialog")
    ui.js("detailDialog.close()")
    _until_closed(ui, "detailDialog")
    ui.bridge._api.editor.change("app", lambda jobs, j: j.update(interview_prep={}), render=False)
    _open_job(ui, "app")
    assert ui.find("prepDocument").property("visible") is False                 # nothing prepared yet: no document button
    ui.js("detailDialog.close()")
    _until_closed(ui, "detailDialog")
    assert ui.warnings == []


def test_show_in_folder_buttons_and_the_laid_out_description(ui, monkeypatch):
    shown = []
    monkeypatch.setattr(ui_api.Api, "show_in_folder", lambda self, kind, job_id: shown.append((kind, job_id)) or {"ok": True, "file": "CV.pdf"})
    (config.TAILORED_CV_DIR / "cv_app.docx").write_bytes(b"x")
    ui.bridge._api.editor.change("app", lambda jobs, j: j.update(tailored_cv_file="cv_app.docx",
                           description="We build tools. Requirements Strong Python Benefits Flexible hours"), render=False)
    _open_job(ui, "app")
    assert ui.find("showCv").property("visible") is True and ui.find("showLetter").property("visible") is False
    ui.js('detailDialog.showFile("cv")')
    assert shown == [("cv", "app")]
    assert ui.find("descText").property("text") == "We build tools.\n\nREQUIREMENTS\nStrong Python\n\nBENEFITS\nFlexible hours"
    ui.js("detailDialog.close()")
    _until_closed(ui, "detailDialog")
    ui.bridge._api.editor.change("app", lambda jobs, j: j.update(tailored_cv_file="", description="We use Python, Docker and Kubernetes."), render=False)
    assert ui.warnings == []


def test_more_about_your_work_is_saved_and_checked(ui):
    dialog = ui.find("aboutDialog")
    start = ui.bridge.options()
    ui.bridge.saveOptions({**start, "where": "Prague", "about_experience": "At Acme I wrote the import service in Python."})
    ui.js("aboutDialog.show()")
    _until_open(ui, "aboutDialog")
    field = ui.find("generalField")
    assert field.property("text") == "At Acme I wrote the import service in Python."
    assert ui.find("notesCount").property("text") == "45 of 4000 characters, all notes together"
    ui.js('generalField.text = "' + "x" * 4010 + '"', dialog)
    assert ui.find("notesCount").property("text") == "10 characters too long together (the limit is 4000)"
    ui.js("aboutDialog.save()", dialog)
    assert "too long together" in ui.find("aboutErr").property("text") and dialog.property("visible") is True   # refused, kept open
    ui.js("confirmClose = true; close()", dialog)
    _until_closed(ui, "aboutDialog")
    ui.bridge.saveOptions(start)
    assert ui.warnings == []


def test_preparation_window_groups_questions_by_who_asks(ui):
    assert ui.js('prepDialog.who({ audience: "hr", kind: "technical" })') == "hr"
    assert ui.js('prepDialog.who({ kind: "technical" })') == "technical" and ui.js('prepDialog.who({ kind: "practical" })') == "hr"   # older saved ones
    source = qml_source()
    assert "If you speak with HR or a recruiter: plain words, no jargon" in source
    assert "If you speak with someone who does this work or the hiring manager: precise answers" in source


def test_preparation_window_shows_what_you_can_ask_each_listener(ui):
    ui.bridge._api.editor.change("app", lambda jobs, j: j.update(interview_prep={"made": "2026-10-08", "summary": "s", "questions": [
        {"audience": "hr", "kind": "about you", "question": "Tell me about yourself.", "why": "", "answer": "I am a junior developer.", "check": ""}],
        "ask_hr": ["What are the next steps?"], "ask_technical": ["What would I own first?"], "ask_them": []}), render=False)
    ui.js('prepDialog.show("app")')
    _until_open(ui, "prepDialog")
    assert ui.find("askHr").property("visible") is True and ui.find("askTech").property("visible") is True
    assert ui.js("prepDialog.prep.ask_hr[0]") == "What are the next steps?"
    ui.js("prepDialog.close()")
    _until_closed(ui, "prepDialog")
    ui.bridge._api.editor.change("app", lambda jobs, j: j.update(interview_prep={}), render=False)
    assert ui.warnings == []


def test_wanted_and_unwanted_roles_are_edited_in_options(ui):
    dialog = ui.find("optDialog")
    start = ui.bridge.options()
    ui.bridge.saveOptions({**start, "where": "Prague", "wanted_roles": ["backend developer"], "unwanted_roles": ["economist"]})
    ui.js("optsData = bridge.options()")
    _open(ui, dialog, 450)
    assert ui.js("wantedTags.all().join(',')") == "backend developer" and ui.js("unwantedTags.all().join(',')") == "economist"
    assert ui.js('unwantedTags.add("x")') is False and "job title" in ui.js("unwantedTags.err")       # wrong input is explained
    source = qml_source()
    assert "is under both the roles you want and the roles you do not want" in source
    note = ui.find("fieldsNote")                                                                      # your profession follows the roles as typed
    assert note.property("text").startswith("Your profession, as the app reads it from these roles: software and IT.")
    assert ui.js('wantedTags.add("registered nurse")') is True
    assert "these roles: software and IT, care and health." in note.property("text")
    ui.js('wantedTags.setAll(["beekeeper"])')
    assert note.property("text").startswith("The app cannot tell your profession from these roles or your CV") \
        or "from your CV" in note.property("text")
    assert "a title that names one is skipped before scoring" in source

    ui.js('wantedTags.setAll(["backend developer"])')                                                 # … or you pick it yourself
    names = ui.js("fieldsBox.read.all.join('|')").split("|")
    assert "software and IT" in names and "care and health" in names and ui.find("fieldsAuto").property("visible") is False
    assert note.property("text").endswith("Pick below to set it yourself.") and ui.js("optDialog.changed()") is False
    ui.js('optDialog.toggleField("law", fieldsBox.read.all); optDialog.toggleField("finance", fieldsBox.read.all)')
    assert ui.js("optDialog.pickedFields.join('|')") == "finance|law"                                 # kept in the app's order
    assert note.property("text").startswith("Your profession, as you picked it: finance, law.") and ui.find("fieldsAuto").property("visible") is True
    assert ui.js("optDialog.changed()") is True and ui.js("optDialog.collect().options.my_fields.join('|')") == "finance|law"
    ui.js('optDialog.toggleField("law", fieldsBox.read.all)')
    assert ui.js("optDialog.pickedFields.join('|')") == "finance"
    ui.js("optDialog.pickedFields = []")                                                              # "Let the app decide"
    assert "as the app reads it from these roles: software and IT" in note.property("text") and ui.js("optDialog.changed()") is False
    ui.js("close()", dialog)
    _until_closed(ui, dialog)
    ui.bridge.saveOptions({**start, "where": "Prague", "wanted_roles": ["backend developer"], "my_fields": ["law", "no such field", "finance"]})
    ui.js("optsData = bridge.options()")
    _open(ui, dialog, 450)
    assert ui.js("optDialog.pickedFields.join('|')") == "finance|law" and ui.js("optDialog.changed()") is False   # an unknown name is dropped on load
    assert "Let the app decide" in source and "Pick it below if you want that." in source
    ui.js("close()", dialog)
    _until_closed(ui, dialog)
    ui.bridge.saveOptions(start)
    assert ui.warnings == []


def test_the_preparation_can_be_edited_and_wrong_steps_are_explained(ui, monkeypatch):
    monkeypatch.setattr(api_core.CvSource, "text", lambda self: "Alex Novak\nPython, Docker.")
    prep = {"made": "2026-10-08", "summary": "They will ask about Python.", "pitch": "I am a junior developer.", "fit": ["Python", "Docker"],
            "questions": [{"audience": "hr", "kind": "about you", "question": "Tell me about yourself.", "why": "", "answer": "I am a junior developer.", "check": ""}],
            "ask_hr": ["What are the next steps?"], "ask_technical": [], "ask_them": [], "checks": []}
    ui.bridge._api.editor.change("app", lambda jobs, j: j.update(interview_prep=dict(prep)), render=False)
    ui.js('prepEditDialog.show("app")')
    _until_open(ui, "prepEditDialog")
    dialog, note = ui.find("prepEditDialog"), ui.find("peNote")
    assert dialog.property("opened") is True and ui.js("peModel.count") == 5
    keys = [ui.js(f"peModel.get({i}).key") for i in range(5)]
    assert keys == ["summary", "fit", "pitch", "questions.0.answer", "ask_hr"]
    # an empty answer is refused, in red, and the window stays open
    ui.js('peModel.setProperty(3, "text", ""); prepEditDialog.dirty = true; prepEditDialog.save()')
    assert "is empty" in note.property("text") and dialog.property("visible") is True
    # too long is said before anything is sent
    ui.js('peModel.setProperty(3, "text", "I am a junior developer."); peModel.setProperty(2, "text", "' + "x" * 3010 + '"); prepEditDialog.save()')
    assert "longer than 3000 characters" in note.property("text")
    # closing with changes asks once
    ui.js('peModel.setProperty(2, "text", "My own pitch."); prepEditDialog.dirty = true; prepEditDialog.tryClose()')
    ui.wait(200)
    assert dialog.property("visible") is True and "did not save" in note.property("text") and ui.js("prepEditDialog.closePolicy === 0") is True
    ui.js("prepEditDialog.save()")
    assert note.property("text") == "Saved." and ui.js("prepEditDialog.dirty") is False
    assert ui.bridge.interviewPrep("app")["pitch"] == "My own pitch."
    ui.js("prepEditDialog.tryClose()")
    ui.wait(450)
    assert dialog.property("visible") is False
    # "Make again" would replace the edit: asked first
    _open_job(ui, "app")
    assert ui.js("detailDialog.job.prep_edited") is True
    source = qml_source()
    assert "Making it again replaces the changes you made in the sheet." in source and 'objectName: "prepEdit"' in source
    ui.js("detailDialog.close()")
    _until_closed(ui, "detailDialog")
    ui.bridge._api.editor.change("app", lambda jobs, j: j.update(interview_prep={}), render=False)
    ui.js('prepEditDialog.show("app")')                                          # nothing to edit: said, no empty window
    ui.wait(250)
    assert dialog.property("visible") is False and "Prepare the interview first" in ui.js("ticker.text")
    assert ui.warnings == []


# ── every text field says when it is full or too long ───────────────────────
def test_text_fields_in_the_job_window_explain_their_limits(ui):
    _open_job(ui, "app")
    line = ui.find("notesLine")
    ui.js('cnContact.text = "' + "c" * 200 + '"; detailDialog.notesInfo = detailDialog.full(cnContact, "“Who you spoke to”"); detailDialog.notesBad = true')
    assert line.property("text") == "“Who you spoke to” takes up to 200 characters."
    ui.js('cnContact.text = "Jana"; cnNotes.text = "' + "n" * 2010 + '"')
    assert ui.find("notesOver").property("visible") is True and ui.find("notesOver").property("text") == "10 characters too long (the limit is 2000)"
    ui.js("detailDialog.saveNotes()")
    assert "10 characters too long" in line.property("text") and ui.js("detailDialog.notesBad") is True      # refused, in red
    from jobhelper import planner
    assert planner.company_notes("Beta")["updated"] == ""                                                    # and not saved cut
    ui.js('cnNotes.text = "Short note."; detailDialog.saveNotes()')
    assert line.property("text") == "Saved for Beta." and ui.js("detailDialog.notesBad") is False
    ui.js('evNote.text = "' + "e" * 200 + '"; detailDialog.eventProblem(detailDialog.full(evNote, "The note"), "note")')
    assert ui.find("eventErr").property("text") == "The note takes up to 200 characters."
    ui.js('detailDialog.eventProblem(""); detailDialog.replyOpen = true; replyText.text = "' + "r" * 8005 + '"')
    assert ui.find("replyHint").property("text").startswith("5 characters too long: paste only the message itself")
    ui.js("detailDialog.fileReply()")
    assert "Paste only the message itself" in ui.js("detailDialog.replyErr") and ui.js("detailDialog.replyBusy") is False
    ui.js('replyText.text = ""; detailDialog.replyOpen = false; cnNotes.text = ""; cnContact.text = ""; detailDialog.saveNotes(); detailDialog.close()')
    _until_closed(ui, "detailDialog")
    assert ui.warnings == []


def test_form_answers_and_preparation_lists_explain_their_limits(ui, monkeypatch):
    ui.bridge.saveFormAnswers("app", [{"question": "Why us?", "answer": "Because.", "check": ""}])
    ui.js('formDialog.show("app")')
    _until_open(ui, "formDialog")
    ui.js('fmModel.setProperty(0, "answer", "' + "a" * 3004 + '"); formDialog.dirty = true; formDialog.save()')
    assert "is longer than 3000 characters" in ui.find("fmNote").property("text")
    assert ui.bridge.formAnswers("app")["items"][0]["answer"] == "Because."       # the long version was not saved
    ui.js("formDialog.dirty = false; formDialog.close()")
    _until_closed(ui, "formDialog")
    assert ui.js("prepEditDialog.listProblem('one\\ntwo')") == ""
    many = "\\n".join(f"line {i}" for i in range(13))
    assert ui.js(f"prepEditDialog.listProblem('{many}')") == "13 lines; the most is 12. Remove some."
    assert ui.js("prepEditDialog.listProblem('" + "q" * 401 + "')") == "Line 1 is 401 characters long; the most is 400."
    ui.bridge.saveFormAnswers("app", [])
    source = qml_source()
    assert "“Right to work” takes up to" in source and "“You can start” takes up to" in source
    assert ui.warnings == []


# ── the older windows: every text field explains itself too ─────────────────
def test_number_fields_say_what_they_accept_instead_of_going_silent(ui):
    dialog = ui.find("optDialog")
    ui.js("optsData = bridge.options()")
    _open(ui, dialog, 450)
    spin = ui.find("scanSpin")
    low, high = ui.bridge._api.OPTION_NUMBERS["scan_days"]                      # the range the back end checks
    start = spin.property("value")
    ui.js(f'scanSpin.contentItem.text = "{high + 1}"; scanSpin.contentItem.textEdited()')
    assert ui.js("scanSpin.note") == f"Between {low} and {high}" and ui.js("scanSpin.noteBad") is True
    assert spin.property("value") == start                                        # a wrong number is not taken
    ui.js("scanSpin.contentItem.editingFinished()")
    assert spin.property("value") == high and ui.js("scanSpin.note") == f"Changed to {high}, the highest allowed"
    assert int("".join(ch for ch in ui.js("scanSpin.contentItem.text") if ch.isdigit())) == high     # shown grouped, like "3 650"
    ui.js('scanSpin.contentItem.text = "0"; scanSpin.contentItem.textEdited(); scanSpin.contentItem.editingFinished()')
    assert spin.property("value") == low and ui.js("scanSpin.note") == f"Changed to {low}, the lowest allowed"
    ui.js('scanSpin.contentItem.text = "45"; scanSpin.contentItem.textEdited()')
    assert spin.property("value") == 45 and ui.js("scanSpin.note") == ""           # a good number: taken at once, nothing said
    ui.js('scanSpin.contentItem.text = ""; scanSpin.contentItem.textEdited(); scanSpin.contentItem.editingFinished()')
    assert spin.property("value") == 45 and ui.js("scanSpin.note") == "Left at 45"
    ui.js(f"scanSpin.value = {start}; scanSpin.note = ''")
    ui.js("close()", dialog)
    _until_closed(ui, dialog)
    assert ui.warnings == []


def test_career_site_name_is_checked(ui):
    dialog = ui.find("optDialog")
    ui.js("optsData = bridge.options()")
    _open(ui, dialog, 450)
    assert ui.js('optDialog.addCompany("' + "n" * 61 + '", "greenhouse", "acme")') is False
    assert ui.js("optDialog.coErr") == "The company name is 61 characters long; the most is 60."
    assert ui.js('optDialog.addCompany("Компания", "greenhouse", "acme")') is False and "Latin letters" in ui.js("optDialog.coErr")
    assert ui.js('optDialog.addCompany("Acme", "greenhouse", "")') is False and "token" in ui.js("optDialog.coErr")
    assert ui.find("coName").property("maximumLength") == 60
    ui.js("close()", dialog)
    _until_closed(ui, dialog)
    assert ui.warnings == []


def test_letter_editor_explains_wrong_steps(ui, monkeypatch):
    from jobhelper import cover_letter
    paras = ["First paragraph of the letter.", "Second one.", "Third one.", "Fourth one.", "Fifth one."]
    monkeypatch.setattr(ui_api.Api, "letter_for_edit", lambda self, job_id: {"ok": True, "id": "lead", "company": "Acme", "title": "Dev",
                        "title_line": "Dev", "date_text": "", "paragraphs": list(paras), "min_words": cover_letter.MIN_WORDS, "max_words": cover_letter.MAX_WORDS})
    saved = []
    monkeypatch.setattr(ui_api.Api, "save_letter", lambda self, job_id, texts: saved.append(list(texts)) or {"ok": True, "warnings": []})
    ui.js('letterDialog.show("lead")')
    _until_open(ui, "letterDialog")
    dialog, err = ui.find("letterDialog"), ui.find("letterErr")
    assert dialog.property("opened") is True
    ui.js('letterDialog.setText(1, "   "); letterDialog.setText(3, "")')
    ui.js("letterDialog.save()")
    assert err.property("text") == "Paragraph 2 and 4 are empty. A letter needs all 5 paragraphs." and saved == []
    ui.js("letterDialog.rewrite(1)")
    assert "Paragraph 2 is empty" in err.property("text") and ui.js("letterDialog.busyIndex") == -1      # nothing to rewrite: Claude is not asked
    ui.js('letterDialog.setText(1, "Second one, changed."); letterDialog.setText(3, "Fourth one."); letterDialog.tryClose()')
    ui.wait(200)
    assert dialog.property("visible") is True and "did not save" in err.property("text")                 # changes are not thrown away on one click
    ui.js("letterDialog.busyIndex = 2; letterDialog.tryClose()")
    assert "still rewriting paragraph 3" in err.property("text") and dialog.property("visible") is True
    ui.js("letterDialog.busyIndex = -1; letterDialog.save()")
    assert saved and saved[0][1] == "Second one, changed."
    assert ui.find("hintField").property("maximumLength") == 300
    ui.js('hintField.text = "' + "h" * 300 + '"; hintField.textEdited()')
    assert err.property("text") == "The note for a rewrite takes up to 300 characters."
    ui.js('hintField.text = ""; letterDialog.err = ""; letterDialog.tryClose()')
    ui.wait(450)
    assert dialog.property("visible") is False                                                           # saved: closes at once
    assert ui.warnings == []


def test_esc_asks_once_before_unsaved_answers_are_thrown_away(ui):
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest
    dialog = ui.find("formDialog")
    ui.js('formDialog.show("app")')
    _until_open(ui, "formDialog")
    ui.win.requestActivate()
    QTest.keyClick(ui.win, Qt.Key_Escape)                                 # nothing changed: Esc closes, as before
    _until_closed(ui, "formDialog")
    ui.js('formDialog.show("app")')
    _until_open(ui, "formDialog")
    ui.js("formDialog.dirty = true")
    ui.win.requestActivate()
    QTest.keyClick(ui.win, Qt.Key_Escape)                                 # changed: Esc used to do nothing at all
    assert _until(lambda: "did not save" in ui.js("formDialog.err")) and dialog.property("visible") is True
    QTest.keyClick(ui.win, Qt.Key_Escape)                                 # asked once: the second Esc closes
    _until_closed(ui, "formDialog")


def test_a_button_that_cannot_be_pressed_looks_it_and_an_error_shows_red_even_with_focus(ui):
    button = ui.find("fmDraft")                                           # enabled: !formDialog.busy
    ui.js('formDialog.show("app")')
    _until_open(ui, "formDialog")
    ui.js("formDialog.busy = true")
    assert button.property("enabled") is False and button.property("opacity") == ui.js("theme.disabledOpacity") == 0.5
    ui.js("formDialog.busy = false")
    assert button.property("enabled") is True and button.property("opacity") == 1
    ui.js("formDialog.close()")
    _until_closed(ui, "formDialog")
    source = qml_source()
    hidden = re.findall(r"border\.color:\s*\(?\s*[\w.]+\.activeFocus \? theme\.accent\s*:[^}]*theme\.crit", source)
    assert hidden == [], "a field's red error border must not hide behind its focus colour (red first, then the focus colour)"
    for name in ("SmallButton", "PillButton"):
        text = (ROOT / "ui_qt" / f"{name}.qml").read_text(encoding="utf-8")
        assert "theme.disabledOpacity" in text and "visualFocus ? (" in text and "primaryStyle ? theme.ink : theme.accent" in text, \
            f"{name}: one paleness, and an outline for Tab that shows on an accent-filled button too"
    letter = (ROOT / "ui_qt" / "LetterDialog.qml").read_text(encoding="utf-8")
    assert 'objectName: "letterX"' in letter and "onClicked: letterDialog.tryClose() }" in letter, "the letter window has a ✕ too"


def _esc(ui):
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest
    ui.win.requestActivate()
    QTest.keyClick(ui.win, Qt.Key_Escape)


def test_esc_in_the_letter_and_the_preparation_asks_once_and_closes_a_clean_window(ui, monkeypatch):
    from jobhelper import cover_letter
    paras = ["First paragraph of the letter.", "Second one.", "Third one.", "Fourth one.", "Fifth one."]
    monkeypatch.setattr(ui_api.Api, "letter_for_edit", lambda self, job_id: {"ok": True, "id": "lead", "company": "Acme", "title": "Dev",
                        "title_line": "Dev", "date_text": "", "paragraphs": list(paras), "min_words": cover_letter.MIN_WORDS, "max_words": cover_letter.MAX_WORDS})
    letter = ui.find("letterDialog")
    ui.js('letterDialog.show("lead")')
    _until_open(ui, "letterDialog")
    _esc(ui)                                                               # clean: Esc closes (it used to do nothing here)
    _until_closed(ui, "letterDialog")
    ui.js('letterDialog.show("lead")')
    _until_open(ui, "letterDialog")
    ui.js('letterDialog.setText(1, "Second one, changed.")')
    _esc(ui)
    assert _until(lambda: "did not save" in ui.find("letterErr").property("text")) and letter.property("visible") is True
    ui.js("letterDialog.busyIndex = 2")
    _esc(ui)                                                               # Claude rewriting: it says to wait, and stays
    assert _until(lambda: "still rewriting" in ui.find("letterErr").property("text")) and letter.property("visible") is True
    ui.js("letterDialog.busyIndex = -1")
    _esc(ui)                                                               # asked once already: now it closes
    _until_closed(ui, "letterDialog")
    prep = {"made": "2026-10-08", "summary": "They will ask about Python.", "pitch": "I am a junior developer.", "fit": ["Python", "Docker"],
            "questions": [{"audience": "hr", "kind": "about you", "question": "Tell me about yourself.", "why": "", "answer": "I am a junior developer.", "check": ""}],
            "ask_hr": ["What are the next steps?"], "ask_technical": [], "ask_them": [], "checks": []}
    ui.bridge._api.editor.change("app", lambda jobs, j: j.update(interview_prep=dict(prep)), render=False)
    ui.js('prepEditDialog.show("app")')
    _until_open(ui, "prepEditDialog")
    ui.js("prepEditDialog.dirty = true")
    _esc(ui)
    assert _until(lambda: "did not save" in ui.js("prepEditDialog.err")) and ui.find("prepEditDialog").property("visible") is True
    _esc(ui)
    _until_closed(ui, "prepEditDialog")


def test_esc_belongs_to_the_window_in_front_whatever_is_open_behind_it(ui):
    """The four stacks the window checker tried with a real Esc: before, the assistant and the job window each stood
    aside while the other was shown, so Esc did nothing in any of them."""
    assistant, job = ui.find("assistantDialog"), ui.find("detailDialog")
    # A: the assistant over the job window, with unsaved details: the assistant asks, the job window stays
    _open_job(ui)
    ui.js("assistantDialog.show()")
    _until_open(ui, "assistantDialog")
    ui.js("assistantDialog.dirty = true")
    _esc(ui)
    assert _until(lambda: "did not save" in (ui.js("assistantDialog.problem.msg") or "")) and assistant.property("visible") and job.property("visible")
    ui.js("assistantDialog.busy = true; assistantDialog.dirty = false; assistantDialog.confirmClose = false")
    _esc(ui)                                                               # B: working: it says so
    assert _until(lambda: "still working" in ui.js("assistantDialog.askErr")) and assistant.property("visible")
    ui.js("assistantDialog.busy = false")
    # an answer typed and not saved holds the window too (it used to close and lose it)
    ui.js('assistantDialog.page = "answers"; assistantDialog.addAnswer("Why do you want this job?")')
    assert ui.js("assistantDialog.unsaved") == 1 and ui.js("assistantDialog.holdsWork") is True
    _esc(ui)
    assert _until(lambda: "not saved" in ui.js("assistantDialog.answersErr")) and assistant.property("visible")
    _esc(ui)                                                               # asked once: now it closes, the job window stays
    _until_closed(ui, "assistantDialog")
    assert job.property("visible") is True
    ui.js("detailDialog.close()")
    _until_closed(ui, "detailDialog")
    # D: the job window opened over a clean assistant, holding a pasted reply: the job window asks, the assistant stays
    ui.js("assistantDialog.show()")
    _until_open(ui, "assistantDialog")
    _open_job(ui)
    ui.js('replyText.text = "Thank you for applying, we would like to meet you."')
    assert ui.js("detailDialog.pending.length") == 1
    _esc(ui)
    assert _until(lambda: ui.js("detailDialog.confirmClose") is True) and job.property("visible") and assistant.property("visible")
    ui.js('replyText.text = ""')
    _esc(ui)                                                               # nothing held now: its own Esc closes it
    _until_closed(ui, "detailDialog")
    assert assistant.property("visible") is True
    ui.js("assistantDialog.close()")
    _until_closed(ui, "assistantDialog")


def test_esc_in_the_assistant_asks_once_and_never_also_closes_the_job_window(ui):
    assistant, job = ui.find("assistantDialog"), ui.find("detailDialog")
    _open_job(ui)
    ui.js("assistantDialog.show()")
    _until_open(ui, "assistantDialog")
    ui.js("assistantDialog.busy = true")
    assert ui.js("detailDialog.pending.length >= 0")
    _esc(ui)                                                               # working: it says so and stays; the job window behind stays too
    assert _until(lambda: "still working" in ui.js("assistantDialog.askErr")) and assistant.property("visible") is True
    assert job.property("visible") is True
    ui.js("assistantDialog.busy = false")
    _esc(ui)                                                               # nothing held: Esc closes the assistant only
    _until_closed(ui, "assistantDialog")
    assert job.property("visible") is True
    ui.js("detailDialog.close()")
    _until_closed(ui, "detailDialog")


def test_esc_keeps_unsaved_answers_when_the_details_are_unsaved_too(ui):
    """One flag stood for both: the first Esc warned about the details, and the second closed and lost the answers."""
    assistant = ui.find("assistantDialog")
    ui.js("assistantDialog.show()")
    _until_open(ui, "assistantDialog")
    ui.js('assistantDialog.dirty = true; assistantDialog.addAnswer("Why do you want this job?")')
    assert ui.js("assistantDialog.unsaved") == 1
    _esc(ui)                                                               # the details first, on their own page
    assert _until(lambda: "did not save" in (ui.js("assistantDialog.problem.msg") or "")) and assistant.property("visible") is True
    assert ui.js("assistantDialog.page") == "details"
    _esc(ui)                                                               # the answers are still held: they are named too
    assert _until(lambda: "not saved" in ui.js("assistantDialog.answersErr")) and assistant.property("visible") is True
    assert ui.js("assistantDialog.page") == "answers" and ui.js("assistantDialog.unsaved") == 1
    ui.js('assistantDialog.setAnswer(0, "answer", "I like the product.")')   # a new change asks again
    _esc(ui)
    assert _until(lambda: "not saved" in ui.js("assistantDialog.answersErr")) and assistant.property("visible") is True
    _esc(ui)                                                               # both asked: now it closes
    _until_closed(ui, "assistantDialog")
    assert ui.warnings == []


def test_a_field_with_an_error_is_red_even_with_the_focus(ui):
    """PR #36 put the red before the focus colour in five fields; five more still hid it, and the time of a date had no red at all."""
    # About you: a pay line that cannot be added
    ui.js("aboutDialog.show()")
    _until_open(ui, "aboutDialog")
    ui.js('aboutDialog.payErr = "Write what kind of role this figure is for."')
    shown, red = _border_with_focus(ui, "aboutDialog", "payRole")
    assert shown == red, "payRole"
    ui.js('aboutDialog.payErr = ""; payRole.text = ""; close()', ui.find("aboutDialog"))
    _until_closed(ui, "aboutDialog")
    # Options: a country with digits in it
    dialog = ui.find("optDialog")
    ui.js("optsData = bridge.options()")
    _open(ui, dialog, 450)
    ui.js('countryField.text = "Germany 2"', dialog)
    shown, red = _border_with_focus(ui, "optDialog", "countryField")
    assert shown == red, "countryField"
    ui.js('countryField.text = ""; confirmClose = true; close()', dialog)
    _until_closed(ui, dialog)
    # the job window: a wrong time is red in the time box, and the date, which is right, is not
    _open_job(ui)
    tomorrow = (date.today() + timedelta(days=1)).isoformat()
    ui.js(f'evDate.text = "{tomorrow}"; evTime.text = "25:99"; detailDialog.saveEvent()')
    assert "HH:MM" in ui.js("detailDialog.eventErr")
    shown, red = _border_with_focus(ui, "detailDialog", "evTime")
    assert shown == red, "evTime"
    assert ui.js("String(evDate.background.border.color)", ui.find("detailDialog")) != red, "the date is right: no red on it"
    ui.js('evDate.text = "next week"; detailDialog.saveEvent()')            # a wrong date: the date is red, the time is not
    shown, red = _border_with_focus(ui, "detailDialog", "evDate")
    assert shown == red and ui.js("String(evTime.background.border.color)", ui.find("detailDialog")) != red
    ui.js('evDate.text = ""; evTime.text = ""; detailDialog.eventProblem(""); detailDialog.close()')
    _until_closed(ui, "detailDialog")
    assert ui.warnings == []


def test_a_message_over_a_window_does_not_come_back_when_a_window_opens_again(ui):
    """"Options saved." was only hidden when Options closed; opened again within four seconds, Options showed it again."""
    dialog, toast = ui.find("optDialog"), ui.find("toast")
    ui.js("optsData = bridge.options()")
    _open(ui, dialog, 450)
    ui.js('ticker.text = "✓ Options saved."')
    assert _until_true(ui, lambda: toast.property("visible") is True)               # over the open window, as before
    ui.js("close()", dialog)
    _until_closed(ui, dialog)
    assert ui.js("toastText") == "" and ui.js("ticker.text") == "✓ Options saved."   # the status line still says it
    _open(ui, dialog, 450)
    assert toast.property("visible") is False, "an old message does not show over a window opened later"
    ui.js("aboutDialog.show()")                                                      # two windows: the message stays ...
    _until_open(ui, "aboutDialog")
    ui.js('ticker.text = "✓ About you saved."')
    ui.js("close()", ui.find("aboutDialog"))
    _until_closed(ui, "aboutDialog")
    assert ui.js("toastText") == "✓ About you saved."                                # ... while one of them is still open
    ui.js("close()", dialog)
    _until_closed(ui, dialog)
    assert ui.js("toastText") == ""
    ui.js('ticker.text = ""')
    assert ui.warnings == []
