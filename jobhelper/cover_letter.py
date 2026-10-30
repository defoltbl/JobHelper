"""Write a cover letter for one job by filling the letter template (.docx).

Follows data/cover_letter_instruction.md: the head block, fonts, spacing and
page setup of data/cover_letter_template.docx stay exactly as they are; only the
title line, the date and the five body paragraphs get new text. We edit
`word/document.xml` inside the template package and copy every other file
unchanged.

The text must sell strengths only: no admission of a gap, and a technology the
master CV cannot back up may be named as the employer's platform or as something
the candidate wants to work with, never as something already done.
"""
from __future__ import annotations

import re
from datetime import date
from typing import Any

from lxml import etree

from . import claims, config, cv_tailor, llm, store
from .claims import ClaimCheck
from .cv_tailor import POSTING_SHOWN, clean, figure_sources
from .word_document import P as _P, T as _T, XML_SPACE, Paragraph, WordDocument, text_of as _text

BODY_PARAGRAPHS = 5
MIN_WORDS, MAX_WORDS = 320, 410        # instruction: about 330-400 words
MAX_CHARS = 2550                       # keeps the letter on one page (example: 2240)
MAX_HINT = 300                         # what you ask to change in one paragraph

# Check 3 of the instruction (a gap admitted), the forward-looking wish (rule 1, option 2) and wording that claims
# hands-on experience: one set of rules for every text in the person's voice (claims.py).
_GAP, _WISH = claims.GAP, claims.WISH

SYSTEM = (
    "You write one candidate's cover letter for a job posting. You sell strengths "
    "only and never mention anything the candidate has not done. You never claim a "
    "technology that is not in the master CV. Write in the first person (I, my): the "
    "candidate is the author, so never refer to them by name or as he/she. Use an en dash (–) with spaces, never "
    "an em dash, and no emojis. You have no tools and no file access here: another "
    "program places your text into the .docx. Reply once, with only the JSON object. " + llm.POSTING_RULE
)

PROMPT = """Write the text of a cover letter for the job posting below. The file handling is
done elsewhere; you only supply the wording.

Rules from the candidate (follow them exactly):
---
{rules}
---

Master CV (the only source of truth for experience):
---
{cv_text}
---
{adapted}
Job posting:
- Title: {title}
- Company: {company}
- Location: {location}
- Description: {description}
{feedback}
Return JSON with exactly these keys:
{{
  "title_line": "<the role applied for, without (m/f/d) or location: used only when the CV has no title line under the name>",
  "paragraphs": ["<hook>", "<strongest evidence>", "<second theme>", "<learning and fit>", "<close>"],
  "role_short": "<the role in 1-3 words, e.g. Middleware Admin>"
}}
Exactly {n} paragraphs, {min_words}-{max_words} words in total. No salutation, no date, no
"Yours sincerely" and no name: those are already in the template."""

GENERIC_RULES = """Sell strengths only. Never mention anything the candidate has not done: no "not
yet", "although", "limited experience", "I lack". A requirement the candidate cannot back up is
left out, or turned into a forward-looking wish ("I would like to bring the same drive to ..."),
or answered with the nearest real evidence. Only technologies in the master CV may appear as
experience.
Five paragraphs, about 330-400 words: (1) hook - role, company, why the work attracts;
(2) strongest evidence for the core requirement - one real project with concrete facts;
(3) a second theme from the posting with one short true story or result; (4) learning and fit,
then the forward-looking wish; (5) close - only the facts listed under "Facts about the candidate"
(location, right to work, start date) and an invitation to talk; a fact that is not stated is left out.
Vary sentence openings, one concrete fact per paragraph, no filler openers, no hedging."""


def load_rules(cv_text: str = "") -> str:
    """Sections 1, 4, 5 and 7 of the instruction (content rules + worked example).
    Without a file the built-in general rules are used. {{placeholders}} in the
    file are filled with the person's own facts (see config.fill_placeholders)."""
    return cv_tailor.instruction_sections(config.LETTER_INSTRUCTION, ("1", "4", "5", "7"), cv_text, GENERIC_RULES)


def rules_with_facts(cv_text: str = "") -> str:
    """The rules, and the only facts about the person a letter may state (location, right to work, start date)."""
    return (load_rules(cv_text) + "\n\nFacts about the candidate (use only these for location, right to work and start date):\n"
            + config.facts_block(cv_text))


def unverified_terms(master_text: str) -> list[str]:
    """The CV's "never claim" list plus the letter instruction's "Anything else (...)" list."""
    terms = cv_tailor.unverified_terms(master_text)
    path = config.LETTER_INSTRUCTION
    if path.exists():
        m = re.search(r"Anything else \(([^)]*)\)", path.read_text(encoding="utf-8"))
        if m:
            low = master_text.lower()
            for raw in m.group(1).split(","):
                t = re.sub(r"\band so on\b", "", raw).strip()
                if t and t not in terms and t.lower() not in low:
                    terms.append(t)
    return terms


_sentences = claims.sentences


def candidate_first_name(cv_text: str) -> str:
    """First word of the CV's first line ("Jane Doe" -> "Jane")."""
    for line in (cv_text or "").splitlines():
        words = re.findall(r"[^\W\d_]+", line)
        if words:
            return words[0]
    return ""


def third_person(paragraphs: list[str], first_name: str = "") -> bool:
    """True if the letter talks about the candidate instead of as the candidate."""
    text = " ".join(paragraphs)
    if claims.names_candidate(text, first_name):
        return True
    return bool(text.strip()) and not claims.FIRST_PERSON.search(text)


def _issues(paragraph: str, check: ClaimCheck) -> list[tuple[str, str, str]]:
    """What breaks the letter rules in one paragraph: ("gap", the words, the sentence) and ("claim", the term, the sentence)."""
    out = []
    for s in claims.sentences(paragraph):
        gap = claims.GAP.search(s)
        if gap:
            out.append(("gap", gap.group(0), s))
        term = check.claimed_in(s)
        if term:
            out.append(("claim", term, s))
    return out


def figure_problems(figures: list[str], what: str = "The letter",
                    advice: str = "Use only the figures the CV or the notes give.") -> list[str]:
    """What is sent back to Claude for each figure no source gives."""
    return [f"{what} states the figure {f}, which none of the candidate's sources gives. {advice}" for f in figures]


def fits_one_page(paragraphs: list[str]) -> bool:
    """Whether the letter's body stays on one page: at most MAX_WORDS words and MAX_CHARS characters."""
    return sum(len(p.split()) for p in paragraphs) <= MAX_WORDS and sum(len(p) for p in paragraphs) <= MAX_CHARS


def content_problems(paragraphs: list[str], unverified: list[str], first_name: str = "") -> list[str]:
    """Why this draft can't be used as is (empty list = fine)."""
    problems = []
    if third_person(paragraphs, first_name):
        problems.append("The letter is written about the candidate in the third person. Rewrite it in the "
                        "first person (I, my); never use the candidate's name or he/his in the body.")
    if len(paragraphs) != BODY_PARAGRAPHS:
        problems.append(f"The letter has {len(paragraphs)} paragraphs; it must have exactly {BODY_PARAGRAPHS}.")
    words = sum(len(p.split()) for p in paragraphs)
    if words < MIN_WORDS:
        problems.append(f"Too short: {words} words; write {MIN_WORDS}-{MAX_WORDS}.")
    if not fits_one_page(paragraphs):
        problems.append(f"Too long for one page: {words} words; write {MIN_WORDS}-{MAX_WORDS}.")
    check = ClaimCheck(unverified)
    for n, para in enumerate(paragraphs, 1):
        for kind, what, s in _issues(para, check):
            problems.append(f"Paragraph {n} admits a gap (\"{what}\"): \"{s[:90]}\". Remove it or state a forward-looking wish instead."
                            if kind == "gap" else
                            f"Paragraph {n} presents '{what}' as experience, but it is not in the "
                            "master CV. Name it only as something the candidate wants to work with.")
    return problems


# Where a sentence turns from the candidate to the reader or to the job: what follows describes the role, not the person.
_TURN = re.compile(r"\b(you|your|yours|the role|this role|the position|this position|the posting|the team|the company)\b", re.IGNORECASE)
# "I" as the one who did something. Not "I'd like", and not the numeral of a job title ("Analyst I opening").
_I_DID = re.compile(r"\bI(?:'ve|'m| have| am)?\b(?!'d| would| will| want| hope| look| opening| role| position| at | in | and | or |\s*[,.;:)–-])")
_LATER = re.compile(r"\b(would|could|will|translate|prepare[ds]? me|ready to|learn|grow|develop into|build on|bring|contribute|apply (this|that|it|them))\b",
                    re.IGNORECASE)


def borrowed_claims(paragraphs: list[str], description: str, backed: str) -> list[str]:
    """Two-word terms of the posting that the letter states as the candidate's own experience although nothing the
    person wrote (the CV, their notes) has either word. The same check as for the adapted CV, read the way a letter
    is written: a letter may speak of the role's terms ("your intercompany postings") and of what the candidate
    wants to do, so a term counts only where the candidate is the one doing it: from "I …" up to where the sentence
    turns to the reader or to the future."""
    found: list[str] = []
    for para in paragraphs:
        for sentence in _sentences(para):
            if not claims.FIRST_PERSON.search(sentence) or claims.WISH.search(sentence):
                continue
            start = _I_DID.search(sentence)                 # "I process …", "I've built …": the candidate as the one who did it
            if not start:
                continue                                   # "the role caught my attention because …" describes the role
            own = sentence[start.start():]
            for stop in (_TURN, _LATER):
                m = stop.search(own)
                if m:
                    own = own[:m.start()]
            for phrase in cv_tailor.borrowed_phrases(own, description, backed):
                if phrase.lower() not in [f.lower() for f in found]:
                    found.append(phrase)
    return found


def enforce(paragraphs: list[str], unverified: list[str]) -> list[str]:
    """Last line of defence: drop sentences that admit a gap or claim unbacked experience."""
    check = ClaimCheck(unverified)
    return [" ".join(s for s in claims.sentences(para) if not claims.GAP.search(s) and not check.claimed_in(s))
            for para in paragraphs]


# ── docx ─────────────────────────────────────────────────────────────────────
def _slots(body) -> dict[str, Any]:
    """Paragraphs of the template by role: name, title, 2 contact lines, date,
    salutation, the body paragraphs, closing, signature."""
    filled = [p for p in body if p.tag == _P and _text(p).strip()]
    if len(filled) < 9:      # name, title, 2 contact lines, date, salutation, body, closing, signature
        raise ValueError("The letter template does not have the expected structure.")
    return {"title": filled[1], "date": filled[4], "body": filled[6:-2]}


def _set_para(p, text: str) -> None:
    para = Paragraph(p)
    if not para.writable:               # its text sits inside a link or a form control: there is nothing to write into
        raise ValueError("A paragraph of the letter template is built in a way the app cannot fill in "
                         "(its text is inside a link or a form control). Retype it as plain text in Word.")
    para.write(text)


APP_MADE = "Made by JobHelper from the head of your CV"      # the template's comment: one the app may bring up to date
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_YEARS = re.compile(r"^(?:(?:19|20)\d{2}[\s()-]*)+$")      # "2019 - 2021": dates, not a phone
_NOT_A_NAME = re.compile(r"\b(cv|curriculum|vitae|resume|résumé|lebenslauf|profile|contact|contacts|summary|skills|experience|"
                         r"education|about|objective|personal|details)\b", re.IGNORECASE)


def template_head(cv_text: str) -> list[str]:
    """The four head lines of a letter template from the top of a CV: name, what you call yourself, and two lines of
    contacts ([] when the CV's top gives no clear name or not two kinds of contact: then the plain letter is written).
    Only the first eight lines are read, so a figure further down ("10 000 000 rows") is never a phone number; a year
    run ("2019 - 2021") is no phone either; and nothing is guessed (no city from the search settings)."""
    head = [_CONTROL.sub("", " ".join(line.split()))[:200] for line in (cv_text or "").splitlines()[:40] if line.strip()][:8]
    name = head[0] if head else ""
    if not name or len(name) > 60 or len(name.split()) > 4 or re.search(r"[@\d|,:]", name) or _NOT_A_NAME.search(name):
        return []
    top = "\n".join(head)
    email = re.search(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", top)
    phone = next((m.group(0).strip() for m in re.finditer(r"\+?\d[\d ()-]{7,}\d", top)
                  if (m.group(0).startswith("+") or len(re.sub(r"\D", "", m.group(0))) >= 9) and not _YEARS.match(m.group(0).strip())), "")
    links = "  ·  ".join(dict.fromkeys(re.findall(r"(?:linkedin\.com/in/|github\.com/)[\w-]+", top)[:2]))
    contacts = [x for x in (email.group(0) if email else "", phone, links) if x]
    if len(contacts) < 2:
        return []                               # one kind of contact: the head has no second line without a guess
    lines = [contacts[0], "  ·  ".join(contacts[1:])] if len(contacts) == 2 else ["  ·  ".join(contacts[:2]), contacts[2]]
    return [name, own_title("\n".join(head)) or "Applicant", *lines]


def make_template(cv_text: str, path) -> bool:
    """Write a plain letter template from the head of your CV when you have none of your own (a new person could never
    use Edit letter before): the head (template_head), the date, the salutation, BODY_PARAGRAPHS body paragraphs, the
    closing and your name again. Marked as made by the app (APP_MADE). True when it was written."""
    import docx
    from docx.shared import Pt
    head = template_head(cv_text)
    if not head:
        return False
    doc = docx.Document()
    for text in (*head, "", letter_date(), "Dear Hiring Team,", *[f"Paragraph {n}." for n in range(1, BODY_PARAGRAPHS + 1)],
                 "Yours sincerely,", head[0]):
        doc.add_paragraph().add_run(text).font.size = Pt(11)
    doc.core_properties.comments = APP_MADE
    doc.save(str(path))
    return True


def app_made_head(path) -> list[str] | None:
    """The four head lines of a template the app made, or None for a template of your own (never changed by the app)."""
    import docx
    doc = docx.Document(str(path))
    if doc.core_properties.comments != APP_MADE:
        return None
    return [p.text for p in doc.paragraphs[:4]]


def letter_date(d: date | None = None) -> str:
    d = d or date.today()
    return f"{d.day} {d:%B %Y}"                      # 6 October 2026


def build_docx(template_path, out_path, title_line: str, paragraphs: list[str], when: date | None = None,
               date_text: str = "") -> None:
    doc = WordDocument(template_path)
    slots = _slots(doc.body)
    if len(slots["body"]) != len(paragraphs):
        raise ValueError(f"The template has {len(slots['body'])} body paragraphs, the letter {len(paragraphs)}.")
    _set_para(slots["title"], title_line)
    _set_para(slots["date"], date_text or letter_date(when))
    for p, text in zip(slots["body"], paragraphs):
        _set_para(p, text)
    doc.save(out_path)


_NOT_A_TITLE = re.compile(r"@|https?:|www\.|\.(?:com|cz|org|net|io|dev)\b|\d{4,}|[.!?]$")


def own_title(cv_text: str) -> str:
    """What you call yourself: the line under your name in your CV ("Junior Software Developer"), or '' when the CV
    has no such line (the second line is then contacts or a heading). The head of a letter says who writes, so it
    carries this line, the same one your CV and your profiles on the web show, not the role applied for."""
    lines = [" ".join(line.split()) for line in (cv_text or "").splitlines() if line.strip()]
    if len(lines) < 2:
        return ""
    line = lines[1]
    if len(line) > 70 or len(line.split()) > 8 or _NOT_A_TITLE.search(line.lower()) or (line.isupper() and len(line.split()) <= 3):
        return ""
    return line if re.search(r"[^\W\d_]{3}", line) else ""


def set_title(path, title: str) -> bool:
    """Put another line under the name of a letter this app built; nothing else in the file changes. False = it already says so."""
    title = clean(str(title)).strip()
    if not title:
        return False
    doc = WordDocument(path)
    slot = _slots(doc.body)["title"]
    if _text(slot).strip() == title:
        return False
    _set_para(slot, title)
    doc.save_over(path)
    return True


def read_letter(path) -> dict[str, Any]:
    """The editable parts of a letter this app built: title line, date, body paragraphs."""
    slots = _slots(WordDocument(path).body)
    return {"title_line": _text(slots["title"]), "date_text": _text(slots["date"]),
            "paragraphs": [_text(p) for p in slots["body"]]}


REWRITE_PROMPT = """Rewrite ONE paragraph of the candidate's cover letter. Everything else stays as it is.

Rules from the candidate (follow them exactly):
---
{rules}
---

Master CV (the only source of truth for experience):
---
{cv_text}
---

Job posting:
- Title: {title}
- Company: {company}
- Description: {description}

The letter now (paragraphs numbered):
{letter}

Rewrite paragraph {number}. Keep its job in the letter and roughly its length ({words} words, at most {max_words}).
{hint}{feedback}
Return JSON only: {{"paragraph": "<the new paragraph {number}>"}}"""


def rewrite_paragraph(job: dict[str, Any], cv_text: str, paragraphs: list[str], index: int, hint: str = "") -> str:
    """A new version of one paragraph (one model call, one retry if it breaks a rule).
    Nothing is saved here: the caller shows it to the user first."""
    if not 0 <= index < len(paragraphs):
        raise ValueError("That paragraph does not exist.")
    unverified, first_name = unverified_terms(cv_text), candidate_first_name(cv_text)
    employers = claims.EmployerCheck(cv_text)
    letter = "\n".join(f"{n}. {p}" for n, p in enumerate(paragraphs, 1))
    words = len(paragraphs[index].split())
    most = max(words + 25, 60)                 # roughly its length: what the prompt asks, and what is checked
    hint_line = f"What the candidate wants changed: {hint.strip()[:MAX_HINT]}\n" if hint.strip() else ""
    feedback, text = "", ""
    for attempt in range(2):
        data = llm.ask_json(REWRITE_PROMPT.format(
            rules=rules_with_facts(cv_text), cv_text=cv_text[:12000], title=job.get("title", ""), company=job.get("company", ""),
            description=llm.untrusted("posting", (job.get("description") or "").strip()[:3000] or "(No description was available.)"),
            letter=letter, number=index + 1, words=words, max_words=most,
            hint=hint_line, feedback=feedback), system=SYSTEM, tag="agent3_letters")
        text = clean(str((data or {}).get("paragraph", ""))).strip() if isinstance(data, dict) else ""
        problems = paragraph_problems(text, unverified, first_name) + _rewrite_length(text, most, paragraphs, index)
        figures = claims.unbacked_figures(text, *figure_sources(job, cv_text, own=(*paragraphs, hint)), brackets_ok=False)
        problems = figure_problems(figures, "It") + problems           # your letter and your hint are yours: their figures count
        misplaced = misplaced_problems(text, employers, "It", job.get("company", "")) if text else []
        if not problems and (not misplaced or attempt == 1):
            return text                     # work at the wrong employer is sent back once, then shown when you save
        problems += misplaced
        feedback = claims.rejected(problems)
    raise ValueError(problems[0] if text else "The model returned no paragraph.")


def _rewrite_length(text: str, most: int, paragraphs: list[str], index: int) -> list[str]:
    """A rewritten paragraph keeps roughly its length and the letter on one page (a letter you made longer by hand
    must at least not grow)."""
    words = len(text.split())
    if words > most:
        return [f"Too long: {words} words; write at most {most}."]
    after = paragraphs[:index] + [text] + paragraphs[index + 1:]
    if not fits_one_page(after) and sum(map(len, after)) > sum(map(len, paragraphs)):
        return ["With this paragraph the letter no longer fits on one page; write it shorter."]
    return []


def misplaced_problems(text: str, employers: claims.EmployerCheck, what: str = "The letter", company: str = "") -> list[str]:
    """Work given to the wrong employer (claims.EmployerCheck), in the words a draft is sent back with."""
    return [f"{what} puts {where}, but the CV names that work only under another position. Keep each piece of work "
            "with the job where it was done." for where in employers.misplaced(text, company)]


def paragraph_problems(text: str, unverified: list[str], first_name: str = "",
                       employers: claims.EmployerCheck | None = None, company: str = "") -> list[str]:
    """The letter rules, applied to a single paragraph."""
    if not text.strip():
        return ["The paragraph is empty."]
    problems = []
    if claims.names_candidate(text, first_name):
        problems.append("Write in the first person (I, my); never use the candidate's name.")
    for kind, what, _s in _issues(text, ClaimCheck(unverified)):
        problems.append(f"It admits a gap (\"{what}\"). Remove it or state a forward-looking wish instead." if kind == "gap"
                        else f"It presents '{what}' as experience, but it is not in the master CV.")
    if employers is not None:
        problems += misplaced_problems(text, employers, "It", company)
    return problems


def letter_warnings(paragraphs: list[str], cv_text: str, company: str = "") -> list[str]:
    """What a hand-edited letter no longer follows (shown to the user, never blocking)."""
    out = []
    words = sum(len(p.split()) for p in paragraphs)
    if words < MIN_WORDS:
        out.append(f"The letter is short: {words} words (the usual range is {MIN_WORDS}-{MAX_WORDS}).")
    if words > MAX_WORDS:
        out.append(f"The letter is long: {words} words (the usual range is {MIN_WORDS}-{MAX_WORDS}); it may not fit on one page.")
    unverified, first_name, employers = unverified_terms(cv_text), candidate_first_name(cv_text), claims.EmployerCheck(cv_text)
    for n, p in enumerate(paragraphs, 1):
        out += [f"Paragraph {n}: {x}" for x in paragraph_problems(p, unverified, first_name, employers, company)]
    return out


def _formatting_only(xml: bytes) -> bytes:
    """document.xml with every piece of text removed (check 2 of the instruction)."""
    root = etree.fromstring(xml)
    for t in root.iter(_T):
        t.text = ""
        t.attrib.pop(XML_SPACE, None)      # whitespace handling of the text, not formatting
    return etree.tostring(root)


def verify(template_path, out_path, unverified: list[str] | None = None) -> list[str]:
    problems = []
    template, out = WordDocument(template_path), WordDocument(out_path)
    for name in template.names():
        if name != WordDocument.TEXT and template.part(name) != out.part(name):
            problems.append(f"{name} differs from the template.")
    if _formatting_only(template.part(WordDocument.TEXT)) != _formatting_only(out.part(WordDocument.TEXT)):
        problems.append("The formatting differs from the template.")
    body = out.body
    paragraphs = [_text(p) for p in _slots(body)["body"]]
    full = "\n".join(_text(p) for p in body if p.tag == _P)
    if "—" in full:
        problems.append("The letter contains an em dash.")
    if cv_tailor.EMOJI.search(full):
        problems.append("The letter contains an emoji.")
    for para in paragraphs:
        gap = _GAP.search(para)
        if gap:
            problems.append(f"The letter admits a gap (\"{gap.group(0)}\").")
    return problems


# ── file name ────────────────────────────────────────────────────────────────
def file_stem(master_path, job: dict[str, Any], role_short: str = "") -> str:
    """<Name>_CoverLetter_<Company>_<RoleShort>, matching the adapted CV's name."""
    cv_stem = cv_tailor.file_stem(master_path, job, role_short)       # ..._CV_<Company>_<Role>
    m = re.match(r"(.+?)_CV_(.+)\.docx$", job.get("tailored_cv_file", "") or "")
    if m:
        return f"{m.group(1)}_CoverLetter_{m.group(2)}"
    return cv_stem.replace("_CV_", "_CoverLetter_", 1)


def _adapted_cv_text(job: dict[str, Any]) -> str:
    name = job.get("tailored_cv_file", "")
    path = config.TAILORED_CV_DIR / name if name else None
    if not path or path.suffix.lower() != ".docx" or not path.exists():
        return ""
    sec = cv_tailor.sections(WordDocument(path).body)
    lines = ([_text(sec["profile"])] if sec["profile"] is not None else []) + [_text(p) for p in sec["rows"]]
    return "\n".join(lines)


# ── entry point ──────────────────────────────────────────────────────────────
def write(job: dict[str, Any], cv_text: str, master_path=None) -> str:
    """Create the cover letter for `job`; returns the file name in LETTERS_DIR."""
    template = config.LETTER_TEMPLATE
    unverified = unverified_terms(cv_text)
    first_name = candidate_first_name(cv_text)
    adapted = _adapted_cv_text(job)
    adapted_block = (f"\nPROFILE and KEY SKILLS of the CV adapted for this same application "
                     f"(stay consistent with it):\n---\n{adapted}\n---\n") if adapted else ""
    description = (job.get("description") or "").strip()[:POSTING_SHOWN] or "(No description was available.)"
    backed = cv_text + "\n" + config.experience_notes()          # everything the person wrote about their own work
    employers = claims.EmployerCheck(cv_text)

    title_line, paragraphs, role_short, feedback = "", [], "", ""
    for _attempt in range(2):
        data = llm.ask_json(PROMPT.format(
            rules=rules_with_facts(cv_text), cv_text=cv_text[:12000], adapted=adapted_block,
            title=job["title"], company=job["company"], location=job.get("location", ""),
            description=llm.untrusted("posting", description), feedback=feedback, n=BODY_PARAGRAPHS,
            min_words=MIN_WORDS + 10, max_words=MAX_WORDS - 10,
        ), system=SYSTEM, tag="agent3_letters")
        if not isinstance(data, dict):
            raise ValueError("The model's answer was not a letter.")
        paragraphs = [clean(str(p)) for p in (data.get("paragraphs") or []) if str(p).strip()]
        title_line = clean(str(data.get("title_line", "")))
        role_short = str(data.get("role_short", ""))
        problems = content_problems(paragraphs, unverified, first_name)
        problems += [f"The letter presents '{phrase}' as the candidate's own experience, but those are the posting's words: nothing in the "
                     "master CV speaks of it. Say what the candidate really did, and name this only as part of the role or as something "
                     "the candidate wants to work with." for phrase in borrowed_claims(paragraphs, description, backed)]
        problems += figure_problems(claims.unbacked_figures("\n".join(paragraphs), *figure_sources(job, cv_text), brackets_ok=False))
        problems += misplaced_problems("\n".join(paragraphs), employers, company=job.get("company", ""))
        if not problems:
            break
        feedback = claims.rejected(problems, "draft")

    if len(paragraphs) != BODY_PARAGRAPHS:
        raise ValueError(f"The model returned {len(paragraphs)} paragraphs instead of {BODY_PARAGRAPHS}.")
    if third_person(paragraphs, first_name):
        raise ValueError("The letter was written in the third person twice; nothing was saved.")
    paragraphs = enforce(paragraphs, unverified)
    if any(not p for p in paragraphs):
        raise ValueError("A paragraph was empty after removing unbacked claims.")
    # What is still stated as experience in the posting's words is not cut out (it can be fair wording for what the CV
    # does show): it is put in front of the person, next to the same note about the adapted CV.
    # work still given to the wrong employer is shown in its own line: the person knows where each piece of work was done
    misplaced = employers.misplaced("\n".join(paragraphs), job.get("company", ""))
    if misplaced:
        job["letter_misplaced"] = misplaced[:8]
    else:
        job.pop("letter_misplaced", None)
    left = borrowed_claims(paragraphs, description, backed)
    if left:
        job["letter_check"] = left[:8]
    else:
        job.pop("letter_check", None)
    figures = claims.unbacked_figures("\n".join(paragraphs), *figure_sources(job, cv_text), brackets_ok=False)   # the final text
    if figures:
        job["letter_figures"] = figures[:8]
    else:
        job.pop("letter_figures", None)
    # The line under the name says who writes: your own title, as your CV has it. Only a CV without such a line gets the role.
    title_line = own_title(cv_text) or title_line or re.sub(r"\s*\(.*?\)", "", job["title"]).strip()

    config.ensure_dirs()
    stem = file_stem(master_path or "Candidate_CV_EN.docx", job, role_short)
    out = config.LETTERS_DIR / f"{stem}.docx"
    if out.exists() and job.get("cover_letter_file") != out.name:
        out = config.LETTERS_DIR / f"{stem}_{store.slugify(job['id'])[-6:]}.docx"
    build_docx(template, out, title_line, paragraphs)
    problems = verify(template, out, unverified)
    if problems:
        out.unlink(missing_ok=True)
        raise ValueError("Cover letter failed its checks: " + "; ".join(problems))
    return out.name
