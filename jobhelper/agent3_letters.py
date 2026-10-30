"""Agent 3 – score each job, write cover letters, and tailor the CV.

For a strong match we (1) write a cover letter and (2) adapt the master CV to
the posting (cv_tailor: PROFILE + KEY SKILLS only, never inventing anything).
Cover letters go to data/cover_letters/, adapted CVs to data/tailored_cvs/.
"""
from __future__ import annotations

import json
import re
from datetime import date
from typing import Any

from . import config, cover_letter, cv_tailor, filters, llm, pdf, prefilter, store, trouble

SYSTEM = (
    "You are an expert career coach who writes concise, specific, non-generic "
    "cover letters in English. You never invent experience the candidate does "
    "not have. Output only valid JSON. " + llm.POSTING_RULE
)

# Step 1 – scoring only (short JSON, never truncates).
_SCORE_RULES = """Rate how well this candidate fits, judging ELIGIBILITY strictly. Give a LOW
score (under 40) if the posting lists hard requirements the candidate clearly
does NOT have, for example: a required certification they lack (e.g. CompTIA A+,
MCP), a minimum number of years of experience beyond what they have, a required
language they don't speak (e.g. Czech), a degree they don't hold, or a mandatory
tech stack absent from their CV. Only score 60+ when the candidate plausibly
meets the core must-have requirements.

Internships, traineeships and practical training are work experience: add up
their months from the profile. A requirement of up to 2 years of experience is
met by a candidate with about that much internship experience; only a stated
requirement of 3 or more years is an unmet hard requirement for them.

Judge only what the posting actually says. Do not assume a requirement that is
not written there: in particular, a posting written in English that does not
ask for Czech counts as an English-speaking role, so never lower the score
because a role "probably" or "typically" needs Czech. Entry-level, trainee and
junior postings do not need prior professional experience unless they state a
number of years. If no description is available, judge from the title alone
and do not invent requirements.
"""

SCORE_PROMPT = """Candidate profile (from their CV):
{profile}

Job posting:
- Title: {title}
- Company: {company}
- Location: {location}
- Description: {description}

{rules}
Return a JSON object only:
{{
  "match_score": <integer 0-100>,
  "match_reason": "<one sentence: why this score, naming any unmet hard requirement>"
}}"""

SCORE_BATCH_PROMPT = """Candidate profile (from their CV):
{profile}

Below are {n} separate job postings, marked J1 to J{n}. Judge each one on its
own, exactly as if it were the only posting.

{postings}

{rules}
Return a JSON object only, with one entry per posting and no others:
{{
  "scores": [
    {{"id": "J1", "match_score": <integer 0-100>, "match_reason": "<one sentence: why this score, naming any unmet hard requirement>"}}
  ]
}}"""

# Step 2 – cover letter (plain text, only for strong matches).
LETTER_SYSTEM = (
    "You write concise, specific, non-generic cover letters in English. You never "
    "invent experience the candidate does not have. Output ONLY the letter body: "
    "start with 'Dear Hiring Team,' and end with 'Yours sincerely,' on its own "
    "line. Do NOT include the candidate's name, address, contact details, date, or "
    "the company address (those are added separately). No em dashes. " + llm.POSTING_RULE
)
LETTER_PROMPT = """Candidate profile (from their CV):
{profile}

Job posting:
- Title: {title}
- Company: {company}
- Location: {location}
- Description: {description}

Write the body of a tailored cover letter, about 250-320 words in 3-4 short
paragraphs, referencing 2-3 concrete requirements from the posting and matching
them to the candidate's real experience. Begin with 'Dear Hiring Team,' and end
with 'Yours sincerely,' on its own final line. Do NOT add the candidate's name
after it. Plain English, no em dashes. Output only the letter body."""

# CV tailoring – re-emphasise real experience to match the posting's keywords.
TAILOR_PROFILE_SYSTEM = (
    "You tailor the PROFILE / summary paragraph of a CV to a specific job for "
    "ATS/HR keyword matching. STRICT RULES: use ONLY facts and skills already in "
    "the candidate's master CV; never invent or exaggerate. Target the specific "
    "role, and mirror the posting's key terminology for things the candidate "
    "genuinely has. 2-3 short, plain sentences. No em dash. Output only the "
    "profile paragraph text (no heading, no markdown, no quotes). " + llm.POSTING_RULE
)
TAILOR_PROFILE_PROMPT = """Master CV (source of truth, do not add anything not here):
---
{cv_text}
---

Target job:
- Title: {title}
- Company: {company}
- Location: {location}
- Description: {description}

Rewrite ONLY the candidate's PROFILE/summary paragraph so it targets this exact
role and surfaces this posting's most important keywords (only skills/experience
the candidate genuinely has, from the master CV above). 2-3 short sentences,
plain English, no em dash. Output only the new profile paragraph."""


def _extract_contact(cv_text: str) -> tuple[str, str, str]:
    """Pull (name, email, phone) from the master CV text for the letter header."""
    import re
    lines = [ln.strip() for ln in (cv_text or "").splitlines() if ln.strip()]
    name = lines[0] if lines else ""
    email = re.search(r"[\w.+-]+@[\w.-]+\.\w+", cv_text or "")
    phone = re.search(r"\+?\d[\d ]{7,}\d", cv_text or "")
    return name, (email.group(0) if email else ""), (phone.group(0).strip() if phone else "")


def _clean_role(title: str) -> str:
    """A concise role label for the letter header (drop location/qualifiers)."""
    import re
    return re.split(r"[|(\[]| - | – ", title)[0].strip() or title


def _save_letter(job: dict[str, Any], letter: str, cv_text: str = "") -> str:
    """Save the cover letter as a properly formatted business letter (.docx) plus
    a plain .md. Header = name / role / contact; body from the model; signature."""
    letter = cv_tailor.en_dashes(_strip_fence(letter))
    config.ensure_dirs()
    base = f"{store.slugify(job['company'])}-{store.slugify(job['title'])}-{job['id']}"
    name, email, phone = _extract_contact(cv_text)
    # under the name stands what the person is (the line under the name in their CV); the role only when the CV has no such line
    role = cover_letter.own_title(cv_text) or _clean_role(job["title"])

    config.LETTERS_DIR.joinpath(f"{base}.md").write_text(letter, encoding="utf-8")

    try:
        import docx
        from docx.shared import Inches, Pt, RGBColor

        black = RGBColor(0, 0, 0)
        doc = docx.Document()
        for s in doc.sections:
            s.left_margin = s.right_margin = Inches(0.83)
            s.top_margin = s.bottom_margin = Inches(0.63)

        def line(text: str, size: int, bold: bool = False):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            run.bold = bold
            run.font.size = Pt(size)
            run.font.color.rgb = black
            return p

        line(name, 16, bold=True)                      # name header
        line(role, 11)                                 # role
        contact = "  ·  ".join(x for x in (email, phone) if x)
        if contact:
            line(contact, 9)                           # email · phone
        doc.add_paragraph()                            # blank line
        for para in (p.strip() for p in letter.split("\n")):
            if para:
                line(para, 11)                         # greeting, body, "Yours sincerely,"
        line(name, 11, bold=True)                      # signature
        doc.save(str(config.LETTERS_DIR / f"{base}.docx"))
    except Exception as e:  # docx is optional; never fail the pipeline for it
        print(f"    (docx skipped for {base}: {e})")

    return f"{base}.md"


def _strip_fence(text: str) -> str:
    """Remove a single wrapping ```-fence the model sometimes adds."""
    import re
    m = re.match(r"^\s*```[a-zA-Z]*\s*\n(.*?)\n\s*```\s*$", text, re.DOTALL)
    return m.group(1).strip() if m else text.strip()


def _save_tailored_cv(job: dict[str, Any], new_profile: str, template_path=None) -> str:
    """Fallback for a non-.docx master: save a Markdown note with the tailored
    profile (a .docx master is adapted by cv_tailor instead)."""
    new_profile = cv_tailor.en_dashes(_strip_fence(new_profile))
    config.ensure_dirs()
    base = f"CV-{store.slugify(job['company'])}-{store.slugify(job['title'])}-{job['id']}"
    md_path = config.TAILORED_CV_DIR / f"{base}.md"
    md_path.write_text(f"# Tailored profile for {job['title']} @ {job['company']}\n\n{new_profile}\n",
                       encoding="utf-8")
    return md_path.name


def _make_tailored_cv(job: dict[str, Any], cv_text: str, cv_path, fields: dict[str, Any]) -> str:
    """Create the adapted CV for one job and return its file name ('' if none).

    With a .docx master this follows data/cv_adaptation_instruction.md (PROFILE +
    KEY SKILLS rewritten, everything else identical). Other master formats fall
    back to a short Markdown note with a tailored profile."""
    previous = job.get("tailored_cv_file", "")
    # a Word CV with a profile section gets a full adapted copy; any other gets the profile note below
    if cv_path and str(cv_path).lower().endswith(".docx") and cv_tailor.adaptable(cv_path):
        name = cv_tailor.tailor(job, cv_text, cv_path)
    else:
        new_profile = llm.ask(TAILOR_PROFILE_PROMPT.format(cv_text=cv_text[:12000], **fields),
                              system=TAILOR_PROFILE_SYSTEM, tag="agent3_letters").strip()
        name = _save_tailored_cv(job, new_profile, template_path=None) if new_profile else ""
    if name and previous and previous != name:      # the superseded file the app wrote: kept in the backup, then removed
        for old in (config.TAILORED_CV_DIR / previous, (config.TAILORED_CV_DIR / previous).with_suffix(".pdf")):
            _retire(old)
    if name.endswith(".docx") and pdf.pdf_for(config.TAILORED_CV_DIR / name) is None:
        pdf.to_pdf(config.TAILORED_CV_DIR / name)   # cv_tailor already made one when it checked the page count
    return name


def _retire(path) -> None:
    """Take a document the app wrote out of the way when a new one replaces it: a copy goes
    to the encrypted backup first (the newest of those are kept), so nothing is simply deleted."""
    from . import vault
    if path.is_file():
        try:
            vault.keep_copy(path, "letter_versions")
        except Exception:  # noqa: BLE001 - without a copy the file is left where it is
            return
        path.unlink(missing_ok=True)


def _generated_letter(job: dict[str, Any]) -> bool:
    """True if the job's current letter is one this app wrote (so it may be
    replaced). The candidate's own imported letters are never touched."""
    name = job.get("cover_letter_file", "")
    return bool(name) and job.get("source") != "manual" and not job.get("own_letter") and (
        (config.LETTERS_DIR / name).exists() or (config.LETTERS_DIR / name).with_suffix(".docx").exists())


def _keep_template(cv_text: str) -> None:
    """A letter template from your CV's head when you have none (so the letter gets every check and Edit letter works),
    brought up to date when your CV's contacts change. A template of your own is never touched, and one that went
    missing is not replaced: the start-up check tells you about it. A template that cannot be made leaves the plain
    letter, as before."""
    from . import integrity
    path = config.LETTER_TEMPLATE
    state = next((f["state"] for f in integrity.check() if f["name"] == "letter_template"), "")
    try:
        if path.exists():
            current = cover_letter.app_made_head(path)
            # only a template still as the app made it: one edited in Word (state "changed", or no check yet) is yours
            if current is None or state != "ok" or current == cover_letter.template_head(cv_text):
                return
        elif state == "missing":
            return
        if path.exists():
            from . import vault
            vault.keep_copy(path, "letter_versions")     # the template as it was, before the app writes it again
        if cover_letter.make_template(cv_text, path):
            integrity.approve(["letter_template"])
    except Exception as e:  # noqa: BLE001 - the plain letter is written instead
        trouble.note("making the letter template", e)


def _make_letter(job: dict[str, Any], cv_text: str, cv_path, fields: dict[str, Any]) -> str:
    """Write the cover letter for one job and return its file name ('' if none).

    With the letter template and a master CV this follows
    data/cover_letter_instruction.md (template head and formatting untouched);
    otherwise it falls back to a plain generated business letter."""
    previous = job.get("cover_letter_file", "") if _generated_letter(job) else ""
    if cv_text.strip():
        _keep_template(cv_text)
    if cv_text.strip() and config.LETTER_TEMPLATE.exists():
        name = cover_letter.write(job, cv_text, cv_path)
    else:
        letter = llm.ask(LETTER_PROMPT.format(**fields), system=LETTER_SYSTEM, tag="agent3_letters").strip()
        name = _save_letter(job, letter, cv_text=cv_text) if letter else ""
    if name and previous and previous != name:       # drop the superseded generated files
        base = config.LETTERS_DIR / previous
        for old in {base, base.with_suffix(".docx"), base.with_suffix(".md"), base.with_suffix(".pdf")}:
            if old.name != name:
                _retire(old)
    if name.endswith(".docx"):
        pdf.to_pdf(config.LETTERS_DIR / name)        # nothing happens without LibreOffice
    return name


LANGUAGE_CAP = 25         # the score a posting gets at most when it requires a language the CV does not show
OTHER_FIELD_CAP, UNWANTED_CAP = 40, 20


def direction_rules(wanted, unwanted) -> str:
    """What the candidate is looking for, as lines for the scoring prompt ('' when nothing is known).
    Without it the scorer rates eligibility only, and an economist role that accepts the degree gets 80."""
    wanted = [str(w).strip() for w in (wanted or []) if str(w).strip()][:12]
    unwanted = [str(w).strip() for w in (unwanted or []) if str(w).strip()][:12]
    lines = []
    if wanted:
        lines.append(
            f"The candidate is looking for these kinds of role: {', '.join(wanted)}. Fit means fit for THAT direction, not only "
            f"eligibility: a role in another field that merely accepts their degree or general skills scores at most {OTHER_FIELD_CAP}, "
            "even when no requirement is unmet. Say in the reason that it is another field.")
    if unwanted:
        lines.append(f"The candidate does NOT want these kinds of role: {', '.join(unwanted)}. A role of one of these kinds scores "
                     f"at most {UNWANTED_CAP}; say so in the reason.")
    return ("\n" + "\n".join(lines) + "\n") if lines else ""


SHORT_TEXT = 400          # a description shorter than this is a summary, not the posting
NO_TEXT_CAP = 40          # the most a posting can score when only its title is known
NO_TEXT = ("The posting's text could not be read, so no letter or CV was written. Open the posting; "
           "you can paste its text in this window and score it again.")


def text_basis(job: dict[str, Any]) -> str:
    """What a score rests on: 'title' (no posting text at all), 'summary' (only a short one) or '' (the posting)."""
    length = len((job.get("description") or "").strip())
    return "title" if length == 0 else "summary" if length < SHORT_TEXT else ""


def apply_score(job: dict[str, Any], result: Any, cv_text: str = "") -> bool:
    """Put one answer from the scorer on a job, then run the app's own check on it.
    False if the answer has no usable number."""
    try:
        job["match_score"] = max(0, min(100, int(result.get("match_score"))))
    except (AttributeError, TypeError, ValueError):
        return False
    job["match_reason"] = str(result.get("match_reason", ""))
    job.pop("score_capped", None)
    # a required language the CV does not show is a hard stop, whatever the model said
    missing = filters.unmet_language(job.get("description") or "", cv_text) if cv_text.strip() else ""
    if missing and job["match_score"] > LANGUAGE_CAP:
        job["score_capped"] = f"Requires {missing}, which your CV does not show."
        job["match_reason"] = (f"{job['score_capped']} The app lowered the score from {job['match_score']} to {LANGUAGE_CAP}. "
                               + job["match_reason"]).strip()
        job["match_score"] = LANGUAGE_CAP
    # a title alone says too little for a high score: such a job is listed, not recommended
    if text_basis(job) == "title" and job["match_score"] > NO_TEXT_CAP:
        job["score_capped"] = "Scored from the title alone: the posting's text could not be read."
        job["match_reason"] = (f"{job['score_capped']} The app lowered the score from {job['match_score']} to {NO_TEXT_CAP}. "
                               + job["match_reason"]).strip()
        job["match_score"] = NO_TEXT_CAP
    job["last_update"] = date.today().isoformat()
    return True


def scoring_rules(profile: dict[str, Any]) -> str:
    """The scoring rules plus the direction from Options (or, when nothing was saved, from the CV)."""
    saved = config.ui_settings()
    wanted = saved.get("wanted_roles")
    return _SCORE_RULES + direction_rules(profile.get("target_roles") if wanted is None else wanted, saved.get("unwanted_roles"))


NO_DESCRIPTION = ("(No description was available for this listing – assess the fit from the title, company and "
                  "location, and keep the cover letter's specific claims about the role modest.)")


class LimitWait:
    """Runs one model step; when Claude's limit is used up it waits through `on_limit` (which returns when the run may
    go on, or raises to stop) and tries again. Without `on_limit` the LimitReached error ends the run at once.
    Scoring and the writing of documents both run their steps through it."""

    def __init__(self, on_limit=None) -> None:
        self._on_limit = on_limit

    def __call__(self, call):
        while True:
            try:
                return call()
            except llm.LimitReached as e:
                if self._on_limit is None:
                    raise
                self._on_limit(e)


class Scorer:
    """Scores postings against one person's profile: several postings in one request (each request has a fixed cost),
    the "J1, J2 …" answers matched back to their postings, and a request of its own for any posting the batch missed.
    Its requests run through a LimitWait. What happens with a score is the caller's: `record(job, result)` keeps it and says
    whether it was usable."""

    BATCH_TEXT, ONE_TEXT = 3000, 4000                # how much of a posting's text one request carries

    def __init__(self, profile: dict[str, Any], on_limit=None, rules: str | None = None) -> None:
        self.blurb = json.dumps({k: profile.get(k) for k in ("summary", "seniority", "years_experience", "internships", "key_skills", "strengths")},
                                ensure_ascii=False)
        self.rules = scoring_rules(profile) if rules is None else rules
        self.ask = on_limit if isinstance(on_limit, LimitWait) else LimitWait(on_limit)

    def fields(self, job: dict[str, Any], limit: int = ONE_TEXT) -> dict[str, str]:
        """A posting as the prompts take it."""
        return dict(profile=self.blurb, title=job.get("title", ""), company=job.get("company", ""), location=job.get("location", ""),
                    description=llm.untrusted("posting", (job.get("description") or "").strip()[:limit] or NO_DESCRIPTION))

    def score(self, chunk: list[dict[str, Any]], record) -> list[dict[str, Any]]:
        """Score these postings; returns the ones that got no usable score."""
        done: set[str] = set()
        if len(chunk) > 1:
            postings = "\n\n".join(
                "### J{n}\n- Title: {title}\n- Company: {company}\n- Location: {location}\n- Description: {description}".format(
                    n=n, **self.fields(job, self.BATCH_TEXT)) for n, job in enumerate(chunk, 1))
            try:
                data = self.ask(lambda: llm.ask_json(SCORE_BATCH_PROMPT.format(profile=self.blurb, postings=postings, rules=self.rules, n=len(chunk)),
                                                     system=SYSTEM, tag="agent3_letters"))
            except llm.STOPS_A_RUN:
                raise                                   # a stuck Claude ends the run: 5 minutes per posting would hold it for hours
            except Exception as e:  # noqa: BLE001 - the ones it missed are asked alone below
                print(f"  ! Scoring {len(chunk)} postings together failed ({e}); asking one by one")
                data = None
            items = data.get("scores") if isinstance(data, dict) else data
            for item in items if isinstance(items, list) else []:
                m = re.fullmatch(r"J(\d+)", str(item.get("id", "")).strip()) if isinstance(item, dict) else None
                if m and 1 <= int(m.group(1)) <= len(chunk):
                    job = chunk[int(m.group(1)) - 1]
                    if job["id"] not in done and record(job, item):
                        done.add(job["id"])
        missed = []
        for job in chunk:
            if job["id"] in done:
                continue
            try:
                result = self.ask(lambda: llm.ask_json(SCORE_PROMPT.format(rules=self.rules, **self.fields(job)), system=SYSTEM, tag="agent3_letters"))
            except llm.STOPS_A_RUN:
                raise
            except Exception as e:  # noqa: BLE001 - one posting that cannot be scored must not stop the others
                print(f"  ! Could not score '{job.get('title', '')}' @ {job.get('company', '')}: {e}")
                missed.append(job)
                continue
            if not record(job, result):
                print(f"  ! Could not score '{job.get('title', '')}' @ {job.get('company', '')}: unusable reply")
                missed.append(job)
        return missed


def rescore(jobs: list[dict[str, Any]], ids: list[str], profile: dict[str, Any], cv_text: str = "",
            progress=None, on_limit=None, status=None) -> dict[str, Any]:
    """Score the given jobs again with the rules as they are now (your direction, the
    language check). Only the score and its reason change: no letter or CV is written
    or removed here. The earlier score is kept as `score_before` so the change shows.
    {scored, changed: [{id, title, company, before, after}], failed}"""
    wanted = [j for j in jobs if j.get("id") in set(ids)]
    scorer = Scorer(profile, on_limit)
    stats: dict[str, Any] = {"scored": 0, "changed": [], "failed": 0}

    def record(job, result) -> bool:
        before = job.get("match_score")
        if not apply_score(job, result, cv_text):
            return False
        stats["scored"] += 1
        if job["match_score"] != before:
            job["score_before"] = before
            stats["changed"].append({"id": job["id"], "title": job.get("title", ""), "company": job.get("company", ""),
                                     "before": before, "after": job["match_score"]})
        if job.get("cover_letter_file"):
            job["match_tier"] = _tier(job["match_score"])
        return True

    size = max(1, int(config.SCORE_BATCH))
    for start in range(0, len(wanted), size):
        chunk = wanted[start:start + size]
        if progress:
            progress(start, len(wanted))
        if status:
            status(f"Scoring again {start + 1}–{start + len(chunk)} of {len(wanted)}…")
        stats["failed"] += len(scorer.score(chunk, record))
    if progress:
        progress(len(wanted), len(wanted))
    return stats


NO_COMPANY = "The posting has no company name, so no letter or CV was written. Open the posting to see who it is."


def _tier(score: int) -> str:
    """For a job that got a letter: strong (60+) | moderate (50-59) | stretch
    (below 50 - written because the user set a lower "write letters above")."""
    if score >= config.STRONG_THRESHOLD:
        return "strong"
    if score >= config.DEFAULT_THRESHOLD:
        return "moderate"
    return "stretch"


def _flag_moderate(job: dict[str, Any]) -> None:
    """Add a review note to moderate (50-59) matches so they aren't sent blindly."""
    if job.get("match_tier") == "moderate":
        note = f"[{date.today().isoformat()}] Moderate match ({job.get('match_score')}%) – review the letter/CV before sending."
        if "Moderate match" not in job.get("notes", ""):
            job["notes"] = (note + "\n" + job.get("notes", "")).strip()
    elif job.get("match_tier") == "stretch":
        note = (f"[{date.today().isoformat()}] Stretch match ({job.get('match_score')}%) – the posting asks for things "
                "your CV does not show. Read \"Why this score\" before you apply.")
        if "Stretch match" not in job.get("notes", ""):
            job["notes"] = (note + "\n" + job.get("notes", "")).strip()


def mark_ready(job: dict[str, Any]) -> None:
    """A job that now has its documents: a found lead becomes letter-ready, and its tier and review note follow its score.
    The search and a request for one job's documents both end this way."""
    if job.get("status") == "found":
        job["status"] = "letter_ready"
    if isinstance(job.get("match_score"), int):
        job["match_tier"] = _tier(job["match_score"])
        _flag_moderate(job)


def generate_for(job: dict[str, Any], cv_text: str = "", cv_path=None,
                 write_letter: bool = True, rewrite_letter: bool = False) -> dict[str, bool]:
    """Generate a cover letter and/or tailored CV for ONE job on demand
    (ignores scoring/skip rules). Used by `status.py retailor`."""
    from datetime import date as _date
    description = (job.get("description") or "").strip()[:Scorer.ONE_TEXT] or NO_DESCRIPTION
    # the profile here is the start of the CV itself (a single job's documents need no CV analysis)
    fields = dict(title=job["title"], company=job["company"],
                  location=job.get("location", ""), description=description,
                  profile=(cv_text[:1500] if cv_text else "(see attached CV)"))
    done = {"letter": False, "cv": False}

    # The CV comes first: the letter is written to match the adapted CV.
    if cv_text.strip():
        try:
            name = _make_tailored_cv(job, cv_text, cv_path, fields)
            if name:
                job["tailored_cv_file"] = name
                done["cv"] = True
        except (*llm.STOPS_A_RUN, llm.PaidApiRefused):
            raise                                       # the reason nothing can be written now: the caller says it
        except Exception as e:
            print(f"  ! tailored CV failed for {job['company']}: {e}")

    missing = not job.get("cover_letter_file")
    if write_letter and (missing or (rewrite_letter and _generated_letter(job))):
        try:
            name = _make_letter(job, cv_text, cv_path, fields)
            if name:
                job["cover_letter_file"] = name
                done["letter"] = True
        except (*llm.STOPS_A_RUN, llm.PaidApiRefused) as e:
            if not done["cv"]:
                raise
            if isinstance(e, llm.CallTimedOut):           # a stuck Claude ends the run; the CV made first stays on the job
                mark_ready(job)
                job["last_update"] = _date.today().isoformat()
                raise
            print(f"  ! letter not written for {job['company']}: {e}")   # the CV that was made is kept
        except Exception as e:
            print(f"  ! letter failed for {job['company']}: {e}")

    if done["letter"] or done["cv"]:
        mark_ready(job)
        job["last_update"] = _date.today().isoformat()
    return done


MAX_DOC_ATTEMPTS = 2      # how many runs may try to write documents for one job


def _give_back_try(job, stop):
    """Claude's limit is no fault of the job, so it keeps its try; a call that hangs every time uses its tries up,
    or it would hold every run for CALL_TIMEOUT and keep the jobs after it waiting."""
    if isinstance(stop, llm.LimitReached):
        job["doc_attempts"] -= 1


def process(
    jobs: list[dict[str, Any]],
    profile: dict[str, Any],
    threshold: int = 60,
    only_new: bool = True,
    cv_text: str = "",
    cv_path=None,
    progress=None,
    on_limit=None,
    status=None,
    score_batch: int | None = None,
) -> dict[str, int]:
    """Score jobs, write letters, and tailor the CV. Mutates jobs in place.

    progress(done, total), if given, is called before each step (a scoring call,
    then each job's documents); status(text) gets a short line about the step.
    score_batch postings are scored per model call.
    on_limit(error), if given, is called when Claude's limit is used up: it should
    block until the run may continue (the step is then retried) or raise to stop.
    Without it the LimitReached error ends the run at once - no job-by-job failures.

    If cv_text (the master CV) is provided, a keyword-tailored CV is generated
    for each strong match alongside the cover letter.
    """
    scorer = Scorer(profile, on_limit)
    stats = {"scored": 0, "letters": 0, "tailored_cvs": 0, "skipped": 0, "prefiltered": 0}

    # Free first look: postings that are plainly another profession are marked
    # and never sent to the model (see prefilter.py).
    if only_new and cv_text.strip():
        level_words = config.ui_settings().get("level_words")
        if level_words is None:
            from . import agent1_cv
            level_words = filters.default_level_words(agent1_cv.level_of(profile))
        saved = config.ui_settings()
        wanted = saved.get("wanted_roles") or profile.get("target_roles") or []
        stats["prefiltered"] = prefilter.apply(
            jobs, cv_text, check_level=bool(level_words), roles=list(wanted) + list(profile.get("adjacent_roles") or []),
            unwanted=saved.get("unwanted_roles") or [], wanted=wanted, fields=saved.get("my_fields"))

    def _skip(job):
        # Applications and anything past the "found" stage are left alone. A found
        # job is still to do when it has no score yet, or when it scored high enough
        # but its documents were never written (a run that was cut short).
        if not only_new:
            return False
        if job.get("status") != "found":
            return True
        score = job.get("match_score")
        if score is None:
            return bool(job.get("prefilter"))
        return score < threshold or bool(job.get("cover_letter_file")) or job.get("doc_attempts", 0) >= MAX_DOC_ATTEMPTS

    def _say(text):
        if status:
            status(text)

    def _record(job, result) -> bool:
        if not apply_score(job, result, cv_text):
            return False
        if job.get("score_capped"):
            stats["capped"] = stats.get("capped", 0) + 1
        stats["scored"] += 1
        return True

    work = [job for job in jobs if not _skip(job)]
    stats["skipped"] = len(jobs) - len(work)
    to_score = [job for job in work if job.get("match_score") is None]
    units, done = len(to_score) + len(work), 0          # one unit per score, one per job for its documents

    # Step 1: score, several postings per call (each call carries a fixed overhead,
    # so five at a time costs far less of the subscription than five calls).
    size = max(1, int(score_batch if score_batch is not None else config.SCORE_BATCH))
    for start in range(0, len(to_score), size):
        chunk = to_score[start:start + size]
        if progress:
            progress(done, units)
        _say(f"Scoring jobs {start + 1}–{start + len(chunk)} of {len(to_score)}…")
        scorer.score(chunk, _record)                   # each posting it could not score is in the log, once
        done += len(chunk)

    # Steps 2 and 3: documents for the ones that scored high enough.
    writing = [job for job in work if (job.get("match_score") or 0) >= threshold and job.get("match_score") is not None]
    for job in work:
        if progress:
            progress(done, units)
        done += 1
        if job.get("match_score") is None:
            continue                                   # could not be scored this time
        if job["match_score"] < threshold:
            print(f"  ·  {job['match_score']:>3}  {job['title']} @ {job['company']}  (below {threshold})")
            continue
        if not (job.get("company") or "").strip():
            from . import agent2_search
            agent2_search.fill_company(job)             # an older record: the name may be in the posting text
        if not (job.get("company") or "").strip():
            # a letter is written to a company; with no name it would be addressed to nobody
            job["hold"] = NO_COMPANY
            print(f"  ?  {job['match_score']:>3}  {job['title']}  (no company name in the posting: no letter written)")
            continue
        if text_basis(job) == "title":
            # without the posting's text a letter can only be generic, and a CV cannot be adapted to anything
            job["hold"] = NO_TEXT
            print(f"  ?  {job['match_score']:>3}  {job['title']} @ {job['company']}  (no posting text: no letter written)")
            continue
        job.pop("hold", None)
        _say(f"Writing CV and letter {writing.index(job) + 1} of {len(writing)}: {job['company'] or job['title']}…")
        fields = scorer.fields(job)
        job["doc_attempts"] = job.get("doc_attempts", 0) + 1
        # Step 2: adapt the CV (PROFILE + KEY SKILLS); the letter is written to match it.
        if cv_text.strip() and not job.get("tailored_cv_file"):
            try:
                name = scorer.ask(lambda: _make_tailored_cv(job, cv_text, cv_path, fields))
            except llm.STOPS_A_RUN as stop:
                _give_back_try(job, stop)
                raise
            except Exception as e:
                print(f"    ! tailored CV failed for {job['company']}: {e}")
                name = ""
            if name:
                job["tailored_cv_file"] = name
                stats["tailored_cvs"] += 1

        # Step 3: write the cover letter.
        try:
            letter_file = scorer.ask(lambda: _make_letter(job, cv_text, cv_path, fields))
        except llm.STOPS_A_RUN as stop:
            _give_back_try(job, stop)
            raise
        except Exception as e:
            print(f"  ! Scored {job['match_score']} but letter failed for {job['company']}: {e}")
            letter_file = ""
        if letter_file:
            job["cover_letter_file"] = letter_file
            mark_ready(job)
            stats["letters"] += 1
            tag = "" if job["match_tier"] == "strong" else f"  ({job['match_tier']} – review)"
            cv_note = " + adapted CV" if job.get("tailored_cv_file") else ""
            print(f"  ✓ {job['match_score']:>3}  {job['title']} @ {job['company']}  → letter{cv_note} saved{tag}")
        else:
            print(f"  ~ {job['match_score']:>3}  {job['title']} @ {job['company']}  (letter failed)")

    return stats
