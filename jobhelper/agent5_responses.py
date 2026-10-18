"""Agent 5 — analyze application responses and build a database.

Reads inbox replies (read-only), classifies each as rejection / interview /
offer / assessment / acknowledgment, extracts a categorised rejection reason
when present, stores everything in a SQLite database (data/responses.db), and
updates the tracked job's status. Then produces analytics — response rate,
rejection-reason breakdown — as a console report and an Excel file.
"""
from __future__ import annotations

import re
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from . import config, db, email_client, llm, store
from .agent4_tracker import header_row, put

SYSTEM = (
    "You analyze emails a job applicant received in reply to their applications. "
    "You extract structured data and categorise rejections. Output only valid JSON."
)

PROMPT = """This email may be a reply to a job application. Company (guessed): "{company}".

From: {sender}
Subject: {subject}
Body:
{body}

Return JSON:
{{
  "is_response": <true if this is a reply about a specific job application; false for newsletters, job alerts, spam, unrelated mail>,
  "stage": "rejection | interview | assessment | offer | acknowledged | other",
  "company": "<the hiring company name as written, or your best guess>",
  "role": "<the role mentioned, or ''>",
  "rejection_reason": "<if a rejection and a reason is given, quote/paraphrase it briefly; else ''>",
  "rejection_category": "one of: {categories}",
  "sentiment": "positive | neutral | negative",
  "next_step": "<any next step requested of the candidate, or ''>",
  "when": "<only if the email fixes the day of an interview, call, test or deadline: 'YYYY-MM-DD HH:MM' in the time zone the email gives (leave out HH:MM if no time is given); otherwise ''>",
  "summary": "<one short sentence>"
}}
If stage is not 'rejection', set rejection_reason='' and rejection_category='no_reason_given'."""


def _ask_waiting(on_limit, call):
    """Run one model step; on a hit limit wait (if the caller can) and try again."""
    while True:
        try:
            return call()
        except llm.LimitReached as e:
            if on_limit is None:
                raise
            on_limit(e)


STAGES = ("rejection", "interview", "assessment", "offer", "acknowledged", "other")
# A date to remember that an email put on the job (not you): a newer email may move it.
EVENT_FROM_EMAIL = "email"
_TEXTS = ("company", "role", "rejection_reason", "rejection_category", "sentiment", "next_step", "summary", "when")


def tidy(result: dict[str, Any]) -> dict[str, Any]:
    """Make a model's answer about a reply safe to file: every text really a text (a model may answer null, a number
    or a list), and the stage one the app knows ("Rejection", "interview | assessment" and anything else become one
    of STAGES). A reply that was read must be filed, not end the scan. Mutates and returns `result`."""
    for key in _TEXTS:
        value = result.get(key)
        result[key] = " ".join(str(value).split()) if isinstance(value, (str, int, float)) and not isinstance(value, bool) else ""
    said = str(result.get("stage") or "").lower()
    result["stage"] = next((name for name in STAGES if re.search(rf"\b{name}\b", said)), "other")
    if result["rejection_category"] not in db.REJECTION_CATEGORIES:
        result["rejection_category"] = "no_reason_given"
    return result


def set_by_hand_since(job: dict[str, Any], email_date: str) -> bool:
    """Whether you set this job's status yourself after this email arrived. Then the email does not change it:
    a rejection from last week must not undo the interview you entered yesterday, on every scan again.
    An email newer than your change still counts, like any news."""
    from datetime import datetime
    marked = str(job.get("status_by_hand") or "")
    if not marked:
        return False
    try:
        mine = datetime.fromisoformat(marked)
    except ValueError:
        return False
    mine = mine if mine.tzinfo else mine.astimezone()
    return db.received_at(email_date) <= mine


def _file_reply(msg: dict[str, Any], result: dict[str, Any], guess: dict[str, Any] | None, stats: dict[str, Any]) -> str:
    """Store one classified reply and move its application along. Returns the stage."""
    tidy(result)
    stage = result["stage"]
    # Safety net: a mail the model filed as a plain update or confirmation but
    # whose text is an unmistakable rejection.
    worded = email_client.rejection_wording(f"{msg['subject']}\n{msg['body']}")
    if worded and stage in ("other", "acknowledged"):
        stage = "rejection"
        result["rejection_category"] = result.get("rejection_category") or worded
        if result.get("rejection_category") == "no_reason_given":
            result["rejection_category"] = worded
    db.upsert_response({
        "msg_id": msg["msg_id"],
        "received_date": msg["date"],
        "from_addr": msg["sender"],
        "subject": msg["subject"],
        "job_id": guess["id"] if guess else "",
        "company": result.get("company") or (guess["company"] if guess else ""),
        "role": result.get("role", ""),
        "stage": stage,
        "rejection_reason": result.get("rejection_reason", ""),
        "rejection_category": result.get("rejection_category", "no_reason_given"),
        "sentiment": result.get("sentiment", "neutral"),
        "next_step": result.get("next_step", ""),
        "summary": result.get("summary", ""),
    })
    stats["new_responses"] += 1
    if stage in ("interview", "assessment", "offer"):
        stats["invites"].append(str(result.get("company") or (guess["company"] if guess else "") or "a company"))
    if stage == "rejection":
        stats["rejections"] += 1
    if guess and stage in ("interview", "assessment"):
        saved = _save_date(guess, result, "interview" if stage == "interview" else "test task")
        if saved:
            stats.setdefault("dates", []).append(saved)

    # Update the matched job's status (forward-only; rejection from anywhere).
    if guess:
        cur = guess.get("status", "found")
        new_status = email_client.status_after(cur, stage)
        if new_status and not set_by_hand_since(guess, msg["date"]):
            if cur in store.LEAD_STATUSES and not guess.get("date_applied"):
                # the company wrote back, so you applied: the reply's day is the best date known
                arrived = db.received_at(msg["date"])
                guess["date_applied"] = arrived.date().isoformat() if arrived.year > 1 else date.today().isoformat()
                guess["date_applied_source"] = "first email"
            guess["status"] = new_status
            guess["last_update"] = date.today().isoformat()
            note = result["summary"]
            if note:
                guess["notes"] = (f"[{date.today().isoformat()}] {note}\n" + guess.get("notes", "")).strip()
            stats["status_updates"] += 1
    return stage


def _save_date(job: dict[str, Any], result: dict[str, Any], kind: str) -> str:
    """Put the day an email fixes (an interview, a test) on the job as its date to remember.
    A date that came from an earlier email is moved by a newer one (an interview that was put off);
    a coming date you set yourself is never overwritten. Returns "13 Oct 11:30 · Company" or ''."""
    from . import planner
    m = re.fullmatch(r"\s*(\d{4}-\d{2}-\d{2})(?:[ T](\d{1,2}:\d{2}))?.*", str(result.get("when") or ""))
    if not m:
        return ""
    day, clock = m.group(1), (m.group(2) or "")
    if clock and len(clock) == 4:
        clock = "0" + clock
    event = job.get("event") or {}
    current = event.get("date", "")
    by_email = event.get("from") == EVENT_FROM_EMAIL
    if current and current >= date.today().isoformat() and current != day and not by_email:
        return ""                                   # you have your own date on this job: it is not overwritten
    if current == day and not clock:
        clock = event.get("time", "")               # a reminder naming only the day keeps the time already known
    note = str(result.get("next_step") or result.get("summary") or "").strip()[:200]
    if planner.set_event(job, day, clock, kind, note):
        return ""                                   # a past or unreadable date: nothing saved
    job["event"]["from"] = EVENT_FROM_EMAIL         # a later email may move it; an edit by hand drops this mark
    when = date.fromisoformat(day)
    return f"{when.day} {when:%b}" + (f" {clock}" if clock else "") + f" · {job.get('company') or job.get('title', '')}"


def scan_summary(stats: dict[str, Any]) -> str:
    """One line about a finished scan: what arrived, invitations first."""
    n, invites, rejections = stats.get("new_responses", 0), len(stats.get("invites") or []), stats.get("rejections", 0)
    plural = lambda k, one, many: f"{k} {one if k == 1 else many}"                       # noqa: E731
    text = f"Scanned {plural(stats.get('scanned', 0), 'email', 'emails')}"
    if "asked" in stats:            # why "0 new" is right: most were read before, and few were worth a question
        text += f" ({stats.get('already_read', 0)} read before, {stats['asked']} sent to Claude)"
    text += f" · {n} new"
    parts = []
    if invites:
        parts.append(plural(invites, "invitation", "invitations"))
    if rejections:
        parts.append(plural(rejections, "rejection", "rejections"))
    other = n - invites - rejections
    if other > 0 and parts:
        parts.append(f"{other} other")
    text += (": " + ", ".join(parts)) if parts else ""
    text += "."
    if stats.get("dates"):
        text += " Date saved: " + "; ".join(stats["dates"]) + "."
    return text


def scan_detail(stats: dict[str, Any]) -> str:
    """How much mail a scan went through and how much of it Claude was asked about, for the Last run card: "45 · 5 asked"."""
    return f"{stats.get('scanned', 0)} · {stats.get('asked', 0)} asked"


def replies_text(stats: dict[str, Any]) -> str:
    """The same news as scan_summary, as a notification: only what arrived."""
    n = stats.get("new_responses", 0)
    if n <= 0:
        return "No new replies."
    tail = scan_summary(stats).split(" new", 1)[1]              # ": 2 invitations, 1 rejection. Date saved: …" or "."
    return ("1 new reply" if n == 1 else f"{n} new replies") + tail


_RULE_STAGES = [
    ("interview", r"\b(interview|pohovor|schedule (?:a|your) (?:call|meeting)|invite you to|book a (?:slot|time)|video call)\b"),
    ("assessment", r"\b(assessment|test task|coding (?:test|challenge)|take-home|online test|hackerrank|codility)\b"),
    ("offer", r"\b(pleased to offer|job offer|offer of employment|we would like to offer you)\b"),
    ("acknowledged", r"(received your application|thank you for (?:applying|your application)|application (?:has been|was) received|"
                     r"děkujeme za|zařadili jsme|obdrželi jsme|we will review)"),
]


MAX_PASTED = 8000        # a pasted reply longer than this is a whole page, not a message


def classify_by_rules(text: str) -> dict[str, Any]:
    """A reply's kind from its wording alone, for when Claude cannot be asked.
    Rejection wording is checked first; an unclear text is filed as 'other'."""
    worded = email_client.rejection_wording(text)
    if worded:
        return {"is_response": True, "stage": "rejection", "rejection_category": worded, "sentiment": "negative"}
    low = text.lower()
    for stage, pattern in _RULE_STAGES:
        if re.search(pattern, low):
            return {"is_response": True, "stage": stage, "sentiment": "positive" if stage != "acknowledged" else "neutral"}
    return {"is_response": True, "stage": "other", "sentiment": "neutral"}


def add_pasted_reply(jobs: list[dict[str, Any]], job_id: str, text: str, use_model: bool = True) -> dict[str, Any]:
    """File a reply you paste in by hand (a LinkedIn message, a text from a portal)
    under one application, exactly like an email found by the scan.

    Claude reads it when possible (one call); if the limit is used up, or
    use_model is False, the wording rules decide instead. Mutates `jobs`."""
    import hashlib
    from email.utils import format_datetime
    from datetime import datetime, timezone
    text = (text or "").strip()
    job = store.find(jobs, job_id)
    if not job:
        return {"ok": False, "msg": "That job is not in the list any more."}
    if len(text) < 20:
        return {"ok": False, "msg": "Paste the whole message: that is too short to read."}
    if len(text) > MAX_PASTED:
        return {"ok": False, "msg": f"That is {len(text)} characters. Paste only the message itself, not the whole page "
                                    f"(up to {MAX_PASTED} characters)."}
    msg_id = "pasted-" + hashlib.sha1((job_id + "\n" + text).encode("utf-8")).hexdigest()[:16]
    if db.has_response(msg_id):
        return {"ok": False, "msg": "This exact message was already added to this application."}
    first_line = next((ln.strip() for ln in text.splitlines() if ln.strip()), "")[:120]
    msg = {"msg_id": msg_id, "date": format_datetime(datetime.now(timezone.utc).astimezone()),
           "sender": "added by hand", "subject": first_line, "body": text}
    by_rules = not use_model
    result: dict[str, Any] = {}
    if use_model:
        try:
            result = llm.ask_json(PROMPT.format(company=job.get("company") or "unknown", sender="(pasted by the candidate)",
                                                subject=first_line, body=text[:3000],
                                                categories=" | ".join(db.REJECTION_CATEGORIES)),
                                  system=SYSTEM, tag="agent5_responses")
            if not isinstance(result, dict):
                result = {}
        except llm.LimitReached:
            by_rules = True
        except Exception:  # noqa: BLE001 - an unreadable answer: the rules still give a usable result
            by_rules = True
    if by_rules or not result.get("stage"):
        result = {**classify_by_rules(text), "summary": first_line}
        by_rules = True
    result["is_response"] = True                    # you pasted it under this application: it is a reply
    result.setdefault("company", job.get("company", ""))
    result.setdefault("role", job.get("title", ""))
    stats = {"new_responses": 0, "status_updates": 0, "rejections": 0, "invites": []}
    stage = _file_reply(msg, result, job, stats)
    db.mark_seen(msg_id)
    return {"ok": True, "stage": stage, "by_rules": by_rules, "status": job.get("status", ""),
            "summary": str(result.get("summary", ""))}


def process_inbox(jobs: list[dict[str, Any]], lookback_days: int | None = None,
                  progress=None, on_limit=None) -> dict[str, int]:
    """Scan inbox, store responses in the DB, and update job statuses.

    progress(done, total), if given, is called before each email."""
    emails = email_client.fetch_recent(lookback_days)
    stats = {"scanned": len(emails), "new_responses": 0, "status_updates": 0, "rejections": 0, "invites": []}

    stats.update(skipped_unrelated=0, already_read=0, asked=0)
    for n, msg in enumerate(emails):
        if progress:
            progress(n, len(emails))
        if db.is_seen(msg["msg_id"]):
            stats["already_read"] += 1
            continue  # already processed (response or not) — don't re-spend usage
        # Cheap pre-filter: skip obvious non-application mail without an LLM call.
        if not email_client.looks_like_application(msg["sender"], msg["subject"], msg["body"], jobs):
            db.mark_seen(msg["msg_id"])
            stats["skipped_unrelated"] += 1
            continue
        guess = email_client.match_job(msg["sender"], msg["subject"], msg["body"], jobs)
        stats["asked"] += 1
        try:
            result = _ask_waiting(on_limit, lambda: llm.ask_json(
                PROMPT.format(
                    company=guess["company"] if guess else "unknown",
                    sender=msg["sender"], subject=msg["subject"], body=msg["body"][:3000],
                    categories=" | ".join(db.REJECTION_CATEGORIES),
                ),
                system=SYSTEM,
                tag="agent5_responses",
            ))
        except llm.LimitReached:
            raise                      # not this email's fault: pause or end the scan, it stays unread
        except Exception as e:  # noqa: BLE001 - one answer that cannot be read: this email is asked again next scan
            from . import trouble
            trouble.note("reading one email with Claude", RuntimeError(type(e).__name__))
            continue

        db.mark_seen(msg["msg_id"])  # spent the call — don't reclassify next time
        if not isinstance(result, dict) or not result.get("is_response"):
            continue                   # an answer that is not what was asked for must not end the scan
        if guess and str(result.get("role") or "").strip():
            # the model read the role from the body: pick that posting among the company's, not just any of them
            guess = email_client.best_job_for_company_role(guess["company"], str(result["role"]), jobs) or guess

        _file_reply(msg, result, guess, stats)

    stats["jobs_created"] = reconcile_jobs(jobs)
    return stats


def repair_unread_bodies(jobs: list[dict[str, Any]], emails: list[dict[str, Any]]) -> dict[str, int]:
    """Fix what an older version got wrong: it read HTML-only emails as empty.

    `emails` are the messages re-read now (with their text). Without any model
    call: a stored reply filed as 'other' / 'acknowledged' whose text is a clear
    rejection becomes a rejection (and its job is marked rejected); messages
    that were skipped or filed as 'other' only because their text was missing
    are queued for the next scan. Mutates `jobs`."""
    stored = {r["msg_id"]: r for r in db.all_responses()}
    by_id = {j["id"]: j for j in jobs}
    stats = {"rejections_found": 0, "jobs_rejected": 0, "requeued": 0}
    requeue = []
    for msg in emails:
        if not msg.get("was_html_only"):
            continue
        rec = stored.get(msg["msg_id"])
        text = f"{msg['subject']}\n{msg['body']}"
        category = email_client.rejection_wording(text)
        if rec and rec["stage"] in ("other", "acknowledged") and category:
            rec.update(stage="rejection", rejection_category=category, sentiment="negative",
                       summary=("Rejection (found on re-reading the email text): "
                                + re.sub(r"\s+", " ", msg["body"])[:220]).strip())
            db.upsert_response(rec)
            stats["rejections_found"] += 1
            job = by_id.get(rec.get("job_id") or "")
            if job and job.get("status") != "rejected":
                job["status"] = "rejected"
                job["last_update"] = date.today().isoformat()
                job["notes"] = (f"[{date.today().isoformat()}] Rejection found on re-reading an earlier email.\n"
                                + (job.get("notes") or "")).strip()
                stats["jobs_rejected"] += 1
        elif rec and rec["stage"] == "other":
            requeue.append(msg["msg_id"])                 # the model saw no text: let it read it properly
        elif not rec and db.is_seen(msg["msg_id"]) and \
                email_client.looks_like_application(msg["sender"], msg["subject"], msg["body"], jobs):
            requeue.append(msg["msg_id"])                 # skipped only because the text was missing
    stats["requeued"] = db.unmark_seen(requeue)
    return stats


def reconcile_jobs(jobs: list[dict[str, Any]]) -> int:
    """Ensure every company that sent a response exists as a tracked job.

    Applications with a reply but weren't in the tracker (not seeded from a cover
    letter, not found in search) get a job record so they appear in the Excel
    with the right status. Returns how many were created.
    """
    comps: dict[str, list[dict[str, Any]]] = {}
    for r in db.all_responses():
        c = (r.get("company") or "").strip()
        if c:
            comps.setdefault(c, []).append(r)

    # Track what already exists, by id AND by normalized company name, so a
    # token-less company (e.g. "HP Inc.", "NN Group") isn't recreated every run.
    existing_ids = {j["id"] for j in jobs}
    existing_norm = {email_client.norm_company(j.get("company", "")) for j in jobs}

    created = 0
    for company, rs in comps.items():
        toks = email_client.company_tokens(company)
        norm = email_client.norm_company(company)
        new_id = f"email-{store.slugify(company)}"
        tracked = (new_id in existing_ids
                   or (norm and norm in existing_norm)
                   or (toks and any(email_client.same_company(company, j.get("company", "")) for j in jobs)))
        if tracked:
            continue  # already tracked (seeded / found / linked / created)
        stages = [x.get("stage", "other") for x in rs]
        result = max(stages, key=lambda s: _STAGE_ORDER.get(s, -1)) if stages else "other"
        role = next((x.get("role") for x in rs if x.get("role")), "")
        jobs.append(store.new_job(
            id=new_id,
            title=role or "(from email response)",
            company=company,
            source="email",
            status=email_client.STAGE_TO_STATUS.get(result, "waiting"),     # they wrote back, so it is an application
            notes="Created from an email response.",
        ))
        existing_ids.add(new_id)
        existing_norm.add(norm)
        created += 1
    return created


_GENERIC_TITLES = {"", "(existing application)", "(from email response)"}


def backfill_roles(jobs: list[dict[str, Any]]) -> int:
    """Fill in real role titles on generic-titled applications, from the role
    the model extracted from that company's response emails. Returns how many
    were updated."""
    from collections import Counter

    roles_by_company: dict[str, Counter] = {}
    for r in db.all_responses():
        role = (r.get("role") or "").strip()
        company = (r.get("company") or "").strip()
        if role and company and role.lower() not in ("unknown", "n/a"):
            roles_by_company.setdefault(company, Counter())[role] += 1

    updated = 0
    for job in jobs:
        if (job.get("title") or "").strip().lower() not in _GENERIC_TITLES:
            continue
        jtoks = email_client.company_tokens(job.get("company", ""))
        if not jtoks:
            continue
        for company, counter in roles_by_company.items():
            if email_client.same_company(job.get("company", ""), company):
                job["title"] = counter.most_common(1)[0][0]
                updated += 1
                break
    return updated


def analytics(jobs: list[dict[str, Any]]) -> dict[str, Any]:
    """Application-centric analytics from the DB + job store.

    Counts distinct applications (not individual emails), so the response rate is
    meaningful even when one application generates many reply emails.
    """
    applied = sum(1 for j in jobs if j.get("status") in store.APPLICATION_STATUSES)
    by_stage = db.counts_by("stage")
    response_emails = db.total()

    rows = progress_rows(jobs)
    companies_replied = sum(1 for r in rows if r["result"])
    apps_with_interview = sum(1 for r in rows if "interview" in r["path"] or "assessment" in r["path"])
    apps_rejected = sum(1 for r in rows if r["result"] == "rejection")
    apps_offer = sum(1 for r in rows if r["result"] == "offer")
    base = max(applied, companies_replied)  # avoid >100% when `applied` is undercounted

    statuses = {s: sum(1 for j in jobs if j.get("status") == s) for s in ("waiting", "interview", "offer", "rejected", "on_hold")}
    by_hand = sum(1 for j in jobs if j.get("status") == "rejected" and "rejected by you" in (j.get("notes") or ""))
    return {
        "applications_by_status": statuses,           # these add up to applications_tracked
        "rejected_by_hand": by_hand,                  # rejections you entered (no email in the mailbox)
        "applications_tracked": applied,
        "companies_replied": companies_replied,
        "response_rate": round(100 * companies_replied / base, 1) if base else 0.0,
        "response_emails": response_emails,
        "by_stage": by_stage,
        "applications_with_interview": apps_with_interview,
        "applications_rejected": apps_rejected,
        "applications_with_offer": apps_offer,
        "rejection_rate": round(100 * apps_rejected / companies_replied, 1) if companies_replied else 0.0,
        "rejection_reasons": db.rejection_reason_counts(),
    }


def print_report(jobs: list[dict[str, Any]]) -> None:
    a = analytics(jobs)
    print("── Application responses ──")
    print(f"Applications with a reply : {a['companies_replied']}")
    print(f"  → interviews reached : {a['applications_with_interview']}")
    print(f"  → offers             : {a['applications_with_offer']}")
    print(f"  → rejections         : {a['applications_rejected']}")
    print(f"Response emails logged : {a['response_emails']}")
    if a["rejection_reasons"]:
        print(f"Why rejected  (rejection rate {a['rejection_rate']}% of replies):")
        for reason, n in a["rejection_reasons"].items():
            print(f"   {reason:<26} {n}")


_STAGE_ORDER = {"acknowledged": 0, "assessment": 1, "interview": 2, "offer": 3,
                "rejection": 3, "other": -1}


def progress_rows(jobs: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """One row per application (company + role): the stage path + final result.

    Grouping by company AND role keeps different positions at the same company
    separate. The path collapses repeated emails about the same stage so it
    reflects the real progression, not email volume.
    """
    groups: dict[str, dict[str, Any]] = {}
    for r in sorted(db.all_responses(), key=lambda x: db.received_at(x.get("received_date", ""))):
        company = (r.get("company") or "").strip()
        role = (r.get("role") or "").strip()
        # Group all emails linked to the same application together; otherwise fall
        # back to company + role (keeps different positions at one company apart).
        key = f"job:{r['job_id']}" if r.get("job_id") else f"{company.lower()}|{role.lower()}"
        g = groups.setdefault(key, {
            "company": company, "role": role, "job_id": r.get("job_id", ""),
            "stages": [], "next_step": "", "first": r.get("received_date", ""),
            "last": r.get("received_date", ""),
        })
        stage = r.get("stage", "other")
        if stage != "other":
            g["stages"].append(stage)
        g["last"] = r.get("received_date", "") or g["last"]
        if r.get("next_step"):
            g["next_step"] = r["next_step"]
        if not g["job_id"] and r.get("job_id"):
            g["job_id"] = r["job_id"]

    jobs_by_id = {j["id"]: j for j in jobs}
    rows = []
    for g in groups.values():
        job = jobs_by_id.get(g["job_id"]) or email_client.job_of_company(g["company"], jobs)
        result = g["stages"][-1] if g["stages"] else ""  # most recent stage = current outcome
        rows.append({
            "job_id": job["id"] if job else "",
            "company": g["company"] or (job["company"] if job else ""),
            "role": g["role"] or (job["title"] if job else ""),
            "status": job["status"] if job else result or "",
            "path": _numbered_path(g["stages"]),
            "result": result,
            "next_step": g["next_step"],
            "first_contact": g["first"],
            "last_update": g["last"],
        })
    # by the day itself: these are dates as emails write them ("Fri, 9 Oct 2026 …"), which as text sort by the weekday's name
    return sorted(rows, key=lambda x: db.received_at(x["last_update"]), reverse=True)


def _numbered_path(stages: list[str]) -> str:
    """Render the stage progression. Consecutive duplicates (many emails about the
    same stage) are collapsed; genuinely separate repeats are numbered. E.g.
    [ack, interview, interview, rejection] → 'acknowledged → interview → rejection';
    [interview, acknowledged, interview] → 'interview #1 → acknowledged → interview #2'."""
    if not stages:
        return "(no response yet)"
    from collections import Counter
    collapsed: list[str] = []
    for s in stages:
        if not collapsed or collapsed[-1] != s:
            collapsed.append(s)
    totals = Counter(collapsed)
    seen: dict[str, int] = {}
    parts = []
    for s in collapsed:
        if totals[s] > 1:
            seen[s] = seen.get(s, 0) + 1
            parts.append(f"{s} #{seen[s]}")
        else:
            parts.append(s)
    return " → ".join(parts)


def backfill_apply_dates(jobs: list[dict[str, Any]]) -> int:
    """Give applications an apply date from their emails where none is known.

    Used when the date is empty, or is only the day a letter was imported (not a
    real date). The date of the company's FIRST email is taken: when that email
    is the confirmation it is the apply day; when the first email kept is a later
    reply the true day was earlier, so the record is marked "first email" either
    way. A date you entered yourself is never changed. Returns how many were set."""
    first: dict[str, Any] = {}
    for r in db.all_responses():
        when = db.received_at(r.get("received_date", ""))
        jid = r.get("job_id") or ""
        if jid and when.year > 1 and (jid not in first or when < first[jid]):
            first[jid] = when
    done = 0
    for job in jobs:
        if job.get("status") not in store.APPLICATION_STATUSES or job["id"] not in first:
            continue
        imported = store.applied_on_import(job)
        from_email = job.get("date_applied_source") == "first email"
        if job.get("date_applied") and not imported and not from_email:
            continue                                   # a real date: leave it
        day = first[job["id"]].date().isoformat()
        if job.get("date_applied") != day or not from_email:
            job["date_applied"], job["date_applied_source"] = day, "first email"
            done += 1
    return done


def annotate_progress(jobs: list[dict[str, Any]]) -> None:
    """Attach each application's stage path + result onto the job records so the
    main tracker (applications.xlsx) can show all steps, not just the status."""
    from . import merge
    merge.auto_merge(jobs)             # one application that arrived by two routes becomes one row
    backfill_apply_dates(jobs)
    jobs_by_id = {j["id"]: j for j in jobs}
    for row in progress_rows(jobs):
        target = jobs_by_id.get(row["job_id"]) or email_client.job_of_company(row["company"], jobs)
        if target is not None:
            target["stages_path"] = row["path"]
            target["result"] = row["result"]
            target["next_step"] = row["next_step"]
            # Keep the job's status in sync with the latest response stage
            # (forward-only; rejection can apply from anywhere). This is the
            # safety net that moves an applied lead into applications.
            new_status = email_client.status_after(target.get("status", "found"), row["result"])
            if new_status and not set_by_hand_since(target, row["last_update"]):
                target["status"] = new_status


def build_report(jobs: list[dict[str, Any]]) -> None:
    """Write data/responses_report.xlsx: Summary + Progress + raw Responses sheets."""
    a = analytics(jobs)
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    bold = Font(bold=True)
    ws["A1"], ws["A1"].font = "Metric", bold
    ws["B1"], ws["B1"].font = "Value", bold
    rows = [
        ("Applications with a reply", a["companies_replied"]),
        ("Applications reaching interview", a["applications_with_interview"]),
        ("Offers", a["applications_with_offer"]),
        ("Applications rejected", a["applications_rejected"]),
        ("Rejection rate % (of replies)", a["rejection_rate"]),
        ("Response emails logged", a["response_emails"]),
        ("", ""),
        ("Response emails by stage", ""),
        *[(f"  {k}", v) for k, v in a["by_stage"].items()],
        ("", ""),
        ("Why rejected", ""),
        *[(f"  {k}", v) for k, v in a["rejection_reasons"].items()],
    ]
    for i, (k, v) in enumerate(rows, start=2):
        put(ws, i, 1, k)
        put(ws, i, 2, v)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 14

    # Interview progress: one row per application with its stage path + result.
    wsp = wb.create_sheet("Progress")
    p_headers = ["company", "role", "status", "path", "result", "next_step",
                 "first_contact", "last_update"]
    header_row(wsp, [(h, 26 if h in ("path", "next_step") else 18) for h in p_headers])
    for r, row in enumerate(progress_rows(jobs), start=2):
        for c, h in enumerate(p_headers, start=1):
            put(wsp, r, c, row.get(h, ""))

    ws2 = wb.create_sheet("Responses")
    headers = ["received_date", "company", "role", "stage", "rejection_category",
               "rejection_reason", "sentiment", "next_step", "summary", "from_addr", "job_id"]
    header_row(ws2, [(h, 22) for h in headers])
    for r, resp in enumerate(db.all_responses(), start=2):
        for c, h in enumerate(headers, start=1):
            put(ws2, r, c, resp.get(h, ""))

    config.ensure_dirs()
    wb.save(str(config.RESPONSES_REPORT))
